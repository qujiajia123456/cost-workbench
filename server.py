from __future__ import annotations

import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import threading
import time
import uuid
import zipfile
from collections import Counter
from copy import copy
from dataclasses import dataclass
from email.parser import BytesParser
from email.policy import default
from datetime import datetime, timedelta
from math import ceil, floor
from calendar import monthrange
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, quote, urlparse
from urllib.request import urlopen

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
STATIC_DIR = ROOT / "static"
DATA_DIR = ROOT / "data"
UPLOAD_DIR = DATA_DIR / "uploads"
EXPORT_DIR = DATA_DIR / "exports"
CODEX_JOB_DIR = DATA_DIR / "codex_material_jobs"
BLACKLIST_CODEX_JOB_DIR = DATA_DIR / "codex_blacklist_jobs"
BLACKLIST_RESULT_DIR = DATA_DIR / "blacklist_results"
BLACKLIST_TYC_WORKER_SCRIPT = ROOT / "tools" / "blacklist_tianyancha_worker.mjs"
BID_SOURCE_DIR = DATA_DIR / "bid_price_sources"
PRICE_LIBRARY_SOURCE_DIR = DATA_DIR / "price_library_sources"
PM_LABOR_PRICE_TASK_DIR = Path.home() / "Documents" / "大pm平台所有劳务合同的清单项进行汇总提取，包括项目名称、班组名称、工序、单位、单价、施工内容、付款条件、质保金、备注等，目的是为了做价格库，从劳务分包模块中履约管理中的劳务分包合同登记中查找所有的合同，还有施工时间等信息，其他重要信息你一块提取"
PM_PROFESSIONAL_PRICE_TASK_DIRS = [
    Path.home() / "Documents" / "分包价格",
    Path.home() / "Documents" / "大pm平台所有专业分包合同的清单项进行汇总提取，包括项目名称、分包商名称、工序、单位、单价、施工内容、付款条件、质保金、备注等，目的是为了做价格库，从专业分包模块中履约管理中的专业分包合同登记中查找所有的合同，还有施工时间等信息，其他重要信息你一块提取",
]
PM_WARNING_JOB_DIR = DATA_DIR / "pm_warning_jobs"
PM_WARNING_RESULT_DIR = DATA_DIR / "pm_warning_results"
PM_WARNING_LOG_DIR = DATA_DIR / "pm_warning_logs"
DB_PATH = DATA_DIR / "db.json"
PRICE_LIBRARY_DB_PATH = DATA_DIR / "price_library.db"
SETTLEMENT_RECEIVE_LEDGER_PATH = DATA_DIR / "settlement_receive_ledger.json"
SETTLEMENT_OVERDUE_MONTH_LEDGER_PATH = DATA_DIR / "settlement_overdue_month_ledger.json"
TEMPLATE_DIR = ROOT / "templates"
NODE_RUNTIME = Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime" / "dependencies" / "node" / "bin" / ("node.exe" if sys.platform.startswith("win") else "node")
PM_WARNING_TEMPLATE_PATH = TEMPLATE_DIR / "pm-warning-template.xls"
RUNTIME_NODE = Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime" / "dependencies" / "node" / "bin" / "node.exe"
RUNTIME_NODE_MODULES = Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime" / "dependencies" / "node" / "node_modules"
PM_WARNING_CDP_LIST_URL = "http://localhost:9222/json/list"
PM_WARNING_BIG_PM_CDP_LIST_URL = "http://localhost:9333/json/list"
PM_WARNING_OLD_PM_URL = "http://218.56.43.116:2020/yjpm2012/MainNew.aspx"
PM_WARNING_BIG_PM_URL = "http://yanjianpm.glodon.com/Portal/Frame/LayoutC/Default.aspx"
# PM 平台只在用户明确需要抓数时提示连接，启动工作台不得自动打开外部平台窗口。
PM_WARNING_AUTO_LAUNCH_BROWSER = False
PM_WARNING_START_DEBUG_BROWSERS_ON_8899 = False
PM_WARNING_START_DEBUG_PLATFORMS_ON_8899: tuple[str, ...] = ()
PM_WARNING_ACTIVE_STATUSES = {"等待抓数", "连接四版平台", "连接大PM平台", "抓取中", "抓取大PM汇总", "补金额明细", "补责任书目标"}
PM_WARNING_DEBUG_BROWSER_BOOT_LOCK = threading.Lock()
PM_WARNING_DEBUG_BROWSER_BOOTING: set[str] = set()

DEFAULT_BLACKLIST_ENTITIES = [
    {
        "id": "demo-black-001",
        "name": "烟建示例失信工程有限公司",
        "creditCode": "91370000DEMO0001X",
        "riskLevel": "黑名单",
        "status": "限制合作",
        "legalRep": "王某某",
        "shareholders": ["王某某", "李某某"],
        "contacts": ["0535-0000001", "13800000001"],
        "bankAccounts": ["370000000000000001"],
        "address": "山东省烟台市示例路1号",
        "reason": "履约严重违约、结算资料长期拒不配合",
        "source": "系统示例数据",
        "updatedAt": "2026-07-01",
    },
    {
        "id": "demo-related-001",
        "name": "烟建示例建材有限公司",
        "creditCode": "91370000DEMO0002X",
        "riskLevel": "关联关注",
        "status": "需复核",
        "legalRep": "王某某",
        "shareholders": ["赵某某", "烟建示例失信工程有限公司"],
        "contacts": ["0535-0000002"],
        "bankAccounts": ["370000000000000002"],
        "address": "山东省烟台市示例路9号",
        "reason": "与黑名单企业存在同法人及股权关联",
        "source": "系统示例数据",
        "updatedAt": "2026-07-01",
    },
    {
        "id": "demo-related-002",
        "name": "鲁东测试劳务有限公司",
        "creditCode": "91370000DEMO0003X",
        "riskLevel": "关联关注",
        "status": "需复核",
        "legalRep": "刘某某",
        "shareholders": ["刘某某"],
        "contacts": ["13800000001"],
        "bankAccounts": ["370000000000000001"],
        "address": "山东省烟台市测试街6号",
        "reason": "与黑名单企业存在同电话及同银行账户关联",
        "source": "系统示例数据",
        "updatedAt": "2026-07-01",
    },
]

COMPANY_NAME_RULES = [
    ("格瑞特", "格瑞特公司"),
    ("国际", "国际公司"),
    ("济南", "济南公司"),
    ("天津", "济南公司"),
    ("六公司", "六公司"),
    ("第六", "六公司"),
    ("市政路桥", "市政路桥公司"),
    ("四公司", "四公司"),
    ("第四", "四公司"),
    ("七公司", "七公司"),
    ("第七", "七公司"),
    ("三公司", "三公司"),
    ("第三", "三公司"),
    ("五公司", "五公司"),
    ("第五", "五公司"),
    ("设备安装", "设备安装公司"),
    ("安装公司", "设备安装公司"),
    ("青岛", "青岛公司"),
    ("十公司", "十公司"),
    ("第十", "十公司"),
    ("装饰幕墙", "装饰幕墙公司"),
    ("上海", "上海公司"),
    ("马来", "马来公司"),
]

REPORT_COMPANY_ORDER = {
    "risk-level": [
        "三公司",
        "四公司",
        "五公司",
        "六公司",
        "七公司",
        "十公司",
        "市政路桥公司",
        "格瑞特公司",
        "青岛公司",
        "济南公司",
        "设备安装公司",
        "国际公司",
    ],
    "settlement": [
        "三公司",
        "四公司",
        "五公司",
        "六公司",
        "七公司",
        "十公司",
        "青岛公司",
        "济南公司",
        "上海公司",
        "格瑞特公司",
        "市政路桥公司",
        "装饰幕墙公司",
        "设备安装公司",
        "国际公司",
        "马来公司",
    ],
}


def ensure_dirs() -> None:
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    CODEX_JOB_DIR.mkdir(parents=True, exist_ok=True)
    BLACKLIST_CODEX_JOB_DIR.mkdir(parents=True, exist_ok=True)
    BLACKLIST_RESULT_DIR.mkdir(parents=True, exist_ok=True)
    PM_WARNING_JOB_DIR.mkdir(parents=True, exist_ok=True)
    PM_WARNING_RESULT_DIR.mkdir(parents=True, exist_ok=True)
    PM_WARNING_LOG_DIR.mkdir(parents=True, exist_ok=True)
    PRICE_LIBRARY_SOURCE_DIR.mkdir(parents=True, exist_ok=True)
    if not DB_PATH.exists():
        save_db({"uploads": []})
    init_price_library_db()


def load_db() -> dict[str, Any]:
    ensure_dirs()
    with DB_PATH.open("r", encoding="utf-8-sig") as f:
        db = json.load(f)
    db.setdefault("uploads", [])
    db.setdefault("issues", [])
    db.setdefault("materialPriceTasks", [])
    db.setdefault("priceLibraryRows", [])
    db.setdefault("blacklistEntities", DEFAULT_BLACKLIST_ENTITIES)
    db.setdefault("blacklistCodexJobs", [])
    return db


def save_db(db: dict[str, Any]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    tmp = DB_PATH.with_suffix(".tmp")
    with tmp.open("w", encoding="utf-8") as f:
        json.dump(db, f, ensure_ascii=False, indent=2)
    tmp.replace(DB_PATH)


def load_settlement_receive_ledger() -> dict[str, Any]:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not SETTLEMENT_RECEIVE_LEDGER_PATH.exists():
        return {"entries": {}}
    try:
        with SETTLEMENT_RECEIVE_LEDGER_PATH.open("r", encoding="utf-8-sig") as f:
            ledger = json.load(f)
    except (OSError, json.JSONDecodeError):
        return {"entries": {}}
    if not isinstance(ledger, dict):
        return {"entries": {}}
    entries = ledger.get("entries")
    if not isinstance(entries, dict):
        ledger["entries"] = {}
    return ledger


def save_settlement_receive_ledger(ledger: dict[str, Any]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    tmp = SETTLEMENT_RECEIVE_LEDGER_PATH.with_suffix(".tmp")
    with tmp.open("w", encoding="utf-8") as f:
        json.dump(ledger, f, ensure_ascii=False, indent=2)
    tmp.replace(SETTLEMENT_RECEIVE_LEDGER_PATH)


def load_settlement_overdue_month_ledger() -> dict[str, Any]:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not SETTLEMENT_OVERDUE_MONTH_LEDGER_PATH.exists():
        return {"months": {}, "manualQuarters": {}}
    try:
        with SETTLEMENT_OVERDUE_MONTH_LEDGER_PATH.open("r", encoding="utf-8-sig") as f:
            ledger = json.load(f)
    except (OSError, json.JSONDecodeError):
        return {"months": {}, "manualQuarters": {}}
    if not isinstance(ledger, dict):
        return {"months": {}, "manualQuarters": {}}
    if not isinstance(ledger.get("months"), dict):
        ledger["months"] = {}
    if not isinstance(ledger.get("manualQuarters"), dict):
        ledger["manualQuarters"] = {}
    return ledger


def save_settlement_overdue_month_ledger(ledger: dict[str, Any]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ledger.setdefault("months", {})
    ledger.setdefault("manualQuarters", {})
    tmp = SETTLEMENT_OVERDUE_MONTH_LEDGER_PATH.with_suffix(".tmp")
    with tmp.open("w", encoding="utf-8") as f:
        json.dump(ledger, f, ensure_ascii=False, indent=2)
    tmp.replace(SETTLEMENT_OVERDUE_MONTH_LEDGER_PATH)


def settlement_receive_ledger_key(row: dict[str, Any]) -> str:
    company = standard_company_name(row.get("company", "")) or clean_text(row.get("company", ""))
    project = clean_text(row.get("project", ""))
    contract_amount = round(to_number(row.get("contract_amount")), 6)
    settled_date = clean_text(row.get("settled_date", ""))
    settled_amount = round(to_number(row.get("settled_amount")), 6)
    if not company or not project:
        return ""
    return "|".join([company, project, f"{contract_amount:.6f}", settled_date, f"{settled_amount:.6f}"])


def json_default(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value)


def parse_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    if not text:
        return None
    text = text.replace(".", "-").replace("/", "-")
    match = re.search(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if match:
        y, m, d = match.groups()
        return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
    return text


def period_cutoff(period: str) -> str:
    match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
    if not match:
        return "9999-12-31"
    year, month = int(match.group(1)), int(match.group(2))
    return f"{year:04d}-{month:02d}-{monthrange(year, month)[1]:02d}"


def settlement_abc_cutoff(period: str) -> str:
    match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
    if not match:
        return period_cutoff(period)
    return f"{int(match.group(1)):04d}-06-30"


def period_start(period: str) -> str:
    match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
    if not match:
        return "0000-01-01"
    year, month = int(match.group(1)), int(match.group(2))
    return f"{year:04d}-{month:02d}-01"


def period_report_date(period: str) -> str:
    match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
    if not match:
        return "9999-12-31"
    year, month = int(match.group(1)), int(match.group(2))
    return f"{year:04d}-{month:02d}-20"


def report_window_start(period: str) -> str:
    match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
    if not match:
        return "0000-01-01"
    year, month = int(match.group(1)), int(match.group(2))
    if month == 1:
        return f"{year - 1:04d}-12-20"
    return f"{year:04d}-{month - 1:02d}-20"


def report_window_begin(period: str) -> str:
    start = parse_date(report_window_start(period))
    if not start:
        return report_window_start(period)
    start_date = datetime.strptime(start, "%Y-%m-%d") + timedelta(days=1)
    return start_date.strftime("%Y-%m-%d")


def report_window_label(period: str) -> str:
    """Return the monthly settlement window shown in report headers, e.g. 6.21—7.20."""
    begin = parse_date(report_window_begin(period))
    end = parse_date(period_report_date(period))
    if not begin or not end:
        return "本月统计期"
    begin_date = datetime.strptime(begin, "%Y-%m-%d")
    end_date = datetime.strptime(end, "%Y-%m-%d")
    return f"{begin_date.month}.{begin_date.day}—{end_date.month}.{end_date.day}"


def previous_report_period(period: str) -> str:
    match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
    if not match:
        return ""
    year, month = int(match.group(1)), int(match.group(2))
    if month == 1:
        return f"{year - 1:04d}-12"
    return f"{year:04d}-{month - 1:02d}"


def row_is_within_report_cutoff(row: dict[str, Any], period: str) -> bool:
    """Exclude only rows with an explicit future completion/settlement date."""
    report_date = period_report_date(period)
    is_settled = row.get("status") == "已定案" or row.get("segment") == "settled"
    if is_settled:
        settled_date = parse_date(row.get("settled_date"))
        return not settled_date or settled_date <= report_date
    completion_date = parse_date(row.get("completion_date"))
    return not completion_date or completion_date <= report_date


def classify_settlement_abc(row: dict[str, Any], cutoff: str) -> str:
    if row.get("segment") in {"在建", "未开工", "诉讼", "停工"}:
        return ""
    status = row.get("status")
    segment = row.get("segment")
    title = clean_text(row.get("section_title", ""))
    completion_date = row.get("completion_date") or ""

    if status == "已定案" or segment == "settled":
        return "A"
    if completion_date:
        return "C" if completion_date > cutoff else "B"
    if any(text in title for text in ["7月1日", "以后", "之后", "7.1"]):
        return "C"
    if any(text in title for text in ["6月30日", "以前", "之前", "6.30"]):
        return "B"
    return "B"


def include_settlement_summary_row(row: dict[str, Any], period: str) -> bool:
    if row.get("segment") in {"在建", "未开工", "诉讼", "停工"}:
        return False
    if row.get("status") == "已定案":
        return True
    report_date = period_report_date(period)
    submit_due_date = row.get("submit_due_date") or ""
    actual_submit_date = row.get("actual_submit_date") or ""
    submitted_amount = row.get("submitted_amount", 0) or 0
    if submit_due_date and submit_due_date > report_date and not actual_submit_date and not submitted_amount:
        return False
    return True


def counts_as_settlement_project(row: dict[str, Any]) -> bool:
    if clean_text(row.get("project")) == "/":
        return False
    return row.get("count_project", True)


def unique_settlement_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    unique_rows: list[dict[str, Any]] = []
    seen: set[tuple[Any, ...]] = set()
    for row in rows:
        key = (
            standard_company_name(row.get("company", "")) or row.get("company", ""),
            clean_text(row.get("project", "")),
            row.get("contract_amount", 0) or 0,
            row.get("completion_date") or "",
            row.get("settled_date") or "",
            row.get("settled_amount", 0) or 0,
            row.get("status") or "",
            row.get("segment") or "",
        )
        if key in seen:
            continue
        seen.add(key)
        unique_rows.append(row)
    return unique_rows


def to_number(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("，", "")
    range_match = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*[-~－—]\s*(\d+(?:\.\d+)?)\s*", text)
    if range_match:
        low, high = (float(part) for part in range_match.groups())
        return (low + high) / 2
    text = re.sub(r"[^\d.\-]", "", text)
    if not text or text in {"-", ".", "-."}:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def to_formula_number(value: Any) -> float:
    text = clean_text(value)
    if not text.startswith("="):
        return to_number(value)
    expression = text[1:].replace(",", "").replace("，", "")
    if not re.fullmatch(r"[0-9+\-*/().\s]+", expression):
        return 0.0
    try:
        result = eval(expression, {"__builtins__": {}}, {})
    except Exception:
        return 0.0
    return float(result) if isinstance(result, (int, float)) else 0.0


def classify_risk_level(score: float) -> str:
    if score >= 80:
        return "A"
    if score >= 60:
        return "B"
    if score >= 40:
        return "C"
    return "D"
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("，", "")
    text = re.sub(r"[^\d.\-]", "", text)
    if not text or text in {"-", ".", "-."}:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, dict):
        for key in ("value", "text", "name", "label", "key"):
            if value.get(key) not in (None, ""):
                return clean_text(value.get(key))
        return ""
    text = str(value).strip()
    object_value = re.search(r"['\"]value['\"]\s*:\s*['\"]([^'\"]+)['\"]", text)
    if object_value:
        return object_value.group(1).strip()
    return text


def standard_company_name(text: str) -> str:
    cleaned = clean_text(text)
    if not cleaned:
        return ""
    cleaned = re.sub(r"^[*_·\-\s]+|[*_·\-\s]+$", "", cleaned)
    cleaned = cleaned.replace("（", "(").replace("）", ")")
    for keyword, company in COMPANY_NAME_RULES:
        if keyword in cleaned:
            return company
    return cleaned


def company_from_filename(filename: str) -> str:
    stem = Path(filename).stem
    if "_" in stem and re.match(r"^[0-9a-fA-F]{16,}_", stem):
        stem = stem.split("_", 1)[1]
    company = standard_company_name(stem)
    if company == stem and any(text in stem for text in ["评分汇总表", "风险等级", "竣工结算", "定案情况"]):
        return ""
    return company


def convert_xls_to_xlsx(source_path: Path) -> Path:
    if source_path.suffix.lower() == ".xlsx":
        return source_path
    if source_path.suffix.lower() != ".xls":
        raise ValueError("仅支持 .xlsx 和 .xls 文件。")

    target_path = source_path.with_suffix(".converted.xlsx")
    script = r"""
$source = $env:REPORT_HUB_XLS_SOURCE
$target = $env:REPORT_HUB_XLS_TARGET
$excel = $null
$workbook = $null
try {
  $excel = New-Object -ComObject Excel.Application
  $excel.Visible = $false
  $excel.DisplayAlerts = $false
  $workbook = $excel.Workbooks.Open($source)
  $workbook.SaveAs($target, 51)
} finally {
  if ($workbook -ne $null) { $workbook.Close($false) | Out-Null }
  if ($excel -ne $null) {
    $excel.Quit()
    [System.Runtime.InteropServices.Marshal]::ReleaseComObject($excel) | Out-Null
  }
}
"""
    env = dict(**__import__("os").environ)
    env["REPORT_HUB_XLS_SOURCE"] = str(source_path)
    env["REPORT_HUB_XLS_TARGET"] = str(target_path)
    result = subprocess.run(
        ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        env=env,
        timeout=120,
    )
    if result.returncode != 0 or not target_path.exists():
        message = (result.stderr or result.stdout or "转换失败").strip()
        raise RuntimeError(f".xls 转换 .xlsx 失败，请确认本机已安装 Excel。{message}")
    return target_path


def parse_multipart_form(headers: Any, body: bytes) -> dict[str, Any]:
    content_type = headers.get("Content-Type", "")
    if "multipart/form-data" not in content_type or "boundary=" not in content_type:
        return {}
    raw = (
        f"Content-Type: {content_type}\r\n"
        "MIME-Version: 1.0\r\n\r\n"
    ).encode("utf-8") + body
    message = BytesParser(policy=default).parsebytes(raw)
    fields: dict[str, Any] = {}
    for part in message.iter_parts():
        disposition = part.get("Content-Disposition", "")
        if "form-data" not in disposition:
            continue
        name = part.get_param("name", header="Content-Disposition")
        filename = part.get_param("filename", header="Content-Disposition")
        payload = part.get_payload(decode=True) or b""
        if not name:
            continue
        if filename is not None:
            file_value = {"filename": filename, "content": payload}
            if name in fields:
                if not isinstance(fields[name], list):
                    fields[name] = [fields[name]]
                fields[name].append(file_value)
            else:
                fields[name] = file_value
        else:
            fields[name] = payload.decode(part.get_content_charset() or "utf-8", errors="replace")
    return fields


def parse_json_body(headers: Any, body: bytes) -> dict[str, Any]:
    if not body:
        return {}
    try:
        return json.loads(body.decode(headers.get_content_charset() or "utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError):
        return {}


def expected_companies() -> list[str]:
    companies: list[str] = []
    for _keyword, company in COMPANY_NAME_RULES:
        if company not in companies:
            companies.append(company)
    return companies


def company_order(report_type: str) -> list[str]:
    return REPORT_COMPANY_ORDER.get(report_type, expected_companies())


def company_sort_key(company: str, report_type: str) -> tuple[int, str]:
    standard = standard_company_name(company)
    order = company_order(report_type)
    if standard in order:
        return (order.index(standard), "")
    return (len(order), standard)


def looks_like_header(row: tuple[Any, ...], keywords: list[str]) -> bool:
    text = "|".join(clean_text(cell) for cell in row)
    return sum(1 for keyword in keywords if keyword in text) >= 2


def row_is_detail(row: tuple[Any, ...]) -> bool:
    first = clean_text(row[0] if row else "")
    second = clean_text(row[1] if len(row) > 1 else "")
    if not second:
        return False
    if any(word in first + second for word in ["小计", "合计", "总计", "序号"]):
        return False
    return bool(re.match(r"^\d+(\.\d+)?$", first))


def formula_row_refs(formula: Any, column_letter: str) -> set[int]:
    if not isinstance(formula, str) or not formula.startswith("="):
        return set()
    refs: set[int] = set()
    col = re.escape(column_letter.upper())
    for start, end in re.findall(rf"\b{col}(\d+)\s*:\s*{col}(\d+)\b", formula.upper()):
        refs.update(range(int(start), int(end) + 1))
    for row_index in re.findall(rf"\b{col}(\d+)\b", formula.upper()):
        refs.add(int(row_index))
    return refs


@dataclass
class ParseResult:
    rows: list[dict[str, Any]]
    errors: list[str]
    warnings: list[str]
    sheet_names: list[str]


class ReportPlugin:
    key = ""
    name = ""
    description = ""

    def parse(self, file_path: Path, company: str) -> ParseResult:
        raise NotImplementedError

    def aggregate(self, rows: list[dict[str, Any]]) -> dict[str, Any]:
        raise NotImplementedError

    def export(self, rows: list[dict[str, Any]], output_path: Path, period: str) -> None:
        raise NotImplementedError


class SettlementPlugin(ReportPlugin):
    key = "settlement"
    name = "竣工结算报表"
    description = "按分公司明细重新计算定案、未定案与结算造价。"

    def parse(self, file_path: Path, company: str) -> ParseResult:
        wb = load_workbook(file_path, data_only=True)
        formula_wb = load_workbook(file_path, data_only=False)
        rows: list[dict[str, Any]] = []
        errors: list[str] = []
        warnings: list[str] = []
        for ws in self.source_sheets(wb, company):
            formula_ws = formula_wb[ws.title] if ws.title in formula_wb.sheetnames else ws
            columns = self.detail_columns(ws)
            counted_rows = self.counted_project_rows(formula_ws, columns)
            sheet_company = company or self.company_from_sheet(ws)
            current_segment = ""
            current_title = ""
            pending: list[dict[str, Any]] = []

            def flush_pending(segment: str | None = None) -> None:
                nonlocal pending, current_segment
                final_segment = segment or current_segment
                if final_segment == "ignored":
                    pending = []
                    return
                for detail in pending:
                    settled_amount = detail.get("settled_amount", 0) or 0
                    settled_date = detail.get("settled_date")
                    if final_segment in {"在建", "未开工", "诉讼", "停工"}:
                        status = final_segment
                    elif final_segment == "settled":
                        status = "已定案"
                    elif final_segment == "unsettled":
                        status = "未定案"
                    else:
                        status = "已定案" if settled_amount or settled_date else "未定案"
                    detail["segment"] = final_segment
                    detail["status"] = status
                    detail["abc_segment"] = classify_settlement_abc(detail, "2026-06-30")
                    rows.append(detail)
                pending = []

            for excel_cells in ws.iter_rows():
                source_row = excel_cells[0].row if excel_cells else 0
                row = tuple(cell.value for cell in excel_cells)
                row_text = " ".join(clean_text(cell) for cell in row if cell is not None)
                is_detail_row = row_is_detail(row)
                special_segment = ""
                if not is_detail_row:
                    for title, segment in [
                        ("在建工程", "在建"),
                        ("在建项目", "在建"),
                        ("未开工工程", "未开工"),
                        ("未开工项目", "未开工"),
                        ("诉讼工程", "诉讼"),
                        ("诉讼项目", "诉讼"),
                        ("停工工程", "停工"),
                        ("停工项目", "停工"),
                    ]:
                        if title in row_text:
                            special_segment = segment
                            break
                if special_segment:
                    flush_pending()
                    current_segment = special_segment
                    current_title = row_text
                    continue
                if "截止到本月竣工工程合计" in row_text:
                    flush_pending()
                    current_segment = "在建"
                    current_title = "在建项目"
                    continue
                if "定案造价小计" in row_text and "未定案" not in row_text:
                    flush_pending("settled")
                    current_segment = "settled"
                    continue
                if "未定案造价小计" in row_text:
                    flush_pending("unsettled")
                    current_segment = "unsettled"
                    continue
                if "未定案工程" in row_text:
                    flush_pending()
                    current_segment = "unsettled"
                    current_title = row_text
                    continue
                if "以下工程" in row_text and "定案工程" in row_text and "未定案" not in row_text:
                    flush_pending()
                    current_segment = "settled"
                    current_title = row_text
                    continue
                if current_segment == "ignored":
                    continue
                if not is_detail_row:
                    continue
                project = clean_text(row[columns["project"]] if len(row) > columns["project"] else "")
                if not project or project in {"工程名称", "项目名称"}:
                    continue
                settled_amount = to_number(row[columns["settled_amount"]] if len(row) > columns["settled_amount"] else None)
                if not settled_amount and source_row:
                    settled_formula_value = formula_ws.cell(source_row, columns["settled_amount"] + 1).value
                    settled_amount = to_formula_number(settled_formula_value)
                settled_date = parse_date(row[columns["settled_date"]] if len(row) > columns["settled_date"] else None)
                detail = {
                    "company": sheet_company,
                    "sheet": ws.title,
                    "source_row": source_row,
                    "project": project,
                    "building_area": clean_text(row[columns["building_area"]] if len(row) > columns["building_area"] else ""),
                    "contract_amount": to_number(row[columns["contract_amount"]] if len(row) > columns["contract_amount"] else None),
                    "contract_start_date": parse_date(row[4] if len(row) > 4 else None) or clean_text(row[4] if len(row) > 4 else ""),
                    "actual_start_date": parse_date(row[5] if len(row) > 5 else None) or clean_text(row[5] if len(row) > 5 else ""),
                    "contract_completion_date": parse_date(row[6] if len(row) > 6 else None) or clean_text(row[6] if len(row) > 6 else ""),
                    "completion_date": parse_date(row[columns["completion_date"]] if len(row) > columns["completion_date"] else None),
                    "submit_due_date": parse_date(row[columns["submit_due_date"]] if len(row) > columns["submit_due_date"] else None),
                    "actual_submit_date": parse_date(row[columns["actual_submit_date"]] if len(row) > columns["actual_submit_date"] else None),
                    "submitted_amount": to_number(row[columns["submitted_amount"]] if len(row) > columns["submitted_amount"] else None),
                    "settled_date": settled_date,
                    "settled_amount": settled_amount,
                    "status": "",
                    "segment": "",
                    "count_project": True,
                    "section_title": current_title,
                    "settlement_note": clean_text(row[columns["settlement_note"]] if len(row) > columns["settlement_note"] else ""),
                    "manager": clean_text(row[columns["manager"]] if len(row) > columns["manager"] else ""),
                    "settlement_overdue_months": to_number(row[17] if len(row) > 17 else None),
                    "settlement_receive_month": parse_date(row[18] if len(row) > 18 else None) or clean_text(row[18] if len(row) > 18 else ""),
                    "month_settled_count": to_number(row[columns["month_settled_count"]] if len(row) > columns["month_settled_count"] else None),
                    "month_settled_amount": to_number(row[columns["month_settled_amount"]] if len(row) > columns["month_settled_amount"] else None),
                    "has_month_settled_marker": bool(
                        clean_text(row[columns["month_settled_count"]] if len(row) > columns["month_settled_count"] else "")
                        or clean_text(row[columns["month_settled_amount"]] if len(row) > columns["month_settled_amount"] else "")
                    ),
                }
                pending.append(detail)
            flush_pending()
        if not rows:
            errors.append("未识别到结算明细行，请确认是否使用了约定模板。")
        return ParseResult(rows, errors, warnings, wb.sheetnames)

    def source_sheets(self, wb: Any, company: str) -> list[Any]:
        company_name = standard_company_name(company)
        if company_name:
            matched = [
                ws
                for ws in wb.worksheets
                if (standard_company_name(ws.title) or ws.title.strip()) == company_name
            ]
            if matched:
                return matched
        named = [ws for ws in wb.worksheets if not re.fullmatch(r"sheet\d*", ws.title.strip(), re.IGNORECASE)]
        return named[:1] or wb.worksheets[:1]

    def counted_project_rows(self, ws: Any, columns: dict[str, int]) -> set[int]:
        count_col_letter = get_column_letter(columns["building_area"] + 1)
        refs: set[int] = set()
        for row_index in range(1, ws.max_row + 1):
            row_text = " ".join(clean_text(ws.cell(row_index, col_index).value) for col_index in range(1, min(ws.max_column, 6) + 1))
            if "小计" not in row_text:
                continue
            formula = ws.cell(row_index, columns["building_area"] + 1).value
            refs.update(formula_row_refs(formula, count_col_letter))
        return refs

    def detail_columns(self, ws: Any) -> dict[str, int]:
        max_col = ws.max_column
        header_rows = range(1, min(ws.max_row, 6) + 1)

        def header_text(col_index: int) -> str:
            return " ".join(clean_text(ws.cell(row_index, col_index + 1).value) for row_index in header_rows)

        def find_col(*keywords: str, default: int) -> int:
            for col_index in range(max_col):
                text = re.sub(r"\s+", "", header_text(col_index))
                if all(keyword in text for keyword in keywords):
                    return col_index
            return default

        def find_amount_col(default: int) -> int:
            candidates: list[int] = []
            for col_index in range(max_col):
                text = re.sub(r"\s+", "", header_text(col_index))
                if "定案造价" in text and "时间" not in text:
                    candidates.append(col_index)
            if candidates:
                return candidates[-1]
            for col_index in range(max_col):
                text = re.sub(r"\s+", "", header_text(col_index))
                if "造价" in text and "定案" in text and "时间" not in text:
                    candidates.append(col_index)
            return candidates[-1] if candidates else default

        def merged_bounds(row_index: int, col_index: int) -> tuple[int, int]:
            cell_row = row_index + 1
            cell_col = col_index + 1
            for merged_range in ws.merged_cells.ranges:
                if (
                    merged_range.min_row <= cell_row <= merged_range.max_row
                    and merged_range.min_col <= cell_col <= merged_range.max_col
                ):
                    return merged_range.min_col - 1, merged_range.max_col - 1
            return col_index, col_index

        def find_month_settled_cols(default_count: int = 21, default_amount: int = 22) -> tuple[int, int]:
            for row_index in header_rows:
                for col_index in range(max_col):
                    text = clean_text(ws.cell(row_index, col_index + 1).value)
                    if "定案项目" not in text:
                        continue
                    start_col, end_col = merged_bounds(row_index - 1, col_index)
                    if end_col > start_col:
                        return start_col, min(start_col + 1, end_col)
                    below = clean_text(ws.cell(row_index + 1, col_index + 1).value)
                    right_below = clean_text(ws.cell(row_index + 1, col_index + 2).value) if col_index + 1 < max_col else ""
                    if "个数" in below and "造价" in right_below:
                        return col_index, col_index + 1
                    return col_index, min(col_index + 1, max_col - 1)
            return default_count, default_amount

        completion_col = 7
        for col_index in range(max_col):
            if "竣工报告" in header_text(col_index):
                completion_col = col_index
                break
        else:
            for col_index in range(max_col):
                if "实际" not in header_text(col_index):
                    continue
                left = max(0, col_index - 3)
                if any(
                    "竣工时间" in clean_text(ws.cell(row_index, scan_col + 1).value)
                    for scan_col in range(left, col_index + 1)
                    for row_index in header_rows
                ):
                    completion_col = col_index
                    break

        month_settled_count_col, month_settled_amount_col = find_month_settled_cols()

        return {
            "project": find_col("工程名称", default=1),
            "building_area": find_col("建筑面积", default=2),
            "contract_amount": find_col("合同价款", default=3),
            "completion_date": completion_col,
            "submit_due_date": find_col("应递交", default=8),
            "actual_submit_date": find_col("实际递交", default=9),
            "submitted_amount": find_col("结算递交造价", default=10),
            "settled_date": find_col("结算定案时间", default=11),
            "settled_amount": find_amount_col(default=12),
            "settlement_note": 13,
            "manager": find_col("项目经理", default=14),
            "month_settled_count": month_settled_count_col,
            "month_settled_amount": month_settled_amount_col,
        }

    def company_from_sheet(self, ws: Any) -> str:
        for row_index in range(1, min(ws.max_row, 5) + 1):
            for col_index in range(1, min(ws.max_column, 4) + 1):
                text = clean_text(ws.cell(row_index, col_index).value)
                if "单位名称" in text:
                    return text.split("：", 1)[-1].split(":", 1)[-1].strip() or ws.title
        return ws.title

    def aggregate(self, rows: list[dict[str, Any]]) -> dict[str, Any]:
        settled = [row for row in rows if row.get("status") == "已定案"]
        unsettled = [row for row in rows if row.get("status") != "已定案"]
        in_progress = [row for row in rows if row.get("segment") == "在建"]
        completed = [row for row in rows if row.get("status") in {"已定案", "未定案"}]
        unsettled_completed = [row for row in completed if row.get("status") != "已定案"]
        by_company: dict[str, dict[str, Any]] = {}
        for row in rows:
            company = row["company"]
            item = by_company.setdefault(
                company,
                {
                    "company": company,
                    "in_progress_projects": 0,
                    "completed_projects": 0,
                    "settled_projects": 0,
                    "unsettled_projects": 0,
                    "contract_amount": 0,
                    "settled_amount": 0,
                    "submitted_amount": 0,
                },
            )
            if row.get("segment") == "在建":
                item["in_progress_projects"] += 1
            if row.get("status") in {"已定案", "未定案"}:
                item["completed_projects"] += 1 if counts_as_settlement_project(row) else 0
                item["contract_amount"] += row.get("contract_amount", 0) or 0
                item["submitted_amount"] += row.get("submitted_amount", 0) or 0
                item["settled_amount"] += row.get("settled_amount", 0) or 0
            if row.get("status") == "已定案":
                item["settled_projects"] += 1 if counts_as_settlement_project(row) else 0
            elif row.get("status") == "未定案":
                item["unsettled_projects"] += 1 if counts_as_settlement_project(row) else 0
        return {
            "cards": [
                {"label": "在建", "value": len(in_progress), "unit": "项"},
                {"label": "竣工", "value": sum(1 for row in completed if counts_as_settlement_project(row)), "unit": "项"},
                {"label": "已定案", "value": sum(1 for row in settled if counts_as_settlement_project(row)), "unit": "项"},
                {"label": "未定案", "value": sum(1 for row in unsettled_completed if counts_as_settlement_project(row)), "unit": "项"},
            ],
            "table": sorted(by_company.values(), key=lambda item: company_sort_key(item["company"], self.key)),
        }

    def export(self, rows: list[dict[str, Any]], output_path: Path, period: str) -> None:
        summary_template = TEMPLATE_DIR / "settlement-summary-template.xlsx"
        if summary_template.exists():
            wb = load_workbook(summary_template)
            for sheet_name in wb.sheetnames[1:]:
                del wb[sheet_name]
            self.write_summary_sheet(wb.active, rows, period)
        else:
            wb = Workbook()
            wb.active.title = "竣工及未定案工程"
            self.write_summary_sheet(wb.active, rows, period)
        details_tmp = EXPORT_DIR / f"_tmp_settled_details_{uuid.uuid4().hex}.xlsx"
        try:
            self.export_settled_details(rows, details_tmp, period)
            details_wb = load_workbook(details_tmp)
            details_ws = details_wb.active
            details_title = "01各单位竣工结算定案工程一览表"
            if details_title in wb.sheetnames:
                del wb[details_title]
            target_ws = wb.create_sheet(details_title[:31], 1)
            clone_worksheet(details_ws, target_ws)
        finally:
            try:
                details_tmp.unlink()
            except OSError:
                pass
        notice_title = "02关于竣工结算完成情况的通报"
        if notice_title in wb.sheetnames:
            del wb[notice_title]
        notice_ws = wb.create_sheet(notice_title, 2)
        self.write_completion_notice_sheet(notice_ws, rows, period)
        submit_title = "03关于竣工结算递交情况的考核"
        if submit_title in wb.sheetnames:
            del wb[submit_title]
        submit_ws = wb.create_sheet(submit_title, 3)
        self.write_submit_assessment_sheet(submit_ws, rows, period)
        main_business_match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
        main_business_title = (
            f"05建安主业单位{int(main_business_match.group(1))}年{int(main_business_match.group(2))}月项目竣工结算定案情况"
            if main_business_match
            else "05建安主业单位项目竣工结算定案情况"
        )
        main_business_title = main_business_title[:31]
        if main_business_title in wb.sheetnames:
            del wb[main_business_title]
        main_business_ws = wb.create_sheet(main_business_title, 4)
        self.write_main_business_settlement_sheet_v2(main_business_ws, rows, period)
        period_match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
        severe_quarter = ceil(int(period_match.group(2)) / 3) if period_match else ceil(datetime.now().month / 3)
        severe_quarter_name = self.chinese_quarter(severe_quarter)
        severe_notice_title = f"06关于{severe_quarter_name}定案严重逾期工程统计情况的通报"
        # 清理旧版本固定“二季度”的06、07表，避免导出文件同时保留新旧两套季度表。
        for sheet_name in list(wb.sheetnames):
            if re.match(r"^06关于.*定案严重逾期工程统计情况的通报$", sheet_name):
                del wb[sheet_name]
        severe_notice_ws = wb.create_sheet(severe_notice_title, 5)
        self.write_quarter_severe_overdue_notice_sheet(severe_notice_ws, rows, period)
        severe_detail_title = f"07{severe_quarter_name}严重逾期工程明细表"
        for sheet_name in list(wb.sheetnames):
            if re.match(r"^07.*严重逾期工程明细表$", sheet_name):
                del wb[sheet_name]
        severe_detail_ws = wb.create_sheet(severe_detail_title, 6)
        self.write_quarter_severe_overdue_detail_sheet(severe_detail_ws, rows, period)
        source_paths = []
        source_companies: dict[str, str] = {}
        for row in rows:
            source_path = row.get("_source_path")
            if source_path and source_path not in source_paths:
                source_paths.append(source_path)
                source_companies[source_path] = standard_company_name(row.get("company", "")) or row.get("company", "")
        if not source_paths:
            template = TEMPLATE_DIR / "settlement-template.xlsx"
            if template.exists():
                source_paths.append(str(template))
        exported_companies: set[str] = set()
        for source_path in source_paths:
            source_wb = load_workbook(source_path, data_only=False)
            company_title = (source_companies.get(source_path) or "").strip()
            source_sheets = source_wb.worksheets
            if company_title:
                matched_sheets = [
                    ws
                    for ws in source_sheets
                    if (standard_company_name(ws.title) or ws.title.strip()) == company_title
                ]
                source_sheets = matched_sheets or source_sheets[:1]
            else:
                source_sheets = source_sheets[:1]
            if company_title in exported_companies:
                continue
            for source_ws in source_sheets[:1]:
                title = (company_title or source_ws.title).strip()[:31]
                if title in wb.sheetnames:
                    continue
                target_ws = wb.create_sheet(title)
                clone_worksheet(source_ws, target_ws)
                if company_title:
                    exported_companies.add(company_title)
            wb.save(output_path)

    def export_result_package(self, rows: list[dict[str, Any]], period: str) -> Path:
        """Export the working sheets to Excel and the three notice sheets to Word in one package."""
        stamp = datetime.now().strftime("%Y%m%d%H%M%S")
        base_name = f"竣工结算报表_{period or '全部'}_{stamp}"
        excel_path = EXPORT_DIR / f"{base_name}_Excel.xlsx"
        bundle_path = EXPORT_DIR / f"{base_name}_成果包.zip"
        word_dir = EXPORT_DIR / f"_word_{uuid.uuid4().hex}"
        word_dir.mkdir(parents=True, exist_ok=True)
        word_files: list[Path] = []
        try:
            self.export(rows, excel_path, period)
            workbook = load_workbook(excel_path)
            notice_sheets = [
                ("02", "竣工结算完成情况的通报"),
                ("03", "竣工结算递交情况的考核"),
                ("06", "定案严重逾期工程统计情况的通报"),
            ]
            for prefix, keyword in notice_sheets:
                sheet_name = next(
                    (name for name in workbook.sheetnames if name.startswith(prefix) and keyword in name),
                    "",
                )
                if not sheet_name:
                    continue
                word_path = word_dir / f"{sheet_name}.docx"
                worksheet_to_docx(workbook[sheet_name], word_path, prefix)
                word_files.append(word_path)
                del workbook[sheet_name]

            for worksheet in workbook.worksheets:
                title = worksheet.title
                if "竣工及未定案工程" in title:
                    worksheet.title = "01竣工及未定案工程"
                elif title.startswith("01") and "竣工结算定案工程一览表" in title:
                    worksheet.title = "02各单位竣工结算定案工程一览表"
                elif title.startswith("05") and "建安主业单位" in title:
                    worksheet.title = f"03{title[2:]}"
                elif title.startswith("07") and "严重逾期工程明细表" in title:
                    worksheet.title = f"04{title[2:]}"
            workbook.save(excel_path)
            workbook.close()

            with zipfile.ZipFile(bundle_path, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.write(excel_path, arcname=f"Excel/{excel_path.name}")
                for word_path in word_files:
                    archive.write(word_path, arcname=f"Word/{word_path.name}")
            return bundle_path
        finally:
            shutil.rmtree(word_dir, ignore_errors=True)

    @staticmethod
    def summary_contract_amount(row: dict[str, Any]) -> float:
        """Use the original contract amount, retaining the source value's decimal precision."""
        return to_number(row.get("contract_amount", 0) or 0)

    @staticmethod
    def source_display_contract_amount(value: Any) -> float:
        """Contract subtotals in the branch template are displayed as whole 万元."""
        amount = to_number(value)
        return float(floor(amount + 0.5)) if amount >= 0 else float(ceil(amount - 0.5))

    def apply_source_display_contract_overrides(self, rows: list[dict[str, Any]], period: str) -> None:
        """Align summary contract totals with the subtotal values visibly shown in each source file."""
        errors: list[str] = []
        for row in rows:
            row.pop("_summary_contract_amount", None)

        by_source: dict[tuple[str, str], dict[int, dict[str, Any]]] = {}
        for row in rows:
            source_path = clean_text(row.get("_source_path"))
            sheet = clean_text(row.get("sheet"))
            source_row = row.get("source_row")
            if not source_path or not sheet or not isinstance(source_row, int):
                continue
            by_source.setdefault((source_path, sheet), {})[source_row] = row

        cutoff = settlement_abc_cutoff(period)
        for (source_path, sheet_name), parsed_rows in by_source.items():
            try:
                # Header merged cells are needed to identify the source contract column reliably.
                workbook = load_workbook(source_path, data_only=True)
                if sheet_name not in workbook.sheetnames:
                    workbook.close()
                    continue
                worksheet = workbook[sheet_name]
                columns = self.detail_columns(worksheet)
                contract_index = columns["contract_amount"]
                pending: list[dict[str, Any]] = []
                for cells in worksheet.iter_rows():
                    source_row = cells[0].row if cells else 0
                    values = tuple(cell.value for cell in cells)
                    if source_row in parsed_rows and row_is_detail(values):
                        pending.append(parsed_rows[source_row])
                    row_text = " ".join(clean_text(value) for value in values if value is not None)
                    is_subtotal = (
                        "未定案造价小计" in row_text
                        or ("定案造价小计" in row_text and "未定案" not in row_text)
                    )
                    if not is_subtotal:
                        continue
                    included = [
                        row for row in pending
                        if counts_as_settlement_project(row)
                        and row_is_within_report_cutoff(row, period)
                        and classify_settlement_abc(row, cutoff)
                    ]
                    segments = {classify_settlement_abc(row, cutoff) for row in included}
                    # A displayed subtotal may be safely used only when it belongs to one A/B/C category.
                    if included and len(segments) == 1 and len(included) == len(pending):
                        # Some branch subtotal formulas include rows that are not statistical projects.
                        # Rebuild this subtotal from the filtered detail rows, then apply the source's
                        # whole-number display convention once for the group.
                        display_amount = self.source_display_contract_amount(
                            sum(to_number(row.get("contract_amount")) for row in included)
                        )
                        included[0]["_summary_contract_amount"] = display_amount
                        for row in included[1:]:
                            row["_summary_contract_amount"] = 0.0
                    pending = []
                workbook.close()
            except Exception as error:
                # Keep the parsed detail amount as a fallback when a source file cannot be reopened.
                errors.append(f"{Path(source_path).name}: {error}")
                continue
        self._summary_contract_override_errors = errors

    def export_settled_details(self, rows: list[dict[str, Any]], output_path: Path, period: str) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "01各单位竣工结算定案工程一览表"

        headers = [
            ("序号", None),
            ("单位名称", None),
            ("工程名称", None),
            ("建筑面积\n(平米)", None),
            ("合同价款\n(万元)", None),
            ("开工时间", "合同"),
            ("开工时间", "实际"),
            ("竣工时间", "合同"),
            ("竣工时间", "实际"),
            ("结算定案\n时间", None),
            ("结算定案造价\n(元)", None),
            ("定案接收\n时间", None),
            ("定案逾期时间\n(月)", None),
            ("备注", None),
        ]
        period_match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
        period_title = f"{period_match.group(1)}年{int(period_match.group(2))}月" if period_match else period
        period_month_label = f"{int(period_match.group(2))}月" if period_match else ""
        receive_ledger = load_settlement_receive_ledger()
        receive_entries = receive_ledger.setdefault("entries", {})
        receive_ledger_changed = False
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws["A1"] = f"{period_title}各单位竣工结算定案工程一览表"
        ws["A1"].font = Font(name="宋体", size=16, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        for col_index, (top, child) in enumerate(headers, 1):
            cell = ws.cell(2, col_index)
            cell.value = top
            if child:
                ws.cell(3, col_index).value = child
            else:
                ws.merge_cells(start_row=2, start_column=col_index, end_row=3, end_column=col_index)
        ws.merge_cells(start_row=2, start_column=6, end_row=2, end_column=7)
        ws.merge_cells(start_row=2, start_column=8, end_row=2, end_column=9)

        detail_rows = [
            row for row in rows
            if row.get("status") == "已定案"
            and row_is_within_report_cutoff(row, period)
            and classify_settlement_abc(row, settlement_abc_cutoff(period)) == "A"
        ]

        def overdue_sort_value(row: dict[str, Any]) -> float:
            value = self.settlement_overdue_months(row)
            if isinstance(value, (int, float)):
                return float(value)
            match = re.search(r"-?\d+(?:\.\d+)?", clean_text(value))
            return float(match.group(0)) if match else -1.0

        detail_rows.sort(
            key=lambda row: (
                company_sort_key(row.get("company", ""), self.key),
                1 if self.is_month_settled(row, period) else 0,
                -overdue_sort_value(row),
                self.sortable_date(row.get("settled_date")),
                clean_text(row.get("project", "")),
            )
        )

        all_companies = sorted(
            {row.get("company") or "未填单位" for row in rows},
            key=lambda name: company_sort_key(name, self.key),
        )
        grouped: dict[str, list[dict[str, Any]]] = {company: [] for company in all_companies}
        for row in detail_rows:
            grouped.setdefault(row.get("company") or "未填单位", []).append(row)

        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        body_font = Font(name="宋体", size=9)
        subtotal_font = Font(name="宋体", size=9, bold=True)
        header_fill = PatternFill("solid", fgColor="F2F2F2")
        month_fill = PatternFill("solid", fgColor="E2F0D9")
        subtotal_fill = PatternFill("solid", fgColor="F8F8F8")

        for row_index in [2, 3]:
            ws.row_dimensions[row_index].height = 24
            for col_index in range(1, len(headers) + 1):
                cell = ws.cell(row_index, col_index)
                cell.font = Font(name="宋体", size=10, bold=True)
                cell.alignment = center
                cell.border = border
                cell.fill = header_fill

        row_index = 4
        serial = 1
        for company in sorted(grouped, key=lambda name: company_sort_key(name, self.key)):
            company_rows = grouped[company]
            company_start = row_index
            subtotal_contract = 0.0
            subtotal_settled_yuan = 0.0
            subtotal_area = 0.0
            if not company_rows:
                ws.row_dimensions[row_index].height = 28
                ws.cell(row_index, 2).value = company
                for col_index in range(1, len(headers) + 1):
                    cell = ws.cell(row_index, col_index)
                    cell.border = border
                    cell.alignment = center
                    cell.font = body_font
                row_index += 1
            for row in company_rows:
                ws.row_dimensions[row_index].height = 28
                contract_amount = row.get("contract_amount", 0) or 0
                settled_yuan = (row.get("settled_amount", 0) or 0) * 10000
                area = to_number(row.get("building_area")) or 0
                subtotal_contract += self.summary_contract_amount(row)
                subtotal_settled_yuan += settled_yuan
                subtotal_area += area
                is_month_row = self.is_month_settled(row, period)
                ledger_key = settlement_receive_ledger_key(row)
                ledger_entry = receive_entries.get(ledger_key, {}) if ledger_key else {}
                receive_month = clean_text(ledger_entry.get("receiveMonth", "")) if isinstance(ledger_entry, dict) else ""
                if not receive_month and is_month_row and period_month_label:
                    receive_month = period_month_label
                    if ledger_key:
                        receive_entries[ledger_key] = {
                            "company": standard_company_name(row.get("company", "")) or row.get("company", ""),
                            "project": row.get("project", ""),
                            "contractAmount": contract_amount,
                            "settledDate": row.get("settled_date", ""),
                            "settledAmount": row.get("settled_amount", 0) or 0,
                            "receiveMonth": receive_month,
                            "receivePeriod": period,
                            "sourceFile": row.get("_source_file", ""),
                            "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        }
                        receive_ledger_changed = True
                values = [
                    serial,
                    company,
                    row.get("project", ""),
                    row.get("building_area", ""),
                    contract_amount,
                    self.detail_date_text(row.get("contract_start_date")),
                    self.detail_date_text(row.get("actual_start_date")),
                    self.detail_date_text(row.get("contract_completion_date")),
                    self.detail_date_text(row.get("completion_date")),
                    self.detail_date_text(row.get("settled_date")),
                    settled_yuan,
                    receive_month,
                    self.settlement_overdue_months(row),
                    row.get("settlement_note", ""),
                ]
                for col_index, value in enumerate(values, 1):
                    cell = ws.cell(row_index, col_index)
                    cell.value = value
                    cell.border = border
                    cell.alignment = left if col_index in {3, 14} else center
                    cell.font = body_font
                    if col_index not in {1, 2} and is_month_row:
                        cell.fill = month_fill
                ws.cell(row_index, 5).number_format = "#,##0.00"
                ws.cell(row_index, 11).number_format = "#,##0.00"
                serial += 1
                row_index += 1
            if len(company_rows) > 1:
                ws.merge_cells(start_row=company_start, start_column=2, end_row=row_index - 1, end_column=2)
                ws.cell(company_start, 2).alignment = center

            subtotal_values = ["", "", "小计", subtotal_area, subtotal_contract, "", "", "", "", "", subtotal_settled_yuan, "", "", ""]
            ws.row_dimensions[row_index].height = 28
            for col_index, value in enumerate(subtotal_values, 1):
                cell = ws.cell(row_index, col_index)
                cell.value = value
                cell.border = border
                cell.alignment = center
                cell.font = subtotal_font
                cell.fill = subtotal_fill
            ws.cell(row_index, 5).number_format = "#,##0.00"
            ws.cell(row_index, 11).number_format = "#,##0.00"
            row_index += 1

        if row_index == 4:
            ws.cell(4, 1).value = "暂无本年度定案项目明细"
            ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=len(headers))
            ws.cell(4, 1).alignment = center
            row_index = 5

        widths = [6, 12, 36, 12, 12, 13, 13, 13, 13, 13, 15, 12, 12, 18]
        for col_index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col_index)].width = width
        ws.freeze_panes = "A4"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        wb.save(output_path)
        if receive_ledger_changed:
            save_settlement_receive_ledger(receive_ledger)

    def write_completion_notice_sheet(self, ws: Any, rows: list[dict[str, Any]], period: str) -> None:
        period_match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
        year = int(period_match.group(1)) if period_match else datetime.now().year
        month = int(period_match.group(2)) if period_match else datetime.now().month
        period_title = f"{year}年{month}月"

        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        widths = [14, 13, 13, 13, 13, 13, 4]
        for col_index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col_index)].width = width

        title_font = Font(name="宋体", size=14, bold=True, color="FF0000")
        body_font = Font(name="宋体", size=10)
        table_font = Font(name="宋体", size=10)
        table_bold = Font(name="宋体", size=10, bold=True)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws.merge_cells("A1:F1")
        ws["A1"] = f"关于{period_title}竣工结算完成情况的通报"
        ws["A1"].font = title_font
        ws["A1"].alignment = center
        ws.row_dimensions[1].height = 32

        ws.merge_cells("A3:F3")
        ws["A3"] = "基层各单位："
        ws["A3"].font = body_font
        ws["A3"].alignment = left
        ws.merge_cells("A4:F5")
        ws["A4"] = (
            "为加强竣工结算的日常过程管理，加快定案进度，集团公司对各基层单位"
            f"{month}月竣工结算完成情况统计通报如下。"
        )
        ws["A4"].font = body_font
        ws["A4"].alignment = left

        table_start = 7
        ws.merge_cells(start_row=table_start, start_column=1, end_row=table_start, end_column=6)
        ws.cell(table_start, 1).value = "竣工结算完成情况统计表"
        ws.cell(table_start, 1).font = table_bold
        ws.cell(table_start, 1).alignment = center

        header_row = table_start + 1
        subheader_row = table_start + 2
        headers = ["单位名称", "纳入考核", "本月", "", "累计", ""]
        subheaders = ["", "", "计划定案", "实际完成", "定案数量", "定案率"]
        for col_index, value in enumerate(headers, 1):
            cell = ws.cell(header_row, col_index)
            cell.value = value
            cell.font = table_bold
            cell.alignment = center
            cell.border = border
        for col_index, value in enumerate(subheaders, 1):
            cell = ws.cell(subheader_row, col_index)
            cell.value = value
            cell.font = table_bold
            cell.alignment = center
            cell.border = border
        ws.merge_cells(start_row=header_row, start_column=1, end_row=subheader_row, end_column=1)
        ws.merge_cells(start_row=header_row, start_column=2, end_row=subheader_row, end_column=2)
        ws.merge_cells(start_row=header_row, start_column=3, end_row=header_row, end_column=4)
        ws.merge_cells(start_row=header_row, start_column=5, end_row=header_row, end_column=6)

        stats = self.completion_notice_stats(rows, period)
        notice_order = [
            ("三公司", ["三公司"]),
            ("四公司", ["四公司"]),
            ("五公司", ["五公司"]),
            ("六公司", ["六公司"]),
            ("七公司", ["七公司"]),
            ("十/装饰小计", ["十公司", "装饰幕墙公司"]),
            ("其中：十公司", ["十公司"]),
            ("装饰幕墙", ["装饰幕墙公司"]),
            ("市政路桥", ["市政路桥公司"]),
            ("格瑞特", ["格瑞特公司"]),
            ("青岛公司", ["青岛公司"]),
            ("济南公司", ["济南公司"]),
            ("上海公司", ["上海公司"]),
            ("国际公司", ["国际公司"]),
            ("马来公司", ["马来公司"]),
            ("设备安装", ["设备安装公司"]),
        ]

        row_index = subheader_row + 1
        total = {"exam": 0, "plan": 0, "actual": 0, "settled": 0}
        for label, companies in notice_order:
            values = self.sum_notice_stats(stats, companies)
            if not label.startswith("其中") and label != "装饰幕墙":
                total["exam"] += values["exam"]
                total["plan"] += values["plan"]
                total["actual"] += values["actual"]
                total["settled"] += values["settled"]
            self.write_notice_table_row(ws, row_index, label, values, border, table_font, center)
            row_index += 1
        self.write_notice_table_row(ws, row_index, "合计", total, border, table_bold, center)
        total_row = row_index

        for row in range(header_row, total_row + 1):
            ws.row_dimensions[row].height = 22
            for col in range(1, 7):
                ws.cell(row, col).border = border
                ws.cell(row, col).alignment = center

        note_row = total_row + 2
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 2, end_column=6)
        ws.cell(note_row, 1).value = (
            "定案率就是定任务，必须未定案是任务未卡、基层各单位需锁定结算目标，"
            "压实各级责任，严格对结算资料流转与送审的关系，强化过程督导与考核兑现，"
            "全力保障项目结算落地，加快欠款回收和资产条件。"
        )
        ws.cell(note_row, 1).font = body_font
        ws.cell(note_row, 1).alignment = left

        sign_row = note_row + 5
        ws.merge_cells(start_row=sign_row, start_column=5, end_row=sign_row, end_column=6)
        ws.cell(sign_row, 5).value = "成本管理部"
        ws.cell(sign_row, 5).font = body_font
        ws.cell(sign_row, 5).alignment = center
        ws.merge_cells(start_row=sign_row + 1, start_column=5, end_row=sign_row + 1, end_column=6)
        ws.cell(sign_row + 1, 5).value = f"{year}年{month}月{datetime.now().day}日"
        ws.cell(sign_row + 1, 5).font = body_font
        ws.cell(sign_row + 1, 5).alignment = center

    def completion_notice_stats(self, rows: list[dict[str, Any]], period: str) -> dict[str, dict[str, int]]:
        cutoff = settlement_abc_cutoff(period)
        stats: dict[str, dict[str, int]] = {
            company: {"exam": 0, "plan": 0, "actual": 0, "settled": 0}
            for company in company_order(self.key)
        }
        for row in rows:
            company = standard_company_name(row.get("company", "")) or row.get("company", "") or "未填单位"
            item = stats.setdefault(company, {"exam": 0, "plan": 0, "actual": 0, "settled": 0})
            count_project = 1 if counts_as_settlement_project(row) else 0
            if not count_project or not row_is_within_report_cutoff(row, period):
                continue
            abc_segment = classify_settlement_abc(row, cutoff)
            if abc_segment in {"A", "B"}:
                item["exam"] += 1
            if abc_segment == "A":
                item["settled"] += 1
                if self.is_month_settled(row, period):
                    item["actual"] += 1

        # The monthly target is committed in the prior month's branch report, not revised from this month's sheet.
        seen_plans: set[tuple[str, str, float]] = set()
        for row in self.previous_period_plan_rows(period):
            company = standard_company_name(row.get("company", "")) or row.get("company", "") or "未填单位"
            project = clean_text(row.get("project", ""))
            contract = round(to_number(row.get("contract_amount")), 6)
            plan_key = (company, project, contract)
            if plan_key in seen_plans:
                continue
            seen_plans.add(plan_key)
            if not counts_as_settlement_project(row) or not row_is_within_report_cutoff(row, period):
                continue
            abc_segment = classify_settlement_abc(row, cutoff)
            if abc_segment in {"A", "B"} and self.month_value_matches(row.get("settlement_receive_month"), period):
                item = stats.setdefault(company, {"exam": 0, "plan": 0, "actual": 0, "settled": 0})
                item["plan"] += 1
        return stats

    @staticmethod
    def sum_notice_stats(stats: dict[str, dict[str, int]], companies: list[str]) -> dict[str, int]:
        result = {"exam": 0, "plan": 0, "actual": 0, "settled": 0}
        for company in companies:
            item = stats.get(company, {})
            for key in result:
                result[key] += int(item.get(key, 0) or 0)
        return result

    @staticmethod
    def write_notice_table_row(ws: Any, row_index: int, label: str, values: dict[str, int], border: Border, font: Font, alignment: Alignment) -> None:
        exam = values.get("exam", 0) or 0
        settled = values.get("settled", 0) or 0
        rate = settled / exam if exam else 0
        row_values = [label, exam, values.get("plan", 0), values.get("actual", 0), settled, rate]
        for col_index, value in enumerate(row_values, 1):
            cell = ws.cell(row_index, col_index)
            cell.value = value
            cell.border = border
            cell.font = font
            cell.alignment = alignment
        ws.cell(row_index, 6).number_format = "0.0%"


    def write_main_business_settlement_sheet(self, ws: Any, rows: list[dict[str, Any]], period: str) -> None:
        period_match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
        year = int(period_match.group(1)) if period_match else datetime.now().year
        month = int(period_match.group(2)) if period_match else datetime.now().month
        period_title = f"{year}?{month}?"

        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        widths = [8, 12, 10, 12, 12, 10, 12, 12, 12, 12, 10, 12, 10, 12, 10, 12, 10, 10, 12]
        for col_index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col_index)].width = width

        title_font = Font(name="??", size=14, bold=True)
        header_font = Font(name="??", size=9, bold=True)
        body_font = Font(name="??", size=9)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill("solid", fgColor="F2F2F2")
        blue_fill = PatternFill("solid", fgColor="D9EAF7")
        gray_fill = PatternFill("solid", fgColor="D9D9D9")
        total_fill = PatternFill("solid", fgColor="B6D7A8")

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=19)
        ws["A1"] = f"??????{period_title}??????????"
        ws["A1"].font = title_font
        ws["A1"].alignment = center
        ws.row_dimensions[1].height = 28

        # Header rows, matching the provided template layout.
        for col in [1, 2]:
            ws.merge_cells(start_row=2, start_column=col, end_row=4, end_column=col)
        ws.cell(2, 1).value = "??\n??"
        ws.cell(2, 2).value = "????"
        ws.merge_cells(start_row=2, start_column=3, end_row=2, end_column=5)
        ws.cell(2, 3).value = "?????????"
        ws.merge_cells(start_row=2, start_column=6, end_row=2, end_column=9)
        ws.cell(2, 6).value = f"??{month}?20?\n???????"
        ws.merge_cells(start_row=2, start_column=10, end_row=2, end_column=13)
        ws.cell(2, 10).value = "???????"
        ws.merge_cells(start_row=2, start_column=14, end_row=2, end_column=17)
        ws.cell(2, 14).value = f"???????5.21?6.20?"
        ws.merge_cells(start_row=2, start_column=18, end_row=2, end_column=19)
        ws.cell(2, 18).value = "??????"

        subheaders = {
            3: [(3, "??(?)"), (4, "???\n(??)"), (5, "??\n???")],
            6: [(6, "??(?)"), (7, "???\n(??)"), (8, "???\n(???)"), (9, "???\n(????)")],
            10: [(10, "??(?)"), (11, "???\n(??)"), (12, "????????\n??(?)"), (13, "????????\n???(??)")],
            14: [(14, "??\n??(?)"), (15, "??\n???(??)"), (16, "??\n??(?)"), (17, "??\n???(??)")],
            18: [(18, "??(?)"), (19, "???\n(??)")],
        }
        for pairs in subheaders.values():
            for col, value in pairs:
                ws.cell(3, col).value = value
                ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
        for row_index in range(2, 5):
            ws.row_dimensions[row_index].height = 26
            for col_index in range(1, 20):
                cell = ws.cell(row_index, col_index)
                cell.font = header_font
                cell.alignment = center
                cell.border = border
                cell.fill = header_fill

        row_defs = [
            ("??", "???", ["???"]),
            ("??", "?/????", ["???", "??????"]),
            ("??", "??????", ["???"]),
            ("??", "????", ["??????"]),
            ("??", "???", ["???"]),
            ("??", "????", ["????"]),
            ("??", "????", ["????"]),
            ("??", "????", ["????"]),
            ("??", "???", ["???"]),
            ("??", "???", ["???"]),
            ("??", "????", ["??????"]),
            ("??", "???", ["?????"]),
            ("??", "????", ["??????"]),
            ("??", "???", ["???"]),
            ("??", "????", ["????"]),
            ("??", "????", ["????"]),
        ]
        all_companies = sorted({company for _, _, companies in row_defs for company in companies}, key=lambda name: company_sort_key(name, self.key))
        stats = self.main_business_settlement_stats(rows, period, all_companies)

        def add_values(target: dict[str, Any], source: dict[str, Any]) -> None:
            for key in target:
                target[key] += source.get(key, 0) or 0

        def sum_companies(companies: list[str]) -> dict[str, Any]:
            result = {
                "exam": 0, "exam_contract": 0.0, "settled": 0, "settled_amount": 0.0,
                "unsettled": 0, "unsettled_contract": 0.0, "serious": 0, "serious_contract": 0.0,
                "plan": 0, "plan_contract": 0.0, "month_settled": 0, "month_settled_amount": 0.0,
                "next_plan": 0, "next_plan_contract": 0.0,
            }
            for company in companies:
                add_values(result, stats.get(company, {}))
            return result

        row_index = 5
        total = {key: 0 for key in [
            "exam", "exam_contract", "settled", "settled_amount", "unsettled", "unsettled_contract",
            "serious", "serious_contract", "plan", "plan_contract", "month_settled", "month_settled_amount",
            "next_plan", "next_plan_contract"
        ]}
        merge_ranges: dict[str, list[int]] = {}
        for leader, label, companies in row_defs:
            values = sum_companies(companies)
            if not label.startswith("??") and label != "????":
                add_values(total, values)
            exam_rate = values["settled"] / values["exam"] if values["exam"] else 0
            amount_rate = values["settled_amount"] / values["exam_contract"] if values["exam_contract"] else 0
            complete_rate = values["month_settled"] / values["plan"] if values["plan"] else 0
            row_values = [
                leader, label,
                values["exam"], values["exam_contract"], None,
                values["settled"], values["settled_amount"], exam_rate, amount_rate,
                values["unsettled"], values["unsettled_contract"], values["serious"], values["serious_contract"],
                values["plan"], values["plan_contract"], values["month_settled"], values["month_settled_amount"],
                values["next_plan"], values["next_plan_contract"],
            ]
            for col_index, value in enumerate(row_values, 1):
                cell = ws.cell(row_index, col_index)
                cell.value = value
                cell.font = body_font
                cell.alignment = center
                cell.border = border
                cell.fill = blue_fill if row_index % 2 else gray_fill
            merge_ranges.setdefault(leader, []).append(row_index)
            row_index += 1

        for leader, row_numbers in merge_ranges.items():
            if len(row_numbers) > 1:
                ws.merge_cells(start_row=min(row_numbers), start_column=1, end_row=max(row_numbers), end_column=1)
                ws.cell(min(row_numbers), 1).alignment = center

        total_rate = total["settled"] / total["exam"] if total["exam"] else 0
        total_amount_rate = total["settled_amount"] / total["exam_contract"] if total["exam_contract"] else 0
        total_complete_rate = total["month_settled"] / total["plan"] if total["plan"] else 0
        total_values = [
            "", "????",
            total["exam"], total["exam_contract"], None,
            total["settled"], total["settled_amount"], total_rate, total_amount_rate,
            total["unsettled"], total["unsettled_contract"], total["serious"], total["serious_contract"],
            total["plan"], total["plan_contract"], total["month_settled"], total["month_settled_amount"],
            total["next_plan"], total["next_plan_contract"],
        ]
        for col_index, value in enumerate(total_values, 1):
            cell = ws.cell(row_index, col_index)
            cell.value = value
            cell.font = header_font
            cell.alignment = center
            cell.border = border
            cell.fill = total_fill

        for r in range(5, row_index + 1):
            ws.row_dimensions[r].height = 22
            for c in [4, 7, 11, 13, 15, 17, 19]:
                ws.cell(r, c).number_format = "#,##0.00"
            for c in [5, 8, 9, 17]:
                if c in [5, 8, 9, 17]:
                    ws.cell(r, c).number_format = "0.00%"

        note_row = row_index + 2
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 2, end_column=19)
        ws.cell(note_row, 1).value = (
            "???\n"
            "1. ??????????2026?6?30?????2026?7?1??????????\n"
            "2. ??????????????????6????????\n"
            "3. ????????????20?????????????????????"
        )
        ws.cell(note_row, 1).font = body_font
        ws.cell(note_row, 1).alignment = left

    def main_business_settlement_stats(self, rows: list[dict[str, Any]], period: str, companies: list[str]) -> dict[str, dict[str, Any]]:
        cutoff = settlement_abc_cutoff(period)
        period_match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
        next_period = ""
        if period_match:
            year, month = int(period_match.group(1)), int(period_match.group(2))
            if month == 12:
                next_period = f"{year + 1:04d}-01"
            else:
                next_period = f"{year:04d}-{month + 1:02d}"
        stats: dict[str, dict[str, Any]] = {
            company: {
                "exam": 0, "exam_contract": 0.0, "settled": 0, "settled_amount": 0.0,
                "unsettled": 0, "unsettled_contract": 0.0, "serious": 0, "serious_contract": 0.0,
                "plan": 0, "plan_contract": 0.0, "month_settled": 0, "month_settled_amount": 0.0,
                "next_plan": 0, "next_plan_contract": 0.0,
            }
            for company in companies
        }
        for row in rows:
            company = standard_company_name(row.get("company", "")) or row.get("company", "")
            if company not in stats or not counts_as_settlement_project(row) or not row_is_within_report_cutoff(row, period):
                continue
            abc_segment = classify_settlement_abc(row, cutoff)
            if not abc_segment:
                continue
            contract = self.summary_contract_amount(row)
            item = stats[company]
            receive_month = row.get("settlement_receive_month")
            is_settled = row.get("status") == "已定案" or abc_segment == "A"
            if is_settled:
                item["exam"] += 1
                item["exam_contract"] += contract
                item["settled"] += 1
                item["settled_amount"] += row.get("settled_amount", 0) or 0
                if self.is_month_settled(row, period):
                    item["month_settled"] += 1
                    item["month_settled_amount"] += row.get("settled_amount", 0) or 0
            elif abc_segment == "B":
                item["exam"] += 1
                item["exam_contract"] += contract
                item["unsettled"] += 1
                item["unsettled_contract"] += contract
                overdue = self.settlement_overdue_months(row)
                if isinstance(overdue, (int, float)) and overdue >= 6:
                    item["serious"] += 1
                    item["serious_contract"] += contract
            if self.month_value_matches(receive_month, period):
                item["plan"] += 1
                item["plan_contract"] += contract
            if next_period and self.month_value_matches(receive_month, next_period):
                item["next_plan"] += 1
                item["next_plan_contract"] += contract
        return stats

    def write_main_business_settlement_sheet_v2(self, ws: Any, rows: list[dict[str, Any]], period: str) -> None:
        period_match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
        year = int(period_match.group(1)) if period_match else datetime.now().year
        month = int(period_match.group(2)) if period_match else datetime.now().month
        title_period = f"{year}年{month}月"
        report_day_label = f"{month}月20日"
        month_window_label = report_window_label(period)

        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "C5"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        widths = [8, 12, 10, 12, 11, 10, 12, 11, 11, 10, 12, 11, 12, 10, 12, 10, 12, 10, 10, 12]
        for col_index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col_index)].width = width

        title_font = Font(name="宋体", size=14, bold=True)
        header_font = Font(name="宋体", size=9, bold=True)
        body_font = Font(name="宋体", size=9)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill("solid", fgColor="F2F2F2")
        blue_fill = PatternFill("solid", fgColor="D9EAF7")
        no_fill = PatternFill(fill_type=None)
        total_fill = PatternFill("solid", fgColor="B6D7A8")

        ws.merge_cells(start_row=1, start_column=3, end_row=1, end_column=20)
        ws.cell(1, 3).value = f"建安主业单位{title_period}项目竣工结算定案情况"
        ws.cell(1, 3).font = title_font
        ws.cell(1, 3).alignment = center
        ws.row_dimensions[1].height = 28

        for col in (1, 2):
            ws.merge_cells(start_row=2, start_column=col, end_row=4, end_column=col)
        ws.cell(2, 1).value = "督导\n领导"
        ws.cell(2, 2).value = "督导单位"
        groups = [
            (3, 5, "纳入本年度考核项目"),
            (6, 9, f"截至{report_day_label}\n累计已定案项目"),
            (10, 13, "竣工未定案项目"),
            (14, 18, f"本月定案项目（{month_window_label}）"),
        ]
        for start_col, end_col, label in groups:
            ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=end_col)
            ws.cell(2, start_col).value = label
        ws.merge_cells(start_row=2, start_column=19, end_row=3, end_column=20)
        ws.cell(2, 19).value = "下月计划定案"

        vertical_headers = {
            3: "数量(个)",
            4: "合同额\n(万元)",
            5: "目标\n定案率",
            6: "数量(个)",
            7: "定案额\n(万元)",
            8: "定案率\n(按数量)",
            9: "定案率\n(按合同额)",
            10: "总数(个)",
            11: "合同额\n(万元)",
        }
        for col, label in vertical_headers.items():
            ws.cell(3, col).value = label
            ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)

        subgroups = [
            (12, 13, "其中严重逾期项目"),
            (14, 15, "计划"),
            (16, 18, "完成"),
        ]
        for start_col, end_col, label in subgroups:
            ws.merge_cells(start_row=3, start_column=start_col, end_row=3, end_column=end_col)
            ws.cell(3, start_col).value = label

        bottom_headers = {
            12: "数量(个)",
            13: "合同额\n(万元)",
            14: "数量(个)",
            15: "合同额\n(万元)",
            16: "数量(个)",
            17: "定案额\n(万元)",
            18: "完成率",
            19: "数量(个)",
            20: "合同额\n(万元)",
        }
        for col, label in bottom_headers.items():
            ws.cell(4, col).value = label
        for row_index in range(2, 5):
            ws.row_dimensions[row_index].height = 26
            for col_index in range(1, 21):
                cell = ws.cell(row_index, col_index)
                cell.font = header_font
                cell.alignment = center
                cell.border = border
                cell.fill = header_fill

        row_defs = [
            ("文总", "七公司", ["七公司"]),
            ("文总", "十/装饰小计", ["十公司", "装饰幕墙公司"]),
            ("文总", "其中：十公司", ["十公司"]),
            ("文总", "装饰幕墙", ["装饰幕墙公司"]),
            ("王总", "四公司", ["四公司"]),
            ("王总", "青岛公司", ["青岛公司"]),
            ("王总", "济南公司", ["济南公司"]),
            ("王总", "上海公司", ["上海公司"]),
            ("宋总", "五公司", ["五公司"]),
            ("宋总", "六公司", ["六公司"]),
            ("汪总", "市政路桥", ["市政路桥公司"]),
            ("汪总", "格瑞特", ["格瑞特公司"]),
            ("汪总", "设备安装", ["设备安装公司"]),
            ("张总", "三公司", ["三公司"]),
            ("高颖", "国际公司", ["国际公司"]),
            ("高颖", "马来公司", ["马来公司"]),
        ]
        leader_band: dict[str, bool] = {}
        for leader, _, _ in row_defs:
            if leader not in leader_band:
                leader_band[leader] = len(leader_band) % 2 == 0
        all_companies = sorted({company for _, _, companies in row_defs for company in companies}, key=lambda name: company_sort_key(name, self.key))
        stats = self.main_business_settlement_stats_v2(rows, period, all_companies)

        stat_keys = [
            "exam", "exam_contract", "settled", "settled_amount", "unsettled", "unsettled_contract",
            "serious", "serious_contract", "plan", "plan_contract", "month_settled", "month_settled_amount",
            "next_plan", "next_plan_contract",
        ]

        def blank_stats() -> dict[str, Any]:
            return {key: 0 for key in stat_keys}

        def add_values(target: dict[str, Any], source: dict[str, Any]) -> None:
            for key in stat_keys:
                target[key] += source.get(key, 0) or 0

        def sum_companies(companies: list[str]) -> dict[str, Any]:
            result = blank_stats()
            for company in companies:
                add_values(result, stats.get(company, {}))
            return result

        total = blank_stats()
        leader_rows: dict[str, list[int]] = {}
        row_index = 5
        for leader, label, companies in row_defs:
            values = sum_companies(companies)
            if label not in ("其中：十公司", "装饰幕墙"):
                add_values(total, values)
            count_rate = values["settled"] / values["exam"] if values["exam"] else 0
            amount_rate = values["settled_amount"] / values["exam_contract"] if values["exam_contract"] else 0
            row_values = [
                leader, label,
                values["exam"], values["exam_contract"], 0.9,
                values["settled"], values["settled_amount"], count_rate, amount_rate,
                values["unsettled"], values["unsettled_contract"], values["serious"], values["serious_contract"],
                None, None, values["month_settled"], values["month_settled_amount"],
                None, values["next_plan"], values["next_plan_contract"],
            ]
            for col_index, value in enumerate(row_values, 1):
                cell = ws.cell(row_index, col_index)
                cell.value = value
                cell.font = body_font
                cell.alignment = center
                cell.border = border
                if col_index <= 2:
                    cell.fill = no_fill
                elif leader_band.get(leader) or col_index == 18:
                    cell.fill = blue_fill
                else:
                    cell.fill = no_fill
            leader_rows.setdefault(leader, []).append(row_index)
            ws.row_dimensions[row_index].height = 22
            row_index += 1

        for leader, indexes in leader_rows.items():
            if len(indexes) > 1:
                ws.merge_cells(start_row=indexes[0], start_column=1, end_row=indexes[-1], end_column=1)
                ws.cell(indexes[0], 1).value = leader
                ws.cell(indexes[0], 1).alignment = center

        total_count_rate = total["settled"] / total["exam"] if total["exam"] else 0
        total_amount_rate = total["settled_amount"] / total["exam_contract"] if total["exam_contract"] else 0
        total_values = [
            "", "集团总计",
            total["exam"], total["exam_contract"], 0.9,
            total["settled"], total["settled_amount"], total_count_rate, total_amount_rate,
            total["unsettled"], total["unsettled_contract"], total["serious"], total["serious_contract"],
            None, None, total["month_settled"], total["month_settled_amount"],
            None, total["next_plan"], total["next_plan_contract"],
        ]
        for col_index, value in enumerate(total_values, 1):
            cell = ws.cell(row_index, col_index)
            cell.value = value
            cell.font = header_font
            cell.alignment = center
            cell.border = border
            cell.fill = total_fill if col_index >= 3 else no_fill
        ws.row_dimensions[row_index].height = 22

        for r in range(5, row_index + 1):
            for c in (4, 7, 11, 13, 15, 17, 20):
                ws.cell(r, c).number_format = "#,##0.00"
            ws.cell(r, 5).number_format = "0%"
            for c in (8, 9):
                ws.cell(r, c).number_format = "0.0%"
                ws.cell(r, c).font = header_font
            ws.cell(r, 18).number_format = "0%"
            ws.cell(r, 18).font = header_font

        note_row = row_index + 2
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 2, end_column=20)
        ws.cell(note_row, 1).value = (
            "备注：\n"
            "1. 纳入本年度考核项目为本年度已定案工程和6月30日及以前实际竣工、截至本月未定案工程。\n"
            "2. 竣工未定案项目按6月30日及以前实际竣工、截至本月未定案工程统计。\n"
            "3. 本月定案项目统计截至本月20日；下月计划定案统计计划定案月份为下月的项目。"
        )
        ws.cell(note_row, 1).font = body_font
        ws.cell(note_row, 1).alignment = left

    def main_business_settlement_stats_v2(self, rows: list[dict[str, Any]], period: str, companies: list[str]) -> dict[str, dict[str, Any]]:
        cutoff = settlement_abc_cutoff(period)
        period_match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
        next_period = ""
        if period_match:
            year, month = int(period_match.group(1)), int(period_match.group(2))
            next_period = f"{year + 1:04d}-01" if month == 12 else f"{year:04d}-{month + 1:02d}"
        stats: dict[str, dict[str, Any]] = {
            company: {
                "exam": 0, "exam_contract": 0.0, "settled": 0, "settled_amount": 0.0,
                "unsettled": 0, "unsettled_contract": 0.0, "serious": 0, "serious_contract": 0.0,
                "plan": 0, "plan_contract": 0.0, "month_settled": 0, "month_settled_amount": 0.0,
                "next_plan": 0, "next_plan_contract": 0.0,
            }
            for company in companies
        }
        for row in rows:
            company = standard_company_name(row.get("company", "")) or row.get("company", "")
            if company not in stats or not counts_as_settlement_project(row) or not row_is_within_report_cutoff(row, period):
                continue
            abc_segment = classify_settlement_abc(row, cutoff)
            if not abc_segment:
                continue
            item = stats[company]
            contract = self.summary_contract_amount(row)
            settled_amount = row.get("settled_amount", 0) or 0
            receive_month = row.get("settlement_receive_month")
            is_settled = row.get("status") == "已定案" or abc_segment == "A"
            if is_settled:
                item["exam"] += 1
                item["exam_contract"] += contract
                if not settled_amount:
                    settled_amount = self.source_settled_amount(row)
                item["settled"] += 1
                item["settled_amount"] += settled_amount
                if self.is_month_settled(row, period):
                    item["month_settled"] += 1
                    item["month_settled_amount"] += settled_amount
            elif abc_segment == "B":
                item["exam"] += 1
                item["exam_contract"] += contract
                item["unsettled"] += 1
                item["unsettled_contract"] += contract
                if self.is_seriously_overdue_by_completion(row, period):
                    item["serious"] += 1
                    item["serious_contract"] += contract
            if self.month_value_matches(receive_month, period):
                item["plan"] += 1
                item["plan_contract"] += contract
            if next_period and self.month_value_matches(receive_month, next_period):
                item["next_plan"] += 1
                item["next_plan_contract"] += contract
        return stats

    def source_settled_amount(self, row: dict[str, Any]) -> float:
        source_path = clean_text(row.get("_source_path"))
        project = clean_text(row.get("project"))
        if not source_path or not project:
            return 0.0
        cache = getattr(self, "_settled_amount_source_cache", None)
        if cache is None:
            cache = {}
            setattr(self, "_settled_amount_source_cache", cache)
        cache_key = (source_path, clean_text(row.get("sheet")))
        if cache_key not in cache:
            values: dict[str, float] = {}
            try:
                wb = load_workbook(source_path, data_only=False)
                sheet_name = cache_key[1]
                worksheets = [wb[sheet_name]] if sheet_name in wb.sheetnames else wb.worksheets
                for ws in worksheets:
                    columns = self.detail_columns(ws)
                    project_col = columns["project"] + 1
                    settled_col = columns["settled_amount"] + 1
                    for cells in ws.iter_rows():
                        source_project = clean_text(cells[project_col - 1].value if len(cells) >= project_col else "")
                        if not source_project:
                            continue
                        formula_value = ws.cell(cells[0].row, settled_col).value
                        amount = to_formula_number(formula_value)
                        if amount:
                            values[source_project] = amount
            except Exception:
                values = {}
            cache[cache_key] = values
        return cache[cache_key].get(project, 0.0)

    @classmethod
    def is_seriously_overdue_by_completion(cls, row: dict[str, Any], period: str) -> bool:
        completion_date = cls.date_value(row.get("completion_date"))
        report_date = cls.date_value(period_report_date(period))
        if not completion_date or not report_date:
            return False
        months = (report_date.year - completion_date.year) * 12 + report_date.month - completion_date.month
        if report_date.day < completion_date.day:
            months -= 1
        return months > 12

    def write_quarter_severe_overdue_notice_sheet(self, ws: Any, rows: list[dict[str, Any]], period: str) -> None:
        period_match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
        year = int(period_match.group(1)) if period_match else datetime.now().year
        month = int(period_match.group(2)) if period_match else datetime.now().month
        quarter = ceil(month / 3)
        # 季度名称按统计月份所属季度显示；取数只截至当前报表日，不能预先计算未来月份。
        report_cutoff = period_report_date(period)

        ws.sheet_view.showGridLines = False
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        widths = [8, 13, 12, 10, 11, 12, 13]
        for col_index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col_index)].width = width

        title_font = Font(name="宋体", size=12, bold=True, color="FF0000")
        body_font = Font(name="宋体", size=10)
        header_font = Font(name="宋体", size=9, bold=True)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill("solid", fgColor="F2F2F2")

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)
        ws["A2"] = f"关于{year}年{self.chinese_quarter(quarter)}定案严重逾期工程统计情况的通报"
        ws["A2"].font = title_font
        ws["A2"].alignment = center
        ws.row_dimensions[2].height = 26

        intro_lines = [
            (4, "基层各单位："),
            (5, "为进一步强化竣工结算管理，提高竣工结算效率，根据《工程竣工结算管理办法》（烟建〔2026〕14号）的相关要求，"),
            (6, "集团公司对基层相关单位定案严重逾期工程进行统计及处罚，现将本季度结果通报如下。"),
        ]
        for row_index, text in intro_lines:
            ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=7)
            ws.cell(row_index, 1).value = text
            ws.cell(row_index, 1).font = body_font
            ws.cell(row_index, 1).alignment = left

        table_start = 8
        ws.merge_cells(start_row=table_start, start_column=1, end_row=table_start, end_column=7)
        ws.cell(table_start, 1).value = "定案严重逾期工程统计表"
        ws.cell(table_start, 1).font = header_font
        ws.cell(table_start, 1).alignment = center

        header_row = table_start + 1
        for col in (1, 2, 3, 7):
            ws.merge_cells(start_row=header_row, start_column=col, end_row=header_row + 1, end_column=col)
        ws.cell(header_row, 1).value = "序号"
        ws.cell(header_row, 2).value = "单位名称"
        ws.cell(header_row, 3).value = "上季度"
        ws.merge_cells(start_row=header_row, start_column=4, end_row=header_row, end_column=6)
        ws.cell(header_row, 4).value = "本季度严重逾期明细"
        ws.cell(header_row, 7).value = "处罚金额\n(元)"
        for col, label in [(4, "合计"), (5, "12-24个月"), (6, "24个月以上")]:
            ws.cell(header_row + 1, col).value = label

        for row_index in range(header_row, header_row + 2):
            ws.row_dimensions[row_index].height = 24
            for col_index in range(1, 8):
                cell = ws.cell(row_index, col_index)
                cell.font = header_font
                cell.alignment = center
                cell.border = border
                cell.fill = header_fill

        row_defs = [
            ("三公司", ["三公司"]),
            ("四公司", ["四公司"]),
            ("五公司", ["五公司"]),
            ("六公司", ["六公司"]),
            ("七公司", ["七公司"]),
            ("十/装饰小计", ["十公司", "装饰幕墙公司"]),
            ("其中：十公司", ["十公司"]),
            ("装饰幕墙", ["装饰幕墙公司"]),
            ("青岛公司", ["青岛公司"]),
            ("济南公司", ["济南公司"]),
            ("上海公司", ["上海公司"]),
            ("格瑞特", ["格瑞特公司"]),
            ("市政路桥", ["市政路桥公司"]),
            ("设备安装", ["设备安装公司"]),
            ("国际公司", ["国际公司"]),
            ("马来公司", ["马来公司"]),
        ]
        all_companies = sorted({company for _, companies in row_defs for company in companies}, key=lambda name: company_sort_key(name, self.key))
        current_period = f"{year:04d}-{month:02d}"
        current_month_stats = self.quarter_severe_overdue_stats(rows, all_companies, period_report_date(current_period))
        self.record_settlement_overdue_month_stats(current_period, current_month_stats, all_companies)

        stats_from_months, current_missing_months = self.quarter_overdue_stats_from_months(
            year, quarter, all_companies, current_period, rows
        )
        stats = (
            stats_from_months
            if not current_missing_months
            else self.quarter_severe_overdue_stats(rows, all_companies, report_cutoff)
        )

        previous_year, previous_q = self.previous_quarter(year, quarter)
        previous_stats, previous_missing_months = self.quarter_overdue_stats_from_months(
            previous_year, previous_q, all_companies, current_period, rows
        )
        if previous_missing_months:
            manual_previous_stats = self.manual_quarter_overdue_stats(previous_year, previous_q, all_companies)
            if manual_previous_stats is not None:
                previous_stats = manual_previous_stats
                previous_missing_months = []
        previous_ready = not previous_missing_months

        data_row = header_row + 2
        total = {"total": 0, "m12_24": 0, "m24": 0, "penalty": 0}
        previous_total = 0
        for index, (label, companies) in enumerate(row_defs, 1):
            values = self.sum_quarter_overdue_stats(stats, companies)
            previous_values = self.sum_quarter_overdue_stats(previous_stats, companies)
            if label not in {"其中：十公司", "装饰幕墙"}:
                for key in total:
                    total[key] += values[key]
                previous_total += previous_values["total"] if previous_ready else 0
            row_values = [
                index,
                label,
                previous_values["total"] if previous_ready else None,
                values["total"],
                values["m12_24"],
                values["m24"],
                values["penalty"],
            ]
            for col_index, value in enumerate(row_values, 1):
                cell = ws.cell(data_row, col_index)
                cell.value = value
                cell.font = body_font
                cell.alignment = center
                cell.border = border
            ws.row_dimensions[data_row].height = 20
            data_row += 1

        ws.merge_cells(start_row=data_row, start_column=1, end_row=data_row, end_column=2)
        total_values = {
            1: "合计",
            3: previous_total if previous_ready else None,
            4: total["total"],
            5: total["m12_24"],
            6: total["m24"],
            7: total["penalty"],
        }
        for col_index, value in total_values.items():
            cell = ws.cell(data_row, col_index)
            cell.value = value
            cell.font = header_font
            cell.alignment = center
            cell.border = border
        for col_index in range(1, 8):
            ws.cell(data_row, col_index).border = border
        data_row += 2

        footer_lines = [
            "集团处罚款项在通知下发后五日内上缴集团公司财务部。",
            "各单位应加快严重逾期工程的定案进度，并加强竣工结算的过程管理，全面提升竣工结算的定案效率和质量。",
        ]
        if previous_missing_months:
            ws.cell(header_row, 3).comment = Comment(
                f"上季度列取自月度严重逾期台账；缺少{', '.join(previous_missing_months)}，本次暂留空。",
                "ReportHub",
            )
        if current_missing_months:
            ws.cell(header_row, 4).comment = Comment(
                f"本季度月度台账缺少{', '.join(current_missing_months)}，本次本季度明细暂按当前报表生成。",
                "ReportHub",
            )
        for text in footer_lines:
            ws.merge_cells(start_row=data_row, start_column=1, end_row=data_row, end_column=7)
            ws.cell(data_row, 1).value = text
            ws.cell(data_row, 1).font = body_font
            ws.cell(data_row, 1).alignment = left
            data_row += 1

        data_row += 2
        ws.merge_cells(start_row=data_row, start_column=5, end_row=data_row, end_column=7)
        ws.cell(data_row, 5).value = "成本管理部"
        ws.cell(data_row, 5).font = body_font
        ws.cell(data_row, 5).alignment = center
        ws.merge_cells(start_row=data_row + 1, start_column=5, end_row=data_row + 1, end_column=7)
        ws.cell(data_row + 1, 5).value = f"{year}年{month}月25日"
        ws.cell(data_row + 1, 5).font = body_font
        ws.cell(data_row + 1, 5).alignment = center

    @staticmethod
    def chinese_quarter(quarter: int) -> str:
        return {1: "一季度", 2: "二季度", 3: "三季度", 4: "四季度"}.get(quarter, "本季度")

    @staticmethod
    def quarter_key(year: int, quarter: int) -> str:
        return f"{year:04d}-Q{quarter}"

    @staticmethod
    def previous_quarter(year: int, quarter: int) -> tuple[int, int]:
        if quarter <= 1:
            return year - 1, 4
        return year, quarter - 1

    @staticmethod
    def quarter_month_periods(year: int, quarter: int) -> list[str]:
        start_month = (quarter - 1) * 3 + 1
        return [f"{year:04d}-{month:02d}" for month in range(start_month, start_month + 3)]

    @staticmethod
    def empty_quarter_overdue_stats(companies: list[str]) -> dict[str, dict[str, int]]:
        return {company: {"total": 0, "m12_24": 0, "m24": 0, "penalty": 0} for company in companies}

    @staticmethod
    def normalized_quarter_overdue_stats(
        stats: dict[str, Any] | None,
        companies: list[str],
    ) -> dict[str, dict[str, int]]:
        normalized = SettlementPlugin.empty_quarter_overdue_stats(companies)
        if not isinstance(stats, dict):
            return normalized
        for company in companies:
            source = stats.get(company, {})
            if not isinstance(source, dict):
                continue
            for key in ("total", "m12_24", "m24", "penalty"):
                normalized[company][key] = int(to_number(source.get(key)))
        return normalized

    def settlement_rows_for_period(self, period: str) -> list[dict[str, Any]]:
        period_rows: list[dict[str, Any]] = []
        try:
            uploads = load_db().get("uploads", [])
        except Exception:
            uploads = []
        for upload in uploads:
            if upload.get("reportType") == self.key and upload.get("period") == period:
                period_rows.extend(upload.get("rows") or [])
        return unique_settlement_rows(period_rows)

    def record_settlement_overdue_month_stats(
        self,
        period: str,
        stats: dict[str, dict[str, int]],
        companies: list[str],
    ) -> None:
        match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
        if not match:
            return
        year, month = int(match.group(1)), int(match.group(2))
        quarter = ceil(month / 3)
        ledger = load_settlement_overdue_month_ledger()
        ledger.setdefault("months", {})
        ledger["months"][f"{year:04d}-{month:02d}"] = {
            "period": f"{year:04d}-{month:02d}",
            "reportDate": period_report_date(period),
            "quarter": self.quarter_key(year, quarter),
            "generatedAt": datetime.now().isoformat(timespec="seconds"),
            "companies": self.normalized_quarter_overdue_stats(stats, companies),
        }
        save_settlement_overdue_month_ledger(ledger)

    def settlement_overdue_month_stats(
        self,
        period: str,
        companies: list[str],
        current_period: str = "",
        current_rows: list[dict[str, Any]] | None = None,
    ) -> dict[str, dict[str, int]] | None:
        period_rows = current_rows if period == current_period and current_rows is not None else self.settlement_rows_for_period(period)
        if period_rows:
            stats = self.quarter_severe_overdue_stats(period_rows, companies, period_report_date(period))
            self.record_settlement_overdue_month_stats(period, stats, companies)
            return stats
        ledger = load_settlement_overdue_month_ledger()
        month_entry = ledger.get("months", {}).get(period)
        if isinstance(month_entry, dict):
            companies_entry = month_entry.get("companies")
            if isinstance(companies_entry, dict):
                return self.normalized_quarter_overdue_stats(companies_entry, companies)
        return None

    def quarter_overdue_stats_from_months(
        self,
        year: int,
        quarter: int,
        companies: list[str],
        current_period: str = "",
        current_rows: list[dict[str, Any]] | None = None,
    ) -> tuple[dict[str, dict[str, int]], list[str]]:
        result = self.empty_quarter_overdue_stats(companies)
        missing: list[str] = []
        for month_period in self.quarter_month_periods(year, quarter):
            month_stats = self.settlement_overdue_month_stats(month_period, companies, current_period, current_rows)
            if month_stats is None:
                missing.append(month_period)
                continue
            for company in companies:
                for key in ("total", "m12_24", "m24", "penalty"):
                    result[company][key] += month_stats.get(company, {}).get(key, 0)
        return result, missing

    def manual_quarter_overdue_stats(
        self,
        year: int,
        quarter: int,
        companies: list[str],
    ) -> dict[str, dict[str, int]] | None:
        ledger = load_settlement_overdue_month_ledger()
        quarter_entry = ledger.get("manualQuarters", {}).get(self.quarter_key(year, quarter))
        if not isinstance(quarter_entry, dict):
            return None
        companies_entry = quarter_entry.get("companies")
        if not isinstance(companies_entry, dict):
            return None
        return self.normalized_quarter_overdue_stats(companies_entry, companies)

    def quarter_severe_overdue_stats(self, rows: list[dict[str, Any]], companies: list[str], quarter_end: str) -> dict[str, dict[str, int]]:
        stats = {company: {"total": 0, "m12_24": 0, "m24": 0, "penalty": 0} for company in companies}
        report_date = self.date_value(quarter_end)
        if not report_date:
            return stats
        for row in rows:
            company = standard_company_name(row.get("company", "")) or row.get("company", "")
            if company not in stats or not counts_as_settlement_project(row):
                continue
            if row.get("segment") in {"在建", "未开工", "诉讼", "停工"} or row.get("status") in {"在建", "未开工", "诉讼", "停工"}:
                continue
            if row.get("status") == "已定案" or classify_settlement_abc(row, quarter_end) == "A":
                continue
            completion_date = self.date_value(row.get("completion_date"))
            if not completion_date:
                continue
            months = (report_date.year - completion_date.year) * 12 + report_date.month - completion_date.month
            if report_date.day < completion_date.day:
                months -= 1
            if months <= 12:
                continue
            item = stats[company]
            item["total"] += 1
            if months <= 24:
                item["m12_24"] += 1
                item["penalty"] += months * 1000
            else:
                item["m24"] += 1
                item["penalty"] += months * 5000
        return stats

    @staticmethod
    def sum_quarter_overdue_stats(stats: dict[str, dict[str, int]], companies: list[str]) -> dict[str, int]:
        result = {"total": 0, "m12_24": 0, "m24": 0, "penalty": 0}
        for company in companies:
            for key in result:
                result[key] += stats.get(company, {}).get(key, 0)
        return result

    def write_quarter_severe_overdue_detail_sheet(self, ws: Any, rows: list[dict[str, Any]], period: str) -> None:
        period_match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
        year = int(period_match.group(1)) if period_match else datetime.now().year
        month = int(period_match.group(2)) if period_match else datetime.now().month
        quarter = ceil(month / 3)
        # 明细与06统计表使用同一实际统计截止日，避免把季度未来月份提前计入。
        report_cutoff = period_report_date(period)

        details = self.quarter_severe_overdue_details(rows, report_cutoff)

        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        widths = [6, 11, 34, 11, 11, 11, 11, 11, 12, 12, 12, 12, 13, 14, 11, 12, 12, 12, 12]
        for col_index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col_index)].width = width

        title_font = Font(name="宋体", size=14, bold=True)
        note_font = Font(name="宋体", size=10)
        header_font = Font(name="宋体", size=9, bold=True)
        body_font = Font(name="宋体", size=9)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        header_fill = PatternFill("solid", fgColor="F2F2F2")
        total_fill = PatternFill("solid", fgColor="E2F0D9")

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=19)
        ws["A1"] = f"{year}年{self.chinese_quarter(quarter)}严重逾期工程明细表"
        ws["A1"].font = title_font
        ws["A1"].alignment = center
        ws.row_dimensions[1].height = 28

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=19)
        ws["A2"] = (
            "说明：逾期12个月至24个月（含24个月）的工程，对基层单位处以每工程1000元/月罚款；"
            "逾期24个月以上的工程，对基层单位处以每工程5000元/月罚款"
        )
        ws["A2"].font = note_font
        ws["A2"].alignment = left
        ws.row_dimensions[2].height = 24

        header_row = 3
        for col in [1, 2, 3, 4, 5, 12, 13, 14, 15, 16, 17, 18, 19]:
            ws.merge_cells(start_row=header_row, start_column=col, end_row=header_row + 1, end_column=col)
        ws.cell(header_row, 1).value = "序号"
        ws.cell(header_row, 2).value = "单位"
        ws.cell(header_row, 3).value = "工程名称"
        ws.cell(header_row, 4).value = "建筑面积\n(平米)"
        ws.cell(header_row, 5).value = "合同价款\n(万元)"
        ws.merge_cells(start_row=header_row, start_column=6, end_row=header_row, end_column=7)
        ws.cell(header_row, 6).value = "开工时间"
        ws.merge_cells(start_row=header_row, start_column=8, end_row=header_row, end_column=9)
        ws.cell(header_row, 8).value = "竣工时间"
        ws.merge_cells(start_row=header_row, start_column=10, end_row=header_row, end_column=11)
        ws.cell(header_row, 10).value = "结算递交时间"
        ws.cell(header_row, 12).value = "结算递交\n造价(万元)"
        ws.cell(header_row, 13).value = "结算\n定案情况"
        ws.cell(header_row, 14).value = "项目\n经理"
        ws.cell(header_row, 15).value = "结算\n负责人"
        ws.cell(header_row, 16).value = "分管\n副总经理"
        ws.cell(header_row, 17).value = "定案逾期时间（月）"
        ws.cell(header_row, 18).value = "计划\n定案时间"
        ws.cell(header_row, 19).value = "罚款金额（元）"
        for col, label in [(6, "合同"), (7, "实际"), (8, "合同"), (9, "竣工报告"), (10, "应递交"), (11, "实际递交")]:
            ws.cell(header_row + 1, col).value = label

        for row_index in range(header_row, header_row + 2):
            ws.row_dimensions[row_index].height = 26
            for col_index in range(1, 20):
                cell = ws.cell(row_index, col_index)
                cell.font = header_font
                cell.alignment = center
                cell.border = border
                cell.fill = header_fill

        data_row = header_row + 2
        total_area = 0.0
        total_contract = 0.0
        total_submitted = 0.0
        total_penalty = 0
        for index, item in enumerate(details, 1):
            row = item["row"]
            source_extra = self.source_people_values(row)
            area = to_number(row.get("building_area"))
            contract = row.get("contract_amount", 0) or 0
            submitted = row.get("submitted_amount", 0) or 0
            penalty = item["penalty"]
            total_area += area
            total_contract += contract
            total_submitted += submitted
            total_penalty += penalty
            row_values = [
                index,
                item["company_label"],
                row.get("project", ""),
                area or row.get("building_area", ""),
                contract,
                self.detail_date_text(row.get("contract_start_date")),
                self.detail_date_text(row.get("actual_start_date")),
                self.detail_date_text(row.get("contract_completion_date")),
                self.detail_date_text(row.get("completion_date")),
                self.detail_date_text(row.get("submit_due_date")),
                self.detail_date_text(row.get("actual_submit_date")),
                submitted,
                row.get("status", ""),
                row.get("manager", "") or source_extra.get("manager", ""),
                source_extra.get("settlement_responsible", ""),
                source_extra.get("deputy_manager", ""),
                item["months"],
                self.detail_month_text(row.get("settlement_receive_month")),
                penalty,
            ]
            for col_index, value in enumerate(row_values, 1):
                cell = ws.cell(data_row, col_index)
                cell.value = value
                cell.font = body_font
                cell.alignment = left if col_index == 3 else center
                cell.border = border
            for col_index in [4, 5, 12]:
                ws.cell(data_row, col_index).number_format = "#,##0.00"
            ws.cell(data_row, 19).number_format = "#,##0"
            ws.row_dimensions[data_row].height = 32
            data_row += 1

        ws.merge_cells(start_row=data_row, start_column=1, end_row=data_row, end_column=3)
        total_values = {
            1: "小计",
            4: total_area,
            5: total_contract,
            12: total_submitted,
            17: f"平均{round(sum(item['months'] for item in details) / len(details))}个月" if details else "",
            19: total_penalty,
        }
        for col_index in range(1, 20):
            cell = ws.cell(data_row, col_index)
            if col_index not in {2, 3}:
                cell.value = total_values.get(col_index, "")
            cell.font = header_font
            cell.alignment = center
            cell.border = border
            cell.fill = total_fill
        for col_index in [4, 5, 12]:
            ws.cell(data_row, col_index).number_format = "#,##0.00"
        ws.cell(data_row, 19).number_format = "#,##0"

    def quarter_severe_overdue_details(self, rows: list[dict[str, Any]], quarter_end: str) -> list[dict[str, Any]]:
        report_date = self.date_value(quarter_end)
        if not report_date:
            return []
        company_order = [
            "三公司", "四公司", "五公司", "六公司", "七公司", "十公司",
            "青岛公司", "格瑞特公司", "市政路桥公司", "装饰幕墙公司", "设备安装公司",
        ]
        label_map = {
            "格瑞特公司": "格瑞特",
            "市政路桥公司": "市政路桥",
            "装饰幕墙公司": "装饰幕墙",
            "设备安装公司": "设备安装",
        }
        details: list[dict[str, Any]] = []
        for row in rows:
            company = standard_company_name(row.get("company", "")) or row.get("company", "")
            if company not in company_order or not counts_as_settlement_project(row):
                continue
            if row.get("segment") in {"在建", "未开工", "诉讼", "停工"} or row.get("status") in {"在建", "未开工", "诉讼", "停工"}:
                continue
            if row.get("status") == "已定案" or classify_settlement_abc(row, quarter_end) == "A":
                continue
            completion_date = self.date_value(row.get("completion_date"))
            if not completion_date:
                continue
            months = (report_date.year - completion_date.year) * 12 + report_date.month - completion_date.month
            if report_date.day < completion_date.day:
                months -= 1
            if months <= 12:
                continue
            details.append({
                "row": row,
                "company": company,
                "company_label": label_map.get(company, company),
                "months": months,
                "bucket": "12-24个月" if months <= 24 else "24个月以上",
                "penalty": months * (1000 if months <= 24 else 5000),
            })
        details.sort(key=lambda item: (
            company_order.index(item["company"]),
            -item["months"],
            clean_text(item["row"].get("project")),
        ))
        return details

    def source_people_values(self, row: dict[str, Any]) -> dict[str, str]:
        source_path = clean_text(row.get("_source_path"))
        project = clean_text(row.get("project"))
        if not source_path or not project:
            return {}
        cache = getattr(self, "_source_people_cache", None)
        if cache is None:
            cache = {}
            setattr(self, "_source_people_cache", cache)
        cache_key = (source_path, clean_text(row.get("sheet")))
        if cache_key not in cache:
            values: dict[str, dict[str, str]] = {}
            try:
                wb = load_workbook(source_path, data_only=True)
                sheet_name = cache_key[1]
                worksheets = [wb[sheet_name]] if sheet_name in wb.sheetnames else wb.worksheets
                for ws in worksheets:
                    columns = self.detail_columns(ws)
                    project_col = columns["project"] + 1
                    manager_col = columns.get("manager", 14) + 1
                    for cells in ws.iter_rows():
                        source_project = clean_text(cells[project_col - 1].value if len(cells) >= project_col else "")
                        if not source_project:
                            continue
                        row_index = cells[0].row
                        values[source_project] = {
                            "manager": clean_text(ws.cell(row_index, manager_col).value),
                            "settlement_responsible": clean_text(ws.cell(row_index, manager_col + 1).value),
                            "deputy_manager": clean_text(ws.cell(row_index, manager_col + 2).value),
                        }
            except Exception:
                values = {}
            cache[cache_key] = values
        return cache[cache_key].get(project, {})

    def write_submit_assessment_sheet(self, ws: Any, rows: list[dict[str, Any]], period: str) -> None:
        period_match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
        year = int(period_match.group(1)) if period_match else datetime.now().year
        month = int(period_match.group(2)) if period_match else datetime.now().month
        period_title = f"{year}年{month}月"

        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        widths = [8, 13, 36, 13, 13]
        for col_index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col_index)].width = width

        title_font = Font(name="宋体", size=14, bold=True)
        body_font = Font(name="宋体", size=10)
        table_font = Font(name="宋体", size=10)
        table_bold = Font(name="宋体", size=10, bold=True)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws.merge_cells("A2:E2")
        ws["A2"] = f"关于{period_title}竣工结算递交情况的考核"
        ws["A2"].font = title_font
        ws["A2"].alignment = center
        ws.row_dimensions[2].height = 30

        ws.merge_cells("A4:E4")
        ws["A4"] = "各基层单位及相关部门："
        ws["A4"].font = body_font
        ws["A4"].alignment = left
        ws.merge_cells("A5:E9")
        ws["A5"] = (
            "依据《工程竣工结算管理办法》（烟建〔2026〕14号）第5.4条“未在规定时间内向发包人递交竣工结算书的，"
            "对基层单位以每个工程1000元/月罚款；逾期3个月未递交的，对基层单位以每个工程2000元/月罚款；"
            "逾期6个月仍未递交的，对基层单位以每个工程4000元/月罚款，集团公司对基层单位负责人、分管副总进行同等处罚。"
            f"截至{month}月20日，下列基层单位及项目的竣工结算应递交而未递交，具体单位及工程如下："
        )
        ws["A5"].font = body_font
        ws["A5"].alignment = left

        table_start = 10
        headers = ["序号", "单位\n名称", "工程名称", "合同额\n(万元)", "处罚金额\n(元)"]
        for col_index, value in enumerate(headers, 1):
            cell = ws.cell(table_start, col_index)
            cell.value = value
            cell.font = table_bold
            cell.alignment = center
            cell.border = border
        ws.row_dimensions[table_start].height = 28

        items = self.submit_assessment_items(rows, period)
        display_count = max(len(items), 5)
        row_index = table_start + 1
        total_contract = 0.0
        total_penalty = 0
        for offset in range(display_count):
            item = items[offset] if offset < len(items) else None
            values = [
                offset + 1,
                item["company"] if item else "",
                item["project"] if item else "",
                item["contract_amount"] if item else "",
                item["penalty"] if item else "",
            ]
            if item:
                total_contract += item["contract_amount"]
                total_penalty += item["penalty"]
            for col_index, value in enumerate(values, 1):
                cell = ws.cell(row_index, col_index)
                cell.value = value
                cell.border = border
                cell.font = table_font
                cell.alignment = left if col_index == 3 else center
            ws.cell(row_index, 4).number_format = "#,##0.00"
            ws.cell(row_index, 5).number_format = "#,##0"
            ws.row_dimensions[row_index].height = 26
            row_index += 1

        ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=3)
        ws.cell(row_index, 1).value = "合计"
        ws.cell(row_index, 4).value = total_contract if items else ""
        ws.cell(row_index, 5).value = total_penalty if items else ""
        for col_index in range(1, 6):
            cell = ws.cell(row_index, col_index)
            cell.border = border
            cell.font = table_bold
            cell.alignment = center
        ws.cell(row_index, 4).number_format = "#,##0.00"
        ws.cell(row_index, 5).number_format = "#,##0"
        ws.row_dimensions[row_index].height = 26

        note_row = row_index + 2
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 1, end_column=5)
        ws.cell(note_row, 1).value = "基层单位在流程结束后五日内将罚款上缴集团公司财务部。"
        ws.cell(note_row, 1).font = body_font
        ws.cell(note_row, 1).alignment = center

        sign_row = note_row + 4
        ws.merge_cells(start_row=sign_row, start_column=4, end_row=sign_row, end_column=5)
        ws.cell(sign_row, 4).value = "成本管理部"
        ws.cell(sign_row, 4).font = body_font
        ws.cell(sign_row, 4).alignment = center
        ws.merge_cells(start_row=sign_row + 1, start_column=4, end_row=sign_row + 1, end_column=5)
        ws.cell(sign_row + 1, 4).value = f"{year}年{month}月{datetime.now().day}日"
        ws.cell(sign_row + 1, 4).font = body_font
        ws.cell(sign_row + 1, 4).alignment = center

    def submit_assessment_items(self, rows: list[dict[str, Any]], period: str) -> list[dict[str, Any]]:
        window_start = self.date_value(report_window_start(period))
        report_date = self.date_value(period_report_date(period))
        if not report_date or not window_start:
            return []
        items: list[dict[str, Any]] = []
        for row in rows:
            if row.get("status") != "未定案":
                continue
            due_date = self.date_value(row.get("submit_due_date"))
            if not due_date:
                continue
            actual_submit = self.date_value(row.get("actual_submit_date"))
            if actual_submit:
                continue
            if due_date > report_date:
                continue
            check_date = report_date
            days_overdue = (check_date - due_date).days
            months_overdue = max(1, ceil(days_overdue / 30))
            if months_overdue >= 6:
                rate = 4000
            elif months_overdue >= 3:
                rate = 2000
            else:
                rate = 1000
            items.append(
                {
                    "company": row.get("company") or "",
                    "project": row.get("project") or "",
                    "contract_amount": row.get("contract_amount", 0) or 0,
                    "penalty": months_overdue * rate,
                    "months_overdue": months_overdue,
                }
            )
        items.sort(key=lambda item: (company_sort_key(item["company"], self.key), item["project"]))
        return items

    @staticmethod
    def detail_date_text(value: Any) -> str:
        text = clean_text(value)
        match = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", text)
        if match:
            return f"{match.group(1)}/{int(match.group(2))}/{int(match.group(3))}"
        return text

    @staticmethod
    def detail_month_text(value: Any) -> str:
        text = clean_text(value)
        match = re.match(r"^(\d{4})-(\d{1,2})(?:-\d{1,2})?$", text)
        if match:
            return f"{int(match.group(2))}月"
        match = re.match(r"^(\d{4})/(\d{1,2})(?:/\d{1,2})?$", text)
        if match:
            return f"{int(match.group(2))}月"
        return text

    @staticmethod
    def month_value_matches(value: Any, period: str) -> bool:
        text = clean_text(value)
        period_match = re.match(r"^(\d{4})-(\d{1,2})$", period or "")
        if not text or not period_match:
            return False
        period_year, period_month = int(period_match.group(1)), int(period_match.group(2))
        match = re.match(r"^(\d{4})[-/年.](\d{1,2})(?:[-/月.]\d{1,2})?", text)
        if match:
            return int(match.group(1)) == period_year and int(match.group(2)) == period_month
        match = re.match(r"^(\d{1,2})月$", text)
        if match:
            return int(match.group(1)) == period_month
        return False

    @staticmethod
    def sortable_date(value: Any) -> str:
        text = clean_text(value)
        match = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", text)
        if match:
            return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
        return text or "9999-12-31"

    @staticmethod
    def date_value(value: Any) -> datetime | None:
        text = clean_text(value)
        match = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", text)
        if not match:
            match = re.match(r"^(\d{4})/(\d{1,2})/(\d{1,2})$", text)
        if not match:
            return None
        year, month, day = (int(part) for part in match.groups())
        try:
            return datetime(year, month, day)
        except ValueError:
            return None

    @classmethod
    def settlement_overdue_months(cls, row: dict[str, Any]) -> int | str:
        settled_date = cls.date_value(row.get("settled_date"))
        report_date = cls.date_value(row.get("actual_submit_date")) or cls.date_value(row.get("completion_date"))
        if not settled_date or not report_date:
            source_value = row.get("settlement_overdue_months", "")
            if source_value is None:
                return ""
            if isinstance(source_value, float) and source_value.is_integer():
                return int(source_value)
            return source_value
        days = (settled_date - report_date).days
        if days <= 0:
            return 0
        return ceil(days / 30)

    @staticmethod
    def is_settled_in_current_window(row: dict[str, Any], period: str) -> bool:
        settled_date = parse_date(row.get("settled_date"))
        return bool(settled_date and report_window_start(period) < settled_date <= period_report_date(period))

    def previous_period_settlement_state(self, period: str) -> dict[str, Any]:
        """Return the prior report's latest branch rows for one-time monthly backfill checks."""
        previous_period = previous_report_period(period)
        if not previous_period:
            return {"companies": set(), "settledProjects": set()}
        try:
            db_mtime = DB_PATH.stat().st_mtime_ns
        except OSError:
            db_mtime = 0
        cache_key = (period, db_mtime)
        cache = getattr(self, "_previous_period_settlement_cache", {})
        if cache_key in cache:
            return cache[cache_key]

        latest_by_company: dict[str, dict[str, Any]] = {}
        for upload in load_db().get("uploads", []):
            if upload.get("reportType") != self.key or upload.get("period") != previous_period:
                continue
            company = standard_company_name(upload.get("company", "")) or clean_text(upload.get("company", ""))
            if not company:
                continue
            current = latest_by_company.get(company)
            if current is None or clean_text(upload.get("uploadedAt")) >= clean_text(current.get("uploadedAt")):
                latest_by_company[company] = upload

        settled_projects: set[tuple[str, str]] = set()
        for company, upload in latest_by_company.items():
            for previous_row in upload.get("rows", []):
                project = clean_text(previous_row.get("project", ""))
                was_settled = previous_row.get("status") == "已定案" or previous_row.get("segment") == "settled"
                if project and was_settled:
                    settled_projects.add((company, project))
        result = {"companies": set(latest_by_company), "settledProjects": settled_projects}
        cache[cache_key] = result
        # Keep only the current cache entry so a long-running service does not retain old upload states.
        self._previous_period_settlement_cache = {cache_key: result}
        return result

    def previous_period_plan_rows(self, period: str) -> list[dict[str, Any]]:
        """Use the latest uploaded branch sheet from the prior report period for 02 monthly targets."""
        previous_period = previous_report_period(period)
        if not previous_period:
            return []
        latest_by_company: dict[str, dict[str, Any]] = {}
        for upload in load_db().get("uploads", []):
            if upload.get("reportType") != self.key or upload.get("period") != previous_period:
                continue
            company = standard_company_name(upload.get("company", "")) or clean_text(upload.get("company", ""))
            if not company:
                continue
            existing = latest_by_company.get(company)
            if existing is None or clean_text(upload.get("uploadedAt")) >= clean_text(existing.get("uploadedAt")):
                latest_by_company[company] = upload
        return [
            row
            for upload in latest_by_company.values()
            for row in upload.get("rows", [])
        ]

    def is_previous_period_late_reported_settlement(self, row: dict[str, Any], period: str) -> bool:
        """Include a project once when its prior-period settlement was unavailable in the prior report."""
        settled_date = parse_date(row.get("settled_date"))
        previous_period = previous_report_period(period)
        if not settled_date or not previous_period:
            return False
        if not (report_window_start(previous_period) < settled_date <= period_report_date(previous_period)):
            return False
        company = standard_company_name(row.get("company", "")) or clean_text(row.get("company", ""))
        project = clean_text(row.get("project", ""))
        if not company or not project:
            return False
        previous_state = self.previous_period_settlement_state(period)
        # No previous branch upload, or the project was still unsettled then: include it once as a late report.
        return company not in previous_state["companies"] or (company, project) not in previous_state["settledProjects"]

    def is_month_settled(self, row: dict[str, Any], period: str) -> bool:
        return self.is_settled_in_current_window(row, period) or self.is_previous_period_late_reported_settlement(row, period)

    def month_settlement_marker_conflict(self, row: dict[str, Any], period: str) -> str:
        marker = bool(row.get("has_month_settled_marker"))
        in_current_window = self.is_settled_in_current_window(row, period)
        is_late_report = self.is_previous_period_late_reported_settlement(row, period)
        in_period = in_current_window or is_late_report
        if not marker and not in_period:
            return ""
        company = standard_company_name(row.get("company", "")) or clean_text(row.get("company", ""))
        source_file = clean_text(row.get("_source_file"))
        source_row = clean_text(row.get("source_row"))
        if source_file and source_row:
            location = f"{source_file}第{source_row}行"
        elif source_file:
            location = source_file
        elif source_row:
            location = f"源表第{source_row}行"
        else:
            location = "源表"
        settled_date = parse_date(row.get("settled_date"))
        project = clean_text(row.get("project"))
        prefix = f"{company}《{project}》{location}："
        if is_late_report:
            return ""
        if marker and not settled_date:
            return prefix + "V/W等“本月新增定案”辅助列有标记，但L列结算定案时间为空或无法解析；本月统计未纳入，请核对源表。"
        if marker and not in_period:
            return prefix + f"辅助列有本月新增定案标记，但L列结算定案时间为{settled_date}，不在{report_window_begin(period)}至{period_report_date(period)}的统计区间内；本月统计未纳入。"
        if in_period and not marker:
            return prefix + f"L列结算定案时间为{settled_date}，落在{report_window_begin(period)}至{period_report_date(period)}统计区间内，但V/W等辅助列无本月新增标记；程序已按日期纳入，请核对辅助列。"
        return ""

    def sync_month_settlement_marker_issues(self, rows: list[dict[str, Any]], period: str) -> int:
        if not period or period == "全部":
            return 0
        conflicts: list[dict[str, str]] = []
        seen: set[str] = set()
        for row in rows:
            if row.get("status") != "已定案" and not row.get("has_month_settled_marker"):
                continue
            description = self.month_settlement_marker_conflict(row, period)
            if not description:
                continue
            company = standard_company_name(row.get("company", "")) or clean_text(row.get("company", ""))
            project = clean_text(row.get("project"))
            source_row = clean_text(row.get("source_row"))
            auto_key = f"month-settled-marker|{period}|{company}|{project}|{source_row}"
            if auto_key in seen:
                continue
            seen.add(auto_key)
            conflicts.append({"company": company, "description": description, "autoKey": auto_key})

        db = load_db()
        existing = db.get("issues", [])
        db["issues"] = [
            issue for issue in existing
            if not (
                issue.get("reportType") == self.key
                and issue.get("period") == period
                and issue.get("autoSource") == "month-settled-marker"
            )
        ]
        for conflict in reversed(conflicts):
            db.setdefault("issues", []).insert(0, {
                "id": uuid.uuid4().hex,
                "reportType": self.key,
                "period": period,
                "issueType": "日期与辅助列冲突",
                "company": conflict["company"],
                "description": conflict["description"],
                "status": "待处理",
                "autoSource": "month-settled-marker",
                "autoKey": conflict["autoKey"],
                "createdAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            })
        save_db(db)
        return len(conflicts)

    def write_summary_sheet(self, ws: Any, rows: list[dict[str, Any]], period: str) -> None:
        start = period_start(period)
        cutoff = settlement_abc_cutoff(period)
        report_date = parse_date(period_report_date(period))
        report_day_label = f"{datetime.strptime(report_date, '%Y-%m-%d').month}月{datetime.strptime(report_date, '%Y-%m-%d').day}日" if report_date else "本月20日"
        month_window_label = report_window_label(period)
        # Keep all displayed dates in the template synchronized with the selected report month.
        ws["B3"] = f"截至{report_day_label}竣工未定案工程"
        ws["J3"] = f"截至{report_day_label}竣工工程"
        ws["P3"] = f"截至{report_day_label}竣工定案工程"
        ws["P4"] = f"本月（{month_window_label}）"
        by_company: dict[str, dict[str, Any]] = {}
        for row in rows:
            company = row.get("company") or "未填单位"
            item = by_company.setdefault(
                company,
                {
                    "company": company,
                    "unsettled_exam_count": 0,
                    "unsettled_exam_amount": 0.0,
                    "unsettled_future_count": 0,
                    "unsettled_future_amount": 0.0,
                    "completed_exam_count": 0,
                    "completed_exam_amount": 0.0,
                    "completed_future_count": 0,
                    "completed_future_amount": 0.0,
                    "settled_month_count": 0,
                    "settled_month_contract": 0.0,
                    "settled_month_amount": 0.0,
                    "settled_total_count": 0,
                    "settled_total_contract": 0.0,
                    "settled_total_amount": 0.0,
                },
            )
            contract_amount = self.summary_contract_amount(row)
            settled_amount = row.get("settled_amount", 0) or 0
            if not settled_amount:
                settled_amount = self.source_settled_amount(row)
            settled_date = row.get("settled_date") or ""
            month_settled_count = row.get("month_settled_count", 0) or 0
            month_settled_amount = row.get("month_settled_amount", 0) or 0
            abc_segment = classify_settlement_abc(row, cutoff)
            if not abc_segment or not row_is_within_report_cutoff(row, period):
                continue
            count_project = 1 if counts_as_settlement_project(row) else 0

            if abc_segment == "A":
                item["completed_exam_count"] += count_project
                item["completed_exam_amount"] += contract_amount
                item["settled_total_count"] += count_project
                item["settled_total_contract"] += contract_amount
                item["settled_total_amount"] += settled_amount
                if self.is_month_settled(row, period):
                    item["settled_month_count"] += count_project or month_settled_count or 1
                    item["settled_month_contract"] += contract_amount
                    item["settled_month_amount"] += settled_amount or month_settled_amount
            elif abc_segment == "C":
                item["completed_future_count"] += count_project
                item["completed_future_amount"] += contract_amount
                item["unsettled_future_count"] += count_project
                item["unsettled_future_amount"] += contract_amount
            else:
                item["completed_exam_count"] += count_project
                item["completed_exam_amount"] += contract_amount
                item["unsettled_exam_count"] += count_project
                item["unsettled_exam_amount"] += contract_amount

        companies = sorted(by_company.values(), key=lambda item: company_sort_key(item["company"], self.key))
        detail_start = 6
        total_row = detail_start + len(companies)
        if len(companies) > 7:
            ws.insert_rows(13, len(companies) - 7)
            for row_index in range(13, total_row):
                copy_row_style(ws, 12, row_index, 21)
        elif len(companies) < 7:
            copy_row_style(ws, 13, total_row, 21)

        for row_index in range(detail_start, total_row + 1):
            for col_index in range(1, 22):
                ws.cell(row_index, col_index).value = None

        total_unsettled_exam = sum(item["unsettled_exam_amount"] for item in companies)
        total_unsettled_all = sum(item["unsettled_exam_amount"] + item["unsettled_future_amount"] for item in companies)
        for offset, item in enumerate(companies):
            row_index = detail_start + offset
            ws.cell(row_index, 1).value = item["company"]
            ws.cell(row_index, 2).value = item["unsettled_exam_count"]
            ws.cell(row_index, 3).value = item["unsettled_exam_amount"]
            ws.cell(row_index, 4).value = item["unsettled_exam_amount"] / total_unsettled_exam if total_unsettled_exam else 0
            ws.cell(row_index, 5).value = item["unsettled_future_count"]
            ws.cell(row_index, 6).value = item["unsettled_future_amount"]
            ws.cell(row_index, 7).value = item["unsettled_exam_count"] + item["unsettled_future_count"]
            ws.cell(row_index, 8).value = item["unsettled_exam_amount"] + item["unsettled_future_amount"]
            ws.cell(row_index, 9).value = (
                (item["unsettled_exam_amount"] + item["unsettled_future_amount"]) / total_unsettled_all
                if total_unsettled_all else 0
            )
            ws.cell(row_index, 10).value = item["completed_exam_count"]
            ws.cell(row_index, 11).value = item["completed_exam_amount"]
            ws.cell(row_index, 12).value = item["completed_future_count"]
            ws.cell(row_index, 13).value = item["completed_future_amount"]
            ws.cell(row_index, 14).value = item["completed_exam_count"] + item["completed_future_count"]
            ws.cell(row_index, 15).value = item["completed_exam_amount"] + item["completed_future_amount"]
            ws.cell(row_index, 16).value = item["settled_month_count"]
            ws.cell(row_index, 17).value = item["settled_month_contract"]
            ws.cell(row_index, 18).value = item["settled_month_amount"]
            ws.cell(row_index, 19).value = item["settled_total_count"]
            ws.cell(row_index, 20).value = item["settled_total_contract"]
            ws.cell(row_index, 21).value = item["settled_total_amount"]

        ws.cell(total_row, 1).value = "小计"
        for col_index in [2, 3, 5, 6, 7, 8, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21]:
            ws.cell(total_row, col_index).value = sum(ws.cell(row_index, col_index).value or 0 for row_index in range(detail_start, total_row))
        ws.cell(total_row, 4).value = 1 if total_unsettled_exam else 0
        ws.cell(total_row, 9).value = 1 if total_unsettled_all else 0
        for row_index in range(detail_start, total_row + 1):
            for col_index in (3, 6, 8, 11, 13, 15, 17, 18, 20, 21):
                ws.cell(row_index, col_index).number_format = "#,##0.00"
        if len(companies) < 7:
            ws.delete_rows(total_row + 1, 7 - len(companies))
        if period:
            month_text = report_day_label
            exam_completed_count = ws.cell(total_row, 10).value or 0
            exam_completed_amount = ws.cell(total_row, 11).value or 0
            month_settled_count = ws.cell(total_row, 16).value or 0
            month_settled_contract = ws.cell(total_row, 17).value or 0
            total_settled_count = ws.cell(total_row, 19).value or 0
            total_settled_contract = ws.cell(total_row, 20).value or 0
            unsettled_count = ws.cell(total_row, 7).value or 0
            unsettled_amount = ws.cell(total_row, 8).value or 0
            amount_to_yi = lambda value: round((value or 0) / 10000, 2)
            ws["A2"] = (
                f"截至{month_text}："
                f"纳入考核的竣工工程{int(exam_completed_count)}个，合同额{amount_to_yi(exam_completed_amount)}亿元；"
                f"本月定案工程{int(month_settled_count)}个，合同额{amount_to_yi(month_settled_contract)}亿元；"
                f"累计定案工程{int(total_settled_count)}个，合同额{amount_to_yi(total_settled_contract)}亿元；"
                f"未定案工程{int(unsettled_count)}个，合同额{amount_to_yi(unsettled_amount)}亿元。"
            )


class RiskLevelPlugin(ReportPlugin):
    key = "risk-level"
    name = "成本风险等级评定表"
    description = "汇总项目风险等级、类别、责任人与整改状态。"

    header_keywords = ["项目", "风险", "等级", "责任", "整改"]

    def parse(self, file_path: Path, company: str) -> ParseResult:
        wb = load_workbook(file_path, data_only=True)
        rows: list[dict[str, Any]] = []
        errors: list[str] = []
        warnings: list[str] = []
        for ws in wb.worksheets:
            sheet_company = company or standard_company_name(clean_text(ws["B2"].value))
            block_start = 5
            while block_start <= ws.max_column:
                project = clean_text(ws.cell(2, block_start).value)
                block_company = sheet_company or self.company_for_project_block(ws, block_start)
                if not sheet_company and standard_company_name(project) == project:
                    project = clean_text(ws.cell(3, block_start).value)
                if not project:
                    project = clean_text(ws.cell(3, block_start).value)
                if not project or project in {"项目情形", "得分", "乘权重\n得分"}:
                    block_start += 3
                    continue
                dimensions: list[dict[str, Any]] = []
                total = 0.0
                for row_index in range(4, 16):
                    weight = to_number(ws.cell(row_index, 3).value)
                    score = to_number(ws.cell(row_index, block_start + 1).value)
                    weighted = to_number(ws.cell(row_index, block_start + 2).value)
                    if not weighted and score and weight:
                        weighted = round(score * weight, 4)
                    total += weighted
                    dimensions.append(
                        {
                            "index": ws.cell(row_index, 1).value,
                            "name": clean_text(ws.cell(row_index, 2).value),
                            "weight": weight,
                            "standard": clean_text(ws.cell(row_index, 4).value),
                            "situation": clean_text(ws.cell(row_index, block_start).value),
                            "score": score,
                            "weighted": weighted,
                        }
                    )
                level = clean_text(ws.cell(17, block_start + 2).value) or classify_risk_level(total)
                rows.append(
                    {
                        "company": block_company or ws.title,
                        "sheet": ws.title,
                        "project": project,
                        "total": round(total, 4),
                        "risk_level": level,
                        "dimensions": dimensions,
                    }
                )
                block_start += 3
        if not rows:
            errors.append("未识别到风险明细行，请确认是否使用了约定模板。")
        return ParseResult(rows, errors, warnings, wb.sheetnames)

    def company_for_project_block(self, ws: Any, block_start: int) -> str:
        for col_index in range(block_start, 4, -1):
            text = standard_company_name(clean_text(ws.cell(2, col_index).value))
            if text:
                return text
        return standard_company_name(clean_text(ws["B2"].value))

    def aggregate(self, rows: list[dict[str, Any]]) -> dict[str, Any]:
        high = [row for row in rows if row.get("risk_level") == "A"]
        by_level: dict[str, int] = {}
        by_company: dict[str, dict[str, Any]] = {}
        for row in rows:
            level = row.get("risk_level") or "未评级"
            by_level[level] = by_level.get(level, 0) + 1
            item = by_company.setdefault(row["company"], {"company": row["company"], "projects": 0, "A": 0, "B": 0, "C": 0, "D": 0})
            item["projects"] += 1
            if level in {"A", "B", "C", "D"}:
                item[level] += 1
        return {
            "cards": [
                {"label": "风险项目", "value": len(rows), "unit": "项"},
                {"label": "A级风险", "value": len(high), "unit": "项"},
                {"label": "最高评分", "value": round(max([row.get("total", 0) for row in rows] or [0]), 2), "unit": "分"},
                {"label": "涉及单位", "value": len(by_company), "unit": "个"},
            ],
            "table": sorted(by_company.values(), key=lambda item: company_sort_key(item["company"], self.key)),
            "levels": by_level,
        }

    def export(self, rows: list[dict[str, Any]], output_path: Path, period: str) -> None:
        rows = sorted(rows, key=lambda row: company_sort_key(row.get("company", ""), self.key))
        template = TEMPLATE_DIR / "risk-level-template.xlsx"
        wb = load_workbook(template) if template.exists() else Workbook()
        ws = wb.active
        ws.title = f"{period}汇总" if period else "风险等级汇总"
        if period:
            year, month = period.split("-", 1)
            ws["A1"] = f"项目成本风险等级评分汇总表（{year}年{int(month)}月）"

        for merged in list(ws.merged_cells.ranges):
            if merged.min_row >= 2 and merged.max_col >= 5:
                ws.unmerge_cells(str(merged))

        max_col = max(ws.max_column, 4 + max(len(rows), 1) * 3)
        for row_index in range(2, 19):
            for col_index in range(5, max_col + 1):
                style_from(ws, row_index, 5 + (col_index - 5) % 3, row_index, col_index)

        for row_index in range(2, 19):
            for col_index in range(5, max_col + 1):
                ws.cell(row_index, col_index).value = None

        company_ranges: dict[str, list[int]] = {}
        for project_index, row in enumerate(rows):
            start_col = 5 + project_index * 3
            end_col = start_col + 2
            company_ranges.setdefault(row.get("company", ""), []).extend([start_col, end_col])
            ws.merge_cells(start_row=3, start_column=start_col, end_row=3, end_column=end_col)
            ws.cell(3, start_col).value = row.get("project")
            ws.cell(4, start_col).value = "项目情形"
            ws.cell(4, start_col + 1).value = "得分"
            ws.cell(4, start_col + 2).value = "乘权重\n得分"
            for offset, dimension in enumerate(row.get("dimensions", []), start=5):
                ws.cell(offset, start_col).value = dimension.get("situation")
                ws.cell(offset, start_col + 1).value = dimension.get("score")
                ws.cell(offset, start_col + 2).value = dimension.get("weighted")
            ws.cell(17, start_col).value = "乘权重\n得分合计"
            ws.cell(17, start_col + 2).value = row.get("total")
            ws.cell(18, start_col).value = "风险等级"
            ws.cell(18, start_col + 2).value = row.get("risk_level")
            apply_cell_box(ws, 17, 18, start_col, end_col)
            ws.merge_cells(start_row=17, start_column=start_col, end_row=17, end_column=start_col + 1)
            ws.merge_cells(start_row=18, start_column=start_col, end_row=18, end_column=start_col + 1)
            ws.cell(17, start_col).value = "乘权重\n得分合计"
            ws.cell(18, start_col).value = "风险等级"
            for col_index in range(start_col, end_col + 1):
                ws.cell(17, col_index).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.cell(18, col_index).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        company_fills = [
            PatternFill("solid", fgColor="E2F0D9"),
            PatternFill("solid", fgColor="D9EAF7"),
        ]
        for company_index, (company, bounds) in enumerate(company_ranges.items()):
            start_col, end_col = min(bounds), max(bounds)
            company_fill = company_fills[company_index % len(company_fills)]
            apply_cell_box(ws, 2, 2, start_col, end_col)
            for col_index in range(start_col, end_col + 1):
                cell = ws.cell(2, col_index)
                cell.fill = copy(company_fill)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if start_col < end_col:
                ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=end_col)
            ws.cell(2, start_col).value = company
            ws.cell(2, start_col).fill = copy(company_fill)
            ws.cell(2, start_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_index in range(5, max_col + 1):
            ws.column_dimensions[get_column_letter(col_index)].width = 13
        self.format_footer(ws, max_col)
        wb.save(output_path)

    def format_footer(self, ws: Any, max_col: int) -> None:
        for merged in list(ws.merged_cells.ranges):
            if (merged.min_col == 1 and merged.max_col <= 4 and merged.min_row <= 18 and merged.max_row >= 17) or merged.min_row == 19:
                ws.unmerge_cells(str(merged))

        risk_text = "风险等级说明：\n评分结果≥80，A级\n60≤评分结果＜80，B级\n40≤评分结果＜60，C级\n评分结果＜40，D级"
        note_text = (
            "填表说明：\n"
            "1.本表按发布之日新开工项目（含合作施工项目）每月更新填报一次，于每月30日前返回。\n"
            "2.结合评分标准，按项目实际情况填报项目情形列内容，并按得分和乘权重得分分别打分；项目情形相近的，须避免得分偏差过大。\n"
            "3.最终评分结果按乘权重得分合计计算，并根据风险等级说明确定风险等级。\n"
            "4.按基层单位评分和集团部门评分取高值，得最终评分并确定风险等级。"
        )

        apply_cell_box(ws, 17, 18, 1, 4)
        for row_index in (17, 18):
            for col_index in range(1, 5):
                ws.cell(row_index, col_index).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=17, start_column=1, end_row=18, end_column=4)
        ws["A17"] = risk_text
        ws["A17"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for col_index in range(1, max_col + 1):
            ws.cell(19, col_index).border = Border()
            ws.cell(19, col_index).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=19, start_column=1, end_row=19, end_column=max_col)
        ws["A19"] = note_text
        ws["A19"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[17].height = 30
        ws.row_dimensions[18].height = 30
        ws.row_dimensions[19].height = 75


def infer_risk_level(score: float) -> str:
    if score >= 80:
        return "高风险"
    if score >= 60:
        return "中风险"
    if score > 0:
        return "低风险"
    return "未评级"


def write_sheet(ws: Any, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    header_fill = PatternFill("solid", fgColor="DCE8F2")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="1F2933")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def autosize(ws: Any) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        width = max(len(clean_text(cell.value)) for cell in column_cells) + 2
        ws.column_dimensions[letter].width = min(max(width, 10), 36)


MATERIAL_PRICE_HEADERS = {
    "material": ["材料名称", "材料", "物资名称", "名称"],
    "spec": ["规格型号", "规格", "型号"],
    "unit": ["单位"],
    "price": ["采购单价", "单价", "含税单价", "价格"],
    "quantity": ["数量", "采购数量"],
    "amount": ["金额", "采购金额", "含税金额"],
    "purchaseDate": ["采购日期", "采购时间", "下单日期", "订单日期"],
    "project": ["项目", "项目名称"],
    "supplier": ["供应商", "供应单位"],
    "company": ["分公司", "采购单位", "单位名称"],
    "platformOrder": ["订单编号", "采购单号", "合同编号"],
}

BID_COMPANY_COLUMNS = [
    ("三公司", ["三公司", "3gs"]),
    ("四公司", ["四公司", "4gs"]),
    ("五公司", ["五公司", "5gs"]),
    ("六公司", ["六公司", "6gs"]),
    ("七公司", ["七公司", "7gs"]),
    ("十公司", ["十公司", "10gs"]),
    ("市政路桥/格瑞特", ["市政路桥", "市政路桥公司", "格瑞特", "格瑞特公司", "szlqgs"]),
    ("青岛公司", ["青岛公司", "qdgs"]),
    ("济南公司", ["济南公司", "jngs"]),
    ("国际公司", ["国际公司", "gjgs"]),
    ("设备安装", ["设备安装", "设备安装公司", "sbazgs"]),
    ("上海公司", ["上海公司", "shgs"]),
    ("东泰物流", ["东泰", "东泰物流", "dtgs"]),
    ("其他公司", ["其他公司", "othergs"]),
]

BID_QUERY_ALIASES = {
    "listNo": ["清单编号", "清单编码"],
    "material": ["材料名称", "清单名称", "物资名称"],
    "spec": ["规格/项目特征", "规格型号", "规格", "项目特征"],
    "unit": ["单位"],
    "company": ["公司名称", "需求单位", "采购单位"],
    "project": ["项目名称"],
    "process": ["采购过程名称", "招标名称", "任务名称"],
    "price": ["含税中标价格", "中标价格", "含税单价", "含税单价(元)", "中标单价"],
    "notaxPrice": ["无税中标价格", "不含税中标价格", "不含税单价"],
    "quantity": ["中标数量", "数量", "采购数量"],
    "amount": ["含税金额", "中标金额", "采购金额", "金额"],
    "bidDate": ["中标日期"],
    "brand": ["品牌/厂家", "品牌", "厂家"],
    "remark": ["备注"],
    "explanation": ["报价备注", "原因说明"],
    "payment": ["付款方式"],
    "paymentDesc": ["付款方式说明", "付款说明"],
    "taskNo": ["任务编号", "招标任务编号", "采购过程编号"],
    "winner": ["中标单位", "中标供应商", "供应商名称"],
    "region": ["地区", "地域"],
}

BID_COMPARISON_ALIASES = {
    "listNo": ["清单编号"],
    "material": ["清单名称", "材料名称", "物资名称"],
    "spec": ["规格/项目特征", "规格型号", "规格", "项目特征"],
    "unit": ["单位"],
    "groupAvg": ["集团平均中标价格", "groupavgprice"],
}

BID_LEDGER_ALIASES = {
    "taskNo": ["任务编号", "招标任务编号"],
    "winner": ["中标单位", "中标供应商", "供应商名称"],
}

BID_PM_ALIASES = {
    "department": ["施工项目部", "项目名称"],
    "project": ["工程项目名称", "项目名称"],
    "region": ["地域", "地区"],
}


def normalize_material_key(row: dict[str, Any]) -> str:
    return "|".join([
        clean_text(row.get("material")).lower(),
        clean_text(row.get("spec")).lower(),
        clean_text(row.get("unit")).lower(),
    ])


def parse_material_price_date(value: Any) -> str:
    parsed = parse_date(value)
    return parsed or clean_text(value)


def date_in_range(value: Any, start_date: str, end_date: str) -> bool:
    text = parse_material_price_date(value)
    if not text:
        return True
    if start_date and text < start_date:
        return False
    if end_date and text > end_date:
        return False
    return True


def material_source_paths() -> list[Path]:
    return [
        DATA_DIR / "material_price_source.xlsx",
        DATA_DIR / "material_price_source.json",
    ]


def material_price_companies() -> list[str]:
    companies = [
        company
        for company in company_order("settlement")
        if "装饰幕墙" not in company
    ]
    if "东泰物流" not in companies:
        companies.append("东泰物流")
    return ["集团"] + companies


def material_price_task_name(task: dict[str, Any]) -> str:
    company = clean_text(task.get("company")) or "集团"
    if company == "全集团":
        company = "集团"
    keyword = clean_text(task.get("keyword"))
    suffix = f"（{keyword}）" if keyword else ""
    return f"{company}材料采购价格横向对比{suffix}"


def material_price_codex_job_path(task_id: str) -> Path:
    safe_id = re.sub(r"[^0-9a-zA-Z_-]", "", clean_text(task_id))
    return CODEX_JOB_DIR / f"{safe_id}.json"


def write_material_price_codex_job(task: dict[str, Any]) -> Path:
    job_path = material_price_codex_job_path(task.get("id", ""))
    payload = {
        "id": task.get("id"),
        "type": "material-price-collection",
        "status": "pending",
        "createdAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "title": task.get("name") or material_price_task_name(task),
        "instruction": "请复用已登录集采平台，进入数据中心-项目定制报表，按筛选时间段导出中标价格对比分析、中标价格查询，并在交易报表导出招标结果台帐；采集完成后由小程序按照材料价格技能生成四张表。",
        "filters": {
            "company": task.get("company") or "集团",
            "startDate": task.get("startDate") or "",
            "endDate": task.get("endDate") or "",
            "keyword": task.get("keyword") or "",
        },
        "excludeCompanies": ["装饰幕墙公司"],
        "requiredExports": [
            {"platform": "集采平台", "report": "中标价格对比分析", "dateField": "招标结果发布时间"},
            {"platform": "集采平台", "report": "中标价格查询", "dateField": "中标日期"},
            {"platform": "集采平台", "report": "招标结果台帐", "dateField": "结果发布日期"},
        ],
        "preferredWriteBack": "将集采平台导出的源表保存到当前任务 sourceDir，并回写 sourcePaths。",
        "businessRules": [
            "集团只表示查询范围，表三公司名称必须写平台中的实际采购单位名称。",
            "删除采购过程名称或需求单位含“测试”的行。",
            "付款方式 CASH/FQ/DFW 分别显示为 现金/分期/抵房/物。",
            "按材料价格技能生成表一、表二、表三、表四；不足两家有效采购单位的材料不进入横向成果。",
        ],
        "afterWriteBack": "回写完成后，小程序刷新材料价格任务即可自动按材料价格技能成表并开放下载。",
    }
    with job_path.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    return job_path


def list_material_price_codex_jobs() -> list[dict[str, Any]]:
    jobs: list[dict[str, Any]] = []
    for path in sorted(CODEX_JOB_DIR.glob("*.json"), key=lambda item: item.stat().st_mtime, reverse=True):
        try:
            with path.open("r", encoding="utf-8-sig") as f:
                job = json.load(f)
        except (OSError, json.JSONDecodeError):
            continue
        job["_path"] = str(path)
        jobs.append(job)
    return jobs


def read_json_file(path: Path, fallback: Any) -> Any:
    try:
        with path.open("r", encoding="utf-8-sig") as f:
            return json.load(f)
    except (OSError, json.JSONDecodeError):
        return fallback


def pm_warning_safe_platform(platform: str) -> str:
    value = clean_text(platform) or "old-pm"
    return value if value in {"old-pm", "big-pm", "all"} else "old-pm"


def pm_warning_job_path(job_id: str) -> Path:
    safe_id = re.sub(r"[^0-9a-zA-Z_-]", "", clean_text(job_id))
    return PM_WARNING_JOB_DIR / f"{safe_id}.json"


def pm_warning_latest_path(platform: str) -> Path:
    return PM_WARNING_RESULT_DIR / f"latest_{pm_warning_safe_platform(platform)}.json"


def pm_warning_debug_profile_dir(platform: str) -> Path:
    profile_root = Path.home() / "AppData/Local/CodexPmProfiles"
    if platform == "old-pm":
        return profile_root / "edge9222-profile"
    return profile_root / "chrome-bigpm-debug-profile"


def find_browser_executable(browser_name: str) -> str:
    is_chrome = browser_name == "chrome"
    candidates = ["chrome", "chrome.exe"] if is_chrome else ["msedge", "msedge.exe"]
    for name in candidates:
        found = shutil.which(name)
        if found:
            return found
    if is_chrome:
        browser_paths = [
            Path("C:/Program Files/Google/Chrome/Application/chrome.exe"),
            Path("C:/Program Files (x86)/Google/Chrome/Application/chrome.exe"),
            Path.home() / "AppData/Local/Google/Chrome/Application/chrome.exe",
        ]
    else:
        browser_paths = [
            Path("C:/Program Files/Microsoft/Edge/Application/msedge.exe"),
            Path("C:/Program Files (x86)/Microsoft/Edge/Application/msedge.exe"),
            Path.home() / "AppData/Local/Microsoft/Edge/Application/msedge.exe",
        ]
    for candidate in browser_paths:
        if candidate.exists():
            return str(candidate)
    return "chrome.exe" if is_chrome else "msedge.exe"


def pm_warning_cdp_open(cdp_url: str) -> bool:
    try:
        with urlopen(cdp_url, timeout=1.5) as response:
            return getattr(response, "status", 200) == 200
    except Exception:
        return False


def stop_edge_startup_boost_processes() -> None:
    if not sys.platform.startswith("win"):
        return
    command = (
        "Get-CimInstance Win32_Process -Filter \"Name='msedge.exe'\" | "
        "Where-Object { $_.CommandLine -match '--no-startup-window' } | "
        "ForEach-Object { Stop-Process -Id $_.ProcessId -Force -ErrorAction SilentlyContinue }"
    )
    try:
        subprocess.run(
            ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", command],
            cwd=str(ROOT),
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=5,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
    except Exception:
        pass


def launch_pm_warning_debug_browser(platform: str) -> tuple[bool, str]:
    safe_platform = pm_warning_safe_platform(platform)
    is_big_pm = safe_platform == "big-pm"
    port = "9333" if is_big_pm else "9222"
    url = PM_WARNING_BIG_PM_URL if is_big_pm else PM_WARNING_OLD_PM_URL
    profile_dir = pm_warning_debug_profile_dir(safe_platform)
    profile_dir.mkdir(parents=True, exist_ok=True)
    browser = find_browser_executable("chrome" if is_big_pm else "edge")
    if not is_big_pm:
        stop_edge_startup_boost_processes()
    args = [
        browser,
        f"--remote-debugging-port={port}",
        "--remote-debugging-address=127.0.0.1",
        f"--user-data-dir={profile_dir}",
        "--new-window",
        "--no-first-run",
        "--no-default-browser-check",
        "--disable-features=msStartupBoost",
        url,
    ]
    try:
        startupinfo = None
        if sys.platform.startswith("win"):
            startupinfo = subprocess.STARTUPINFO()
            startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
            startupinfo.wShowWindow = 1
        subprocess.Popen(
            args,
            cwd=str(ROOT),
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            startupinfo=startupinfo,
            creationflags=subprocess.CREATE_NO_WINDOW if sys.platform.startswith("win") else 0,
        )
        return True, f"已自动打开{pm_warning_platform_name(safe_platform)}调试浏览器（端口 {port}）。"
    except Exception as exc:
        return False, f"自动打开{pm_warning_platform_name(safe_platform)}调试浏览器失败：{exc}"


def start_pm_warning_debug_browser_async(platform: str = "old-pm") -> None:
    safe_platform = pm_warning_safe_platform(platform)
    cdp_url = PM_WARNING_BIG_PM_CDP_LIST_URL if safe_platform == "big-pm" else PM_WARNING_CDP_LIST_URL
    if pm_warning_cdp_open(cdp_url):
        return
    with PM_WARNING_DEBUG_BROWSER_BOOT_LOCK:
        if safe_platform in PM_WARNING_DEBUG_BROWSER_BOOTING:
            return
        PM_WARNING_DEBUG_BROWSER_BOOTING.add(safe_platform)

    def worker() -> None:
        try:
            if pm_warning_cdp_open(cdp_url):
                return
            launched, message = launch_pm_warning_debug_browser(safe_platform)
            print(message)
            if launched:
                for _ in range(16):
                    time.sleep(0.5)
                    if pm_warning_cdp_open(cdp_url):
                        print(f"{pm_warning_platform_name(safe_platform)}调试端口已就绪。")
                        break
        finally:
            with PM_WARNING_DEBUG_BROWSER_BOOT_LOCK:
                PM_WARNING_DEBUG_BROWSER_BOOTING.discard(safe_platform)

    threading.Thread(target=worker, name=f"pm-warning-debug-browser-{safe_platform}", daemon=True).start()


def start_pm_warning_debug_browsers_for_server(port: int) -> list[str]:
    if port != 8899 or not PM_WARNING_START_DEBUG_BROWSERS_ON_8899:
        return []
    messages: list[str] = []
    cdp_urls = {
        "old-pm": PM_WARNING_CDP_LIST_URL,
        "big-pm": PM_WARNING_BIG_PM_CDP_LIST_URL,
    }
    for platform in PM_WARNING_START_DEBUG_PLATFORMS_ON_8899:
        cdp_url = cdp_urls.get(platform, PM_WARNING_CDP_LIST_URL)
        platform_name = pm_warning_platform_name(platform)
        debug_port = "9333" if platform == "big-pm" else "9222"
        if pm_warning_cdp_open(cdp_url):
            check = check_pm_warning_big_pm_cdp() if platform == "big-pm" else check_pm_warning_cdp()
            if check.get("ok"):
                messages.append(f"{platform_name}调试端口 {debug_port} 已就绪。")
                continue
        launched, message = launch_pm_warning_debug_browser(platform)
        if not launched:
            messages.append(message)
            continue
        ready = False
        for _ in range(16):
            time.sleep(0.5)
            if pm_warning_cdp_open(cdp_url):
                ready = True
                break
        if ready:
            messages.append(f"{platform_name}调试端口 {debug_port} 已随 8899 服务启动。")
        else:
            messages.append(f"{message} 但暂未检测到端口 {debug_port}，请确认浏览器窗口没有被安全软件拦截。")
    return messages


def ensure_pm_warning_debug_browser(platform: str) -> tuple[bool, dict[str, Any]]:
    safe_platform = pm_warning_safe_platform(platform)
    cdp_url = PM_WARNING_BIG_PM_CDP_LIST_URL if safe_platform == "big-pm" else PM_WARNING_CDP_LIST_URL
    check = check_pm_warning_big_pm_cdp() if safe_platform == "big-pm" else check_pm_warning_cdp()
    if check.get("ok"):
        return True, check
    if not pm_warning_cdp_open(cdp_url):
        if not PM_WARNING_AUTO_LAUNCH_BROWSER:
            port = "9333" if safe_platform == "big-pm" else "9222"
            browser_name = "Chrome" if safe_platform == "big-pm" else "Edge"
            return False, {
                "ok": False,
                "autoLaunched": False,
                "message": f"未检测到{pm_warning_platform_name(safe_platform)}调试端口 {port}。为避免影响你当前工作，程序不会自动打开浏览器；请先手动打开专用{browser_name}调试窗口并登录后，再点击“开始生成”。",
                "previous": check,
            }
        launched, launch_message = launch_pm_warning_debug_browser(safe_platform)
        for _ in range(12):
            time.sleep(0.5)
            check = check_pm_warning_big_pm_cdp() if safe_platform == "big-pm" else check_pm_warning_cdp()
            if check.get("ok"):
                check["autoLaunched"] = launched
                check["launchMessage"] = launch_message
                return True, check
        message = (
            f"{launch_message} 请在新打开的浏览器里完成登录，登录完成后再点击“开始生成”。"
            if launched
            else launch_message
        )
        return False, {"ok": False, "autoLaunched": launched, "message": message, "previous": check}
    check["message"] = f"{check.get('message', '')} 如页面未登录，请在这个调试浏览器中登录后再点击“开始生成”。"
    return False, check


def check_pm_warning_cdp(cdp_url: str = PM_WARNING_CDP_LIST_URL) -> dict[str, Any]:
    try:
        with urlopen(cdp_url, timeout=2) as response:
            status = getattr(response, "status", 200)
            if status != 200:
                return {
                    "ok": False,
                    "message": f"四版 PM 浏览器调试端口返回异常：HTTP {status}。请确认已登录的浏览器是用 9222 调试端口打开的。",
                }
            pages = json.loads(response.read().decode("utf-8", errors="replace"))
    except Exception:
        return {
            "ok": False,
            "message": f"未连接到四版 PM 浏览器调试端口 9222（{cdp_url}）。请先用带 --remote-debugging-port=9222 的已登录 Edge/Chrome 打开四版 PM 页面，再点击“四版平台”。",
        }

    if not isinstance(pages, list):
        return {
            "ok": False,
            "message": "四版 PM 浏览器调试端口返回内容异常，请重启带 9222 调试端口的已登录浏览器后重试。",
        }

    page = next((
        item for item in pages
        if isinstance(item, dict)
        and item.get("type") == "page"
        and (
            "/yjpm2012/MainNew.aspx" in clean_text(item.get("url"))
            or "/yjpm2012/DefaultMainNew.aspx" in clean_text(item.get("url"))
            or "烟建综合项目管理系统" in clean_text(item.get("title"))
            or "烟建PM平台" in clean_text(item.get("title"))
        )
    ), None)
    if not page:
        open_pages = "；".join(
            " ".join(value for value in [clean_text(item.get("title")), clean_text(item.get("url"))] if value)
            for item in pages
            if isinstance(item, dict) and item.get("type") == "page"
        )
        return {
            "ok": False,
            "message": f"已连接到浏览器调试端口，但没有找到已登录的四版 PM 主页面。请在该浏览器中打开并登录四版 PM 后重试。当前页面：{open_pages[:300] or '无'}",
        }

    return {"ok": True, "message": "四版 PM 浏览器连接正常。", "page": {"title": page.get("title", ""), "url": page.get("url", "")}}


def check_pm_warning_big_pm_cdp(cdp_url: str = PM_WARNING_BIG_PM_CDP_LIST_URL) -> dict[str, Any]:
    try:
        with urlopen(cdp_url, timeout=2) as response:
            status = getattr(response, "status", 200)
            if status != 200:
                return {
                    "ok": False,
                    "message": f"大PM Chrome 调试端口返回异常：HTTP {status}。请确认已登录的大PM Chrome 使用 9333 调试端口打开。",
                }
            pages = json.loads(response.read().decode("utf-8", errors="replace"))
    except Exception:
        return {
            "ok": False,
            "message": f"未连接到大PM Chrome 调试端口 9333（{cdp_url}）。请先打开并登录大PM调试 Chrome，再点击“大PM平台”。",
        }

    if not isinstance(pages, list):
        return {
            "ok": False,
            "message": "大PM Chrome 调试端口返回内容异常，请重启带 9333 调试端口的已登录 Chrome 后重试。",
        }

    page = next((
        item for item in pages
        if isinstance(item, dict)
        and item.get("type") == "page"
        and "yanjianpm.glodon.com" in clean_text(item.get("url"))
        and "login" not in clean_text(item.get("url")).lower()
    ), None)
    if not page:
        open_pages = "；".join(
            " ".join(value for value in [clean_text(item.get("title")), clean_text(item.get("url"))] if value)
            for item in pages
            if isinstance(item, dict) and item.get("type") == "page"
        )
        return {
            "ok": False,
            "message": f"已连接到 9333 调试端口，但没有找到已登录的大PM页面。请在该 Chrome 中打开并登录大PM后重试。当前页面：{open_pages[:300] or '无'}",
        }

    return {"ok": True, "message": "大PM Chrome 浏览器连接正常。", "page": {"title": page.get("title", ""), "url": page.get("url", "")}}


def pm_warning_fallback_data() -> dict[str, Any]:
    return {
        "source": {"displaySource": "尚未执行本次抓数"},
        "companies": [],
        "listChecks": [],
        "projects": [],
        "totals": {"companies": 0, "inProgress": 0, "normal": 0, "blue": 0, "red": 0, "detailProjects": 0},
        "captureCompleteness": {"ok": True, "expected": 0, "actual": 0, "missing": 0, "incompleteCount": 0, "incompleteChecks": []},
    }


def pm_warning_platform_name(platform: str) -> str:
    return {"old-pm": "四版平台", "big-pm": "大PM平台", "all": "两个平台汇总"}.get(platform, platform)


def pm_warning_result_is_publishable(platform: str, data: dict[str, Any]) -> bool:
    if platform != "big-pm":
        return True
    capture = data.get("capture") if isinstance(data.get("capture"), dict) else {}
    source = data.get("source") if isinstance(data.get("source"), dict) else {}
    return clean_text(capture.get("status") or source.get("status")).lower() == "completed"


def pm_warning_capture_completeness_from_list_checks(data: dict[str, Any]) -> dict[str, Any]:
    checks = data.get("listChecks") if isinstance(data, dict) else []
    expected = 0
    actual = 0
    incomplete_checks: list[dict[str, Any]] = []
    for check in checks or []:
        if not isinstance(check, dict):
            continue
        item_expected = pm_warning_int(check.get("expected"))
        item_actual = pm_warning_int(check.get("actual"))
        if item_expected <= 0:
            continue
        expected += item_expected
        actual += item_actual
        missing = max(0, item_expected - item_actual)
        if missing:
            incomplete_checks.append({
                "reportPeriod": clean_text(check.get("reportPeriod")),
                "company": clean_text(check.get("company")),
                "type": clean_text(check.get("type")),
                "status": clean_text(check.get("status")),
                "expected": item_expected,
                "actual": item_actual,
                "missing": missing,
                "attempts": pm_warning_int(check.get("attempts")) or 1,
                "message": clean_text(check.get("message")),
                "sourcePlatform": clean_text(check.get("sourcePlatform")),
            })
    missing_total = sum(pm_warning_int(item.get("missing")) for item in incomplete_checks)
    return {
        "ok": missing_total == 0,
        "expected": expected,
        "actual": actual,
        "missing": missing_total,
        "checked": sum(1 for check in checks or [] if isinstance(check, dict) and pm_warning_int(check.get("expected")) > 0),
        "incompleteCount": len(incomplete_checks),
        "incompleteChecks": incomplete_checks[:200],
    }


def load_pm_warning_single_platform(platform: str) -> dict[str, Any]:
    safe_platform = pm_warning_safe_platform(platform)
    latest_path = pm_warning_latest_path(safe_platform)
    data = read_json_file(latest_path, None)
    if data is not None and not pm_warning_result_is_publishable(safe_platform, data):
        data = None
    if data is None:
        data = pm_warning_fallback_data()
        data.setdefault("source", {})["displaySource"] = "尚未执行本次抓数"
    else:
        data.setdefault("source", {})["displaySource"] = f"{pm_warning_platform_name(safe_platform)}实时抓数"
    data.setdefault("source", {})["platform"] = safe_platform
    for item in data.get("companies") or []:
        if isinstance(item, dict):
            item["sourcePlatform"] = safe_platform
    for item in data.get("projects") or []:
        if isinstance(item, dict):
            item["sourcePlatform"] = safe_platform
    if not isinstance(data.get("captureCompleteness"), dict):
        data["captureCompleteness"] = pm_warning_capture_completeness_from_list_checks(data)
    return data


def pm_warning_empty_reason_consistency() -> dict[str, Any]:
    return {
        "sourcePlatform": "all",
        "platforms": [],
        "totalProjects": 0,
        "eligibleProjects": 0,
        "skippedMissingPath": 0,
        "checked": 0,
        "conflicts": 0,
        "reviews": 0,
        "ok": 0,
        "errors": [],
        "checks": [],
    }


def pm_warning_merge_reason_consistency(target: dict[str, Any], source: Any, platform: str) -> None:
    if not isinstance(source, dict):
        return
    target.setdefault("platforms", []).append(
        {
            "platform": platform,
            "checked": pm_warning_int(source.get("checked")),
            "conflicts": pm_warning_int(source.get("conflicts")),
            "reviews": pm_warning_int(source.get("reviews")),
            "ok": pm_warning_int(source.get("ok")),
            "skippedMissingPath": pm_warning_int(source.get("skippedMissingPath")),
        }
    )
    for key in ("totalProjects", "eligibleProjects", "skippedMissingPath", "checked", "conflicts", "reviews", "ok"):
        target[key] = pm_warning_int(target.get(key)) + pm_warning_int(source.get(key))
    for item in source.get("errors") or []:
        if isinstance(item, dict):
            target.setdefault("errors", []).append({**item, "sourcePlatform": item.get("sourcePlatform") or platform})
        else:
            target.setdefault("errors", []).append({"sourcePlatform": platform, "message": clean_text(item)})
    if clean_text(source.get("error")):
        target.setdefault("errors", []).append({"sourcePlatform": platform, "message": clean_text(source.get("error"))})
    for check in source.get("checks") or []:
        if isinstance(check, dict):
            target.setdefault("checks", []).append({**check, "sourcePlatform": check.get("sourcePlatform") or platform})


def merge_pm_warning_platform_data(platforms: list[str]) -> dict[str, Any]:
    merged = pm_warning_fallback_data()
    merged["source"] = {"platform": "all", "displaySource": "两个平台汇总"}
    merged["companies"] = []
    merged["listChecks"] = []
    merged["projects"] = []
    merged["totals"] = {"companies": 0, "inProgress": 0, "normal": 0, "blue": 0, "red": 0, "detailProjects": 0}
    merged["detailHealth"] = {"total": 0, "real": 0, "placeholder": 0, "complete": 0, "missing": 0, "contract": 0, "budget": 0, "actual": 0, "reduction": 0}
    merged["captureCompleteness"] = {"ok": True, "expected": 0, "actual": 0, "missing": 0, "incompleteCount": 0, "incompleteChecks": []}
    merged["reasonConsistency"] = pm_warning_empty_reason_consistency()
    for platform in platforms:
        data = load_pm_warning_single_platform(platform)
        merged["companies"].extend(data.get("companies") or [])
        merged["listChecks"].extend(data.get("listChecks") or [])
        merged["projects"].extend(data.get("projects") or [])
        pm_warning_merge_reason_consistency(merged["reasonConsistency"], data.get("reasonConsistency"), platform)
        totals = data.get("totals") or {}
        for key in ("companies", "inProgress", "normal", "blue", "red", "detailProjects"):
            merged["totals"][key] += int(totals.get(key) or 0)
        detail_health = data.get("detailHealth") or {}
        for key in ("total", "real", "placeholder", "complete", "missing", "contract", "budget", "actual", "reduction"):
            merged["detailHealth"][key] += int(detail_health.get(key) or 0)
        completeness = data.get("captureCompleteness") if isinstance(data.get("captureCompleteness"), dict) else {}
        merged["captureCompleteness"]["expected"] += int(completeness.get("expected") or 0)
        merged["captureCompleteness"]["actual"] += int(completeness.get("actual") or 0)
        merged["captureCompleteness"]["missing"] += int(completeness.get("missing") or 0)
        merged["captureCompleteness"]["incompleteCount"] += int(completeness.get("incompleteCount") or 0)
        merged["captureCompleteness"]["incompleteChecks"].extend(completeness.get("incompleteChecks") or [])
    merged["captureCompleteness"]["ok"] = merged["captureCompleteness"]["missing"] == 0
    merged["reasonConsistency"]["hasData"] = bool(merged["reasonConsistency"].get("platforms"))
    return merged


def pm_warning_missing_value(value: Any) -> bool:
    return value is None or clean_text(value) == ""


def pm_warning_number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = clean_text(value).replace(",", "").replace("%", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def pm_warning_has_value(value: Any) -> bool:
    return not pm_warning_missing_value(value)


def pm_warning_period_window(period: str) -> list[str]:
    match = re.match(r"^(\d{4})-(\d{2})$", clean_text(period))
    if not match:
        now = datetime.now()
        base_year = now.year
        base_month = now.month
    else:
        base_year = int(match.group(1))
        base_month = int(match.group(2))
    periods: list[str] = []
    for offset in (-2, -1, 0):
        month_index = (base_year * 12 + base_month - 1) + offset
        year = month_index // 12
        month = month_index % 12 + 1
        periods.append(f"{year:04d}-{month:02d}")
    return periods


def pm_warning_is_zero(value: Any) -> bool:
    number = pm_warning_number(value)
    return number is not None and abs(number) <= 0.000001


def pm_warning_valid_no_budget_profit_zero(project: dict[str, Any]) -> bool:
    if not pm_warning_missing_value(project.get("budget")):
        return False
    if pm_warning_missing_value(project.get("actual")):
        return False
    decision = clean_text(project.get("warningDecision"))
    if decision == "warning-no-budget-profit-zero":
        return True
    return pm_warning_is_zero(project.get("reduceActual")) or pm_warning_is_zero(project.get("profitRate"))


def pm_warning_amount_row_ready(project: dict[str, Any]) -> bool:
    has_contract = pm_warning_has_value(project.get("contract")) or bool(project.get("contractNoDataVerified"))
    has_actual = pm_warning_has_value(project.get("actual"))
    has_budget = pm_warning_has_value(project.get("budget"))
    has_reduction = pm_warning_has_value(project.get("reduction"))
    return has_contract and has_budget and has_actual and has_reduction


def pm_warning_recompute_detail_health(data: dict[str, Any]) -> None:
    projects = [item for item in (data.get("projects") or []) if isinstance(item, dict)]
    placeholder_count = sum(1 for item in projects if item.get("placeholder"))
    detail_health = data.setdefault("detailHealth", {})
    detail_health["total"] = len(projects)
    detail_health["real"] = max(0, len(projects) - placeholder_count)
    detail_health["placeholder"] = placeholder_count
    detail_health["contract"] = sum(
        1 for item in projects
        if pm_warning_has_value(item.get("contract")) or bool(item.get("contractNoDataVerified"))
    )
    detail_health["budget"] = sum(1 for item in projects if pm_warning_has_value(item.get("budget")))
    detail_health["actual"] = sum(1 for item in projects if pm_warning_has_value(item.get("actual")))
    detail_health["reduction"] = sum(1 for item in projects if pm_warning_has_value(item.get("reduction")))
    detail_health["noBudgetProfitZero"] = sum(1 for item in projects if pm_warning_valid_no_budget_profit_zero(item))
    detail_health["complete"] = sum(1 for item in projects if pm_warning_amount_row_ready(item))
    detail_health["missing"] = max(0, detail_health["total"] - detail_health["complete"])
    totals = data.setdefault("totals", {})
    totals["detailProjects"] = detail_health["total"]
    totals["amountComplete"] = detail_health["complete"]
    totals["amountMissing"] = detail_health["missing"]


def pm_warning_detail_type(project: dict[str, Any]) -> str:
    warning_type = clean_text(project.get("type")).lower()
    status = clean_text(project.get("status") or project.get("displayStatus") or project.get("warningStatus"))
    if warning_type in ("red", "ks") or status == "红色":
        return "red"
    if warning_type in ("blue", "wwc") or status == "蓝色":
        return "blue"
    return ""


def pm_warning_detail_company_key(value: Any) -> str:
    text = clean_text(value)
    return text.replace("分公司", "公司")


def pm_warning_int(value: Any) -> int:
    try:
        return int(float(clean_text(value).replace(",", "")))
    except (TypeError, ValueError):
        return 0


def pm_warning_status_rank(status: Any) -> int:
    text = clean_text(status)
    if text == "红色":
        return 2
    if text == "蓝色":
        return 1
    return 0


def pm_warning_rank_status(rank: int) -> str:
    if rank >= 2:
        return "红色"
    if rank == 1:
        return "蓝色"
    return ""


def pm_warning_accounting_key(value: Any) -> str:
    text = re.sub(r"[（(].*?[）)]", "", clean_text(value))
    text = re.sub(r"工程|专业|施工", "", text)
    text = re.sub(r"\s+", "", text).lower()
    if "土建" in text:
        return "土建"
    if "安装" in text:
        return "安装"
    if "装饰" in text or "幕墙" in text:
        return "装饰幕墙"
    return text


def pm_warning_platform_rank(value: Any) -> int:
    platform = clean_text(value)
    if platform == "big-pm":
        return 2
    if platform == "old-pm":
        return 1
    return 0


def pm_warning_project_group_key(project: dict[str, Any]) -> str:
    company = pm_warning_detail_company_key(project.get("company") or project.get("companyName"))
    name = re.sub(r"\s+", "", clean_text(project.get("name") or project.get("project") or project.get("projectName"))).lower()
    major = pm_warning_accounting_key(project.get("hsdx") or project.get("hsdxName") or project.get("major"))
    pid = clean_text(project.get("pid") or project.get("projectId"))
    if not name and pid:
        name = pid
    return "|".join([company, name, major])


def pm_warning_project_period(project: dict[str, Any], default_period: str) -> str:
    for key in ("reportPeriod", "period", "queryPeriod", "month"):
        value = clean_text(project.get(key))
        if re.match(r"^\d{4}-\d{2}$", value):
            return value
    return default_period


def pm_warning_project_monthly_statuses(project: dict[str, Any], default_period: str) -> dict[str, str]:
    statuses: dict[str, str] = {}
    raw = project.get("monthlyStatuses")
    if isinstance(raw, dict):
        for key, value in raw.items():
            period = clean_text(key)
            status = clean_text(value)
            if re.match(r"^\d{4}-\d{2}$", period) and pm_warning_status_rank(status):
                statuses[period] = status
    period = pm_warning_project_period(project, default_period)
    status = clean_text(project.get("status") or project.get("displayStatus") or project.get("warningStatus"))
    if re.match(r"^\d{4}-\d{2}$", period) and pm_warning_status_rank(status):
        statuses[period] = status
    return statuses


def pm_warning_project_has_amount(project: dict[str, Any]) -> bool:
    return any(pm_warning_has_value(project.get(key)) for key in ("contract", "budget", "actual", "reduction"))


def pm_warning_copy_project_amount_fields(target: dict[str, Any], source: dict[str, Any]) -> None:
    for key in (
        "contract",
        "contractNoDataVerified",
        "budget",
        "actual",
        "reduction",
        "reduceDuty",
        "responsibilityTargetDisplay",
        "targetSignStatus",
        "reduceActual",
        "profitRate",
        "detailUrl",
        "warningDecision",
        "sourceRaw",
    ):
        if key in source:
            target[key] = source[key]


def pm_warning_recompute_totals_from_projects(data: dict[str, Any]) -> None:
    projects = [item for item in (data.get("projects") or []) if isinstance(item, dict)]
    totals = data.setdefault("totals", {})
    red_count = sum(1 for item in projects if pm_warning_status_rank(item.get("status")) >= 2)
    blue_count = sum(1 for item in projects if pm_warning_status_rank(item.get("status")) == 1)
    totals["red"] = red_count
    totals["blue"] = blue_count
    totals["detailProjects"] = len(projects)
    in_progress = pm_warning_int(totals.get("inProgress"))
    if in_progress:
        totals["normal"] = max(0, in_progress - red_count - blue_count)

    company_seed: dict[str, dict[str, Any]] = {}
    for company in data.get("companies") or []:
        if not isinstance(company, dict):
            continue
        company_name = clean_text(company.get("company") or company.get("name") or company.get("companyName"))
        if not company_name:
            continue
        key = pm_warning_detail_company_key(company_name)
        company_seed.setdefault(key, dict(company))
        company_seed[key]["company"] = company_name
        company_seed[key]["red"] = 0
        company_seed[key]["blue"] = 0

    for project in projects:
        company_name = clean_text(project.get("company") or project.get("companyName"))
        if not company_name:
            continue
        key = pm_warning_detail_company_key(company_name)
        if key not in company_seed:
            company_seed[key] = {"company": company_name, "total": 0, "normal": 0, "blue": 0, "red": 0}
        rank = pm_warning_status_rank(project.get("status"))
        if rank >= 2:
            company_seed[key]["red"] = pm_warning_int(company_seed[key].get("red")) + 1
        elif rank == 1:
            company_seed[key]["blue"] = pm_warning_int(company_seed[key].get("blue")) + 1

    for company in company_seed.values():
        total = pm_warning_int(company.get("total"))
        red = pm_warning_int(company.get("red"))
        blue = pm_warning_int(company.get("blue"))
        if total <= 0:
            total = red + blue
            company["total"] = total
        company["normal"] = max(0, total - red - blue)
    if company_seed:
        data["companies"] = list(company_seed.values())


def pm_warning_merge_monthly_projects(data: dict[str, Any], selected_period: str = "") -> None:
    source = data.setdefault("source", {})
    base_period = clean_text(selected_period or data.get("_selectedPeriod") or data.get("reportPeriod") or source.get("reportPeriod"))
    periods = data.get("reportPeriods")
    if not isinstance(periods, list) or len(periods) != 3:
        periods = pm_warning_period_window(base_period)
    periods = [clean_text(item) for item in periods if re.match(r"^\d{4}-\d{2}$", clean_text(item))]
    if len(periods) != 3:
        periods = pm_warning_period_window(base_period)
    data["reportPeriods"] = periods
    source["reportPeriods"] = periods
    if base_period:
        data["reportPeriod"] = base_period
        source["reportPeriod"] = base_period

    projects = [item for item in (data.get("projects") or []) if isinstance(item, dict)]
    grouped: dict[str, dict[str, Any]] = {}
    latest_amount_period: dict[str, str] = {}
    target_period = periods[-1] if periods else base_period
    for project in projects:
        period = pm_warning_project_period(project, target_period)
        project["reportPeriod"] = period
        statuses = pm_warning_project_monthly_statuses(project, target_period)
        project["monthlyStatuses"] = statuses
        key = pm_warning_project_group_key(project)
        if key not in grouped:
            grouped[key] = dict(project)
            grouped[key]["monthlyStatuses"] = dict(statuses)
            grouped[key]["_primaryPlatformRank"] = pm_warning_platform_rank(project.get("sourcePlatform"))
            latest_amount_period[key] = period if pm_warning_project_has_amount(project) else ""
            continue
        merged = grouped[key]
        merged_rank = int(merged.get("_primaryPlatformRank") or pm_warning_platform_rank(merged.get("sourcePlatform")))
        project_rank = pm_warning_platform_rank(project.get("sourcePlatform"))
        if project_rank > merged_rank:
            grouped[key] = dict(project)
            grouped[key]["monthlyStatuses"] = dict(statuses)
            grouped[key]["_primaryPlatformRank"] = project_rank
            latest_amount_period[key] = period if pm_warning_project_has_amount(project) else ""
            continue
        if project_rank < merged_rank:
            continue
        merged_statuses = merged.setdefault("monthlyStatuses", {})
        for status_period, status in statuses.items():
            if pm_warning_status_rank(status) >= pm_warning_status_rank(merged_statuses.get(status_period)):
                merged_statuses[status_period] = status
        merged["sourcePlatform"] = merged.get("sourcePlatform") or project.get("sourcePlatform")
        merged["company"] = merged.get("company") or project.get("company")
        merged["name"] = merged.get("name") or project.get("name")
        merged["major"] = merged.get("major") or project.get("major") or project.get("hsdx")
        merged["manager"] = merged.get("manager") or project.get("manager")
        merged["placeholder"] = bool(merged.get("placeholder")) and bool(project.get("placeholder"))
        current_amount_period = latest_amount_period.get(key, "")
        should_replace_amount = False
        if period == target_period and pm_warning_project_has_amount(project):
            should_replace_amount = True
        elif not pm_warning_project_has_amount(merged) and pm_warning_project_has_amount(project):
            should_replace_amount = True
        elif pm_warning_project_has_amount(project) and period > current_amount_period and current_amount_period != target_period:
            should_replace_amount = True
        if should_replace_amount:
            pm_warning_copy_project_amount_fields(merged, project)
            latest_amount_period[key] = period

    merged_projects = []
    for project in grouped.values():
        statuses = project.get("monthlyStatuses") if isinstance(project.get("monthlyStatuses"), dict) else {}
        max_rank = max((pm_warning_status_rank(statuses.get(period)) for period in periods), default=0)
        if max_rank <= 0:
            max_rank = pm_warning_status_rank(project.get("status") or project.get("displayStatus") or project.get("warningStatus"))
        status = pm_warning_rank_status(max_rank)
        project["status"] = status
        project["displayStatus"] = status
        if max_rank >= 2:
            project["type"] = "red"
            project["t"] = "ks"
        elif max_rank == 1:
            project["type"] = "blue"
            project["t"] = "wwc"
        project["monthlyStatuses"] = {period: statuses.get(period, "") for period in periods}
        project.pop("_primaryPlatformRank", None)
        merged_projects.append(project)
    data["projects"] = merged_projects
    pm_warning_recompute_totals_from_projects(data)


def pm_warning_pad_missing_warning_projects(data: dict[str, Any]) -> None:
    projects = data.setdefault("projects", [])
    if not isinstance(projects, list):
        projects = []
        data["projects"] = projects

    reliable_list_counts: dict[tuple[str, str, str], int] = {}
    for check in data.get("listChecks") or []:
        if not isinstance(check, dict):
            continue
        message = clean_text(check.get("message"))
        actual = pm_warning_int(check.get("actual"))
        total = pm_warning_int(check.get("total"))
        if message or total != actual:
            continue
        platform = clean_text(check.get("sourcePlatform") or (data.get("source") or {}).get("platform") or "old-pm")
        company_key = pm_warning_detail_company_key(check.get("company"))
        detail_type = clean_text(check.get("type"))
        if company_key and detail_type:
            reliable_list_counts[(platform, company_key, detail_type)] = actual

    existing: dict[tuple[str, str, str], int] = {}
    for project in projects:
        if not isinstance(project, dict):
            continue
        platform = clean_text(project.get("sourcePlatform") or (data.get("source") or {}).get("platform") or "old-pm")
        company_key = pm_warning_detail_company_key(project.get("company") or project.get("companyName"))
        detail_type = pm_warning_detail_type(project)
        if not company_key or not detail_type:
            continue
        key = (platform, company_key, detail_type)
        existing[key] = existing.get(key, 0) + 1

    companies = data.get("companies") or []
    sequence = len(projects) + 1
    for company in companies:
        if not isinstance(company, dict):
            continue
        platform = clean_text(company.get("sourcePlatform") or (data.get("source") or {}).get("platform") or "old-pm")
        company_name = clean_text(company.get("company") or company.get("name") or company.get("companyName"))
        company_key = pm_warning_detail_company_key(company_name)
        if not company_key:
            continue
        for detail_type, status_text, t_value, expected_key in (
            ("red", "红色", "ks", "red"),
            ("blue", "蓝色", "wwc", "blue"),
        ):
            expected = pm_warning_int(company.get(expected_key))
            key = (platform, company_key, detail_type)
            if key in reliable_list_counts:
                expected = reliable_list_counts[key]
            missing = max(0, expected - existing.get(key, 0))
            for index in range(missing):
                projects.append(
                    {
                        "sequence": sequence,
                        "sourcePlatform": platform,
                        "company": company_name,
                        "name": f"{pm_warning_platform_name(platform)} {company_name}{status_text}汇总占位 {index + 1}（明细待补抓）",
                        "major": "汇总占位",
                        "manager": "非项目明细",
                        "projectStatus": "在建",
                        "status": status_text,
                        "displayStatus": status_text,
                        "type": detail_type,
                        "t": t_value,
                        "contract": "",
                        "budget": "",
                        "actual": "",
                        "reduction": "",
                        "reduceActual": "",
                        "profitRate": "",
                        "reduceDuty": "明细待补抓",
                        "responsibilityTargetDisplay": "明细待补抓",
                        "targetSignStatus": "明细待补抓",
                        "warningDecision": "placeholder-missing-detail",
                        "placeholder": True,
                        "remark": "该行由平台汇总数补齐生成，不是平台真实项目明细；补抓项目列表后会被真实明细替换。",
                    }
                )
                sequence += 1
            existing[key] = existing.get(key, 0) + missing


def pm_warning_round(value: float) -> float:
    return round(value + 0, 2)


def pm_warning_fill_derived(project: dict[str, Any]) -> int:
    actual = pm_warning_number(project.get("actual"))
    budget = pm_warning_number(project.get("budget"))
    rate_source = project.get("reduceActual")
    if pm_warning_missing_value(rate_source):
        rate_source = project.get("profitRate")
    actual_rate = pm_warning_number(rate_source)
    has_actual_rate = not pm_warning_missing_value(rate_source) and actual_rate is not None
    eps = 0.000001

    if has_actual_rate and (
        budget is None
        or actual is None
        or (abs(budget) <= eps and abs(actual) <= eps)
    ):
        project["amountLogicIssue"] = "actual-rate-without-budget-actual"
    else:
        project.pop("amountLogicIssue", None)

    return 0


def pm_warning_clear_auto_contract_no_data(data: dict[str, Any]) -> None:
    for project in data.get("projects") or []:
        if not isinstance(project, dict):
            continue
        project.pop("contractNoData", None)
        raw = project.get("sourceRaw")
        if isinstance(raw, dict):
            raw.pop("contractNoData", None)


def pm_warning_match_with_amount(match: Any) -> bool:
    if not isinstance(match, dict):
        return False
    if pm_warning_number(match.get("amountYuan")) not in (None, 0):
        return True
    return pm_warning_int(match.get("withAmount")) > 0


def pm_warning_mark_verified_contract_no_data(data: dict[str, Any]) -> None:
    for project in data.get("projects") or []:
        if not isinstance(project, dict):
            continue
        if pm_warning_has_value(project.get("contract")):
            project.pop("contractNoDataVerified", None)
            continue
        raw = project.get("sourceRaw")
        if not isinstance(raw, dict):
            project.pop("contractNoDataVerified", None)
            continue
        contract_match = raw.get("contractMatch")
        basic_match = raw.get("contractBasicMatch")
        both_paths_checked = isinstance(contract_match, dict) and isinstance(basic_match, dict)
        has_amount = pm_warning_match_with_amount(contract_match) or pm_warning_match_with_amount(basic_match)
        if both_paths_checked and not has_amount:
            project["contractNoDataVerified"] = True
        else:
            project.pop("contractNoDataVerified", None)


def normalize_pm_warning_report_fields(data: dict[str, Any], selected_period: str = "") -> dict[str, Any]:
    pm_warning_clear_auto_contract_no_data(data)
    pm_warning_mark_verified_contract_no_data(data)
    pm_warning_pad_missing_warning_projects(data)
    for project in data.get("projects") or []:
        if isinstance(project, dict):
            pm_warning_fill_derived(project)
            if pm_warning_missing_value(project.get("reduceDuty")):
                project["responsibilityTargetDisplay"] = "待签"
                project["targetSignStatus"] = "待签"
            elif pm_warning_missing_value(project.get("responsibilityTargetDisplay")):
                project["responsibilityTargetDisplay"] = project.get("reduceDuty")
    pm_warning_merge_monthly_projects(data, selected_period)
    pm_warning_recompute_detail_health(data)
    return data


def load_pm_warning_data(platform: str = "old-pm", period: str = "") -> dict[str, Any]:
    if clean_text(platform) == "all":
        data = merge_pm_warning_platform_data(["old-pm", "big-pm"])
    else:
        data = load_pm_warning_single_platform(platform)
    return normalize_pm_warning_report_fields(data, period)


def clear_pm_warning_data(platform: str = "old-pm") -> dict[str, Any]:
    ensure_dirs()
    if clean_text(platform) == "all":
        removed: list[str] = []
        for item in ("old-pm", "big-pm"):
            result = clear_pm_warning_data(item)
            removed.extend(result.get("removed") or [])
        data = pm_warning_fallback_data()
        data.setdefault("source", {})["displaySource"] = "已清空，等待重新抓数"
        data.setdefault("source", {})["platform"] = "all"
        return {"ok": True, "removed": removed, "data": data}
    safe_platform = pm_warning_safe_platform(platform)
    targets = [
        pm_warning_latest_path(safe_platform),
        PM_WARNING_RESULT_DIR / "red_blue_company_summary_old-pm.csv",
        PM_WARNING_RESULT_DIR / "red_blue_projects_old-pm.csv",
    ]
    removed: list[str] = []
    for target in targets:
        try:
            if target.exists() and target.is_file():
                target.unlink()
                removed.append(str(target))
        except OSError:
            continue

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for path in PM_WARNING_JOB_DIR.glob("*.json"):
        job = read_json_file(path, None)
        if not isinstance(job, dict):
            continue
        if pm_warning_safe_platform(job.get("platform", "old-pm")) != safe_platform:
            continue
        if job.get("status") in PM_WARNING_ACTIVE_STATUSES:
            job["status"] = "已清空"
            job["message"] = "红蓝预警网页数据已清空。"
            job["updatedAt"] = now
            with path.open("w", encoding="utf-8") as f:
                json.dump(job, f, ensure_ascii=False, indent=2)

    data = pm_warning_fallback_data()
    data.setdefault("source", {})["displaySource"] = "已清空，等待重新抓数"
    return {"ok": True, "removed": removed, "data": data}


def write_pm_warning_job(job: dict[str, Any]) -> Path:
    ensure_dirs()
    job_path = pm_warning_job_path(job.get("id", ""))
    with job_path.open("w", encoding="utf-8") as f:
        json.dump(job, f, ensure_ascii=False, indent=2)
    return job_path


def parse_pm_warning_job_datetime(value: Any) -> datetime | None:
    text = clean_text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S.%fZ", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(text.replace("+08:00", ""), fmt)
        except ValueError:
            continue
    return None


def refresh_pm_warning_job_status(job: dict[str, Any], path: Path) -> dict[str, Any]:
    if job.get("status") not in PM_WARNING_ACTIVE_STATUSES:
        return job
    created_at = parse_pm_warning_job_datetime(job.get("createdAt"))
    updated_at = parse_pm_warning_job_datetime(job.get("updatedAt")) or created_at
    result_path_text = clean_text(job.get("resultPath"))
    result_path = Path(result_path_text) if result_path_text else None
    if result_path and result_path.exists():
        result_time = datetime.fromtimestamp(result_path.stat().st_mtime)
        if created_at is None or result_time >= created_at - timedelta(seconds=5):
            data = read_json_file(result_path, {}) or {}
            totals = data.get("totals") if isinstance(data, dict) else {}
            detail_health = data.get("detailHealth") if isinstance(data, dict) else {}
            capture_completeness = data.get("captureCompleteness") if isinstance(data.get("captureCompleteness"), dict) else {}
            in_progress = clean_text((totals or {}).get("inProgress"))
            red = clean_text((totals or {}).get("red"))
            blue = clean_text((totals or {}).get("blue"))
            job["status"] = "已完成"
            platform_name = clean_text(job.get("platformName")) or pm_warning_platform_name(clean_text(job.get("platform")))
            warning_total = pm_warning_int((totals or {}).get("red")) + pm_warning_int((totals or {}).get("blue"))
            detail_total = pm_warning_int((detail_health or {}).get("total") or (totals or {}).get("detailProjects"))
            pending_detail = pm_warning_int((capture_completeness or {}).get("missing")) or max(0, warning_total - detail_total)
            if clean_text(job.get("platform")) == "big-pm" and pending_detail:
                job["message"] = f"{platform_name}汇总抓取完成：在建 {in_progress or 0} 项，红色 {red or 0} 项，蓝色 {blue or 0} 项；项目明细仍有 {pending_detail} 项待补抓。"
            else:
                job["message"] = f"{platform_name}抓数完成：在建 {in_progress or 0} 项，红色 {red or 0} 项，蓝色 {blue or 0} 项。"
            job["totals"] = totals or {}
            job["detailHealth"] = detail_health or {}
            job["captureCompleteness"] = capture_completeness or {}
            job["completedAt"] = result_time.strftime("%Y-%m-%d %H:%M:%S")
            job["updatedAt"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            write_pm_warning_job(job)
            return job
    timeout_minutes = 45 if isinstance(job.get("periods"), list) and len(job.get("periods") or []) >= 3 else 8
    if updated_at and datetime.now() - updated_at > timedelta(minutes=timeout_minutes):
        job["status"] = "失败"
        job["message"] = "后台抓数任务长时间没有更新，已自动结束。请确认四版 PM 页面仍登录后重新点击“四版平台”。"
        job["failedAt"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        job["updatedAt"] = job["failedAt"]
        write_pm_warning_job(job)
    return job


def list_pm_warning_jobs(job_id: str = "") -> list[dict[str, Any]]:
    ensure_dirs()
    jobs: list[dict[str, Any]] = []
    for path in sorted(PM_WARNING_JOB_DIR.glob("*.json"), key=lambda item: item.stat().st_mtime, reverse=True):
        job = read_json_file(path, None)
        if not isinstance(job, dict):
            continue
        job = refresh_pm_warning_job_status(job, path)
        if job_id and job.get("id") != job_id:
            continue
        job["_path"] = str(path)
        jobs.append(job)
    return jobs


def start_pm_warning_worker(job: dict[str, Any], job_path: Path) -> tuple[bool, str]:
    platform = pm_warning_safe_platform(job.get("platform"))
    base_script_path = ROOT / "tools" / ("pm_warning_big_pm_summary_worker.mjs" if platform == "big-pm" else "pm_warning_old_pm_worker.mjs")
    script_path = ROOT / "tools" / "pm_warning_multi_period_worker.mjs"
    if not script_path.exists():
        script_path = base_script_path
    if not base_script_path.exists():
        return False, f"{pm_warning_platform_name(platform)}红蓝预警 worker 不存在。"
    node_path = RUNTIME_NODE if RUNTIME_NODE.exists() else Path("node")
    log_path = PM_WARNING_LOG_DIR / f"{job.get('id', 'pm-warning')}.log"
    try:
        import os

        env = os.environ.copy()
        period = clean_text(job.get("period")) or datetime.now().strftime("%Y-%m")
        periods = job.get("periods")
        if not isinstance(periods, list) or len(periods) != 3:
            periods = pm_warning_period_window(period)
        env["PM_WARNING_PERIOD"] = period
        env["PM_WARNING_PERIODS"] = ",".join(clean_text(item) for item in periods)
        env["PM_WARNING_BASE_WORKER_SCRIPT"] = str(base_script_path)
        if platform == "big-pm":
            env.setdefault("PM_WARNING_ALLOW_PAGE_NAVIGATION", "1")
            env.setdefault("PM_WARNING_BIG_PM_PATCH_DETAILS", "1")
    except Exception:
        env = {}
    try:
        with log_path.open("a", encoding="utf-8") as log:
            log.write(f"\n[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] start PM warning worker {job_path}\n")
            subprocess.Popen(
                [str(node_path), str(script_path), str(job_path)],
                cwd=str(ROOT),
                stdout=log,
                stderr=log,
                env=env or None,
                creationflags=subprocess.CREATE_NO_WINDOW if sys.platform.startswith("win") else 0,
            )
    except Exception as exc:
        return False, f"四版平台抓数启动失败：{exc}"
    job["logPath"] = str(log_path)
    return True, f"{pm_warning_platform_name(platform)}抓数任务已启动。"


def export_pm_warning_report(platform: str, period: str, output_path: Path) -> tuple[bool, str]:
    if not PM_WARNING_TEMPLATE_PATH.exists():
        return False, "红蓝预警报表模板不存在。"
    script_path = ROOT / "tools" / "export_pm_warning_report.ps1"
    if not script_path.exists():
        return False, "红蓝预警报表导出脚本不存在。"
    data = load_pm_warning_data(platform, period)
    detail_health = data.get("detailHealth") or {}
    detail_total = int(detail_health.get("total") or 0)
    core_complete = int(detail_health.get("complete") or 0)
    if False and detail_total and core_complete < detail_total:
        return (
            False,
            f"PM金额明细核心列未取全（{core_complete}/{detail_total}），为避免导出旧模板偏差数据，请先重新执行四版平台抓数并确认合同额、实际成本、预算/降低额或无预算0利润率规则均已取全。",
        )
    input_path = PM_WARNING_RESULT_DIR / f"export_input_{datetime.now().strftime('%Y%m%d%H%M%S')}_{uuid.uuid4().hex[:8]}.json"
    PM_WARNING_RESULT_DIR.mkdir(parents=True, exist_ok=True)
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    with input_path.open("w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    powershell = Path("powershell.exe")
    result = subprocess.run(
        [
            str(powershell),
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(script_path),
            "-JsonPath",
            str(input_path),
            "-OutputPath",
            str(output_path),
            "-TemplatePath",
            str(PM_WARNING_TEMPLATE_PATH),
            "-Period",
            period,
        ],
        cwd=str(ROOT),
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        message = (result.stderr or result.stdout or "红蓝预警报表导出失败。").strip()
        return False, message
    return True, result.stdout.strip()


def start_material_platform_worker(task: dict[str, Any], job_path: Path) -> tuple[bool, str]:
    script_path = ROOT / "tools" / "material_platform_worker.mjs"
    if not script_path.exists():
        return False, "材料价格平台采集 worker 不存在。"
    node_path = RUNTIME_NODE if RUNTIME_NODE.exists() else Path("node")
    log_dir = DATA_DIR / "material_worker_logs"
    log_dir.mkdir(parents=True, exist_ok=True)
    log_path = log_dir / f"{task.get('id', 'material')}.log"
    env = dict()
    try:
        import os

        env = os.environ.copy()
    except Exception:
        env = {}
    if RUNTIME_NODE_MODULES.exists():
        pnpm_modules = RUNTIME_NODE_MODULES / ".pnpm" / "node_modules"
        env["NODE_PATH"] = str(pnpm_modules) + (";" + str(RUNTIME_NODE_MODULES) if sys.platform.startswith("win") else ":" + str(RUNTIME_NODE_MODULES))
    try:
        with log_path.open("a", encoding="utf-8") as log:
            log.write(f"\n[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] start worker {job_path}\n")
            subprocess.Popen(
                [str(node_path), str(script_path), str(job_path)],
                cwd=str(ROOT),
                stdout=log,
                stderr=log,
                env=env or None,
                creationflags=subprocess.CREATE_NO_WINDOW if sys.platform.startswith("win") else 0,
            )
    except Exception as exc:
        return False, f"后台采集启动失败：{exc}"
    task["materialWorkerLog"] = str(log_path)
    return True, f"已启动后台平台采集：{log_path.name}"


def blacklist_codex_job_path(job_id: str) -> Path:
    safe_id = re.sub(r"[^0-9a-zA-Z_-]", "", clean_text(job_id))
    return BLACKLIST_CODEX_JOB_DIR / f"{safe_id}.json"


def blacklist_result_path(job_id: str) -> Path:
    safe_id = re.sub(r"[^0-9a-zA-Z_-]", "", clean_text(job_id))
    return BLACKLIST_RESULT_DIR / f"{safe_id}.json"


def estimate_blacklist_finish(created_at: str) -> str:
    try:
        base = datetime.strptime(clean_text(created_at), "%Y-%m-%d %H:%M:%S")
    except ValueError:
        base = datetime.now()
    return (base + timedelta(minutes=10)).strftime("%Y-%m-%d %H:%M:%S")


BLACKLIST_ROUTE_KEYS = (
    "routes",
    "upwardRoutes",
    "downwardRoutes",
    "branchRoutes",
    "personCrossRoutes",
    "sameContactRoutes",
    "naturalPersonRoutes",
    "otherShareholderRoutes",
)

BLACKLIST_TRAVERSAL_GROUPS = [
    "股东向上",
    "自然人股东层",
    "向下投资及控制",
    "分支机构",
    "主要人员交叉",
    "同电话/同地址/同账户",
    "其他自然人股东",
    "风险记录关联",
]


def write_blacklist_codex_job(job: dict[str, Any]) -> Path:
    BLACKLIST_CODEX_JOB_DIR.mkdir(parents=True, exist_ok=True)
    BLACKLIST_RESULT_DIR.mkdir(parents=True, exist_ok=True)
    job_path = blacklist_codex_job_path(job.get("id", ""))
    result_path = blacklist_result_path(job.get("id", ""))
    payload = {
        "id": job.get("id"),
        "type": "blacklist-tyc-deep-search",
        "status": "pending",
        "createdAt": job.get("createdAt"),
        "estimatedAt": job.get("estimatedAt") or estimate_blacklist_finish(job.get("createdAt", "")),
        "title": f"天眼查层层穿透：{job.get('query')}",
        "query": job.get("query"),
        "instruction": "请使用天眼查或已授权企业信息数据源，对查询对象进行企业关系层层穿透。参照泰山集团关系图的查询思路：以查询对象为中心，把股东向上、自然人股东层、向下投资及控制、分支机构、主要人员交叉、同电话/同地址/同账户、其他自然人股东、风险记录关联等分类逐项展开。必须覆盖所有可见且授权可访问的关联公司，不只摘取前几条；遇到“展开更多”、分页、下一层控股/参股企业时继续展开，直到没有新增主体、达到最大层级或受到授权/验证码/付费限制。每个已展开节点和未能继续展开的节点都要在 coverage 中记录完成状态和停止原因。并按围标串标线索筛查口径整理主体、人员、联系方式、地址、账户、股权控制、对外投资和风险记录。输出应区分已核事实、异常线索、重点核查和待补证事项，不得因单一关联直接认定围标串标。仅使用已授权访问方式，不绕过验证码、登录、付费或反爬限制。",
        "searchScope": [
            "企业基础信息",
            "法定代表人",
            "股东及出资关系",
            "向上穿透：直接股东、控股股东、实际控制人、最终上层主体",
            "对外投资与被投资企业",
            "向下穿透：控股子公司、参股公司、分支机构、关联经营主体",
            "主要人员",
            "联系方式、注册地址、经营地址",
            "失信、被执行、限制高消费、行政处罚、经营异常",
            "围标串标线索：同法人/同股东/同高管/同联系人/同电话/同地址/同账户/历史投标异常",
        ],
        "maxDepth": 4,
        "traversalPolicy": {
            "goal": "覆盖所有可见且授权可访问的关联公司，并记录完整路径。",
            "expandUntil": "无新增主体、达到 maxDepth、或遇到授权/验证码/付费/反爬限制。",
            "requiredGroups": BLACKLIST_TRAVERSAL_GROUPS,
            "nodeRule": "每个企业节点都要继续检查股东、对外投资、分支机构、主要人员和联系方式交叉；自然人节点至少检查任职、投资、控制或担任股东的企业。",
            "routeRule": "每条路线必须包含最上层和最下层，中间层不要省略；无法确认的关系写入 missingEvidence。",
        },
        "relationGraphStyle": {
            "layout": "center-radial",
            "center": "query",
            "relationGroups": BLACKLIST_TRAVERSAL_GROUPS,
            "note": "routes/upwardRoutes/downwardRoutes 等路线请保留关系类别，便于小程序渲染成类似泰山集团的中心企业向外发散关系图。",
        },
        "preferredOutput": str(result_path),
        "requiredOutputSchema": {
            "query": "原查询对象",
            "source": "天眼查或其他授权来源",
            "checkedAt": "yyyy-mm-dd HH:MM:SS",
            "entities": [
                {
                    "id": "稳定编号，可用信用代码或自定义ID",
                    "name": "企业名称",
                    "creditCode": "统一社会信用代码",
                    "riskLevel": "黑名单/高风险/关联关注/正常",
                    "status": "限制合作/需复核/可关注等",
                    "legalRep": "法定代表人",
                    "shareholders": ["股东名称"],
                    "contacts": ["电话"],
                    "bankAccounts": ["如有授权资料再填写"],
                    "address": "注册地址或经营地址",
                    "reason": "风险或关联原因",
                    "source": "信息来源说明",
                    "updatedAt": "yyyy-mm-dd",
                }
            ],
            "allRelatedCompanies": [
                {
                    "name": "所有已发现关联公司名称",
                    "relationGroup": "股东向上/自然人股东层/向下投资及控制/分支机构/主要人员交叉/同电话同地址同账户/其他自然人股东/风险记录关联",
                    "depth": "距离查询对象的层级",
                    "path": ["最上层主体", "中间主体", "最下层主体"],
                    "expanded": "true/false",
                    "stopReason": "无新增主体/达到最大层级/授权限制/验证码限制/付费限制/页面未披露",
                }
            ],
            "coverage": {
                "status": "complete/partial",
                "completedDepth": 0,
                "completedGroups": BLACKLIST_TRAVERSAL_GROUPS,
                "expandedNodeCount": 0,
                "relatedCompanyCount": 0,
                "unexpandedNodes": [
                    {"name": "未继续穿透的主体", "reason": "停止原因", "group": "所属关系分类"}
                ],
                "stopReasons": ["整体停止原因"],
                "summary": "说明本次是否已覆盖所有可见关联公司，以及哪些地方需要人工补查。",
            },
            "clueMatrix": [
                {
                    "id": "CL-001",
                    "objects": "涉及对象",
                    "indicator": "异常指标，如同法人、同电话、同地址、同账户、股权控制、人员交叉",
                    "verifiedFact": "已核事实和天眼查定位",
                    "riskExplanation": "风险解释，使用异常线索/重点核查等表述",
                    "possibleExplanation": "合理解释或反证",
                    "missingEvidence": "待补资料",
                    "priority": "高/中/低",
                }
            ],
            "evidenceCatalog": [
                {
                    "id": "E-001",
                    "name": "证据或页面名称",
                    "source": "天眼查/授权来源",
                    "location": "页面模块、截图编号或导出文件位置",
                    "relatedClue": "CL-001",
                    "sensitiveHandling": "脱敏或受控底稿说明",
                }
            ],
            "routes": [
                {
                    "path": ["最上层控股主体/实际控制人", "中间层主体", "查询对象或最下层关联企业"],
                    "relations": ["股东/控制/任职/同电话/同地址", "对外投资/控制/关联"],
                    "source": "兼容字段：查询对象或上层企业/人员",
                    "target": "兼容字段：已穿透到的下层企业/人员",
                    "label": "兼容字段：股东/法定代表人/对外投资/同电话/同地址/同账户",
                    "sourceType": "root/black/related/person",
                    "targetType": "black/related/person/reason",
                }
            ],
            "upwardRoutes": [
                {
                    "path": ["最终上层主体/实际控制人", "控股股东", "查询对象"],
                    "relations": ["控制/持股", "控制/持股"],
                    "note": "用于展示向上穿透的上级控股公司或实际控制链",
                }
            ],
            "downwardRoutes": [
                {
                    "path": ["查询对象", "控股/参股企业", "下层关联企业"],
                    "relations": ["对外投资/控制", "对外投资/控制"],
                    "note": "用于展示向下穿透的投资和被投资链",
                }
            ],
            "branchRoutes": [{"path": ["查询对象", "分支机构"], "relations": ["设立/分支"], "note": "分支机构路线"}],
            "personCrossRoutes": [{"path": ["查询对象", "主要人员", "关联企业"], "relations": ["任职/高管", "任职/投资"], "note": "主要人员交叉路线"}],
            "sameContactRoutes": [{"path": ["查询对象", "同电话/同地址/同账户", "关联企业"], "relations": ["同电话/同地址/同账户", "关联"], "note": "联系方式或账户交叉路线"}],
            "naturalPersonRoutes": [{"path": ["查询对象", "自然人股东", "自然人控制/投资企业"], "relations": ["股东", "投资/任职/控制"], "note": "自然人股东层路线"}],
            "otherShareholderRoutes": [{"path": ["查询对象", "其他股东", "其他股东关联企业"], "relations": ["股东", "投资/任职/控制"], "note": "其他自然人股东或法人股东路线"}],
        },
        "afterWriteBack": "结果写入 preferredOutput 后，在小程序中刷新或重新查询，即可生成关系图和 Excel。",
    }
    with job_path.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    return job_path


def start_blacklist_tianyancha_worker(job_id: str) -> bool:
    if not BLACKLIST_TYC_WORKER_SCRIPT.exists():
        return False
    node = NODE_RUNTIME if NODE_RUNTIME.exists() else Path("node")
    try:
        env = os.environ.copy()
        env.setdefault("TYC_SECOND_LAYER_LIMIT", "24")
        env.setdefault("TYC_SECOND_LAYER_CONCURRENCY", "4")
        env.setdefault("TYC_DETAIL_SCROLLS", "5")
        subprocess.Popen(
            [str(node), str(BLACKLIST_TYC_WORKER_SCRIPT), "--job", job_id],
            cwd=str(ROOT),
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            env=env,
            creationflags=subprocess.CREATE_NO_WINDOW if sys.platform.startswith("win") else 0,
        )
        return True
    except Exception:
        return False


def refresh_blacklist_codex_jobs(db: dict[str, Any]) -> bool:
    changed = False
    for job in db.get("blacklistCodexJobs", []):
        result_path = Path(job.get("resultPath", ""))
        if result_path.exists() and load_blacklist_entities_from_path(result_path):
            if job.get("status") != "已回写":
                job["status"] = "已回写"
                job["message"] = "天眼查穿透结果已回写，可重新查询并导出。"
                changed = True
        elif result_path.exists() and job.get("status") != "待填写结果":
            job["status"] = "待填写结果"
            job["message"] = "回写模板已生成，请填入天眼查穿透结果。"
            changed = True
    return changed


def list_blacklist_codex_jobs() -> list[dict[str, Any]]:
    db = load_db()
    if refresh_blacklist_codex_jobs(db):
        save_db(db)
    jobs = []
    for job in db.get("blacklistCodexJobs", []):
        item = dict(job)
        item["status"] = clean_text(item.get("status")).replace("Codex", "天眼查")
        item["estimatedAt"] = item.get("estimatedAt") or estimate_blacklist_finish(item.get("createdAt", ""))
        result_path = Path(item.get("resultPath", ""))
        if result_path.exists():
            item["resultStatus"] = "已有结果" if load_blacklist_entities_from_path(result_path) or load_blacklist_routes_from_path(result_path) else "模板待填写"
            summary = load_blacklist_coverage_summary(result_path)
            item.update(summary)
        else:
            item["resultStatus"] = "未回写"
            item["coverageStatus"] = ""
            item["completedDepth"] = ""
            item["unexpandedCount"] = ""
        jobs.append(item)
    return jobs


def load_blacklist_coverage_summary(path: Path) -> dict[str, Any]:
    try:
        with path.open("r", encoding="utf-8-sig") as f:
            data = json.load(f)
    except (OSError, json.JSONDecodeError):
        return {"coverageStatus": "", "completedDepth": "", "unexpandedCount": ""}
    coverage = data.get("coverage", {}) if isinstance(data, dict) else {}
    unexpanded = coverage.get("unexpandedNodes", []) if isinstance(coverage, dict) else []
    return {
        "coverageStatus": clean_text(coverage.get("status") if isinstance(coverage, dict) else ""),
        "completedDepth": coverage.get("completedDepth", "") if isinstance(coverage, dict) else "",
        "unexpandedCount": len(unexpanded) if isinstance(unexpanded, list) else "",
    }


def load_blacklist_result_entities() -> list[dict[str, Any]]:
    entities: list[dict[str, Any]] = []
    for path in sorted(BLACKLIST_RESULT_DIR.glob("*.json")):
        entities.extend(load_blacklist_entities_from_path(path))
    return entities


def load_blacklist_result_routes(query: str) -> list[dict[str, Any]]:
    routes: list[dict[str, Any]] = []
    normalized_query = normalize_lookup_text(query)
    for path in sorted(BLACKLIST_RESULT_DIR.glob("*.json")):
        try:
            with path.open("r", encoding="utf-8-sig") as f:
                data = json.load(f)
        except (OSError, json.JSONDecodeError):
            continue
        if normalized_query and normalized_query not in normalize_lookup_text(data.get("query", "")):
            entity_names = [normalize_lookup_text(row.get("name")) for row in data.get("entities", []) if isinstance(row, dict)]
            if not any(normalized_query in name for name in entity_names):
                continue
        for route_key in BLACKLIST_ROUTE_KEYS:
            for route in data.get(route_key, []):
                if isinstance(route, dict):
                    route.setdefault("direction", route_key)
                    routes.append(route)
        for relation in data.get("relations", []):
            if isinstance(relation, dict):
                routes.append({
                    "source": relation.get("source") or relation.get("from"),
                    "target": relation.get("target") or relation.get("to"),
                    "label": relation.get("label") or relation.get("relation") or "关联",
                    "sourceType": relation.get("sourceType", ""),
                    "targetType": relation.get("targetType", ""),
                })
    return routes


def load_blacklist_routes_from_path(path: Path) -> list[dict[str, Any]]:
    try:
        with path.open("r", encoding="utf-8") as f:
            data = json.load(f)
    except (OSError, json.JSONDecodeError):
        return []
    routes = []
    for route_key in BLACKLIST_ROUTE_KEYS:
        for route in data.get(route_key, []):
            if isinstance(route, dict):
                route.setdefault("direction", route_key)
                routes.append(route)
    relations = [
        {
            "source": relation.get("source") or relation.get("from"),
            "target": relation.get("target") or relation.get("to"),
            "label": relation.get("label") or relation.get("relation") or "关联",
            "sourceType": relation.get("sourceType", ""),
            "targetType": relation.get("targetType", ""),
        }
        for relation in data.get("relations", [])
        if isinstance(relation, dict)
    ]
    return [route for route in routes + relations if route_edges_from_route(route)]


def load_blacklist_entities_from_path(path: Path) -> list[dict[str, Any]]:
    try:
        with path.open("r", encoding="utf-8-sig") as f:
            data = json.load(f)
    except (OSError, json.JSONDecodeError):
        return []
    rows = data if isinstance(data, list) else data.get("entities", [])
    if not isinstance(rows, list):
        return []
    entities: list[dict[str, Any]] = []
    for row in rows:
        if not isinstance(row, dict) or not clean_text(row.get("name")):
            continue
        row.setdefault("id", row.get("creditCode") or f"codex-{uuid.uuid4().hex[:10]}")
        row.setdefault("source", data.get("source", "Codex回写"))
        row.setdefault("updatedAt", data.get("checkedAt", "")[:10])
        entities.append(row)
    return entities


def route_edges_from_route(route: dict[str, Any]) -> list[dict[str, str]]:
    edges: list[dict[str, str]] = []
    path_values = route.get("path")
    if isinstance(path_values, list) and len(path_values) >= 2:
        labels = [clean_text(value) for value in path_values if clean_text(value)]
        relation_labels = route.get("relations") if isinstance(route.get("relations"), list) else []
        for index in range(len(labels) - 1):
            label = clean_text(relation_labels[index] if index < len(relation_labels) else route.get("label")) or "穿透"
            edges.append({"source": labels[index], "target": labels[index + 1], "label": label, "basis": "天眼查穿透路线"})
        return edges
    source = clean_text(route.get("source") or route.get("from"))
    target = clean_text(route.get("target") or route.get("to"))
    if source and target:
        edges.append({
            "source": source,
            "target": target,
            "label": clean_text(route.get("label") or route.get("relation")) or "关联",
            "basis": "天眼查穿透关系",
        })
    return edges


def blacklist_relation_edges() -> list[dict[str, str]]:
    edges: list[dict[str, str]] = []
    for route in load_blacklist_result_routes(""):
        edges.extend(route_edges_from_route(route))
    entities = load_db().get("blacklistEntities", []) + load_blacklist_result_entities()
    for left_index, left in enumerate(entities):
        for right in entities[left_index + 1:]:
            reasons = relation_reasons(left, right)
            if not reasons:
                continue
            edges.append({
                "source": clean_text(left.get("name")),
                "target": clean_text(right.get("name")),
                "label": "；".join(reasons[:3]),
                "basis": "本地企业字段交叉比对",
            })
    return [edge for edge in edges if edge.get("source") and edge.get("target")]


def company_name_matches(name: str, query: str) -> bool:
    normalized_name = normalize_lookup_text(name)
    normalized_query = normalize_lookup_text(query)
    return bool(normalized_query and (normalized_query in normalized_name or normalized_name in normalized_query))


def find_relation_node(nodes: set[str], query: str) -> str:
    normalized_query = normalize_lookup_text(query)
    if not normalized_query:
        return ""
    exact = [node for node in nodes if normalize_lookup_text(node) == normalized_query]
    if exact:
        return exact[0]
    fuzzy = [node for node in nodes if company_name_matches(node, query)]
    return sorted(fuzzy, key=len)[0] if fuzzy else ""


def judge_company_relation(left: str, right: str) -> dict[str, Any]:
    left = clean_text(left)
    right = clean_text(right)
    if not left or not right:
        return {"related": False, "message": "请输入两个企业名称。", "path": [], "steps": []}
    edges = blacklist_relation_edges()
    nodes = {edge["source"] for edge in edges} | {edge["target"] for edge in edges}
    left_node = find_relation_node(nodes, left)
    right_node = find_relation_node(nodes, right)
    if not left_node or not right_node:
        missing = []
        if not left_node:
            missing.append(left)
        if not right_node:
            missing.append(right)
        return {
            "related": False,
            "status": "missing-data",
            "message": f"本地穿透库缺少：{'、'.join(missing)}。请先对缺失企业发起天眼查穿透，回写后再判断。",
            "leftMatched": left_node,
            "rightMatched": right_node,
            "missing": missing,
            "path": [],
            "steps": [],
        }
    adjacency: dict[str, list[dict[str, str]]] = {}
    for edge in edges:
        adjacency.setdefault(edge["source"], []).append(edge)
        adjacency.setdefault(edge["target"], []).append({
            "source": edge["target"],
            "target": edge["source"],
            "label": edge["label"],
            "basis": edge["basis"],
        })
    queue: list[tuple[str, list[dict[str, str]]]] = [(left_node, [])]
    seen = {left_node}
    found_steps: list[dict[str, str]] = []
    while queue:
        current, path = queue.pop(0)
        if current == right_node:
            found_steps = path
            break
        if len(path) >= 6:
            continue
        for edge in adjacency.get(current, []):
            target = edge["target"]
            if target in seen:
                continue
            seen.add(target)
            queue.append((target, path + [edge]))
    if not found_steps and left_node != right_node:
        return {
            "related": False,
            "status": "no-path",
            "message": "已找到两个企业，但当前数据中未发现可连通的关系路径。",
            "leftMatched": left_node,
            "rightMatched": right_node,
            "path": [],
            "steps": [],
        }
    path_nodes = [left_node]
    for step in found_steps:
        path_nodes.append(step["target"])
    return {
        "related": True,
        "status": "related",
        "message": "发现关联路径。",
        "leftMatched": left_node,
        "rightMatched": right_node,
        "distance": len(found_steps),
        "path": path_nodes,
        "steps": found_steps,
    }


def write_blacklist_result_template(job: dict[str, Any]) -> Path:
    output_path = Path(job.get("resultPath") or blacklist_result_path(job.get("id", "")))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        return output_path
    payload = {
        "query": job.get("query", ""),
        "source": "天眼查/授权企业信息数据源",
        "checkedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "entities": [
            {
                "id": "",
                "name": "",
                "creditCode": "",
                "riskLevel": "关联关注",
                "status": "需复核",
                "legalRep": "",
                "shareholders": [],
                "contacts": [],
                "bankAccounts": [],
                "address": "",
                "reason": "",
                "source": "天眼查",
                "updatedAt": datetime.now().strftime("%Y-%m-%d"),
            }
        ],
        "allRelatedCompanies": [
            {
                "name": "",
                "relationGroup": "股东向上/自然人股东层/向下投资及控制/分支机构/主要人员交叉/同电话同地址同账户/其他自然人股东/风险记录关联",
                "depth": "",
                "path": [job.get("query", "")],
                "expanded": False,
                "stopReason": "",
            }
        ],
        "coverage": {
            "status": "partial",
            "completedDepth": 0,
            "completedGroups": [],
            "expandedNodeCount": 0,
            "relatedCompanyCount": 0,
            "unexpandedNodes": [
                {"name": job.get("query", ""), "reason": "模板待填写，尚未完成天眼查穿透", "group": "根节点"}
            ],
            "stopReasons": [],
            "summary": "按泰山集团关系图思路逐项展开：股东向上、自然人股东层、向下投资及控制、分支机构、主要人员交叉、同电话/同地址/同账户、其他自然人股东、风险记录关联。完成后把 status 改为 complete；若受授权或页面限制，保留 partial 并写明原因。",
        },
        "routes": [
            {
                "path": ["最终上层控股主体/实际控制人", "中间层控股股东", job.get("query", "")],
                "relations": ["控制/持股", "控制/持股"],
                "source": job.get("query", ""),
                "target": "",
                "label": "向上穿透上级控股公司",
                "sourceType": "root",
                "targetType": "related",
            }
        ],
        "upwardRoutes": [
            {
                "path": ["最终上层控股主体/实际控制人", "中间层控股股东", job.get("query", "")],
                "relations": ["控制/持股", "控制/持股"],
                "note": "向上穿透：填写查询对象的股东、控股股东、实际控制人或最终上层主体。",
            }
        ],
        "downwardRoutes": [
            {
                "path": [job.get("query", ""), "控股/参股企业", "下层关联企业"],
                "relations": ["对外投资/控制", "对外投资/控制"],
                "note": "向下穿透：填写查询对象投资、控股或参股的下层企业。",
            }
        ],
        "branchRoutes": [
            {
                "path": [job.get("query", ""), "分支机构"],
                "relations": ["设立/分支"],
                "note": "分支机构路线：列出所有可见分支机构，并继续检查其负责人、地址、联系方式和风险记录。",
            }
        ],
        "personCrossRoutes": [
            {
                "path": [job.get("query", ""), "主要人员", "关联企业"],
                "relations": ["任职/高管", "任职/投资"],
                "note": "主要人员交叉路线：法定代表人、董监高、关键人员投资或任职的企业。",
            }
        ],
        "sameContactRoutes": [
            {
                "path": [job.get("query", ""), "同电话/同地址/同账户", "关联企业"],
                "relations": ["同电话/同地址/同账户", "关联"],
                "note": "联系方式交叉路线：同电话、同注册地址、同经营地址、同账户等授权可见线索。",
            }
        ],
        "naturalPersonRoutes": [
            {
                "path": [job.get("query", ""), "自然人股东", "自然人控制/投资企业"],
                "relations": ["股东", "投资/任职/控制"],
                "note": "自然人股东层路线：自然人股东名下投资、任职或控制企业。",
            }
        ],
        "otherShareholderRoutes": [
            {
                "path": [job.get("query", ""), "其他股东", "其他股东关联企业"],
                "relations": ["股东", "投资/任职/控制"],
                "note": "其他股东关联路线：除控股股东外其他法人或自然人股东的关联企业。",
            }
        ],
        "clueMatrix": [
            {
                "id": "CL-001",
                "objects": "",
                "indicator": "",
                "verifiedFact": "",
                "riskExplanation": "",
                "possibleExplanation": "",
                "missingEvidence": "",
                "priority": "中",
            }
        ],
        "evidenceCatalog": [
            {
                "id": "E-001",
                "name": "",
                "source": "天眼查",
                "location": "",
                "relatedClue": "CL-001",
                "sensitiveHandling": "敏感信息脱敏展示，完整信息保留受控底稿。",
            }
        ],
    }
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return output_path


def load_material_price_rows(task: dict[str, Any]) -> list[dict[str, Any]]:
    source_xlsx = DATA_DIR / "material_price_source.xlsx"
    source_json = DATA_DIR / "material_price_source.json"
    rows: list[dict[str, Any]] = []
    for source_path in bid_source_paths_for_task(task):
        if source_path.suffix.lower() != ".json":
            continue
        try:
            with source_path.open("r", encoding="utf-8-sig") as f:
                data = json.load(f)
            candidate_rows = data if isinstance(data, list) else data.get("rows", [])
            if isinstance(candidate_rows, list) and candidate_rows:
                rows = candidate_rows
                break
        except Exception:
            continue
    if not rows and source_xlsx.exists():
        wb = load_workbook(source_xlsx, data_only=True)
        for ws in wb.worksheets:
            rows.extend(read_material_sheet(ws))
    elif not rows and source_json.exists():
        with source_json.open("r", encoding="utf-8-sig") as f:
            data = json.load(f)
        rows = data if isinstance(data, list) else data.get("rows", [])
    start_date = clean_text(task.get("startDate"))
    end_date = clean_text(task.get("endDate"))
    keyword = clean_text(task.get("keyword")).lower()
    selected_company = clean_text(task.get("company"))
    if selected_company == "全集团":
        selected_company = "集团"
    filtered = []
    for row in rows:
        row = normalize_material_row(row)
        if not row.get("material"):
            continue
        row_company = standard_company_name(row.get("company", "")) or row.get("company", "")
        if "装饰幕墙" in row_company:
            continue
        if selected_company and selected_company != "集团" and row_company != selected_company:
            continue
        if keyword and keyword not in (row.get("material", "") + row.get("spec", "")).lower():
            continue
        if not date_in_range(row.get("purchaseDate"), start_date, end_date):
            continue
        filtered.append(row)
    return filtered


def read_material_sheet(ws: Any) -> list[dict[str, Any]]:
    header_row = None
    header_map: dict[str, int] = {}
    for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True), start=1):
        values = [clean_text(cell) for cell in row]
        matched: dict[str, int] = {}
        for field, aliases in MATERIAL_PRICE_HEADERS.items():
            for col_index, value in enumerate(values):
                if any(alias in value for alias in aliases):
                    matched[field] = col_index
                    break
        if "material" in matched and ("price" in matched or "amount" in matched):
            header_row = row_index
            header_map = matched
            break
    if not header_row:
        return []
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        item = {field: row[index] if index < len(row) else "" for field, index in header_map.items()}
        rows.append(item)
    return rows


def read_material_price_file(file_path: Path) -> list[dict[str, Any]]:
    suffix = file_path.suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        parse_path = convert_xls_to_xlsx(file_path) if suffix == ".xls" else file_path
        wb = load_workbook(parse_path, data_only=True)
        rows: list[dict[str, Any]] = []
        for ws in wb.worksheets:
            rows.extend(read_material_sheet(ws))
        return rows
    if suffix == ".json":
        with file_path.open("r", encoding="utf-8-sig") as f:
            data = json.load(f)
        if isinstance(data, list):
            return data
        if isinstance(data, dict):
            return data.get("rows", [])
    raise ValueError("仅支持 .xlsx、.xls 和 .json 数据文件。")


def sheet_rows_with_headers(file_path: Path) -> list[dict[str, Any]]:
    suffix = file_path.suffix.lower()
    if suffix == ".json":
        raw = file_path.read_bytes()
        data = None
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                data = json.loads(raw.decode(encoding))
                break
            except (UnicodeDecodeError, json.JSONDecodeError):
                continue
        if data is None:
            return []
        rows = data if isinstance(data, list) else data.get("rows", [])
        return rows if isinstance(rows, list) else []
    parse_path = convert_xls_to_xlsx(file_path) if suffix == ".xls" else file_path
    wb = load_workbook(parse_path, data_only=True)
    all_rows: list[dict[str, Any]] = []
    company_header_aliases = [
        clean_text(alias)
        for _company, aliases in BID_COMPANY_COLUMNS
        for alias in aliases
    ]
    for ws in wb.worksheets:
        header_row = None
        headers: list[str] = []
        data_start_row = None
        for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 25), values_only=True), start=1):
            values = [clean_text(cell) for cell in row]
            useful = [value for value in values if value]
            if len(useful) >= 3 and any("清单编号" in value or "清单编码" in value or "任务编号" in value or "采购过程编号" in value or "施工项目部" in value for value in values):
                header_row = row_index
                headers = values
                data_start_row = header_row + 1
                if row_index < ws.max_row:
                    next_values = [
                        clean_text(cell)
                        for cell in next(ws.iter_rows(min_row=row_index + 1, max_row=row_index + 1, values_only=True))
                    ]
                    has_company_header_row = any(
                        value and any(alias and alias in value for alias in company_header_aliases)
                        for value in next_values
                    )
                    if has_company_header_row:
                        headers = [
                            next_values[index] if index < len(next_values) and next_values[index] else header
                            for index, header in enumerate(headers)
                        ]
                        data_start_row = header_row + 2
                break
        if not header_row:
            continue
        for row in ws.iter_rows(min_row=data_start_row or header_row + 1, values_only=True):
            item = {headers[index]: row[index] if index < len(row) else "" for index in range(len(headers)) if headers[index]}
            if any(clean_text(value) for value in item.values()):
                item["_sheet"] = ws.title
                all_rows.append(item)
    return all_rows


def value_by_alias(row: dict[str, Any], aliases: list[str]) -> Any:
    normalized = {clean_text(key).lower(): value for key, value in row.items()}
    for alias in aliases:
        alias_text = clean_text(alias).lower()
        if alias_text in normalized:
            return normalized[alias_text]
    for key, value in row.items():
        key_text = clean_text(key).lower()
        if any(clean_text(alias).lower() in key_text for alias in aliases):
            return value
    return ""


def map_alias_row(row: dict[str, Any], aliases: dict[str, list[str]]) -> dict[str, Any]:
    return {field: value_by_alias(row, names) for field, names in aliases.items()}


def classify_bid_source_rows(rows: list[dict[str, Any]], filename: str) -> str:
    text = filename + "|" + "|".join(clean_text(key) for row in rows[:5] for key in row.keys())
    if "项目基本信息" in text or "施工项目部" in text:
        return "pm"
    has_company_prices = any(any(clean_text(alias) in text for alias in aliases) for _name, aliases in BID_COMPANY_COLUMNS)
    if "中标价格对比分析" in text or ("集团平均中标价格" in text and has_company_prices):
        return "comparison"
    if "中标价格查询" in text or "含税中标价格" in text or "付款方式" in text:
        return "query"
    if "招标结果台帐" in text or "中标单位" in text:
        return "ledger"
    return ""


def valid_price(value: Any) -> float:
    number = to_number(value)
    return number if number > 0 else 0.0


def clean_bid_company_name(value: Any, project: Any = "") -> str:
    company = clean_text(value)
    if "->" in company:
        company = company.split("->", 1)[0].strip()
    if "-" in company:
        company = company.split("-", 1)[0].strip()
    if company == "烟建集团有限公司":
        company = clean_text(project)
    return company


def bid_company_bucket(company: Any) -> str:
    text = clean_text(company)
    if "东泰" in text:
        return "东泰物流"
    if "市政路桥" in text or "格瑞特" in text:
        return "市政路桥/格瑞特"
    standard = standard_company_name(text)
    if standard == "设备安装公司":
        return "设备安装"
    if standard in {"市政路桥公司", "格瑞特公司"}:
        return "市政路桥/格瑞特"
    if standard == "东泰物流":
        return "东泰物流"
    if standard:
        return standard
    for bucket, aliases in BID_COMPANY_COLUMNS:
        if any(clean_text(alias) and clean_text(alias) in text for alias in aliases):
            return bucket
    return text


def infer_region_from_project_name(project: Any) -> str:
    text = clean_text(project)
    if not text:
        return ""
    direct_city_markers = [
        "烟台", "青岛", "济南", "威海", "潍坊", "淄博", "临沂", "日照", "东营", "泰安", "济宁", "德州", "聊城", "滨州", "菏泽", "枣庄",
        "北京", "上海", "天津", "重庆", "南京", "苏州", "杭州", "宁波", "广州", "深圳", "成都", "西安", "郑州", "合肥", "武汉",
    ]
    search_texts = [text]
    if "公司" in text:
        tail = text.split("公司", 1)[1]
        if tail:
            search_texts.insert(0, tail)
    for search_text in search_texts:
        for marker in direct_city_markers:
            if marker in search_text:
                return marker if marker.endswith("市") else f"{marker}市"
    province_markers = ["山东", "江苏", "浙江", "广东", "河南", "河北", "安徽", "湖北", "四川", "陕西", "辽宁", "吉林", "黑龙江", "福建", "江西", "湖南", "山西", "云南", "贵州", "广西", "海南", "甘肃", "新疆", "内蒙古", "宁夏", "青海", "西藏"]
    for marker in province_markers:
        if marker in text:
            return marker if marker.endswith(("省", "区")) else f"{marker}省"
    return ""


def normalize_region(value: Any, project: Any = "") -> str:
    text = clean_text(value)
    if not text:
        return infer_region_from_project_name(project)
    if text == "全国":
        return infer_region_from_project_name(project) or text
    city_markers = [
        "烟台", "青岛", "济南", "威海", "潍坊", "淄博", "临沂", "日照", "东营", "泰安", "济宁", "德州", "聊城", "滨州", "菏泽", "枣庄",
        "北京", "上海", "天津", "重庆", "南京", "苏州", "杭州", "宁波", "广州", "深圳", "成都", "西安", "郑州", "合肥", "武汉",
    ]
    for marker in city_markers:
        if text == marker:
            return marker if marker.endswith("市") else f"{marker}市"
    province_markers = ["山东", "江苏", "浙江", "广东", "河南", "河北", "安徽", "湖北", "四川", "陕西", "辽宁", "吉林", "黑龙江", "福建", "江西", "湖南", "山西", "云南", "贵州", "广西", "海南", "甘肃", "新疆", "内蒙古", "宁夏", "青海", "西藏"]
    for marker in province_markers:
        if text == marker:
            return marker if marker.endswith(("省", "区")) else f"{marker}省"
    if text.endswith(("市", "省", "自治区", "特别行政区")):
        return text
    return text


def normalize_payment(value: Any) -> str:
    text = clean_text(value).upper()
    mapping = {"CASH": "现金", "FQ": "分期", "TFK": "分期", "DFW": "抵房/物"}
    return mapping.get(text, clean_text(value))


def infer_price_library_region(row: dict[str, Any]) -> str:
    text = " ".join(
        clean_text(row.get(field))
        for field in ("project", "supplier", "contractNo", "source", "remark", "item", "spec")
    )
    city_map = [
        ("\u5305\u5934", "\u5185\u8499\u53e4\u5305\u5934"),
        ("\u70df\u53f0", "\u5c71\u4e1c\u7701\u70df\u53f0\u5e02"),
        ("\u9752\u5c9b", "\u5c71\u4e1c\u7701\u9752\u5c9b\u5e02"),
        ("\u6d4e\u5357", "\u5c71\u4e1c\u7701\u6d4e\u5357\u5e02"),
        ("\u5a01\u6d77", "\u5c71\u4e1c\u7701\u5a01\u6d77\u5e02"),
        ("\u6f4d\u574a", "\u5c71\u4e1c\u7701\u6f4d\u574a\u5e02"),
        ("\u6dc4\u535a", "\u5c71\u4e1c\u7701\u6dc4\u535a\u5e02"),
        ("\u4e34\u6c82", "\u5c71\u4e1c\u7701\u4e34\u6c82\u5e02"),
        ("\u65e5\u7167", "\u5c71\u4e1c\u7701\u65e5\u7167\u5e02"),
        ("\u4e1c\u8425", "\u5c71\u4e1c\u7701\u4e1c\u8425\u5e02"),
        ("\u6cf0\u5b89", "\u5c71\u4e1c\u7701\u6cf0\u5b89\u5e02"),
        ("\u6d4e\u5b81", "\u5c71\u4e1c\u7701\u6d4e\u5b81\u5e02"),
        ("\u5fb7\u5dde", "\u5c71\u4e1c\u7701\u5fb7\u5dde\u5e02"),
        ("\u804a\u57ce", "\u5c71\u4e1c\u7701\u804a\u57ce\u5e02"),
        ("\u6ee8\u5dde", "\u5c71\u4e1c\u7701\u6ee8\u5dde\u5e02"),
        ("\u83cf\u6cfd", "\u5c71\u4e1c\u7701\u83cf\u6cfd\u5e02"),
        ("\u67a3\u5e84", "\u5c71\u4e1c\u7701\u67a3\u5e84\u5e02"),
    ]
    for marker, region in city_map:
        if marker in text:
            return region
    return ""


def infer_price_library_project_type(row: dict[str, Any]) -> str:
    text = " ".join(
        clean_text(row.get(field))
        for field in ("project", "item", "spec", "remark", "source")
    )
    if any(word in text for word in (
        "\u5382\u623f", "\u8f66\u95f4", "\u4ed3\u5e93", "\u5e93\u623f",
        "\u751f\u4ea7\u7ebf", "\u4ea7\u4e1a\u56ed", "\u5de5\u4e1a", "\u5149\u4f0f", "\u5236\u9020",
    )):
        return "\u5382\u623f\u3001\u4ed3\u5e93"
    if any(word in text for word in ("\u5e02\u653f", "\u9053\u8def", "\u7ba1\u7f51", "\u6c61\u6c34", "\u6865\u6881", "\u7ba1\u5eca")):
        return "\u5e02\u653f"
    if any(word in text for word in (
        "\u5b66\u6821", "\u533b\u9662", "\u529e\u516c", "\u5546\u4e1a", "\u6559\u5b66",
        "\u79d1\u7814", "\u9152\u5e97", "\u516c\u5bd3", "\u6d88\u9632\u7ad9", "\u6d3e\u51fa\u6240",
    )):
        return "\u516c\u7528\u5efa\u7b51"
    if any(word in text for word in ("\u4f4f\u5b85", "\u5c0f\u533a", "\u5b89\u7f6e", "\u57ce\u4e2d\u6751")):
        return "\u4f4f\u5b85"
    return ""


def build_bid_source_sets(paths: list[Path]) -> dict[str, Any]:
    sources: dict[str, Any] = {"comparison": [], "query": [], "ledger": [], "pm": [], "unmatched": []}
    for path in paths:
        rows = sheet_rows_with_headers(path)
        kind = classify_bid_source_rows(rows, path.name)
        if not kind:
            kind = classify_bid_source_rows([], path.name)
        if kind:
            sources[kind].extend(rows or [{"_sourceFile": path.name, "_emptySource": True}])
        else:
            sources["unmatched"].append(path.name)
    return sources


def missing_bid_source_labels(sources: dict[str, Any]) -> list[str]:
    required = [
        ("comparison", "中标价格对比分析"),
        ("query", "中标价格查询"),
        ("ledger", "招标结果台帐"),
    ]
    return [label for key, label in required if not sources.get(key)]


def bid_source_paths_for_task(task: dict[str, Any]) -> list[Path]:
    paths: list[Path] = []
    for value in task.get("sourcePaths") or []:
        path = Path(clean_text(value))
        if path.exists():
            paths.append(path)
    source_dir_text = clean_text(task.get("sourceDir"))
    if source_dir_text:
        source_dir = Path(source_dir_text)
        if source_dir.exists():
            for path in sorted(source_dir.glob("*")):
                if path.suffix.lower() in {".xlsx", ".xls", ".json"} and path not in paths:
                    paths.append(path)
    return paths


def has_material_bid_info_source(task: dict[str, Any]) -> bool:
    for path in bid_source_paths_for_task(task):
        if path.name == "material_bid_info_source.json":
            return True
        if path.suffix.lower() == ".json":
            try:
                with path.open("r", encoding="utf-8-sig") as f:
                    data = json.load(f)
                if clean_text(data.get("source") if isinstance(data, dict) else "") in {"平台-材料中标信息", "平台材料中标信息"}:
                    return True
            except Exception:
                continue
    return False


def build_bid_workbook_from_task_sources(task: dict[str, Any]) -> dict[str, Any]:
    source_paths = bid_source_paths_for_task(task)
    if not source_paths:
        return {"built": False, "missing": ["中标价格对比分析", "中标价格查询", "招标结果台帐"]}
    sources = build_bid_source_sets(source_paths)
    missing = missing_bid_source_labels(sources)
    if missing:
        return {"built": False, "missing": missing, "sources": sources}
    output_path = EXPORT_DIR / f"bid_price_auto_{task.get('startDate','')}_{task.get('endDate','')}_{task.get('id','')[:8]}.xlsx"
    stats = write_bid_price_workbook(task, sources, output_path)
    return {"built": True, "stats": stats, "outputPath": output_path, "sources": sources}


def comparison_company_prices(row: dict[str, Any]) -> dict[str, float]:
    prices: dict[str, float] = {}
    for company, aliases in BID_COMPANY_COLUMNS:
        price = valid_price(value_by_alias(row, aliases))
        if price:
            prices[company] = price
    return prices


def build_bid_detail_rows(sources: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    comparison_rows = sources.get("comparison", [])
    query_rows = sources.get("query", [])
    ledger_rows = sources.get("ledger", [])
    pm_rows = sources.get("pm", [])

    ledger_by_task = {
        clean_text(map_alias_row(row, BID_LEDGER_ALIASES).get("taskNo") or row.get("tendercode") or row.get("taskcode") or row.get("code") or row.get("procurementcode")):
        clean_text(map_alias_row(row, BID_LEDGER_ALIASES).get("winner") or row.get("suppliername") or row.get("winbidder") or row.get("biddername"))
        for row in ledger_rows
        if clean_text(map_alias_row(row, BID_LEDGER_ALIASES).get("taskNo") or row.get("tendercode") or row.get("taskcode") or row.get("code") or row.get("procurementcode"))
    }
    region_by_project: dict[str, str] = {}
    for row in pm_rows:
        mapped = map_alias_row(row, BID_PM_ALIASES)
        region = normalize_region(mapped.get("region"), mapped.get("project") or mapped.get("department"))
        for key in [mapped.get("department"), mapped.get("project")]:
            text = clean_text(key)
            if text and region:
                region_by_project[text] = region

    if any("history" in row for row in query_rows):
        return expand_verified_bid_details(query_rows, ledger_by_task, region_by_project), comparison_rows

    details: list[dict[str, Any]] = []
    for row in query_rows:
        mapped = map_alias_row(row, BID_QUERY_ALIASES)
        list_no = clean_text(mapped.get("listNo"))
        project = clean_text(mapped.get("project"))
        process = clean_text(mapped.get("process"))
        company = clean_bid_company_name(mapped.get("company"), project)
        if "测试" in process or "测试" in company or "测试" in project:
            continue
        price = valid_price(mapped.get("price"))
        if not list_no or not price:
            continue
        quantity = to_number(mapped.get("quantity"))
        amount = to_number(mapped.get("amount"))
        if not amount and quantity:
            amount = price * quantity
        if not quantity and amount:
            quantity = amount / price if price else 0.0
        task_no = clean_text(mapped.get("taskNo"))
        region = clean_text(mapped.get("region")) or region_by_project.get(project, "")
        if not region:
            for pm_project, pm_region in region_by_project.items():
                if pm_project and (pm_project in project or project in pm_project):
                    region = pm_region
                    break
        region = normalize_region(region, project)
        details.append({
            "listNo": list_no,
            "material": clean_text(mapped.get("material")),
            "spec": clean_text(mapped.get("spec")),
            "unit": clean_text(mapped.get("unit")),
            "company": company,
            "companyBucket": bid_company_bucket(company),
            "project": project,
            "process": process,
            "price": price,
            "notaxPrice": to_number(mapped.get("notaxPrice")),
            "quantity": quantity,
            "amount": amount,
            "bidDate": parse_date(mapped.get("bidDate")) or clean_text(mapped.get("bidDate")),
            "brand": clean_text(mapped.get("brand")),
            "remark": clean_text(mapped.get("remark")),
            "explanation": clean_text(mapped.get("explanation")),
            "payment": normalize_payment(mapped.get("payment")),
            "paymentDesc": clean_text(mapped.get("paymentDesc")),
            "taskNo": task_no,
            "winner": clean_text(mapped.get("winner")) or ledger_by_task.get(task_no, ""),
            "region": region,
        })
    details.sort(key=lambda row: (
        clean_text(row.get("listNo")),
        clean_text(row.get("material")),
        clean_text(row.get("spec")),
        clean_text(row.get("unit")),
        clean_text(row.get("company")),
        clean_text(row.get("project")),
        clean_text(row.get("bidDate")),
    ))
    return details, comparison_rows


def expand_verified_bid_details(query_rows: list[dict[str, Any]], ledger_by_task: dict[str, str], region_by_project: dict[str, str]) -> list[dict[str, Any]]:
    details: list[dict[str, Any]] = []
    seen: set[tuple[Any, ...]] = set()
    for row in query_rows:
        history = row.get("history")
        if isinstance(history, str):
            try:
                history = json.loads(history)
            except Exception:
                history = []
        if not isinstance(history, list):
            history = []
        for item in history:
            price = valid_price(item.get("price"))
            if not price:
                continue
            project = clean_text(item.get("projectname") or row.get("projectname"))
            raw_org = clean_text(item.get("orgname") or row.get("orgname"))
            process = clean_text(item.get("tendername") or row.get("tendername") or row.get("taskname"))
            demand = clean_bid_company_name(raw_org, project)
            if "测试" in process or "测试" in demand or "测试" in project:
                continue
            task_no = clean_text(item.get("tendercode") or row.get("tendercode") or row.get("taskcode"))
            company = clean_bid_company_name(raw_org, project)
            region = clean_text(item.get("areaname") or row.get("areaname")) or region_by_project.get(project, "")
            if not region:
                for pm_project, pm_region in region_by_project.items():
                    if pm_project and (pm_project in project or project in pm_project):
                        region = pm_region
                        break
            detail = {
                "listNo": clean_text(item.get("materialcode") or row.get("materialcode")),
                "material": clean_text(item.get("materialname") or row.get("materialname")),
                "spec": clean_text(item.get("specification") or row.get("specification")),
                "unit": clean_text(item.get("unit") or row.get("unit")),
                "company": company,
                "companyBucket": bid_company_bucket(company),
                "project": project,
                "process": process,
                "price": price,
                "notaxPrice": to_number(item.get("notaxprice")),
                "quantity": to_number(item.get("quantity")),
                "amount": 0,
                "bidDate": parse_date(item.get("resultpublishdate") or row.get("resultpublishdate")) or clean_text(item.get("resultpublishdate") or row.get("resultpublishdate")),
                "brand": clean_text(item.get("brand")),
                "remark": clean_text(item.get("remark")),
                "explanation": clean_text(item.get("explanation")),
                "payment": normalize_payment(item.get("paymethodcode") or row.get("paymethodcode")),
                "paymentDesc": clean_text(item.get("paymethoddesc") or row.get("paymethoddesc")),
                "taskNo": task_no,
                "winner": clean_text(item.get("suppliername") or row.get("suppliername")) or ledger_by_task.get(task_no, ""),
                "region": normalize_region(region, project),
            }
            key = (
                detail["taskNo"],
                detail["listNo"],
                detail["material"],
                detail["spec"],
                detail["unit"],
                detail["company"],
                detail["project"],
                detail["price"],
                detail["bidDate"],
            )
            if key in seen:
                continue
            seen.add(key)
            details.append(detail)
    details.sort(key=lambda row: (
        clean_text(row.get("listNo")),
        clean_text(row.get("material")),
        clean_text(row.get("spec")),
        clean_text(row.get("unit")),
        clean_text(row.get("company")),
        clean_text(row.get("project")),
        clean_text(row.get("bidDate")),
    ))
    return details


def comparable_bid_details(details: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[tuple[str, str, str, str], list[dict[str, Any]]] = {}
    for detail in details:
        groups.setdefault(bid_group_key(detail), []).append(detail)
    keep_list_nos = {
        key[0]
        for key, group_rows in groups.items()
        if len({clean_text(row.get("company")) for row in group_rows if row.get("price")}) >= 2
    }
    return [row for row in details if clean_text(row.get("listNo")) in keep_list_nos]


def bid_group_key(row: dict[str, Any]) -> tuple[str, str, str, str]:
    return (
        clean_text(row.get("listNo")),
        clean_text(row.get("material")),
        clean_text(row.get("spec")),
        clean_text(row.get("unit")),
    )


def region_summary(rows: list[dict[str, Any]]) -> str:
    regions = sorted({normalize_region(row.get("region"), row.get("project")) for row in rows if clean_text(row.get("region"))})
    if not regions:
        return ""
    province_by_city = {
        "烟台市": "山东省", "青岛市": "山东省", "济南市": "山东省", "威海市": "山东省", "潍坊市": "山东省", "淄博市": "山东省", "临沂市": "山东省", "日照市": "山东省", "东营市": "山东省", "泰安市": "山东省", "济宁市": "山东省", "德州市": "山东省", "聊城市": "山东省", "滨州市": "山东省", "菏泽市": "山东省", "枣庄市": "山东省",
        "南京市": "江苏省", "苏州市": "江苏省", "杭州市": "浙江省", "宁波市": "浙江省", "广州市": "广东省", "深圳市": "广东省", "成都市": "四川省", "西安市": "陕西省", "郑州市": "河南省", "合肥市": "安徽省", "武汉市": "湖北省",
    }
    city_by_area = {
        "海阳市": "烟台市", "莱阳市": "烟台市", "莱州市": "烟台市", "龙口市": "烟台市", "招远市": "烟台市", "栖霞市": "烟台市", "蓬莱区": "烟台市", "莱山区": "烟台市", "牟平区": "烟台市", "芝罘区": "烟台市", "高新区": "烟台市", "开发区": "烟台市",
        "黄岛区": "青岛市", "即墨区": "青岛市", "胶州市": "青岛市", "平度市": "青岛市", "莱西市": "青岛市",
    }
    municipalities = {"北京市", "上海市", "天津市", "重庆市"}

    def province_city(region: str) -> tuple[str, str | None]:
        if region == "全国":
            return "全国", None
        if region in municipalities:
            return region, region
        if region in province_by_city:
            return province_by_city[region], region
        if region in city_by_area:
            city = city_by_area[region]
            return province_by_city.get(city, ""), city
        if region.endswith("省"):
            return region, None
        for area, city in city_by_area.items():
            if area and area in region:
                return province_by_city.get(city, ""), city
        for city, province in province_by_city.items():
            if city and city in region:
                return province, city
        return region, None

    provinces = set()
    cities = set()
    for region in regions:
        province, city = province_city(region)
        if province == "全国":
            return "全国"
        if province:
            provinces.add(province)
        if city:
            cities.add(city)
    if len(provinces) > 1:
        return "全国"
    if len(cities) == 1:
        return next(iter(cities))
    if len(provinces) == 1:
        return next(iter(provinces))
    return regions[0]


def build_bid_average_rows(details: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str, str, str], list[dict[str, Any]]] = {}
    for row in details:
        grouped.setdefault(bid_group_key(row), []).append(row)
    rows: list[dict[str, Any]] = []
    for key, group_rows in grouped.items():
        prices = [row["price"] for row in group_rows if row.get("price")]
        if not prices:
            continue
        item = {
            "清单编号": key[0],
            "材料名称": key[1],
            "规格/项目特征": key[2],
            "单位": key[3],
            "地区": region_summary(group_rows),
            "集团平均中标价格": round(sum(prices) / len(prices), 2),
        }
        for company, _aliases in BID_COMPANY_COLUMNS:
            company_prices = [row["price"] for row in group_rows if row.get("companyBucket") == company]
            item[company] = round(sum(company_prices) / len(company_prices), 2) if company_prices else ""
        rows.append(item)
    return rows


def price_ratio(high: float, low: float) -> float:
    return (high - low) / low if low else 0.0


def build_bid_high_rows(details: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str, str, str], list[dict[str, Any]]] = {}
    for row in details:
        grouped.setdefault(bid_group_key(row), []).append(row)
    high_rows: list[dict[str, Any]] = []
    for key, group_rows in grouped.items():
        sorted_rows = sorted(group_rows, key=lambda row: row.get("price") or 0)
        if len(sorted_rows) < 2:
            continue
        low = sorted_rows[0]
        high = sorted_rows[-1]
        dongtai_rows = [row for row in sorted_rows if row.get("companyBucket") == "东泰物流"]
        include = False
        reason = ""
        dt_high = dt_low = None
        if dongtai_rows:
            dt_low = min(dongtai_rows, key=lambda row: row.get("price") or 0)
            dt_high = max(dongtai_rows, key=lambda row: row.get("price") or 0)
            other_rows = [row for row in sorted_rows if row.get("companyBucket") != "东泰物流"]
            if other_rows:
                other_low = other_rows[0]
                other_high = other_rows[-1]
                if price_ratio(dt_high["price"], other_low["price"]) > 0.5 or price_ratio(other_high["price"], dt_low["price"]) > 0.5:
                    include = True
                    reason = "东泰物流与其他公司价差超过50%"
            if not include and price_ratio(dt_high["price"], dt_low["price"]) > 0.1:
                include = True
                reason = "东泰物流自身高低价差超过10%"
        elif price_ratio(high["price"], low["price"]) > 0.2:
            include = True
            reason = "其他公司高低价差超过20%"
        if not include:
            continue
        high_rows.append({
            "清单编号": key[0],
            "材料名称": key[1],
            "规格/项目特征": key[2],
            "单位": key[3],
            "最高价公司": high.get("company", ""),
            "最高价": round(high.get("price", 0), 2),
            "最低价公司": low.get("company", ""),
            "最低价": round(low.get("price", 0), 2),
            "价差率": round(price_ratio(high.get("price", 0), low.get("price", 0)), 4),
            "东泰物流最高价": round(dt_high.get("price", 0), 2) if dt_high else "",
            "东泰物流最低价": round(dt_low.get("price", 0), 2) if dt_low else "",
            "原因": reason,
            "_highDetail": high,
            "_lowDetail": low,
            "_dtHighDetail": dt_high,
            "_dtLowDetail": dt_low,
        })
    return high_rows


def build_bid_dongtai_rows(comparison_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row in comparison_rows:
        price = valid_price(value_by_alias(row, ["东泰", "东泰物流", "dtgs"]))
        if not price:
            continue
        mapped = map_alias_row(row, BID_COMPARISON_ALIASES)
        rows.append({
            "清单编号": clean_text(mapped.get("listNo")),
            "清单名称": clean_text(mapped.get("material")),
            "规格/项目特征": clean_text(mapped.get("spec")),
            "单位": clean_text(mapped.get("unit")),
            "东泰物流平均中标价格": price,
        })
    return rows


def save_material_price_source(rows: list[dict[str, Any]], source: str) -> Path:
    output_path = DATA_DIR / "material_price_source.json"
    payload = {
        "source": source,
        "importedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "rows": rows,
    }
    with output_path.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    return output_path


def fix_bid_comment_vml(output_path: Path) -> None:
    script_path = ROOT / "_skill_yanjian_bid_price_comparison" / "yanjian-bid-price-comparison" / "scripts" / "fix_comment_vml.py"
    if not script_path.exists():
        return
    temp_path = output_path.with_name(f"{output_path.stem}.vmltmp{output_path.suffix}")
    try:
        subprocess.run([sys.executable, str(script_path), str(output_path), str(temp_path)], check=True, capture_output=True, text=True)
        temp_path.replace(output_path)
    except Exception:
        if temp_path.exists():
            temp_path.unlink()


def validate_bid_workbook(output_path: Path, expected_dongtai_rows: int) -> list[str]:
    errors: list[str] = []
    wb = load_workbook(output_path, data_only=False)
    required_sheets = ["中标价格高的材料", "平均中标价格横向对比", "中标价格横向对比数据明细", "东泰物流平均中标价格"]
    for sheet in required_sheets:
        if sheet not in wb.sheetnames:
            errors.append(f"缺少工作表：{sheet}")
    if errors:
        return errors
    ws1 = wb["中标价格高的材料"]
    ws3 = wb["中标价格横向对比数据明细"]
    ws4 = wb["东泰物流平均中标价格"]
    if ws3.max_column < 20:
        errors.append("表三未扩展到A:T。")
    if ws3["Q3"].value != "付款方式":
        errors.append("表三Q列不是付款方式。")
    valid_payments = {"", "现金", "分期", "抵房/物"}
    for row_index in range(4, ws3.max_row + 1):
        payment = clean_text(ws3.cell(row_index, 17).value)
        if payment not in valid_payments:
            errors.append(f"表三Q{row_index}付款方式未标准化：{payment}")
            break
        if "测试" in clean_text(ws3.cell(row_index, 7).value) or "测试" in clean_text(ws3.cell(row_index, 5).value):
            errors.append(f"表三第{row_index}行含测试数据。")
            break
    actual_dongtai_rows = max(ws4.max_row - 2, 0)
    if actual_dongtai_rows != expected_dongtai_rows:
        errors.append(f"表四东泰行数不匹配：应为{expected_dongtai_rows}，实际{actual_dongtai_rows}。")
    for row_index in range(5, ws1.max_row + 1):
        for col_index in range(6, 20):
            cell = ws1.cell(row_index, col_index)
            if cell.value not in ("", None) and not cell.hyperlink:
                errors.append(f"表一{cell.coordinate}缺少跳转到表三J列的超链接。")
                return errors
            color = getattr(getattr(cell, "font", None), "color", None)
            color_value = clean_text(getattr(color, "rgb", "")) if color and getattr(color, "type", "") == "rgb" else ""
            if cell.value not in ("", None) and color_value.upper().endswith("FF0000") and not cell.comment:
                errors.append(f"表一{cell.coordinate}最高价缺少批注。")
                return errors
    return errors


def clear_bid_template_sheet(ws: Any, start_row: int) -> None:
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)


def bid_date_text(task: dict[str, Any]) -> str:
    def fmt(value: Any) -> str:
        text = clean_text(value)
        try:
            year, month, day = text.split("-")[:3]
            return f"{int(year)}年{int(month)}月{int(day)}日"
        except Exception:
            return text

    return f"{fmt(task.get('startDate', ''))}—{fmt(task.get('endDate', ''))}"


def bid_grouped_details(details: list[dict[str, Any]]) -> dict[tuple[str, str, str, str], list[dict[str, Any]]]:
    groups: dict[tuple[str, str, str, str], list[dict[str, Any]]] = {}
    for row in details:
        groups.setdefault(bid_group_key(row), []).append(row)
    return dict(sorted(groups.items(), key=lambda item: tuple(str(value or "") for value in item[0])))


def bid_table1_trigger(items: list[dict[str, Any]]) -> bool:
    prices = [row["price"] for row in items if row.get("price") not in (None, 0)]
    if len(prices) < 2 or min(prices) <= 0:
        return False
    overall = price_ratio(max(prices), min(prices))
    if overall > 0.50:
        return True
    dongtai_prices = [row["price"] for row in items if row.get("companyBucket") == "东泰物流" and row.get("price")]
    other_prices = [row["price"] for row in items if row.get("companyBucket") != "东泰物流" and row.get("price")]
    if dongtai_prices:
        if other_prices:
            for dongtai_price in dongtai_prices:
                for other_price in other_prices:
                    if min(dongtai_price, other_price) > 0 and price_ratio(max(dongtai_price, other_price), min(dongtai_price, other_price)) > 0.50:
                        return True
        return len(dongtai_prices) >= 2 and min(dongtai_prices) > 0 and price_ratio(max(dongtai_prices), min(dongtai_prices)) > 0.10
    return len(other_prices) >= 2 and min(other_prices) > 0 and price_ratio(max(other_prices), min(other_prices)) > 0.20


def bid_record_for_value(items: list[dict[str, Any]], company: str, value: float, prefer_high: bool = True) -> dict[str, Any] | None:
    candidates = [row for row in items if row.get("companyBucket") == company and row.get("price") == value]
    if not candidates:
        candidates = [row for row in items if row.get("price") == value]
    if not candidates:
        return None
    return sorted(candidates, key=lambda row: clean_text(row.get("bidDate")), reverse=prefer_high)[0]


def write_verified_bid_details(ws: Any, details: list[dict[str, Any]]) -> dict[tuple[Any, ...], int]:
    detail_row_by_identity: dict[tuple[Any, ...], int] = {}
    def blank_if_none(value: Any) -> Any:
        return "" if value is None else value
    for seq, detail in enumerate(sorted(details, key=lambda row: (row.get("listNo", ""), row.get("material", ""), row.get("spec", ""), row.get("company", ""), row.get("project", ""))), start=1):
        ws.append([
            seq,
            detail.get("listNo", ""),
            detail.get("material", ""),
            detail.get("spec", ""),
            detail.get("unit", ""),
            detail.get("company", ""),
            detail.get("project", ""),
            blank_if_none(detail.get("quantity")),
            blank_if_none(detail.get("notaxPrice")),
            round(detail.get("price", 0), 2) if detail.get("price") else "",
            detail.get("bidDate", ""),
            detail.get("brand", ""),
            detail.get("remark", ""),
            detail.get("explanation", ""),
            detail.get("taskNo", ""),
            detail.get("process", ""),
            detail.get("payment", ""),
            detail.get("paymentDesc", ""),
            detail.get("winner", ""),
            detail.get("region", ""),
        ])
        detail_row_by_identity[bid_detail_identity(detail)] = ws.max_row
    return detail_row_by_identity


def write_verified_bid_table2(ws: Any, groups: dict[tuple[str, str, str, str], list[dict[str, Any]]]) -> int:
    seq = 1
    for key, group_rows in groups.items():
        prices = [row["price"] for row in group_rows if row.get("price")]
        if not prices:
            continue
        code, name, spec, unit = key
        row_index = ws.max_row + 1
        ws.cell(row_index, 1, seq)
        ws.cell(row_index, 2, "物资材料库")
        ws.cell(row_index, 5, code)
        ws.cell(row_index, 6, name)
        ws.cell(row_index, 7, spec)
        ws.cell(row_index, 8, unit)
        ws.cell(row_index, 9, region_summary(group_rows))
        group_avg = sum(prices) / len(prices)
        ws.cell(row_index, 10, group_avg)
        for company_index, (company, _aliases) in enumerate(BID_COMPANY_COLUMNS, start=11):
            if company == "其他公司":
                known = {item[0] for item in BID_COMPANY_COLUMNS if item[0] != "其他公司"}
                company_prices = [row["price"] for row in group_rows if row.get("companyBucket") not in known and row.get("price")]
            else:
                company_prices = [row["price"] for row in group_rows if row.get("companyBucket") == company and row.get("price")]
            if company_prices:
                ws.cell(row_index, company_index, sum(company_prices) / len(company_prices))
        seq += 1
    return seq - 1


def write_verified_bid_table1(ws1: Any, ws3: Any, groups: dict[tuple[str, str, str, str], list[dict[str, Any]]], detail_row_by_identity: dict[tuple[Any, ...], int]) -> int:
    red_font = Font(name="宋体", size=10, bold=True, color="FFFF0000")
    black_font = Font(name="宋体", size=10, bold=True, color="FF000000")
    yellow_fill = PatternFill("solid", fgColor="FFFF00")
    seq = 1
    company_columns = [(company, col) for col, (company, _aliases) in enumerate(BID_COMPANY_COLUMNS[:12], start=6)]
    for key, items in groups.items():
        if not bid_table1_trigger(items):
            continue
        code, name, spec, unit = key
        prices = [row["price"] for row in items if row.get("price")]
        if not prices:
            continue
        overall_high = max(prices)
        overall_low = min(prices)
        row_index = ws1.max_row + 1
        ws1.cell(row_index, 1, seq)
        ws1.cell(row_index, 2, code)
        ws1.cell(row_index, 3, name)
        ws1.cell(row_index, 4, spec)
        ws1.cell(row_index, 5, unit)
        high_records: list[dict[str, Any] | None] = []
        for company, col_index in company_columns:
            company_prices = [row["price"] for row in items if row.get("companyBucket") == company and row.get("price")]
            if not company_prices:
                continue
            company_max = max(company_prices)
            company_min = min(company_prices)
            if abs(company_max - overall_high) < 0.0001:
                value = company_max
                is_high = True
            elif abs(company_min - overall_low) < 0.0001:
                value = company_min
                is_high = False
            else:
                continue
            cell = ws1.cell(row_index, col_index, round(value, 2))
            detail = bid_record_for_value(items, company, value, is_high)
            cell.font = red_font if is_high else black_font
            add_verified_bid_link_and_comment(cell, detail, detail_row_by_identity, high=is_high)
            if is_high:
                high_records.append(detail)
        dongtai_rows = [row for row in items if row.get("companyBucket") == "东泰物流" and row.get("price")]
        if dongtai_rows:
            dt_high = max(row["price"] for row in dongtai_rows)
            dt_low = min(row["price"] for row in dongtai_rows)
            high_cell = ws1.cell(row_index, 18, round(dt_high, 2))
            low_cell = ws1.cell(row_index, 19, round(dt_low, 2))
            high_detail = bid_record_for_value(items, "东泰物流", dt_high, True)
            low_detail = bid_record_for_value(items, "东泰物流", dt_low, False)
            is_high = abs(dt_high - overall_high) < 0.0001
            high_cell.font = red_font if is_high else black_font
            low_cell.font = black_font
            add_verified_bid_link_and_comment(high_cell, high_detail, detail_row_by_identity, high=is_high)
            add_verified_bid_link_and_comment(low_cell, low_detail, detail_row_by_identity, high=False)
            if is_high:
                high_records.append(high_detail)
        for detail in high_records:
            target_row = detail_row_by_identity.get(bid_detail_identity(detail or {}))
            if target_row:
                ws3.cell(target_row, 10).fill = yellow_fill
        seq += 1
    return seq - 1


def add_verified_bid_link_and_comment(cell: Any, detail: dict[str, Any] | None, detail_row_by_identity: dict[tuple[Any, ...], int], high: bool = False) -> None:
    if not detail:
        return
    target_row = detail_row_by_identity.get(bid_detail_identity(detail))
    if target_row:
        cell.hyperlink = f"#'中标价格横向对比数据明细'!J{target_row}"
    if high:
        comment = Comment(
            f"1.项目名称：{detail.get('project', '')}\n"
            f"2.项目地区：{detail.get('region', '')}\n"
            f"3.付款方式：{detail.get('paymentDesc') or detail.get('payment', '')}\n"
            "4.原因说明：",
            "Codex",
        )
        comment.width = 340
        comment.height = 227
        cell.comment = comment


def write_verified_bid_table4(ws: Any, dongtai_rows: list[dict[str, Any]]) -> int:
    for seq, item in enumerate(dongtai_rows, start=1):
        ws.append([
            seq,
            item.get("清单编号", ""),
            item.get("清单名称", ""),
            item.get("规格/项目特征", ""),
            item.get("单位", ""),
            item.get("东泰物流平均中标价格", ""),
        ])
    return len(dongtai_rows)


def apply_verified_bid_template_format(wb: Workbook) -> None:
    thin = Side(style="thin", color="FF000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    red_font = Font(name="宋体", size=10, bold=True, color="FFFF0000")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        autosize(ws)
    ws2 = wb["平均中标价格横向对比"]
    for row in ws2.iter_rows(min_row=5, max_row=ws2.max_row, min_col=11, max_col=min(ws2.max_column, 24)):
        group_avg = to_number(ws2.cell(row[0].row, 10).value)
        for cell in row:
            if to_number(cell.value) > group_avg > 0:
                cell.font = red_font
    for ws, start_row in [(wb["中标价格高的材料"], 5), (ws2, 5), (wb["中标价格横向对比数据明细"], 4), (wb["东泰物流平均中标价格"], 3)]:
        for row in ws.iter_rows(min_row=start_row):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0" if cell.column == 1 else "0.00"


def write_bid_price_workbook(task: dict[str, Any], sources: dict[str, Any], output_path: Path) -> dict[str, int]:
    source_details, comparison_rows = build_bid_detail_rows(sources)
    details = comparable_bid_details(source_details)
    dongtai_rows = build_bid_dongtai_rows(comparison_rows)
    template_path = TEMPLATE_DIR / "bid-price-template.xlsx"
    wb = load_workbook(template_path) if template_path.exists() else Workbook()
    if not template_path.exists():
        raise FileNotFoundError("缺少已验证样表 templates/bid-price-template.xlsx")
    ws1 = wb["中标价格高的材料"]
    ws2 = wb["平均中标价格横向对比"]
    ws3 = wb["中标价格横向对比数据明细"]
    ws4 = wb["东泰物流平均中标价格"]
    for ws, start_row in [(ws1, 5), (ws2, 5), (ws3, 4), (ws4, 3)]:
        clear_bid_template_sheet(ws, start_row)
    date_text = bid_date_text(task)
    ws1.cell(1, 1).value = f"2026年集采平台近三个月中标价格对比（{date_text}）"
    ws2.cell(1, 1).value = f"2026年集采平台近三个月中标价格对比（{date_text}）"
    ws3.cell(1, 1).value = f"2026年集采平台近三个月中标价格横向对比数据明细（{date_text}）"
    ws4.cell(1, 1).value = f"2026年集采平台中标价格（{date_text}）"
    groups = bid_grouped_details(details)
    detail_row_by_identity = write_verified_bid_details(ws3, details)
    average_count = write_verified_bid_table2(ws2, groups)
    high_count = write_verified_bid_table1(ws1, ws3, groups, detail_row_by_identity)
    dongtai_count = write_verified_bid_table4(ws4, dongtai_rows)
    apply_verified_bid_template_format(wb)
    wb.save(output_path)
    fix_bid_comment_vml(output_path)
    validation_errors = validate_bid_workbook(output_path, dongtai_count)
    return {
        "details": len(details),
        "sourceDetails": len(source_details),
        "average": average_count,
        "high": high_count,
        "dongtai": dongtai_count,
        "validationErrors": validation_errors,
    }


def bid_detail_identity(row: dict[str, Any]) -> tuple[Any, ...]:
    return (
        row.get("listNo"),
        row.get("material"),
        row.get("spec"),
        row.get("unit"),
        row.get("company"),
        row.get("project"),
        row.get("price"),
        row.get("taskNo"),
    )


def apply_bid_workbook_format(wb: Workbook, high_rows: list[dict[str, Any]], detail_row_by_identity: dict[tuple[Any, ...], int]) -> None:
    thin = Side(style="thin", color="B8C2CC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="DCE8F2")
    yellow_fill = PatternFill("solid", fgColor="FFF2CC")
    alt_fill_a = PatternFill("solid", fgColor="D9EAF7")
    alt_fill_b = PatternFill("solid", fgColor="FCE4D6")
    red_font = Font(name="宋体", color="FF0000", bold=True)
    body_font = Font(name="宋体", size=10)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.font = body_font
        max_col = ws.max_column
        if max_col > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
            if not ws.title.startswith("表四"):
                ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
        ws.row_dimensions[1].height = 49 if not ws.title.startswith("表三") else 37
        ws.cell(1, 1).font = Font(name="宋体" if not ws.title.startswith("表三") else "Microsoft YaHei UI", size=18, bold=True)
        ws.cell(2, 1).font = Font(name="宋体" if not ws.title.startswith("表三") else "Microsoft YaHei UI", size=14, bold=True, color="1F4E79")
        ws.cell(2, 1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        header_rows = [4] if ws.title.startswith(("表一", "表二")) else [3] if ws.title.startswith("表三") else [2]
        for row_index in header_rows:
            if row_index <= ws.max_row:
                ws.row_dimensions[row_index].height = 36 if not ws.title.startswith("表三") else 24
                for cell in ws[row_index]:
                    cell.fill = header_fill
                    cell.font = Font(name="宋体" if not ws.title.startswith("表三") else "Microsoft YaHei UI", size=12 if not ws.title.startswith("表三") else 9, bold=True)
        if ws.title.startswith(("表一", "表二")) and ws.max_row >= 3:
            ws.row_dimensions[3].height = 36
            for cell in ws[3]:
                cell.fill = header_fill
                cell.font = Font(name="宋体", size=12, bold=True)
        data_start = 5 if ws.title.startswith(("表一", "表二")) else 4 if ws.title.startswith("表三") else 3
        for row_index in range(data_start, ws.max_row + 1):
            ws.row_dimensions[row_index].height = 24 if ws.title.startswith("表三") else 18
            for cell in ws[row_index]:
                cell.font = Font(name="Microsoft YaHei UI" if ws.title.startswith("表三") else "宋体", size=9 if ws.title.startswith("表三") else 10)
        autosize(ws)
        ws.freeze_panes = "A5" if ws.title.startswith(("表一", "表二")) else "A4" if ws.title.startswith("表三") else "A3"
    ws1 = wb["表一 中标价格高的材料"]
    ws3 = wb["表三 中标价格横向对比数据明细"]
    for index, item in enumerate(high_rows, start=5):
        for col in (6, 10):
            if ws1.cell(index, col).value not in ("", None):
                ws1.cell(index, col).font = red_font
        high_detail = item.get("_highDetail") or {}
        low_detail = item.get("_lowDetail") or {}
        dt_high_detail = item.get("_dtHighDetail") or {}
        dt_low_detail = item.get("_dtLowDetail") or {}
        for col, detail in [(6, high_detail), (8, low_detail), (10, dt_high_detail), (11, dt_low_detail)]:
            target_row = detail_row_by_identity.get(bid_detail_identity(detail))
            if not target_row:
                continue
            ws1.cell(index, col).hyperlink = f"#'表三 中标价格横向对比数据明细'!J{target_row}"
            ws1.cell(index, col).style = "Hyperlink"
            ws3.cell(target_row, 10).fill = yellow_fill
            if col in (6, 10):
                comment = Comment(
                    f"1.项目名称：{detail.get('project', '')}\n"
                    f"2.项目地区：{detail.get('region', '')}\n"
                    f"3.付款方式：{detail.get('payment', '')}\n"
                    "4.原因说明：",
                    "Codex",
                )
                comment.width = 340
                comment.height = 226
                ws1.cell(index, col).comment = comment
        for col in (9,):
            ws1.cell(index, col).number_format = "0.00%"
    ws2 = wb["表二 平均中标价格横向对比"]
    for row in ws2.iter_rows(min_row=5, min_col=11, max_col=min(ws2.max_column, 24)):
        group_avg = to_number(ws2.cell(row[0].row, 6).value)
        for cell in row:
            if to_number(cell.value) > group_avg > 0:
                cell.font = red_font
    for row in ws2.iter_rows(min_row=5, max_row=ws2.max_row):
        row[9].number_format = "0.00%"
    seen_list_no: dict[str, PatternFill] = {}
    for row_index in range(4, ws3.max_row + 1):
        list_no = clean_text(ws3.cell(row_index, 1).value)
        if list_no not in seen_list_no:
            seen_list_no[list_no] = alt_fill_a if len(seen_list_no) % 2 == 0 else alt_fill_b
        fill = seen_list_no[list_no]
        for col_index in range(1, min(ws3.max_column, 20) + 1):
            if ws3.cell(row_index, col_index).fill == yellow_fill:
                continue
            ws3.cell(row_index, col_index).fill = fill
        ws3.cell(row_index, 10).number_format = "0.00"
        ws3.cell(row_index, 17).value = normalize_payment(ws3.cell(row_index, 17).value)


def write_platform_material_skill_workbook(task: dict[str, Any], rows: list[dict[str, Any]], output_path: Path) -> dict[str, int]:
    date_text = f"{task.get('startDate', '')}—{task.get('endDate', '')}"
    raw_details: list[dict[str, Any]] = []
    for row in rows:
        normalized = normalize_material_row(row)
        price = to_number(normalized.get("price"))
        if not normalized.get("material") or price <= 0:
            continue
        company = clean_bid_company_name(normalized.get("company"), normalized.get("project"))
        if not company or company == "集团":
            company = clean_text(row.get("company")) or clean_text(row.get("project"))
        if "装饰幕墙" in company or "测试" in company or "测试" in normalized.get("project", ""):
            continue
        list_no = clean_text(row.get("materialCode") or row.get("listNo"))
        if not list_no:
            list_no = f"PLAT-{uuid.uuid5(uuid.NAMESPACE_DNS, normalize_material_key(normalized)).hex[:10].upper()}"
        raw_details.append({
            "listNo": list_no,
            "material": normalized.get("material", ""),
            "spec": normalized.get("spec", ""),
            "unit": normalized.get("unit", ""),
            "company": company,
            "companyBucket": bid_company_bucket(company),
            "project": normalized.get("project", ""),
            "process": clean_text(row.get("process")) or "平台材料中标信息",
            "price": price,
            "bidDate": normalized.get("purchaseDate", ""),
            "payment": normalize_payment(row.get("payment")),
            "taskNo": clean_text(row.get("taskNo") or row.get("platformOrder")),
            "winner": clean_text(row.get("supplier") or row.get("winner")),
            "region": normalize_region(row.get("region"), normalized.get("project")),
        })

    grouped: dict[tuple[str, str, str, str], list[dict[str, Any]]] = {}
    for detail in raw_details:
        grouped.setdefault(bid_group_key(detail), []).append(detail)
    details = [
        detail
        for group_rows in grouped.values()
        if len({clean_text(item.get("company")) for item in group_rows if clean_text(item.get("company"))}) >= 2
        for detail in group_rows
    ]
    comparison_rows = [
        {
            "清单编号": key[0],
            "清单名称": key[1],
            "规格/项目特征": key[2],
            "单位": key[3],
            **{
                company: round(sum(item["price"] for item in group_rows if item.get("companyBucket") == company) / len([item for item in group_rows if item.get("companyBucket") == company]), 2)
                for company, _aliases in BID_COMPANY_COLUMNS
                if [item for item in group_rows if item.get("companyBucket") == company]
            },
        }
        for key, group_rows in grouped.items()
    ]
    average_rows = build_bid_average_rows(details)
    high_rows = build_bid_high_rows(details)
    dongtai_rows = build_bid_dongtai_rows(comparison_rows)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "表一 中标价格高的材料"
    ws2 = wb.create_sheet("表二 平均中标价格横向对比")
    ws3 = wb.create_sheet("表三 中标价格横向对比数据明细")
    ws4 = wb.create_sheet("表四 东泰物流平均中标价格")

    headers1 = [
        "清单编号", "材料名称", "规格/项目特征", "单位", "最高价公司", "最高价", "最低价公司", "最低价", "价差率",
        "东泰物流最高价", "东泰物流最低价", "原因", "最高价项目", "最高价地区", "最高价付款方式",
        "最低价项目", "最低价地区", "最低价付款方式", "任务编号",
    ]
    headers2 = [
        "清单编号", "材料名称", "规格/项目特征", "单位", "地区", "集团平均中标价格", "价格记录数", "最低价", "最高价", "价差率",
    ] + [company for company, _aliases in BID_COMPANY_COLUMNS]
    headers3 = [
        "清单编号", "材料名称", "规格/项目特征", "单位", "公司名称", "项目名称", "采购过程名称", "中标单位", "地区", "含税中标价格",
        "中标日期", "预留1", "预留2", "预留3", "预留4", "预留5", "付款方式", "任务编号", "预留6", "预留7",
    ]
    headers4 = ["清单编号", "清单名称", "规格/项目特征", "单位", "东泰物流平均中标价格"]

    ws1.append([f"烟建集团有限公司中标价格高的材料（{date_text}）"])
    ws1.append(["红色字体为高价，点击价格可跳转至表三明细。"])
    ws1.append([""] * len(headers1))
    ws1.append(headers1)
    ws2.append([f"烟建集团有限公司平均中标价格横向对比（{date_text}）"])
    ws2.append(["红色字体为高于集团平均中标价格。"])
    ws2.append([""] * len(headers2))
    ws2.append(headers2)
    ws3.append([f"烟建集团有限公司中标价格横向对比数据明细（{date_text}）"])
    ws3.append(["按清单编号排序，剔除测试数据，付款方式已标准化。"])
    ws3.append(headers3)
    ws4.append([f"烟建集团有限公司东泰物流平均中标价格（{date_text}）"])
    ws4.append(headers4)

    for item in high_rows:
        high_detail = item.get("_highDetail") or {}
        low_detail = item.get("_lowDetail") or {}
        ws1.append([
            item.get("清单编号", ""),
            item.get("材料名称", ""),
            item.get("规格/项目特征", ""),
            item.get("单位", ""),
            item.get("最高价公司", ""),
            item.get("最高价", ""),
            item.get("最低价公司", ""),
            item.get("最低价", ""),
            item.get("价差率", ""),
            item.get("东泰物流最高价", ""),
            item.get("东泰物流最低价", ""),
            item.get("原因", ""),
            high_detail.get("project", ""),
            high_detail.get("region", ""),
            high_detail.get("payment", ""),
            low_detail.get("project", ""),
            low_detail.get("region", ""),
            low_detail.get("payment", ""),
            high_detail.get("taskNo", "") or low_detail.get("taskNo", ""),
        ])
    for item in average_rows:
        group_key = (item.get("清单编号"), item.get("材料名称"), item.get("规格/项目特征"), item.get("单位"))
        group_prices = [detail["price"] for detail in details if bid_group_key(detail) == group_key and detail.get("price")]
        min_price = min(group_prices) if group_prices else 0
        max_price = max(group_prices) if group_prices else 0
        payload = {
            **item,
            "价格记录数": len(group_prices),
            "最低价": round(min_price, 2) if min_price else "",
            "最高价": round(max_price, 2) if max_price else "",
            "价差率": round(price_ratio(max_price, min_price), 4) if min_price else "",
        }
        ws2.append([payload.get(header, "") for header in headers2])
    detail_row_by_identity: dict[tuple[Any, ...], int] = {}
    for detail in sorted(details, key=lambda row: (row.get("listNo", ""), row.get("material", ""), row.get("company", ""))):
        ws3.append([
            detail.get("listNo", ""),
            detail.get("material", ""),
            detail.get("spec", ""),
            detail.get("unit", ""),
            detail.get("company", ""),
            detail.get("project", ""),
            detail.get("process", ""),
            detail.get("winner", ""),
            detail.get("region", ""),
            round(detail.get("price", 0), 2),
            detail.get("bidDate", ""),
            "",
            "",
            "",
            "",
            "",
            detail.get("payment", ""),
            detail.get("taskNo", ""),
            "",
            "",
        ])
        detail_row_by_identity[bid_detail_identity(detail)] = ws3.max_row
    for item in dongtai_rows:
        ws4.append([item.get(header, "") for header in headers4])

    apply_bid_workbook_format(wb, high_rows, detail_row_by_identity)
    wb.save(output_path)
    fix_bid_comment_vml(output_path)
    return {"high": len(high_rows), "average": len(average_rows), "details": len(details), "dongtai": len(dongtai_rows)}


def apply_platform_skill_workbook_format(wb: Workbook, high_rows: list[dict[str, Any]], detail_row_by_identity: dict[tuple[str, str, float, str], int]) -> None:
    thin = Side(style="thin", color="B8C2CC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="DCE8F2")
    red_font = Font(name="宋体", color="FF0000", bold=True)
    yellow_fill = PatternFill("solid", fgColor="FFF2CC")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.font = Font(name="宋体", size=10)
        ws.cell(1, 1).font = Font(name="宋体", size=16, bold=True)
        for row_index in (2, 3):
            if row_index <= ws.max_row:
                for cell in ws[row_index]:
                    cell.fill = header_fill if row_index == 3 or ws.title.startswith("表四") and row_index == 2 else PatternFill(fill_type=None)
                    cell.font = Font(name="宋体", size=11, bold=True)
        ws.freeze_panes = "A4" if ws.max_row >= 4 else None
        autosize(ws)
    ws1 = wb["表一 中标价格高的材料"]
    ws3 = wb["表三 中标价格横向对比数据明细"]
    for item in high_rows:
        row_index = item["rowIndex"]
        ws1.cell(row_index, 6).font = red_font
        ws1.cell(row_index, 9).number_format = "0.00%"
        for col, project, price, order in [
            (6, item.get("highProject", ""), item.get("highPrice", 0), item.get("highOrder", "")),
            (8, item.get("lowProject", ""), item.get("lowPrice", 0), item.get("lowOrder", "")),
        ]:
            target_row = detail_row_by_identity.get((item["key"], project, price, order))
            if target_row:
                ws1.cell(row_index, col).hyperlink = f"#'表三 中标价格横向对比数据明细'!J{target_row}"
                ws1.cell(row_index, col).style = "Hyperlink"
                ws3.cell(target_row, 10).fill = yellow_fill
    for row in wb["表二 平均中标价格横向对比"].iter_rows(min_row=4):
        row[9].number_format = "0.00%"


def normalize_material_row(row: dict[str, Any]) -> dict[str, Any]:
    quantity = to_number(row.get("quantity"))
    amount = to_number(row.get("amount"))
    price = to_number(row.get("price"))
    if not price and amount and quantity:
        price = amount / quantity
    if not amount and price and quantity:
        amount = price * quantity
    return {
        "material": clean_text(row.get("material")),
        "spec": clean_text(row.get("spec")),
        "unit": clean_text(row.get("unit")),
        "price": round(price, 4) if price else 0,
        "quantity": round(quantity, 4) if quantity else 0,
        "amount": round(amount, 2) if amount else 0,
        "purchaseDate": parse_material_price_date(row.get("purchaseDate")),
        "project": clean_text(row.get("project")),
        "supplier": clean_text(row.get("supplier")),
        "company": clean_text(row.get("company")),
        "platformOrder": clean_text(row.get("platformOrder")),
    }


PRICE_LIBRARY_ALIASES = {
    "category": ["类型", "类型库", "类型库\n（下拉选择）", "类别", "价格类型", "业务类型"],
    "company": ["分公司", "基层单位", "公司", "单位", "采购单位", "所属单位"],
    "major": ["专业类别", "一级专业", "分包专业", "成本科目", "价格类型"],
    "discipline": ["专业", "二级专业", "建筑安装", "建筑/安装"],
    "trade": ["工种", "班组", "作业工种"],
    "process": ["工序", "施工工序", "工艺", "作业内容"],
    "item": ["分项名称", "分包项名称", "工序/费用项名称", "费用项名称", "施工内容", "工作内容", "清单名称", "材料名称", "名称"],
    "spec": ["规格型号", "规格/做法", "规格", "型号", "做法", "项目特征", "层高", "施工部位"],
    "unit": ["单位", "计量单位", "计量单位"],
    "price": ["含税单价", "不含税单价", "控制价", "控制价\n不含税", "实际采购价", "采购价", "单价", "价格", "综合单价", "合同单价", "结算单价", "采购单价", "变更后单价", "含税指导价"],
    "taxRate": ["税率", "税率(%)", "增值税率"],
    "quantity": ["工程量", "数量", "变更后数量"],
    "amount": ["金额", "合价", "合同金额", "结算金额", "变更后金额"],
    "region": ["地区", "区域", "城市", "所在地域"],
    "projectType": ["项目类型", "工程类型", "项目业态", "工程类型\n（下拉选择）"],
    "project": ["项目", "项目部", "项目名称", "项目部/项目名称", "项目基本信息匹配名称", "工程名称"],
    "supplier": ["供应商", "分包商", "分包商/班组", "班组名称", "分供商", "单位名称", "中标单位", "合同相对方"],
    "contractNo": ["合同编号", "结算编号", "招采编号", "订单编号", "单据编号"],
    "priceDate": ["日期", "价格日期", "定标日期", "合同日期", "结算日期", "采购日期", "中标日期", "开工日期"],
    "source": ["来源", "数据来源", "依据"],
    "remark": ["备注", "付款条件", "付款条件总结", "付款条件原文", "清单备注", "说明", "适用条件"],
}

PRICE_LIBRARY_STANDARD_FIELDS = [
    {"key": "category", "header": "类型", "required": True, "type": "枚举", "rule": "只能填写劳务分包或专业分包。", "options": "劳务分包、专业分包", "source": "PM合同分类、Excel类型库"},
    {"key": "company", "header": "分公司", "required": False, "type": "文本", "rule": "使用统一分公司简称，如三公司、四公司。", "options": "", "source": "公司、基层单位、编制机构"},
    {"key": "major", "header": "专业类别", "required": False, "type": "文本", "rule": "建议固定为土建、安装、装饰、市政等成本口径。", "options": "土建、安装、装饰、市政", "source": "成本科目、价格类型、专业类别"},
    {"key": "discipline", "header": "专业", "required": False, "type": "枚举", "rule": "优先填写建筑或安装，空白时系统按分项名称自动识别。", "options": "建筑、安装", "source": "专业、建筑安装、系统识别"},
    {"key": "trade", "header": "工种", "required": False, "type": "文本", "rule": "如木工、瓦工、油工、电工、水工；空白时系统自动识别。", "options": "", "source": "班组、作业工种、系统识别"},
    {"key": "process", "header": "工序", "required": False, "type": "文本", "rule": "如模板支设、砌筑、电气配管配线、给排水管道安装。", "options": "", "source": "工序、施工工序、系统识别"},
    {"key": "item", "header": "分项名称", "required": True, "type": "文本", "rule": "价格对应的最小可比施工内容，不写平均值、合计行。", "options": "", "source": "分包项名称、工序/费用项名称、工作内容"},
    {"key": "spec", "header": "规格/做法", "required": False, "type": "文本", "rule": "填写规格、项目特征、做法、层高、施工部位等影响价格的条件。", "options": "", "source": "项目特征、层高、施工部位、工程量计算规则"},
    {"key": "unit", "header": "单位", "required": True, "type": "文本", "rule": "统一用m2、m3、m、t、个、项、工日等，不同单位不要混算。", "options": "", "source": "单位、计量单位"},
    {"key": "price", "header": "单价", "required": True, "type": "数值", "rule": "填写可比较单价；区间价可写3.22-16.72，系统按均值入库并保留原值说明。", "options": "", "source": "单价、控制价不含税、合同单价、结算单价"},
    {"key": "taxRate", "header": "税率", "required": False, "type": "数值/百分比", "rule": "可写9%、9或0.09，系统统一为小数。", "options": "", "source": "税率、税率(%)"},
    {"key": "quantity", "header": "工程量", "required": False, "type": "数值", "rule": "可空；用于反算金额或校核样本规模。", "options": "", "source": "数量、工程量、变更后数量"},
    {"key": "amount", "header": "金额", "required": False, "type": "数值", "rule": "可空；有单价和工程量时系统可反算。", "options": "", "source": "金额、合价、变更后金额"},
    {"key": "region", "header": "地区", "required": False, "type": "文本", "rule": "建议到城市或区域，如山东省烟台市、山东-青岛。", "options": "", "source": "地区、所在地域、城市"},
    {"key": "projectType", "header": "项目类型", "required": False, "type": "文本", "rule": "如多层住宅、小高层住宅、高层住宅、公用建筑、厂房仓库。", "options": "", "source": "工程类型、项目类型、工作表名称"},
    {"key": "project", "header": "项目", "required": False, "type": "文本", "rule": "项目名称或项目部名称，用于项目维度筛选。", "options": "", "source": "项目名称、项目部、使用项目"},
    {"key": "supplier", "header": "供应商/分包商", "required": False, "type": "文本", "rule": "劳务班组、分包商或专业分包单位名称。", "options": "", "source": "分包商/班组、分包商名称、合同相对方"},
    {"key": "contractNo", "header": "合同/结算编号", "required": False, "type": "文本", "rule": "用于追溯来源，尽量保留原始编号。", "options": "", "source": "合同编号、结算编号、参考单据编号"},
    {"key": "priceDate", "header": "价格日期", "required": False, "type": "日期", "rule": "统一为YYYY-MM-DD，可取定标日期、合同日期、结算日期或开工日期。", "options": "", "source": "定标日期、合同日期、结算日期、开工日期"},
    {"key": "source", "header": "来源", "required": False, "type": "文本", "rule": "标识数据来源，如PM劳务价格抓取、手工控制价库、合同清单。", "options": "", "source": "系统自动补充或Excel来源列"},
    {"key": "remark", "header": "备注", "required": False, "type": "文本", "rule": "付款条件、质保金、价格口径、特殊做法、风险等级等全部放这里追溯。", "options": "", "source": "备注、付款条件、质保金、施工内容"},
]

PRICE_LIBRARY_TEMPLATE_HEADERS = [field["header"] for field in PRICE_LIBRARY_STANDARD_FIELDS]

PRICE_LIBRARY_REQUIRED_FIELDS = {
    field["header"] for field in PRICE_LIBRARY_STANDARD_FIELDS if field["required"]
}

PRICE_LIBRARY_TEMPLATE_GUIDE = [
    (field["header"], f"{'必填' if field['required'] else '可选'}；{field['rule']}")
    for field in PRICE_LIBRARY_STANDARD_FIELDS
]


def price_library_connect() -> sqlite3.Connection:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(PRICE_LIBRARY_DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_price_library_db() -> None:
    with price_library_connect() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS price_library (
                id TEXT PRIMARY KEY,
                category TEXT NOT NULL,
                company TEXT,
                major TEXT,
                discipline TEXT,
                trade TEXT,
                process TEXT,
                item TEXT NOT NULL,
                spec TEXT,
                unit TEXT NOT NULL,
                price REAL NOT NULL DEFAULT 0,
                tax_rate REAL NOT NULL DEFAULT 0,
                quantity REAL NOT NULL DEFAULT 0,
                amount REAL NOT NULL DEFAULT 0,
                region TEXT,
                project_type TEXT,
                project TEXT,
                supplier TEXT,
                contract_no TEXT,
                price_date TEXT,
                source TEXT,
                remark TEXT,
                source_file TEXT,
                created_at TEXT,
                updated_at TEXT
            )
            """
        )
        existing_columns = {row["name"] for row in conn.execute("PRAGMA table_info(price_library)")}
        for column_name in ("discipline", "trade", "process"):
            if column_name not in existing_columns:
                conn.execute(f"ALTER TABLE price_library ADD COLUMN {column_name} TEXT")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_price_library_item ON price_library(item, spec, unit)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_price_library_filters ON price_library(category, company, region, price_date)")
        count = conn.execute("SELECT COUNT(*) FROM price_library").fetchone()[0]
    if count == 0:
        migrate_price_library_from_json()


def migrate_price_library_from_json() -> None:
    if not DB_PATH.exists():
        return
    try:
        with DB_PATH.open("r", encoding="utf-8-sig") as f:
            legacy_rows = json.load(f).get("priceLibraryRows", [])
    except (OSError, json.JSONDecodeError, AttributeError):
        return
    if not legacy_rows:
        return
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with price_library_connect() as conn:
        for legacy in legacy_rows:
            if not isinstance(legacy, dict):
                continue
            row = normalize_price_library_row(legacy)
            if not row.get("item") or not row.get("unit") or not row.get("price"):
                continue
            row["id"] = clean_text(legacy.get("id")) or uuid.uuid4().hex
            row["createdAt"] = clean_text(legacy.get("createdAt")) or now
            row["updatedAt"] = clean_text(legacy.get("updatedAt")) or row["createdAt"]
            insert_price_library_row(conn, row)


PRICE_LIBRARY_CLASSIFICATION_RULES = [
    {"discipline": "建筑", "trade": "木工", "process": "模板支设", "keywords": ["模板", "支模", "拆模", "木方", "对拉螺栓"]},
    {"discipline": "建筑", "trade": "钢筋工", "process": "钢筋制作安装", "keywords": ["钢筋", "绑扎", "套筒", "直螺纹", "植筋"]},
    {"discipline": "建筑", "trade": "混凝土工", "process": "混凝土浇筑", "keywords": ["混凝土", "砼", "浇筑", "振捣", "泵送"]},
    {"discipline": "建筑", "trade": "瓦工", "process": "砌筑", "keywords": ["砌筑", "砌块", "砖墙", "加气块", "二次结构"]},
    {"discipline": "建筑", "trade": "瓦工", "process": "抹灰找平", "keywords": ["抹灰", "找平", "砂浆", "地面", "楼地面"]},
    {"discipline": "建筑", "trade": "油工", "process": "腻子涂料", "keywords": ["腻子", "涂料", "乳胶漆", "油漆", "喷涂"]},
    {"discipline": "建筑", "trade": "防水工", "process": "防水施工", "keywords": ["防水", "卷材", "涂膜", "聚氨酯", "SBS"]},
    {"discipline": "建筑", "trade": "架子工", "process": "脚手架搭拆", "keywords": ["脚手架", "架体", "盘扣", "悬挑架", "外架", "满堂架"]},
    {"discipline": "建筑", "trade": "保温工", "process": "保温施工", "keywords": ["保温", "岩棉", "挤塑板", "聚苯板", "真石漆基层"]},
    {"discipline": "建筑", "trade": "土方工", "process": "土方开挖回填", "keywords": ["土方", "开挖", "回填", "外运", "场平"]},
    {"discipline": "建筑", "trade": "门窗幕墙工", "process": "门窗幕墙安装", "keywords": ["门窗", "幕墙", "玻璃", "铝合金", "百叶"]},
    {"discipline": "安装", "trade": "电工", "process": "电气配管配线", "keywords": ["电气", "配管", "桥架", "电缆", "电线", "配电箱", "灯具", "开关", "插座"]},
    {"discipline": "安装", "trade": "水工", "process": "给排水管道安装", "keywords": ["给水", "排水", "雨水", "污水", "PPR", "PVC", "管道", "阀门", "洁具"]},
    {"discipline": "安装", "trade": "暖通工", "process": "暖通空调安装", "keywords": ["暖通", "通风", "风管", "空调", "采暖", "散热器", "新风"]},
    {"discipline": "安装", "trade": "消防工", "process": "消防系统安装", "keywords": ["消防", "喷淋", "消火栓", "报警", "烟感", "防排烟"]},
    {"discipline": "安装", "trade": "弱电工", "process": "弱电智能化安装", "keywords": ["弱电", "智能化", "监控", "网络", "综合布线", "门禁", "广播"]},
]


def classify_price_library_work(row: dict[str, Any]) -> dict[str, str]:
    text = " ".join(clean_text(row.get(field)) for field in ("item", "spec", "remark", "major", "process"))
    for rule in PRICE_LIBRARY_CLASSIFICATION_RULES:
        if any(keyword.lower() in text.lower() for keyword in rule["keywords"]):
            return {
                "discipline": rule["discipline"],
                "trade": rule["trade"],
                "process": rule["process"],
            }
    return {"discipline": "", "trade": "", "process": ""}


def normalize_price_library_category(value: Any) -> str:
    text = clean_text(value)
    if "劳务" in text:
        return "劳务分包"
    if "专业" in text or "分供" in text or "材料" in text or "供应" in text:
        return "专业分包"
    if text == "分包" or not text:
        return "劳务分包"
    return text


def normalize_price_library_row(row: dict[str, Any]) -> dict[str, Any]:
    price = to_number(row.get("price"))
    quantity = to_number(row.get("quantity"))
    amount = to_number(row.get("amount"))
    if not amount and price and quantity:
        amount = price * quantity
    if not price and amount and quantity:
        price = amount / quantity
    tax_rate = to_number(row.get("taxRate"))
    if tax_rate > 1:
        tax_rate = tax_rate / 100
    price_date = parse_date(row.get("priceDate")) or clean_text(row.get("priceDate"))
    classified = classify_price_library_work(row)
    region = clean_text(row.get("region")) or infer_price_library_region(row)
    project_type = clean_text(row.get("projectType")) or infer_price_library_project_type(row)
    return {
        "category": normalize_price_library_category(row.get("category")),
        "company": standard_company_name(row.get("company")) or clean_text(row.get("company")),
        "major": clean_text(row.get("major")),
        "discipline": clean_text(row.get("discipline")) or classified.get("discipline", ""),
        "trade": clean_text(row.get("trade")) or classified.get("trade", ""),
        "process": clean_text(row.get("process")) or classified.get("process", ""),
        "item": clean_text(row.get("item")),
        "spec": clean_text(row.get("spec")),
        "unit": clean_text(row.get("unit")),
        "price": round(price, 4) if price else 0,
        "taxRate": round(tax_rate, 4) if tax_rate else 0,
        "quantity": round(quantity, 4) if quantity else 0,
        "amount": round(amount, 2) if amount else 0,
        "region": region,
        "projectType": project_type,
        "project": clean_text(row.get("project")),
        "supplier": clean_text(row.get("supplier")),
        "contractNo": clean_text(row.get("contractNo")),
        "priceDate": price_date,
        "source": clean_text(row.get("source")),
        "remark": clean_text(row.get("remark")),
    }


def price_library_project_context_key(row: dict[str, Any]) -> str:
    project = clean_text(row.get("project"))
    if not project:
        return ""
    return re.sub(r"[\s\u3000()（）【】\[\]_-]+", "", project).lower()


def enrich_price_library_project_fields(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    project_types: dict[str, Counter[str]] = {}
    regions: dict[str, Counter[str]] = {}
    for row in rows:
        key = price_library_project_context_key(row)
        if not key:
            continue
        project_type = clean_text(row.get("projectType"))
        region = clean_text(row.get("region"))
        if project_type:
            project_types.setdefault(key, Counter()).update([project_type])
        if region:
            regions.setdefault(key, Counter()).update([region])

    project_type_fill: dict[str, str] = {}
    for key, counts in project_types.items():
        total = sum(counts.values())
        if not total:
            continue
        value, count = counts.most_common(1)[0]
        if len(counts) == 1 or count / total >= 0.7:
            project_type_fill[key] = value

    region_fill: dict[str, str] = {
        key: counts.most_common(1)[0][0]
        for key, counts in regions.items()
        if counts
    }

    enriched: list[dict[str, Any]] = []
    for row in rows:
        key = price_library_project_context_key(row)
        if not key:
            enriched.append(row)
            continue
        next_row = dict(row)
        if not clean_text(next_row.get("projectType")) and project_type_fill.get(key):
            next_row["projectType"] = project_type_fill[key]
        if not clean_text(next_row.get("region")) and region_fill.get(key):
            next_row["region"] = region_fill[key]
        enriched.append(next_row)
    return enriched


def price_library_sql_row_to_dict(row: sqlite3.Row) -> dict[str, Any]:
    return {
        "id": row["id"],
        "category": row["category"],
        "company": row["company"],
        "major": row["major"],
        "discipline": row["discipline"],
        "trade": row["trade"],
        "process": row["process"],
        "item": row["item"],
        "spec": row["spec"],
        "unit": row["unit"],
        "price": row["price"],
        "taxRate": row["tax_rate"],
        "quantity": row["quantity"],
        "amount": row["amount"],
        "region": row["region"],
        "projectType": row["project_type"],
        "project": row["project"],
        "supplier": row["supplier"],
        "contractNo": row["contract_no"],
        "priceDate": row["price_date"],
        "source": row["source"],
        "remark": row["remark"],
        "_sourceFile": row["source_file"],
        "createdAt": row["created_at"],
        "updatedAt": row["updated_at"],
    }


def price_library_duplicate_key(row: dict[str, Any]) -> tuple[str, ...]:
    return (
        clean_text(row.get("category")).lower(),
        clean_text(row.get("company")).lower(),
        clean_text(row.get("supplier")).lower(),
        clean_text(row.get("contractNo")).lower(),
        clean_text(row.get("item")).lower(),
        clean_text(row.get("spec")).lower(),
        clean_text(row.get("unit")).lower(),
        str(round(to_number(row.get("price")), 4)),
    )


def find_existing_price_library_row(conn: sqlite3.Connection, row_id: str, normalized: dict[str, Any]) -> dict[str, Any] | None:
    matches = find_existing_price_library_rows(conn, row_id, normalized)
    return matches[0] if matches else None


def find_existing_price_library_rows(conn: sqlite3.Connection, row_id: str, normalized: dict[str, Any]) -> list[dict[str, Any]]:
    matches: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    if row_id:
        existing = conn.execute("SELECT * FROM price_library WHERE id = ?", (row_id,)).fetchone()
        if existing:
            existing_dict = price_library_sql_row_to_dict(existing)
            matches.append(existing_dict)
            seen_ids.add(existing_dict["id"])
    candidates = [
        price_library_sql_row_to_dict(candidate)
        for candidate in conn.execute(
            """
            SELECT * FROM price_library
            WHERE item = ? AND unit = ? AND ABS(price - ?) < 0.0001
            """,
            (normalized.get("item"), normalized.get("unit"), normalized.get("price") or 0),
        )
    ]
    target_key = price_library_duplicate_key(normalized)
    for candidate in candidates:
        if candidate["id"] in seen_ids:
            continue
        if price_library_duplicate_key(candidate) == target_key:
            matches.append(candidate)
            seen_ids.add(candidate["id"])
    return matches


def merge_unique_text(existing: Any, incoming: Any) -> str:
    existing_parts = [part.strip() for part in clean_text(existing).split("；") if part.strip()]
    incoming_parts = [part.strip() for part in clean_text(incoming).split("；") if part.strip()]
    seen = {part.lower() for part in existing_parts}
    merged = list(existing_parts)
    for part in incoming_parts:
        if part.lower() not in seen:
            merged.append(part)
            seen.add(part.lower())
    return "；".join(merged)


def merge_price_library_rows(existing: dict[str, Any], incoming: dict[str, Any]) -> dict[str, Any]:
    merged = dict(existing)
    text_fields = [
        "category", "company", "major", "discipline", "trade", "process", "item", "spec", "unit",
        "region", "projectType", "project", "supplier", "contractNo", "priceDate",
    ]
    for field in text_fields:
        if not clean_text(merged.get(field)) and clean_text(incoming.get(field)):
            merged[field] = incoming.get(field)
    for field in ("price", "taxRate", "quantity", "amount"):
        if not to_number(merged.get(field)) and to_number(incoming.get(field)):
            merged[field] = incoming.get(field)
    merged["source"] = merge_unique_text(merged.get("source"), incoming.get("source"))
    merged["remark"] = merge_unique_text(merged.get("remark"), incoming.get("remark"))
    merged["_sourceFile"] = merge_unique_text(merged.get("_sourceFile"), incoming.get("_sourceFile"))
    merged["id"] = existing.get("id") or incoming.get("id")
    merged["createdAt"] = existing.get("createdAt") or incoming.get("createdAt")
    merged["updatedAt"] = incoming.get("updatedAt") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    return merged


def insert_price_library_row(conn: sqlite3.Connection, row: dict[str, Any]) -> str:
    normalized = normalize_price_library_row(row)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    row_id = clean_text(row.get("id")) or uuid.uuid4().hex
    existing_rows = find_existing_price_library_rows(conn, row_id, normalized)
    if existing_rows:
        keeper = existing_rows[0]
        merged = merge_price_library_rows(keeper, {**normalized, **row, "id": row_id})
        for duplicate in existing_rows[1:]:
            merged = merge_price_library_rows(merged, duplicate)
        update_price_library_row(conn, keeper["id"], merged)
        for duplicate in existing_rows[1:]:
            conn.execute("DELETE FROM price_library WHERE id = ?", (duplicate["id"],))
        return "updated"
    conn.execute(
        """
        INSERT OR REPLACE INTO price_library (
            id, category, company, major, discipline, trade, process, item, spec, unit, price, tax_rate,
            quantity, amount, region, project_type, project, supplier,
            contract_no, price_date, source, remark, source_file, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            row_id,
            normalized.get("category"),
            normalized.get("company"),
            normalized.get("major"),
            normalized.get("discipline"),
            normalized.get("trade"),
            normalized.get("process"),
            normalized.get("item"),
            normalized.get("spec"),
            normalized.get("unit"),
            normalized.get("price"),
            normalized.get("taxRate"),
            normalized.get("quantity"),
            normalized.get("amount"),
            normalized.get("region"),
            normalized.get("projectType"),
            normalized.get("project"),
            normalized.get("supplier"),
            normalized.get("contractNo"),
            normalized.get("priceDate"),
            normalized.get("source"),
            normalized.get("remark"),
            clean_text(row.get("_sourceFile")),
            clean_text(row.get("createdAt")) or now,
            clean_text(row.get("updatedAt")) or clean_text(row.get("createdAt")) or now,
        ),
    )
    return "inserted"


def can_edit_price_library(_headers: Any) -> bool:
    # 权限后续在这里接入：可按登录用户、角色、审批状态或请求头令牌判断。
    return True


def update_price_library_row(conn: sqlite3.Connection, row_id: str, row: dict[str, Any]) -> dict[str, Any] | None:
    existing = conn.execute("SELECT * FROM price_library WHERE id = ?", (row_id,)).fetchone()
    if not existing:
        return None
    existing_dict = price_library_sql_row_to_dict(existing)
    normalized = normalize_price_library_row({**existing_dict, **row})
    updated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn.execute(
        """
        UPDATE price_library
        SET category = ?, company = ?, major = ?, discipline = ?, trade = ?, process = ?,
            item = ?, spec = ?, unit = ?, price = ?, tax_rate = ?, quantity = ?, amount = ?, region = ?,
            project_type = ?, project = ?, supplier = ?, contract_no = ?,
            price_date = ?, source = ?, remark = ?, source_file = ?, updated_at = ?
        WHERE id = ?
        """,
        (
            normalized.get("category"),
            normalized.get("company"),
            normalized.get("major"),
            normalized.get("discipline"),
            normalized.get("trade"),
            normalized.get("process"),
            normalized.get("item"),
            normalized.get("spec"),
            normalized.get("unit"),
            normalized.get("price"),
            normalized.get("taxRate"),
            normalized.get("quantity"),
            normalized.get("amount"),
            normalized.get("region"),
            normalized.get("projectType"),
            normalized.get("project"),
            normalized.get("supplier"),
            normalized.get("contractNo"),
            normalized.get("priceDate"),
            normalized.get("source"),
            normalized.get("remark"),
            clean_text(row.get("_sourceFile")) or clean_text(existing_dict.get("_sourceFile")),
            updated_at,
            row_id,
        ),
    )
    saved = conn.execute("SELECT * FROM price_library WHERE id = ?", (row_id,)).fetchone()
    return price_library_sql_row_to_dict(saved) if saved else None


def price_library_key(row: dict[str, Any]) -> str:
    return "|".join([
        clean_text(row.get("category")).lower(),
        clean_text(row.get("discipline")).lower(),
        clean_text(row.get("trade")).lower(),
        clean_text(row.get("process")).lower(),
        clean_text(row.get("item")).lower(),
        clean_text(row.get("unit")).lower(),
        clean_text(row.get("region")).lower(),
    ])


def is_price_library_process_price(row: dict[str, Any]) -> bool:
    item = clean_text(row.get("item"))
    unit = clean_text(row.get("unit"))
    price = to_number(row.get("price"))
    if not item or not unit or price <= 0:
        return False
    if price > 1000000:
        return False
    contract_markers = (
        "\u5408\u540c", "\u52b3\u52a1\u5206\u5305\u5408\u540c", "\u4e13\u4e1a\u5206\u5305\u5408\u540c",
        "\u5206\u5305\u5408\u540c", "\u91c7\u8d2d\u5408\u540c",
    )
    if any(marker in item for marker in contract_markers):
        return False
    return True


def price_library_match(row: dict[str, Any], params: dict[str, list[str]]) -> bool:
    keyword = clean_text(params.get("keyword", [""])[0]).lower()
    category = clean_text(params.get("category", [""])[0])
    company = clean_text(params.get("company", [""])[0]).lower()
    major = clean_text(params.get("major", [""])[0]).lower()
    discipline = clean_text(params.get("discipline", [""])[0]).lower()
    project_type = clean_text(params.get("projectType", [""])[0]).lower()
    trade = clean_text(params.get("trade", [""])[0]).lower()
    process = clean_text(params.get("process", [""])[0]).lower()
    region = clean_text(params.get("region", [""])[0]).lower()
    project = clean_text(params.get("project", [""])[0]).lower()
    start_date = clean_text(params.get("startDate", [""])[0])
    end_date = clean_text(params.get("endDate", [""])[0])
    if category and row.get("category") != category:
        return False
    if company and company not in clean_text(row.get("company")).lower():
        return False
    if major:
        row_major = clean_text(row.get("major")).lower()
        row_discipline = clean_text(row.get("discipline")).lower()
        if major in {"土建", "建筑"}:
            if "土建" not in row_major and "建筑" not in row_discipline:
                return False
        elif major == "安装":
            if "安装" not in row_major and "安装" not in row_discipline:
                return False
        elif major not in row_major and major not in row_discipline:
            return False
    if discipline and discipline not in clean_text(row.get("discipline")).lower():
        return False
    if project_type and project_type not in clean_text(row.get("projectType")).lower():
        return False
    if trade and trade not in clean_text(row.get("trade")).lower():
        return False
    if process and process not in clean_text(row.get("process")).lower():
        return False
    if region and region not in clean_text(row.get("region")).lower():
        return False
    if project and project not in clean_text(row.get("project")).lower():
        return False
    if keyword:
        haystack = " ".join(clean_text(row.get(field)).lower() for field in ("item", "spec", "unit", "project", "supplier", "remark", "contractNo", "major", "discipline", "trade", "process"))
        if keyword not in haystack:
            return False
    row_date = clean_text(row.get("priceDate"))
    if start_date and row_date and row_date < start_date:
        return False
    if end_date and row_date and row_date > end_date:
        return False
    return True


def load_price_library_rows(params: dict[str, list[str]] | None = None) -> list[dict[str, Any]]:
    init_price_library_db()
    with price_library_connect() as conn:
        rows = [price_library_sql_row_to_dict(row) for row in conn.execute("SELECT * FROM price_library")]
    normalized = [normalize_price_library_row(row) | {
        "id": row.get("id", uuid.uuid4().hex),
        "createdAt": row.get("createdAt", ""),
        "updatedAt": row.get("updatedAt", ""),
        "_sourceFile": row.get("_sourceFile", ""),
    } for row in rows]
    normalized = enrich_price_library_project_fields(normalized)
    if not params:
        return normalized
    return [row for row in normalized if price_library_match(row, params)]


def summarize_price_library(rows: list[dict[str, Any]]) -> dict[str, Any]:
    prices = [row.get("price", 0) for row in rows if row.get("price", 0) > 0]
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        if is_price_library_process_price(row):
            grouped.setdefault(price_library_key(row), []).append(row)
    spreads = []
    for group_rows in grouped.values():
        group_prices = [row["price"] for row in group_rows if row.get("price")]
        if len(group_prices) < 2:
            continue
        low = min(group_rows, key=lambda row: row.get("price") or float("inf"))
        high = max(group_rows, key=lambda row: row.get("price") or 0)
        low_price = low.get("price") or 0
        high_price = high.get("price") or 0
        spreads.append({
            "category": group_rows[0].get("category"),
            "item": group_rows[0].get("item"),
            "unit": group_rows[0].get("unit"),
            "region": group_rows[0].get("region"),
            "lowPrice": round(low_price, 4),
            "highPrice": round(high_price, 4),
            "spreadRate": round((high_price - low_price) / low_price, 4) if low_price else 0,
            "lowSupplier": low.get("supplier"),
            "highSupplier": high.get("supplier"),
            "records": len(group_rows),
        })
    spreads.sort(key=lambda row: (row.get("spreadRate") or 0, row.get("records") or 0), reverse=True)
    monthly: dict[str, list[float]] = {}
    by_company: dict[str, list[float]] = {}
    by_item: dict[str, list[dict[str, Any]]] = {}
    by_classification: dict[tuple[str, str, str, str], list[float]] = {}
    for row in rows:
        price = row.get("price", 0)
        if price <= 0:
            continue
        date_text = clean_text(row.get("priceDate"))
        month = date_text[:7] if re.match(r"^\d{4}-\d{2}", date_text) else "未填日期"
        monthly.setdefault(month, []).append(price)
        by_company.setdefault(clean_text(row.get("company")) or "未填分公司", []).append(price)
        by_item.setdefault(clean_text(row.get("item")) or "未命名施工内容", []).append(row)
        class_key = (
            clean_text(row.get("category")) or "未分类类型",
            clean_text(row.get("discipline")) or "未分类专业",
            clean_text(row.get("trade")) or "未分类工种",
            clean_text(row.get("process")) or "未分类工序",
        )
        by_classification.setdefault(class_key, []).append(price)
    trend_rows = [
        {
            "month": month,
            "avgPrice": round(sum(values) / len(values), 4),
            "minPrice": round(min(values), 4),
            "maxPrice": round(max(values), 4),
            "records": len(values),
        }
        for month, values in sorted(monthly.items())
    ]
    company_rows = sorted([
        {
            "company": company,
            "avgPrice": round(sum(values) / len(values), 4),
            "minPrice": round(min(values), 4),
            "maxPrice": round(max(values), 4),
            "records": len(values),
        }
        for company, values in by_company.items()
    ], key=lambda row: row["records"], reverse=True)[:20]
    item_rows = []
    for item, item_group in by_item.items():
        item_prices = [row.get("price", 0) for row in item_group if row.get("price", 0) > 0]
        if not item_prices:
            continue
        low_price = min(item_prices)
        high_price = max(item_prices)
        item_rows.append({
            "item": item,
            "avgPrice": round(sum(item_prices) / len(item_prices), 4),
            "minPrice": round(low_price, 4),
            "maxPrice": round(high_price, 4),
            "spreadRate": round((high_price - low_price) / low_price, 4) if low_price else 0,
            "records": len(item_prices),
        })
    item_rows.sort(key=lambda row: (row["records"], row["spreadRate"]), reverse=True)
    classification_rows = sorted([
        {
            "category": key[0],
            "discipline": key[1],
            "trade": key[2],
            "process": key[3],
            "avgPrice": round(sum(values) / len(values), 4),
            "minPrice": round(min(values), 4),
            "maxPrice": round(max(values), 4),
            "records": len(values),
        }
        for key, values in by_classification.items()
    ], key=lambda row: row["records"], reverse=True)
    return {
        "total": len(rows),
        "laborSubcontract": sum(1 for row in rows if row.get("category") == "劳务分包"),
        "professionalSubcontract": sum(1 for row in rows if row.get("category") == "专业分包"),
        "avgPrice": round(sum(prices) / len(prices), 4) if prices else 0,
        "minPrice": round(min(prices), 4) if prices else 0,
        "maxPrice": round(max(prices), 4) if prices else 0,
        "spreadRows": spreads[:20],
        "trendRows": trend_rows[-24:],
        "companyRows": company_rows,
        "itemRows": item_rows[:20],
        "classificationRows": classification_rows[:50],
    }


def price_library_payload(params: dict[str, list[str]] | None = None) -> dict[str, Any]:
    rows = load_price_library_rows(params)
    rows.sort(key=lambda row: (row.get("priceDate") or "", row.get("createdAt") or ""), reverse=True)
    return {"rows": rows, "summary": summarize_price_library(rows)}


def price_library_header_score(row_values: tuple[Any, ...]) -> int:
    headers = [clean_text(value) for value in row_values if clean_text(value)]
    if not headers:
        return 0
    score = 0
    for aliases in PRICE_LIBRARY_ALIASES.values():
        if any(header == alias or alias in header for header in headers for alias in aliases):
            score += 1
    return score


def price_library_major_from_text(*values: Any) -> tuple[str, str]:
    text = " ".join(clean_text(value) for value in values)
    if "安装" in text:
        return "安装", "安装"
    if "土建" in text:
        return "土建", "建筑"
    if "建筑" in text:
        return "土建", "建筑"
    return "", ""


def read_price_library_file(file_path: Path) -> list[dict[str, Any]]:
    suffix = file_path.suffix.lower()
    if suffix == ".json":
        data = read_json_file(file_path, [])
        source_rows = data.get("rows", []) if isinstance(data, dict) else data
        return [normalize_price_library_row(row) for row in source_rows if isinstance(row, dict)]
    workbook = load_workbook(file_path, data_only=True)
    rows: list[dict[str, Any]] = []
    for ws in workbook.worksheets:
        values = list(ws.iter_rows(values_only=True))
        if not values:
            continue
        header_index = 0
        for index, row_values in enumerate(values[:12]):
            if price_library_header_score(row_values) >= 3:
                header_index = index
                break
        sheet_title = clean_text(ws.title)
        title_text = clean_text(values[0][0]) if values and values[0] else ""
        source_major, source_discipline = price_library_major_from_text(file_path.name, sheet_title, title_text)
        headers = [clean_text(value) for value in values[header_index]]
        for raw_values in values[header_index + 1:]:
            raw = {headers[index]: raw_values[index] if index < len(raw_values) else "" for index in range(len(headers)) if headers[index]}
            mapped = map_alias_row(raw, PRICE_LIBRARY_ALIASES)
            mapped = normalize_manual_control_price_row(raw, mapped, ws.title, source_major, source_discipline)
            mapped["_sourceSheet"] = sheet_title
            if mapped.get("_skip"):
                continue
            normalized = normalize_price_library_row(mapped)
            if normalized.get("item") and normalized.get("price"):
                normalized["_sourceSheet"] = sheet_title
                rows.append(normalized)
    return rows


def normalize_manual_control_price_row(raw: dict[str, Any], mapped: dict[str, Any], sheet_name: str = "", source_major: str = "", source_discipline: str = "") -> dict[str, Any]:
    item = clean_text(value_from_any(raw, ["分包项名称"]))
    control_price = value_from_any(raw, ["实际采购价", "采购价", "控制价\n不含税", "控制价不含税", "控制价"])
    if not item and not clean_text(control_price):
        return mapped
    if item.endswith("平均值") or "平均值" in item:
        return {**mapped, "_skip": True}
    construction_content = clean_text(value_from_any(raw, ["施工内容"]))
    position = clean_text(value_from_any(raw, ["施工部位"]))
    height = clean_text(value_from_any(raw, ["层高"]))
    payment_period = clean_text(value_from_any(raw, ["付款周期"]))
    payment_ratio = clean_text(value_from_any(raw, ["付款比例"]))
    payment_method = clean_text(value_from_any(raw, ["付款方式"]))
    risk_level = clean_text(value_from_any(raw, ["房源风险等级"]))
    payment_condition = clean_text(value_from_any(raw, ["付款条件", "付款条件总结", "付款条件原文"]))
    original_price = clean_text(control_price)
    note = clean_text(value_from_any(raw, ["备注"]))
    spec_parts = [
        f"层高：{height}" if height else "",
        f"施工部位：{position}" if position else "",
    ]
    remark_parts = [
        "手工控制价库",
        f"控制价口径：不含税，原值{original_price}" if original_price else "控制价口径：不含税",
        f"付款周期：{payment_period}" if payment_period else "",
        f"付款比例：{payment_ratio}" if payment_ratio else "",
        f"付款方式：{payment_method}" if payment_method else "",
        f"付款条件：{payment_condition}" if payment_condition else "",
        f"房源风险等级：{risk_level}" if risk_level else "",
        f"施工内容：{construction_content}" if construction_content else "",
        note,
    ]
    return {
        **mapped,
        "category": value_from_any(raw, ["类型库\n（下拉选择）", "类型库"]) or mapped.get("category"),
        "company": value_from_any(raw, ["基层单位"]) or mapped.get("company"),
        "major": mapped.get("major") or source_major,
        "discipline": mapped.get("discipline") or source_discipline,
        "item": item or mapped.get("item"),
        "spec": "；".join(part for part in spec_parts if part) or mapped.get("spec"),
        "unit": value_from_any(raw, ["计量单位"]) or mapped.get("unit"),
        "price": control_price or mapped.get("price"),
        "region": value_from_any(raw, ["地区"]) or mapped.get("region"),
        "projectType": value_from_any(raw, ["项目业态", "工程类型\n（下拉选择）", "工程类型"]) or sheet_name or mapped.get("projectType"),
        "project": value_from_any(raw, ["项目部/项目名称", "项目部", "项目基本信息匹配名称", "项目名称"]) or mapped.get("project"),
        "priceDate": value_from_any(raw, ["定标日期"]) or mapped.get("priceDate"),
        "source": f"手工控制价库/{sheet_name}" if sheet_name else "手工控制价库",
        "remark": "；".join(part for part in remark_parts if part),
    }


def stable_uploaded_price_library_id(filename: str, row: dict[str, Any]) -> str:
    parts = [
        clean_text(row.get("category")),
        clean_text(row.get("company")),
        clean_text(row.get("projectType")),
        clean_text(row.get("project")),
        clean_text(row.get("supplier")),
        clean_text(row.get("contractNo")),
        clean_text(row.get("item")),
        clean_text(row.get("spec")),
        clean_text(row.get("unit")),
        clean_text(row.get("price")),
        clean_text(row.get("priceDate")),
        clean_text(row.get("_sourceSheet")),
    ]
    return uuid.uuid5(uuid.NAMESPACE_URL, "|".join(parts)).hex


def write_price_library_template_workbook(output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "导入模板"
    ws.append(PRICE_LIBRARY_TEMPLATE_HEADERS)
    ws.append([
        "劳务分包",
        "三公司",
        "土建",
        "建筑",
        "木工",
        "模板支设",
        "梁板模板支设",
        "含支模、拆模、清理",
        "m2",
        58,
        "3%",
        1200,
        69600,
        "烟台",
        "工业厂房",
        "示例项目",
        "示例劳务班组",
        "LWFBHT-2026-001",
        "2026-07-01",
        "合同清单",
        "付款、质保、赶工等影响价格的条件写在这里",
    ])
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="B00020")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if cell.value in PRICE_LIBRARY_REQUIRED_FIELDS:
            cell.comment = Comment("必填字段", "系统")
    for index, header in enumerate(PRICE_LIBRARY_TEMPLATE_HEADERS, start=1):
        width = max(12, min(24, len(header) * 2 + 8))
        ws.column_dimensions[get_column_letter(index)].width = width

    guide = wb.create_sheet("字段说明")
    guide.append(["字段", "是否必填", "类型", "填写规则", "允许值", "抓取/识别来源"])
    for field in PRICE_LIBRARY_STANDARD_FIELDS:
        guide.append([
            field["header"],
            "必填" if field["required"] else "可选",
            field["type"],
            field["rule"],
            field["options"],
            field["source"],
        ])
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 12
    guide.column_dimensions["C"].width = 14
    guide.column_dimensions["D"].width = 60
    guide.column_dimensions["E"].width = 28
    guide.column_dimensions["F"].width = 38
    for cell in guide[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="263244")

    mapping = wb.create_sheet("抓取映射")
    mapping.append(["标准字段", "PM劳务抓取字段", "手工控制价库字段", "处理规则"])
    mapping_rows = [
        ("类型", "固定为劳务分包", "类型库", "统一为劳务分包/专业分包"),
        ("分公司", "公司、编制机构、DeptName", "基层单位", "统一公司简称"),
        ("专业类别", "成本科目、价格类型", "固定为安装或原专业类别", "用于专业大类筛选"),
        ("专业", "系统自动识别", "固定为安装", "建筑/安装"),
        ("工种", "系统自动识别", "系统自动识别", "可人工修正"),
        ("工序", "系统自动识别", "系统自动识别", "可人工修正"),
        ("分项名称", "工序/费用项名称、*费用项名称", "分包项名称", "平均值、合计行不入库"),
        ("规格/做法", "项目特征、部位、工程量计算规则", "层高、施工部位", "影响价格的条件放这里"),
        ("单位", "单位", "计量单位", "作为价格可比口径"),
        ("单价", "变更后单价、单价、含税指导价", "控制价不含税", "区间价取均值，原值写备注"),
        ("税率", "税率(%)、TaxRatio", "", "9/9%/0.09统一为0.09"),
        ("工程量", "变更后数量、数量", "", "可空"),
        ("金额", "变更后金额、金额", "", "可空"),
        ("地区", "所在地域、地区", "地区", "建议到城市"),
        ("项目类型", "", "工程类型、工作表名称", "用于工程类型筛选"),
        ("项目", "项目名称、使用项目", "项目部", "用于项目筛选"),
        ("供应商/分包商", "分包商/班组、分包商名称", "", "用于供应商维度"),
        ("合同/结算编号", "合同编号、参考单据编号", "", "追溯依据"),
        ("价格日期", "开工日期、CreateTime、BizDate", "定标日期", "统一YYYY-MM-DD"),
        ("来源", "PM劳务价格抓取/文件名", "手工控制价库/工作表", "系统自动补充"),
        ("备注", "付款条件、质保金、合同备注", "付款周期、付款比例、付款方式、风险等级、施工内容、备注", "上下文和价格口径都放这里"),
    ]
    for row in mapping_rows:
        mapping.append(row)
    for col, width in zip(("A", "B", "C", "D"), (18, 42, 38, 50)):
        mapping.column_dimensions[col].width = width
    for cell in mapping[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="263244")
    wb.save(output_path)


def value_from_any(row: dict[str, Any], names: list[str]) -> Any:
    for name in names:
        value = row.get(name)
        if clean_text(value):
            return value
    return ""


def stable_price_library_import_id(prefix: str, row: dict[str, Any], source_name: str) -> str:
    parts = [
        prefix,
        source_name,
        clean_text(value_from_any(row, ["合同ID", "contractId"])),
        clean_text(value_from_any(row, ["预算ID", "budgetId"])),
        clean_text(value_from_any(row, ["清单编码", "编码", "numberer"])),
        clean_text(value_from_any(row, ["工序/费用项名称", "*费用项名称", "费用项名称"])),
        clean_text(value_from_any(row, ["单位"])),
    ]
    raw = "|".join(parts)
    return uuid.uuid5(uuid.NAMESPACE_URL, raw).hex


def map_pm_price_row(row: dict[str, Any], source_name: str, category: str, source_label: str, id_prefix: str) -> dict[str, Any]:
    item = value_from_any(row, ["工序/费用项名称", "*费用项名称", "费用项名称", "工作内容", "合同工作内容"])
    spec_parts = [
        clean_text(value_from_any(row, ["项目特征"])),
        clean_text(value_from_any(row, ["工程量计算规则"])),
        clean_text(value_from_any(row, ["部位名称", "部位全路径"])),
    ]
    remark_parts = [
        f"合同名称：{clean_text(value_from_any(row, ['合同名称', 'contractName', 'Name']))}",
        f"预算名称：{clean_text(value_from_any(row, ['预算名称', 'YSMC']))}",
        f"付款条件：{clean_text(value_from_any(row, ['付款条件/详细描述', '付款条件或详细描述']))}",
        f"质保金：{clean_text(value_from_any(row, ['质保金(%)', '质保金']))}",
        f"工期：{clean_text(value_from_any(row, ['工期']))}",
        clean_text(value_from_any(row, ["合同施工范围", "施工范围", "SGFW"])),
        clean_text(value_from_any(row, ["清单备注", "合同备注", "Remark"])),
    ]
    return {
        "id": stable_price_library_import_id(id_prefix, row, source_name),
        "category": category,
        "company": value_from_any(row, ["公司", "编制机构", "DeptName"]),
        "major": value_from_any(row, ["成本科目", "价格类型", "专业类别"]),
        "item": item,
        "spec": "；".join(part for part in spec_parts if part),
        "unit": value_from_any(row, ["单位"]),
        "price": value_from_any(row, ["变更后单价", "单价", "含税指导价", "变更前单价"]),
        "taxRate": value_from_any(row, ["税率(%)", "TaxRatio"]),
        "quantity": value_from_any(row, ["变更后数量", "数量", "变更前数量"]),
        "amount": value_from_any(row, ["变更后金额", "金额", "变更前金额"]),
        "region": value_from_any(row, ["所在地域", "地区", "区域"]),
        "projectType": "",
        "project": value_from_any(row, ["项目名称", "使用项目", "CGXM_Name", "总包企业名称"]),
        "supplier": value_from_any(row, ["分包商/班组", "分包商名称", "YF_Name"]),
        "contractNo": value_from_any(row, ["合同编号", "参考单据编号", "Code"]),
        "priceDate": value_from_any(row, ["开工日期", "CreateTime", "BizDate"]),
        "source": f"{source_label}/{source_name}",
        "remark": "；".join(part for part in remark_parts if part and not part.endswith("：")),
    }


def map_pm_labor_price_row(row: dict[str, Any], source_name: str) -> dict[str, Any]:
    return map_pm_price_row(row, source_name, "劳务分包", "PM劳务价格抓取", "pm-labor-price")


def map_pm_professional_price_row(row: dict[str, Any], source_name: str) -> dict[str, Any]:
    return map_pm_price_row(row, source_name, "专业分包", "PM专业分包价格抓取", "pm-professional-price")


def find_pm_labor_price_sources() -> list[Path]:
    if not PM_LABOR_PRICE_TASK_DIR.exists():
        return []
    candidates = [
        path for path in PM_LABOR_PRICE_TASK_DIR.glob("*价格库底表*.json")
        if "含合同条款摘录" in path.name or "逐页" in path.name or "全量" in path.name
    ]
    preferred: dict[str, Path] = {}
    priority_words = [("逐页", 3), ("全量", 2), ("当前页", 1)]
    for path in candidates:
        company = path.name.split("_", 1)[0]
        priority = next((score for word, score in priority_words if word in path.name), 0)
        current = preferred.get(company)
        if not current:
            preferred[company] = path
            continue
        current_priority = next((score for word, score in priority_words if word in current.name), 0)
        if (priority, path.stat().st_mtime) > (current_priority, current.stat().st_mtime):
            preferred[company] = path
    return sorted(preferred.values(), key=lambda path: path.stat().st_mtime, reverse=True)


def read_pm_labor_price_sources() -> tuple[list[dict[str, Any]], list[str]]:
    rows: list[dict[str, Any]] = []
    sources: list[str] = []
    for source_path in find_pm_labor_price_sources():
        data = read_json_file(source_path, [])
        if not isinstance(data, list):
            continue
        source_rows = [
            map_pm_labor_price_row(row, source_path.name)
            for row in data
            if isinstance(row, dict)
        ]
        source_rows = [
            normalize_price_library_row(row) | {"id": row.get("id"), "_sourceFile": source_path.name}
            for row in source_rows
            if clean_text(row.get("item")) and to_number(row.get("price")) > 0 and clean_text(row.get("unit"))
        ]
        if source_rows:
            rows.extend(source_rows)
            sources.append(source_path.name)
    return rows, sources


def find_pm_professional_price_sources() -> list[Path]:
    candidates: list[Path] = []
    for folder in PM_PROFESSIONAL_PRICE_TASK_DIRS:
        if not folder.exists():
            continue
        for suffix in ("*.json", "*.xlsx", "*.xls"):
            for path in folder.rglob(suffix):
                name = clean_text(path.name)
                if "专业" not in name and "专业分包" not in name:
                    continue
                if "价格库" not in name and "底表" not in name and "清单" not in name:
                    continue
                candidates.append(path)
    return sorted(candidates, key=lambda path: path.stat().st_mtime, reverse=True)


def read_pm_professional_price_sources() -> tuple[list[dict[str, Any]], list[str]]:
    rows: list[dict[str, Any]] = []
    sources: list[str] = []
    for source_path in find_pm_professional_price_sources():
        source_rows: list[dict[str, Any]] = []
        if source_path.suffix.lower() == ".json":
            data = read_json_file(source_path, [])
            raw_rows = data.get("rows", []) if isinstance(data, dict) else data
            if isinstance(raw_rows, list):
                source_rows = [
                    map_pm_professional_price_row(row, source_path.name)
                    for row in raw_rows
                    if isinstance(row, dict)
                ]
        elif source_path.suffix.lower() in {".xlsx", ".xls"}:
            parse_path = convert_xls_to_xlsx(source_path) if source_path.suffix.lower() == ".xls" else source_path
            source_rows = [
                {**row, "category": "专业分包", "source": row.get("source") or f"PM专业分包价格抓取/{source_path.name}"}
                for row in read_price_library_file(parse_path)
            ]
        normalized_rows = [
            normalize_price_library_row(row) | {"id": row.get("id") or stable_price_library_import_id("pm-professional-price", row, source_path.name), "_sourceFile": source_path.name}
            for row in source_rows
            if clean_text(row.get("item")) and to_number(row.get("price")) > 0 and clean_text(row.get("unit"))
        ]
        if normalized_rows:
            rows.extend(normalized_rows)
            sources.append(source_path.name)
    return rows, sources


def price_library_remark_attrs(remark: Any) -> dict[str, str]:
    text = clean_text(remark)
    attrs: dict[str, str] = {}
    for part in [item.strip() for item in text.split("；") if item.strip()]:
        if "：" not in part:
            continue
        key, value = part.split("：", 1)
        attrs[clean_text(key)] = clean_text(value)
    return attrs


def price_library_export_context(row: dict[str, Any]) -> dict[str, str]:
    attrs = price_library_remark_attrs(row.get("remark"))
    spec = clean_text(row.get("spec"))
    height = ""
    position = ""
    spec_parts = [part.strip() for part in spec.split("；") if part.strip()]
    remaining_spec = []
    for part in spec_parts:
        if part.startswith("层高："):
            height = part.removeprefix("层高：")
        elif part.startswith("施工部位："):
            position = part.removeprefix("施工部位：")
        else:
            remaining_spec.append(part)
    construction_content = attrs.get("施工内容") or clean_text(row.get("process")) or clean_text(row.get("trade")) or "；".join(remaining_spec)
    return {
        "paymentPeriod": attrs.get("付款周期", ""),
        "paymentRatio": attrs.get("付款比例", ""),
        "paymentMethod": attrs.get("付款方式", ""),
        "paymentCondition": attrs.get("付款条件", "") or "；".join(
            part for part in [attrs.get("付款周期", ""), attrs.get("付款比例", ""), attrs.get("付款方式", "")]
            if part
        ),
        "riskLevel": attrs.get("房源风险等级", ""),
        "height": height,
        "position": position,
        "constructionContent": construction_content,
    }


def price_library_export_date(value: Any) -> Any:
    text = clean_text(value)
    if not text:
        return ""
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m", "%Y/%m"):
        try:
            parsed = datetime.strptime(text[:10] if fmt in ("%Y-%m-%d", "%Y/%m/%d") else text[:7], fmt)
            return parsed
        except ValueError:
            continue
    return text


def price_library_export_row(row: dict[str, Any], index: int, average: bool = False, avg_price: float | None = None) -> list[Any]:
    context = price_library_export_context(row)
    remark = context["paymentCondition"] or clean_text(row.get("remark"))
    if average:
        remark = "现金价平均值"
    return [
        index,
        clean_text(row.get("category")),
        f"{clean_text(row.get('item'))} 平均值" if average else clean_text(row.get("item")),
        clean_text(row.get("unit")),
        round(avg_price, 4) if average and avg_price is not None else row.get("price", ""),
        context["paymentPeriod"],
        context["paymentRatio"],
        "现金" if average else context["paymentMethod"],
        context["riskLevel"],
        clean_text(row.get("region")),
        clean_text(row.get("projectType")),
        context["height"],
        price_library_export_date(row.get("priceDate")),
        "" if average else clean_text(row.get("company")),
        "" if average else clean_text(row.get("project")),
        context["position"],
        context["constructionContent"],
        remark,
    ]


def price_library_average_key(row: dict[str, Any]) -> tuple[str, str, str, str, str]:
    context = price_library_export_context(row)
    return (
        clean_text(row.get("category")),
        clean_text(row.get("item")),
        clean_text(row.get("unit")),
        clean_text(row.get("projectType")),
        context["height"],
    )


def price_library_is_cash_price(row: dict[str, Any]) -> bool:
    context = price_library_export_context(row)
    payment_method = context["paymentMethod"]
    if "现金" not in payment_method:
        return False
    if any(word in payment_method for word in ("抵房", "抵物", "房/物", "物资")):
        return False
    return to_number(row.get("price")) > 0


def price_library_cash_average_rows(rows: list[dict[str, Any]]) -> dict[tuple[str, str, str, str, str], list[dict[str, Any]]]:
    grouped: dict[tuple[str, str, str, str, str], list[dict[str, Any]]] = {}
    for row in rows:
        if not price_library_is_cash_price(row):
            continue
        grouped.setdefault(price_library_average_key(row), []).append(row)
    return grouped


def price_library_project_type_sort_key(project_type: str) -> tuple[int, str]:
    order = ["多层住宅", "小高层住宅", "高层住宅", "公用建筑", "公共建筑", "厂房、仓库", "市政", "未分类"]
    return (order.index(project_type) if project_type in order else len(order), project_type)


def price_library_export_project_type(value: Any, mode: str) -> str:
    text = clean_text(value)
    if text in {"多层住宅", "小高层住宅", "高层住宅", "公用建筑", "厂房、仓库", "市政"}:
        return text
    if text == "公共建筑":
        return "公用建筑"
    if "小高" in text:
        return "小高层住宅"
    if "高层" in text:
        return "高层住宅"
    if "多层" in text or "住宅" in text:
        return "多层住宅"
    if "厂房" in text or "仓库" in text:
        return "厂房、仓库"
    if "市政" in text:
        return "市政"
    return "公用建筑"


PRICE_LIBRARY_PROJECT_TYPES = ["多层住宅", "小高层住宅", "高层住宅", "公用建筑", "厂房、仓库", "市政"]
PRICE_LIBRARY_CIVIL_WIDTHS = [7.77777777777778, 12.6296296296296, 20.6296296296296, 10.6296296296296, 13.0, 13.0, 10.3796296296296, 20.6296296296296, 12.6296296296296, 12.6296296296296, 13.0, 13.6296296296296, 10.6296296296296, 10.6296296296296, 37.1666666666667, 15.0, 20.6296296296296, 15.6296296296296]
PRICE_LIBRARY_INSTALL_WIDTHS = [8.77777777777778, 12.6296296296296, 20.6296296296296, 10.6296296296296, 13.85, 10.6296296296296, 13.0, 20.6296296296296, 12.6296296296296, 12.6296296296296, 13.0, 13.6296296296296, 10.6296296296296, 10.6296296296296, 13.0, 18.72, 33.81, 15.6296296296296]


def price_library_export_mode(rows: list[dict[str, Any]]) -> str:
    installish = False
    civilish = False
    for row in rows:
        major = clean_text(row.get("major"))
        discipline = clean_text(row.get("discipline"))
        if "安装" in major or "安装" in discipline:
            installish = True
        if "土建" in major or "建筑" in discipline or "土建" in discipline:
            civilish = True
    if installish and not civilish:
        return "安装"
    if civilish and not installish:
        return "土建"
    return "混合"


def write_price_library_workbook(rows: list[dict[str, Any]], output_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    export_mode = price_library_export_mode(rows)
    base_project_types = PRICE_LIBRARY_PROJECT_TYPES
    export_widths = PRICE_LIBRARY_INSTALL_WIDTHS if export_mode == "安装" else PRICE_LIBRARY_CIVIL_WIDTHS
    export_headers = [
        "序号", "类型库\n（下拉选择）", "分包项名称", "计量单位", "实际采购价",
        "付款周期", "付款比例", "付款方式", "房源风险等级", "地区",
        "工程类型\n（下拉选择）", "层高", "定标日期", "基层单位", "项目部",
        "施工部位", "施工内容", "付款条件",
    ]
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        export_row = dict(row)
        project_type = price_library_export_project_type(row.get("projectType"), export_mode)
        export_row["projectType"] = project_type
        grouped.setdefault(project_type, []).append(export_row)
    ordered_project_types = base_project_types + [
        project_type for project_type in sorted(grouped, key=price_library_project_type_sort_key)
        if project_type not in base_project_types
    ]
    if not rows:
        ordered_project_types = base_project_types
    for project_type in ordered_project_types:
        group_rows = grouped.get(project_type, [])
        safe_title = re.sub(r"[:\\/?*\[\]]", "_", project_type)[:31] or "未分类"
        ws = wb.create_sheet(safe_title)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(export_headers))
        majors = {clean_text(row.get("major")) for row in group_rows if clean_text(row.get("major"))}
        disciplines = {clean_text(row.get("discipline")) for row in group_rows if clean_text(row.get("discipline"))}
        title_major = "安装" if "安装" in majors or "安装" in disciplines or export_mode == "安装" else "土建" if "土建" in majors or "建筑" in disciplines or export_mode == "土建" else ""
        title = f"{title_major}分包价格库—{project_type}" if title_major else f"分包价格库—{project_type}"
        ws.cell(1, 1).value = title
        ws.row_dimensions[1].height = 60
        ws.cell(1, 1).font = Font(name="宋体", bold=True, size=24)
        ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.append(export_headers)
        group_rows.sort(key=lambda row: (
            clean_text(row.get("item")),
            clean_text(row.get("unit")),
            clean_text(row.get("priceDate")),
            clean_text(row.get("company")),
        ))
        ordered_groups: dict[tuple[str, str, str, str, str], list[dict[str, Any]]] = {}
        for row in group_rows:
            ordered_groups.setdefault(price_library_average_key(row), []).append(row)
        index = 1
        for average_key, item_rows in ordered_groups.items():
            for row in item_rows:
                ws.append(price_library_export_row(row, index))
                ws.cell(ws.max_row, 1).value = f"=ROW(A{ws.max_row})-2"
                index += 1
            cash_rows = [
                row for row in item_rows
                if price_library_is_cash_price(row)
            ]
            if cash_rows:
                row = cash_rows[0]
                avg_price = sum(to_number(item.get("price")) for item in cash_rows) / len(cash_rows)
                ws.append(price_library_export_row(row, index, average=True, avg_price=avg_price))
                ws.cell(ws.max_row, 1).value = f"=ROW(A{ws.max_row})-2"
                index += 1
    for sheet in wb.worksheets:
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for index, width in enumerate(export_widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row_cells in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=18):
            for cell in row_cells:
                cell.font = Font(name="宋体", size=11, bold=(cell.row == 2))
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
                if cell.column == 13 and cell.row >= 3 and cell.value:
                    cell.number_format = "yyyy/m/d"
        if export_mode != "安装":
            sheet.auto_filter.ref = f"A2:R{sheet.max_row}"
    wb.save(output_path)


def build_material_comparison_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        grouped.setdefault(normalize_material_key(row), []).append(row)
    comparison: list[dict[str, Any]] = []
    for group_rows in grouped.values():
        prices = [row.get("price", 0) for row in group_rows if row.get("price", 0) > 0]
        if not prices:
            continue
        min_price = min(prices)
        max_price = max(prices)
        avg_price = sum(prices) / len(prices)
        latest = sorted(group_rows, key=lambda item: item.get("purchaseDate") or "", reverse=True)[0]
        comparison.append({
            "材料名称": latest.get("material", ""),
            "规格型号": latest.get("spec", ""),
            "单位": latest.get("unit", ""),
            "采购次数": len(group_rows),
            "最低价": round(min_price, 4),
            "最高价": round(max_price, 4),
            "平均价": round(avg_price, 4),
            "价差": round(max_price - min_price, 4),
            "价差率": round((max_price - min_price) / min_price, 4) if min_price else 0,
            "最低价供应商": next((row.get("supplier", "") for row in group_rows if row.get("price") == min_price), ""),
            "最高价供应商": next((row.get("supplier", "") for row in group_rows if row.get("price") == max_price), ""),
            "最近采购日期": latest.get("purchaseDate", ""),
            "最近项目": latest.get("project", ""),
        })
    return sorted(comparison, key=lambda item: (item["价差率"], item["价差"]), reverse=True)


def build_material_price_dispersion() -> dict[str, Any]:
    rows = latest_bid_detail_rows_for_dispersion()
    if not rows:
        rows = load_material_price_rows({"company": "集团"})
    material_groups: dict[str, dict[str, list[dict[str, float]]]] = {}
    material_labels: dict[str, dict[str, str]] = {}
    for row in rows:
        row = normalize_material_row(row)
        price = to_number(row.get("price"))
        if price <= 0:
            continue
        material_name = clean_text(row.get("material"))
        if is_non_material_price_item(material_name):
            continue
        company = material_dispersion_company(row)
        if "装饰幕墙" in company:
            continue
        key = normalize_material_key(row)
        quantity = to_number(row.get("quantity"))
        amount = to_number(row.get("amount"))
        if not amount and quantity:
            amount = price * quantity
        if not quantity and amount:
            quantity = amount / price if price else 0.0
        material_groups.setdefault(key, {}).setdefault(company, []).append({
            "price": price,
            "quantity": quantity,
            "amount": amount,
        })
        material_labels.setdefault(key, {
            "material": material_name,
            "spec": clean_text(row.get("spec")),
            "unit": clean_text(row.get("unit")),
        })

    materials: list[dict[str, Any]] = []
    for _key, company_prices in material_groups.items():
        if len(company_prices) < 2:
            continue
        company_avgs = {
            company: (
                sum(item["amount"] for item in records) / sum(item["quantity"] for item in records)
                if sum(item["quantity"] for item in records) > 0
                else sum(item["price"] for item in records) / len(records)
            )
            for company, records in company_prices.items()
            if records
        }
        if len(company_avgs) < 2:
            continue
        group_avg = sum(company_avgs.values()) / len(company_avgs)
        min_avg = min(company_avgs.values())
        max_avg = max(company_avgs.values())
        if group_avg <= 0 or min_avg <= 0:
            continue
        low_company, low_price = min(company_avgs.items(), key=lambda item: item[1])
        high_company, high_price = max(company_avgs.items(), key=lambda item: item[1])
        spread = high_price - low_price
        spread_rate = spread / low_price if low_price else 0.0
        total_quantity = sum(item["quantity"] for records in company_prices.values() for item in records)
        total_amount = sum(item["amount"] for records in company_prices.values() for item in records)
        impact_amount = spread * total_quantity if total_quantity else spread * sum(len(records) for records in company_prices.values())
        labels = material_labels.get(_key, {})
        company_detail = [
            {
                "company": company,
                "avgPrice": round(avg_price, 2),
                "recordCount": len(company_prices.get(company, [])),
                "quantity": round(sum(item["quantity"] for item in company_prices.get(company, [])), 4),
                "amount": round(sum(item["amount"] for item in company_prices.get(company, [])), 2),
            }
            for company, avg_price in sorted(company_avgs.items(), key=lambda item: item[1])
        ]
        materials.append({
            "key": _key,
            "material": labels.get("material", ""),
            "spec": labels.get("spec", ""),
            "unit": labels.get("unit", ""),
            "label": " ".join(value for value in [labels.get("material", ""), labels.get("spec", "")] if value),
            "lowCompany": low_company,
            "lowPrice": round(low_price, 2),
            "highCompany": high_company,
            "highPrice": round(high_price, 2),
            "spread": round(spread, 2),
            "spreadRate": round(spread_rate, 4),
            "totalQuantity": round(total_quantity, 4),
            "totalAmount": round(total_amount, 2),
            "impactAmount": round(impact_amount, 2),
            "companyCount": len(company_avgs),
            "priceCount": sum(len(prices) for prices in company_prices.values()),
            "companies": company_detail,
        })

    materials.sort(key=lambda item: (item["totalAmount"], item["totalQuantity"], item["impactAmount"]), reverse=True)
    bulk_materials = [
        item for item in materials
        if item.get("spreadRate", 0) > 0
        and (item.get("totalAmount", 0) >= 10000 or item.get("totalQuantity", 0) >= 100)
    ]
    display_materials = bulk_materials if len(bulk_materials) >= 5 else materials
    return {
        "type": "material-spread-ranking",
        "metric": "按采购金额优先的相同材料跨分公司价差",
        "materials": display_materials[:30],
        "companies": [],
    }


def is_non_material_price_item(material: Any) -> bool:
    text = clean_text(material)
    if not text:
        return True
    keywords = [
        "人工", "零工", "小工", "大工", "技工", "电工", "焊工", "瓦工", "木工", "钢筋工", "架子工",
        "机械", "租赁", "挖掘机", "洒水车", "货车", "吊车", "泵车", "铲车", "装载机", "运输", "运费",
    ]
    return any(keyword in text for keyword in keywords)


def latest_bid_detail_rows_for_dispersion() -> list[dict[str, Any]]:
    db = load_db()
    for task in db.get("materialPriceTasks", []):
        if not (task.get("sourcePaths") or task.get("sourceDir")):
            continue
        source_paths = bid_source_paths_for_task(task)
        if not source_paths:
            continue
        try:
            sources = build_bid_source_sets(source_paths)
            if missing_bid_source_labels(sources):
                continue
            details, _comparison_rows = build_bid_detail_rows(sources)
        except Exception:
            continue
        rows: list[dict[str, Any]] = []
        for detail in details:
            company = clean_text(detail.get("companyBucket")) or bid_company_bucket(detail.get("company"))
            if not company or "装饰幕墙" in company:
                continue
            rows.append({
                "material": detail.get("material", ""),
                "spec": detail.get("spec", ""),
                "unit": detail.get("unit", ""),
                "price": detail.get("price", 0),
                "quantity": detail.get("quantity", 0),
                "amount": detail.get("amount", 0),
                "purchaseDate": detail.get("bidDate", ""),
                "project": detail.get("project", ""),
                "supplier": detail.get("winner", ""),
                "company": company,
                "platformOrder": detail.get("taskNo", ""),
            })
        if rows:
            return rows
    return []


def material_dispersion_company(row: dict[str, Any]) -> str:
    raw_company = clean_text(row.get("company"))
    project = clean_text(row.get("project"))
    company = clean_bid_company_name(raw_company, project)
    bucket = bid_company_bucket(company)
    if bucket and bucket != company:
        return bucket
    standard = standard_company_name(company)
    if standard and standard != company:
        return standard
    if any(alias in company for _name, aliases in BID_COMPANY_COLUMNS for alias in aliases):
        return bid_company_bucket(company)
    return company or "未识别"


def export_material_price_task(task: dict[str, Any], rows: list[dict[str, Any]], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "横向对比表"
    headers = ["材料名称", "规格型号", "单位", "采购次数", "最低价", "最高价", "平均价", "价差", "价差率", "最低价供应商", "最高价供应商", "最近采购日期", "最近项目"]
    ws["A1"] = "材料采购价格横向对比表"
    ws["A1"].font = Font(name="宋体", size=16, bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws["A2"] = f"任务：{task.get('name', '')}    时间：{task.get('startDate', '')} 至 {task.get('endDate', '')}"
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    comparison_rows = build_material_comparison_rows(rows)
    ws.append(headers)
    for row in comparison_rows:
        ws.append([row.get(header, "") for header in headers])
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(ws.max_row, 3)}"
    for cell in ws[3]:
        cell.fill = PatternFill("solid", fgColor="DCE8F2")
        cell.font = Font(name="宋体", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=4):
        row[8].number_format = "0.00%"
    autosize(ws)

    detail = wb.create_sheet("平台明细")
    detail_headers = ["材料名称", "规格型号", "单位", "采购单价", "数量", "金额", "采购日期", "项目", "供应商", "分公司", "订单编号"]
    detail_rows = [[
        row.get("material", ""),
        row.get("spec", ""),
        row.get("unit", ""),
        row.get("price", 0),
        row.get("quantity", 0),
        row.get("amount", 0),
        row.get("purchaseDate", ""),
        row.get("project", ""),
        row.get("supplier", ""),
        row.get("company", ""),
        row.get("platformOrder", ""),
    ] for row in rows]
    write_sheet(detail, detail_headers, detail_rows)
    autosize(detail)

    info = wb.create_sheet("指令信息")
    info_rows = [
        ["分公司", "集团" if task.get("company") == "全集团" else task.get("company", "集团")],
        ["开始时间", task.get("startDate", "")],
        ["结束时间", task.get("endDate", "")],
        ["材料关键词", task.get("keyword", "")],
        ["数据来源", task.get("source", "")],
        ["状态", task.get("status", "")],
        ["说明", task.get("message", "")],
    ]
    write_sheet(info, ["字段", "内容"], info_rows)
    autosize(info)
    wb.save(output_path)


def clone_cell(source: Any, target: Any) -> None:
    target.value = source.value
    if source.has_style:
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)
    if source.hyperlink:
        target._hyperlink = copy(source.hyperlink)
    if source.comment:
        target.comment = copy(source.comment)


def _set_docx_run_font(run: Any, font_name: str, size: Any, bold: bool, color: Any) -> None:
    """Apply a captured template font, including its East Asian font setting."""
    run.font.name = font_name
    if size:
        run.font.size = size
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    if rpr.rFonts is None:
        from docx.oxml import OxmlElement

        rpr.insert(0, OxmlElement("w:rFonts"))
    rpr.rFonts.set(qn("w:eastAsia"), font_name)


def _write_template_cell(cell: Any, value: Any) -> None:
    """Replace visible text without changing a template cell's borders, fill or layout."""
    paragraph = cell.paragraphs[0]
    sample = next((run for run in paragraph.runs if run.text), None)
    font_name = (sample.font.name if sample and sample.font.name else "宋体")
    size = sample.font.size if sample else Pt(12)
    bold = bool(sample.font.bold) if sample else False
    color = sample.font.color.rgb if sample and sample.font.color.type else None
    alignment = paragraph.alignment
    line_spacing = paragraph.paragraph_format.line_spacing
    # Merged template cells can carry residual paragraphs from their original
    # constituent cells.  Keep the first paragraph and remove the rest before
    # writing, otherwise header words are visibly duplicated.
    for extra_paragraph in cell.paragraphs[1:]:
        extra_paragraph._element.getparent().remove(extra_paragraph._element)
    paragraph.clear()
    run = paragraph.add_run("" if value is None else clean_text(value))
    _set_docx_run_font(run, font_name, size, bold, color)
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = line_spacing


def _write_template_paragraph(paragraph: Any, value: Any) -> None:
    """Replace paragraph text while preserving the reference document's paragraph style."""
    sample = next((run for run in paragraph.runs if run.text), None)
    font_name = (sample.font.name if sample and sample.font.name else "宋体")
    size = sample.font.size if sample else Pt(12)
    bold = bool(sample.font.bold) if sample else False
    color = sample.font.color.rgb if sample and sample.font.color.type else None
    alignment = paragraph.alignment
    line_spacing = paragraph.paragraph_format.line_spacing
    paragraph.clear()
    run = paragraph.add_run("" if value is None else clean_text(value))
    _set_docx_run_font(run, font_name, size, bold, color)
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = line_spacing


def _template_table_fill(source_ws: Any, table: Any) -> None:
    """Fill a Word template table from an Excel notice table, retaining the Word template shape."""
    required_rows = source_ws.max_row
    while len(table.rows) < required_rows:
        table._tbl.append(copy(table.rows[-1]._tr))

    seen_cells: set[int] = set()
    for row_index in range(1, min(source_ws.max_row, len(table.rows)) + 1):
        for col_index in range(1, min(source_ws.max_column, len(table.columns)) + 1):
            source_cell = source_ws.cell(row_index, col_index)
            if source_cell.value is None:
                continue
            target_cell = table.cell(row_index - 1, col_index - 1)
            cell_identity = id(target_cell._tc)
            if cell_identity in seen_cells:
                continue
            seen_cells.add(cell_identity)
            value = source_cell.value
            if isinstance(value, (int, float)) and "%" in (source_cell.number_format or ""):
                value = f"{value * 100:.1f}%"
            _write_template_cell(target_cell, value)


def _worksheet_to_template_docx(source_ws: Any, output_path: Path, notice_kind: str) -> bool:
    template_path = TEMPLATE_DIR / f"settlement-notice-{notice_kind}.docx"
    if not template_path.exists():
        return False

    shutil.copyfile(template_path, output_path)
    doc = Document(output_path)
    if not doc.tables:
        return False

    # The supplied Word files are the format authority.  Match text from the
    # current Excel notice into that preserved document rather than recreate a layout.
    source_title = clean_text(source_ws.cell(2 if notice_kind in {"03", "06"} else 1, 1).value)
    if source_title:
        title = next((p for p in doc.paragraphs if "关于" in p.text and "通报" in p.text or "考核" in p.text), None)
        if title:
            _write_template_paragraph(title, source_title)

    table = doc.tables[0]
    table.autofit = False
    if notice_kind == "02":
        table_start = 8
    elif notice_kind == "03":
        table_start = 10
    else:
        table_start = 9

    paragraph_rows = {
        "02": [(1, 1, 5), (3, 1, 7), (4, 1, 8), (7, 1, 9), (28, 1, 10), (33, 5, 12), (34, 5, 13)],
        "03": [(2, 1, 0), (4, 1, 2), (5, 1, 3), (18, 1, 4), (22, 4, 5), (23, 4, 6)],
        "06": [(2, 1, 5), (4, 1, 7), (5, 1, 8), (8, 1, 9), (29, 1, 10), (30, 1, 11), (33, 5, 13), (34, 5, 14)],
    }
    for source_row, source_col, paragraph_index in paragraph_rows.get(notice_kind, []):
        if paragraph_index >= len(doc.paragraphs):
            continue
        source_text = source_ws.cell(source_row, source_col).value
        if source_text is not None:
            _write_template_paragraph(doc.paragraphs[paragraph_index], source_text)

    # Use a compact worksheet view so the reusable table filler only sees the
    # notice table; ending at the '合计' row avoids appending signature rows.
    table_end = table_start
    for row_index in range(table_start, source_ws.max_row + 1):
        row_values = [clean_text(source_ws.cell(row_index, col).value) for col in range(1, source_ws.max_column + 1)]
        if any(value == "合计" for value in row_values):
            table_end = row_index
            break

    class NoticeTableView:
        def __init__(self, ws: Any, start_row: int, end_row: int) -> None:
            self._ws = ws
            self._start = start_row
            self.max_row = max(1, end_row - start_row + 1)
            self.max_column = ws.max_column

        def cell(self, row: int, column: int) -> Any:
            return self._ws.cell(self._start + row - 1, column)

    _template_table_fill(NoticeTableView(source_ws, table_start, table_end), table)
    doc.save(output_path)
    return True


def worksheet_to_docx(source_ws: Any, output_path: Path, notice_kind: str = "") -> None:
    """Create a notice Word file from its approved template, with a safe fallback."""
    if notice_kind and _worksheet_to_template_docx(source_ws, output_path, notice_kind):
        return
    _worksheet_to_docx_fallback(source_ws, output_path)


def _worksheet_to_docx_fallback(source_ws: Any, output_path: Path) -> None:
    """Fallback used only when an approved Word template is unavailable."""
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE if source_ws.max_column > 8 else WD_ORIENT.PORTRAIT
    if section.orientation == WD_ORIENT.LANDSCAPE:
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.3)
    section.right_margin = Cm(1.3)

    table = doc.add_table(rows=source_ws.max_row, cols=source_ws.max_column)
    table.style = "Table Grid"
    table.autofit = False

    merged_followers: set[tuple[int, int]] = set()
    title_rows: set[int] = set()
    for merged_range in source_ws.merged_cells.ranges:
        start = table.cell(merged_range.min_row - 1, merged_range.min_col - 1)
        end = table.cell(merged_range.max_row - 1, merged_range.max_col - 1)
        start.merge(end)
        if (
            merged_range.min_col == 1
            and merged_range.max_col == source_ws.max_column
            and "关于" in clean_text(source_ws.cell(merged_range.min_row, merged_range.min_col).value)
        ):
            title_rows.add(merged_range.min_row)
        for row_index in range(merged_range.min_row, merged_range.max_row + 1):
            for col_index in range(merged_range.min_col, merged_range.max_col + 1):
                if (row_index, col_index) != (merged_range.min_row, merged_range.min_col):
                    merged_followers.add((row_index, col_index))

    for row_index in range(1, source_ws.max_row + 1):
        for col_index in range(1, source_ws.max_column + 1):
            if (row_index, col_index) in merged_followers:
                continue
            source_cell = source_ws.cell(row_index, col_index)
            target_cell = table.cell(row_index - 1, col_index - 1)
            target_cell.text = "" if source_cell.value is None else clean_text(source_cell.value)
            target_cell.vertical_alignment = 1
            for paragraph in target_cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if source_cell.alignment.horizontal != "left" else WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "宋体"
                    run.font.size = Pt(13 if row_index in title_rows else 8.5)
                    run.font.bold = bool(source_cell.font.bold)
                    if row_index in title_rows:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(output_path)


def clone_worksheet(source_ws: Any, target_ws: Any) -> None:
    for row in source_ws.iter_rows():
        for source_cell in row:
            clone_cell(source_cell, target_ws.cell(source_cell.row, source_cell.column))
    for merged in source_ws.merged_cells.ranges:
        target_ws.merge_cells(str(merged))
    for key, dim in source_ws.column_dimensions.items():
        target_ws.column_dimensions[key].width = dim.width
        target_ws.column_dimensions[key].hidden = dim.hidden
    for key, dim in source_ws.row_dimensions.items():
        target_ws.row_dimensions[key].height = dim.height
        target_ws.row_dimensions[key].hidden = dim.hidden
    target_ws.freeze_panes = source_ws.freeze_panes
    target_ws.sheet_view.showGridLines = source_ws.sheet_view.showGridLines


def style_from(ws: Any, source_row: int, source_col: int, target_row: int, target_col: int) -> None:
    clone_cell(ws.cell(source_row, source_col), ws.cell(target_row, target_col))
    ws.cell(target_row, target_col).value = None


def copy_row_style(ws: Any, source_row: int, target_row: int, max_col: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col_index in range(1, max_col + 1):
        style_from(ws, source_row, col_index, target_row, col_index)


def apply_cell_box(ws: Any, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index in range(min_row, max_row + 1):
        for col_index in range(min_col, max_col + 1):
            cell = ws.cell(row_index, col_index)
            cell.border = border
            cell.alignment = center


BLACKLIST_RELATION_FIELDS = [
    ("legalRep", "同法人"),
    ("shareholders", "同股东"),
    ("contacts", "同联系电话"),
    ("bankAccounts", "同银行账户"),
    ("address", "同地址"),
]


def normalize_lookup_text(value: Any) -> str:
    return re.sub(r"\s+", "", clean_text(value)).lower()


def blacklist_values(entity: dict[str, Any], field: str) -> list[str]:
    value = entity.get(field)
    if value is None:
        return []
    if isinstance(value, list):
        return [clean_text(item) for item in value if clean_text(item)]
    text = clean_text(value)
    return [text] if text else []


def blacklist_matches_query(entity: dict[str, Any], query: str) -> bool:
    if not query:
        return True
    haystack = [
        entity.get("name"),
        entity.get("creditCode"),
        entity.get("legalRep"),
        entity.get("address"),
        entity.get("reason"),
        *blacklist_values(entity, "shareholders"),
        *blacklist_values(entity, "contacts"),
        *blacklist_values(entity, "bankAccounts"),
    ]
    normalized_query = normalize_lookup_text(query)
    return any(normalized_query in normalize_lookup_text(value) for value in haystack)


def relation_reasons(source: dict[str, Any], target: dict[str, Any]) -> list[str]:
    reasons: list[str] = []
    for field, label in BLACKLIST_RELATION_FIELDS:
        source_values = {normalize_lookup_text(value): value for value in blacklist_values(source, field)}
        target_values = {normalize_lookup_text(value): value for value in blacklist_values(target, field)}
        common = [source_values[key] for key in source_values.keys() & target_values.keys() if key]
        if common:
            reasons.append(f"{label}：{'、'.join(common)}")
    if normalize_lookup_text(source.get("name")) in {normalize_lookup_text(value) for value in blacklist_values(target, "shareholders")}:
        reasons.append("目标股东包含命中企业")
    if normalize_lookup_text(target.get("name")) in {normalize_lookup_text(value) for value in blacklist_values(source, "shareholders")}:
        reasons.append("命中企业股东包含目标企业")
    return reasons


def blacklist_result_row(entity: dict[str, Any], hit_type: str, reasons: list[str]) -> dict[str, Any]:
    return {
        "id": entity.get("id", ""),
        "hitType": hit_type,
        "name": entity.get("name", ""),
        "creditCode": entity.get("creditCode", ""),
        "riskLevel": entity.get("riskLevel", ""),
        "status": entity.get("status", ""),
        "legalRep": entity.get("legalRep", ""),
        "contacts": "、".join(blacklist_values(entity, "contacts")),
        "reason": "；".join(reasons) if reasons else entity.get("reason", ""),
        "updatedAt": entity.get("updatedAt", ""),
    }


def route_node_id(label: Any) -> str:
    text = normalize_lookup_text(label) or "unknown"
    return "route-" + re.sub(r"[^0-9a-zA-Z\u4e00-\u9fff_-]", "", text)[:48]


WEAK_BLACKLIST_ROUTE_KEYWORDS = (
    "页面可见",
    "同地区同行业",
    "同行业",
    "同地区",
    "推荐",
    "浏览",
    "搜索",
)


def is_strong_blacklist_relation(label: Any) -> bool:
    text = clean_text(label)
    if not text:
        return False
    if any(keyword in text for keyword in WEAK_BLACKLIST_ROUTE_KEYWORDS):
        return False
    return any(keyword in text for keyword in (
        "股东",
        "控股",
        "投资",
        "出资",
        "分支",
        "分公司",
        "人员",
        "高管",
        "任职",
        "法人",
        "电话",
        "地址",
        "账户",
        "风险",
        "失信",
        "被执行",
        "裁判",
        "处罚",
    ))


def blacklist_mindmap(query: str, rows: list[dict[str, Any]], routes: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    root_label = query or "全部黑名单关联"
    nodes = [{"id": "root", "label": root_label, "type": "root"}]
    links: list[dict[str, str]] = []
    if routes:
        node_ids = {"root"}
        for route in routes:
            path_values = route.get("path")
            if isinstance(path_values, list) and len(path_values) >= 2:
                labels = [clean_text(value) for value in path_values if clean_text(value)]
                relation_labels = route.get("relations") if isinstance(route.get("relations"), list) else []
                for index in range(len(labels) - 1):
                    source_label = labels[index]
                    target_label = labels[index + 1]
                    label = clean_text(relation_labels[index] if index < len(relation_labels) else route.get("label")) or "穿透"
                    if not is_strong_blacklist_relation(label):
                        continue
                    source_id = "root" if index == 0 and normalize_lookup_text(source_label) == normalize_lookup_text(root_label) else route_node_id(source_label)
                    target_id = route_node_id(target_label)
                    if source_id not in node_ids:
                        nodes.append({"id": source_id, "label": source_label, "type": clean_text(route.get("sourceType")) or ("root" if index == 0 else "related")})
                        node_ids.add(source_id)
                    if target_id not in node_ids:
                        node_type = clean_text(route.get("targetType")) or ("related" if index < len(labels) - 2 else "black")
                        nodes.append({"id": target_id, "label": target_label, "type": node_type})
                        node_ids.add(target_id)
                    links.append({"source": source_id, "target": target_id, "label": label})
                continue
            source_label = clean_text(route.get("source") or route.get("from") or root_label)
            target_label = clean_text(route.get("target") or route.get("to"))
            if not target_label:
                continue
            label = clean_text(route.get("label") or route.get("relation") or "关联")
            if not is_strong_blacklist_relation(label):
                continue
            source_id = "root" if normalize_lookup_text(source_label) == normalize_lookup_text(root_label) else route_node_id(source_label)
            target_id = route_node_id(target_label)
            if source_id not in node_ids:
                nodes.append({"id": source_id, "label": source_label, "type": clean_text(route.get("sourceType")) or "related"})
                node_ids.add(source_id)
            if target_id not in node_ids:
                nodes.append({"id": target_id, "label": target_label, "type": clean_text(route.get("targetType")) or "related"})
                node_ids.add(target_id)
            links.append({"source": source_id, "target": target_id, "label": label})
        if links:
            return {"nodes": nodes, "links": links, "mode": "routes"}
        return {"nodes": nodes, "links": links, "mode": "routes", "filteredWeakRelations": True}
    for row in rows:
        entity_id = f"entity-{row.get('id')}"
        nodes.append({
            "id": entity_id,
            "label": row.get("name", ""),
            "type": "black" if row.get("riskLevel") == "黑名单" else "related",
        })
        links.append({"source": "root", "target": entity_id, "label": row.get("hitType", "")})
        reasons = [item for item in clean_text(row.get("reason")).split("；") if item]
        for index, reason in enumerate(reasons[:4], start=1):
            reason_id = f"{entity_id}-reason-{index}"
            nodes.append({"id": reason_id, "label": reason, "type": "reason"})
            links.append({"source": entity_id, "target": reason_id, "label": "依据"})
    return {"nodes": nodes, "links": links, "mode": "evidence"}


def search_blacklist_entities(query: str) -> dict[str, Any]:
    entities = load_db().get("blacklistEntities", []) + load_blacklist_result_entities()
    routes = load_blacklist_result_routes(query)
    direct = [entity for entity in entities if blacklist_matches_query(entity, query)]
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    for entity in direct:
        entity_id = clean_text(entity.get("id")) or clean_text(entity.get("name"))
        seen.add(entity_id)
        rows.append(blacklist_result_row(entity, "直接命中", [clean_text(entity.get("reason"))]))
    for source in direct:
        for target in entities:
            target_id = clean_text(target.get("id")) or clean_text(target.get("name"))
            if target_id in seen:
                continue
            reasons = relation_reasons(source, target)
            if reasons:
                seen.add(target_id)
                rows.append(blacklist_result_row(target, "关联命中", reasons))
    high_risk = sum(1 for row in rows if row.get("riskLevel") == "黑名单")
    related = sum(1 for row in rows if row.get("hitType") == "关联命中")
    return {
        "query": query,
        "total": len(rows),
        "highRisk": high_risk,
        "related": related,
        "rows": rows,
        "mindmap": blacklist_mindmap(query, rows, routes),
    }


def export_blacklist_result(result: dict[str, Any], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "关联查询结果"
    ws.merge_cells("A1:I1")
    ws["A1"] = "黑名单企业关联查询结果"
    ws["A1"].font = Font(bold=True, size=16, color="B71924")
    ws["A2"] = "查询条件"
    ws["B2"] = result.get("query") or "全部"
    ws["D2"] = "命中数量"
    ws["E2"] = result.get("total", 0)
    ws["F2"] = "黑名单"
    ws["G2"] = result.get("highRisk", 0)
    ws["H2"] = "关联命中"
    ws["I2"] = result.get("related", 0)
    headers = ["命中类型", "企业名称", "统一社会信用代码", "风险等级", "状态", "法人", "联系方式", "命中依据", "更新时间"]
    ws.append([])
    ws.append(headers)
    for row in result.get("rows", []):
        ws.append([
            row.get("hitType", ""),
            row.get("name", ""),
            row.get("creditCode", ""),
            row.get("riskLevel", ""),
            row.get("status", ""),
            row.get("legalRep", ""),
            row.get("contacts", ""),
            row.get("reason", ""),
            row.get("updatedAt", ""),
        ])
    header_fill = PatternFill("solid", fgColor="B71924")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[4]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=9):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = [14, 28, 24, 12, 12, 12, 22, 46, 14]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A5"

    relation_ws = wb.create_sheet("关系图数据")
    relation_ws.append(["节点ID", "节点名称", "节点类型", "父节点", "关系"])
    labels = {node.get("id"): node.get("label", "") for node in result.get("mindmap", {}).get("nodes", [])}
    types = {node.get("id"): node.get("type", "") for node in result.get("mindmap", {}).get("nodes", [])}
    relation_ws.append(["root", labels.get("root", ""), types.get("root", ""), "", ""])
    for link in result.get("mindmap", {}).get("links", []):
        target = link.get("target", "")
        relation_ws.append([target, labels.get(target, ""), types.get(target, ""), labels.get(link.get("source", ""), ""), link.get("label", "")])
    for column in range(1, 6):
        relation_ws.column_dimensions[get_column_letter(column)].width = [24, 38, 14, 38, 14][column - 1]
    for cell in relation_ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    wb.save(output_path)


def clamp_percent(value: float) -> int:
    return max(0, min(100, int(round(value))))


def latest_month_key(*values: Any) -> str:
    months: list[str] = []
    for value in values:
        text = clean_text(value)
        if re.match(r"^\d{4}-\d{2}", text):
            months.append(text[:7])
    if months:
        return max(months)
    return datetime.now().strftime("%Y-%m")


def parse_datetime_value(value: Any) -> datetime | None:
    text = clean_text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y-%m-%dT%H:%M:%S.%fZ", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(text.replace("+08:00", ""), fmt)
        except ValueError:
            continue
    return None


def month_end_datetime(year_month: str) -> datetime:
    year, month = [int(part) for part in year_month.split("-")]
    return datetime(year, month, monthrange(year, month)[1], 23, 59, 59)


def month_keys_between(start: str, end: str) -> list[str]:
    start_year, start_month = [int(part) for part in start.split("-")]
    end_year, end_month = [int(part) for part in end.split("-")]
    keys: list[str] = []
    year, month = start_year, start_month
    while (year, month) <= (end_year, end_month):
        keys.append(f"{year:04d}-{month:02d}")
        month += 1
        if month > 12:
            year += 1
            month = 1
    return keys


def event_datetimes(db: dict[str, Any]) -> list[datetime]:
    values: list[datetime] = []
    for upload in db.get("uploads", []):
        dt = parse_datetime_value(upload.get("uploadedAt"))
        if dt:
            values.append(dt)
    for task in db.get("materialPriceTasks", []):
        dt = parse_datetime_value(task.get("updatedAt") or task.get("createdAt"))
        if dt:
            values.append(dt)
    for job in db.get("blacklistCodexJobs", []):
        dt = parse_datetime_value(job.get("updatedAt") or job.get("createdAt"))
        if dt:
            values.append(dt)
    for entity in db.get("blacklistEntities", []):
        dt = parse_datetime_value(entity.get("updatedAt"))
        if dt:
            values.append(dt)
    for path in sorted(PM_WARNING_RESULT_DIR.glob("latest_*.json")):
        try:
            values.append(datetime.fromtimestamp(path.stat().st_mtime))
        except OSError:
            continue
    return sorted(values)


def date_points_between(start: datetime, end: datetime) -> list[datetime]:
    start_day = datetime(start.year, start.month, start.day, 23, 59, 59)
    end_day = datetime(end.year, end.month, end.day, 23, 59, 59)
    points: list[datetime] = []
    cursor = start_day
    while cursor <= end_day:
        points.append(cursor)
        cursor += timedelta(days=1)
    return points


def month_points_between(start: datetime, end: datetime) -> list[datetime]:
    keys = month_keys_between(start.strftime("%Y-%m"), end.strftime("%Y-%m"))
    return [month_end_datetime(key) for key in keys]


def build_workbench_trend(db: dict[str, Any], period: str) -> list[dict[str, Any]]:
    events = event_datetimes(db)
    now = datetime.now()
    if not events:
        events = [now]
    start = min(events)
    span_days = max((now.date() - start.date()).days, 0)
    points = date_points_between(start, now) if span_days <= 45 else month_points_between(start, now)
    return [
        {
            "label": point.strftime("%m-%d") if span_days <= 45 else f"{point.month}月",
            "date": point.strftime("%Y-%m-%d"),
            "value": modules_overall(build_workbench_modules(db, period, point)),
        }
        for point in points
    ]


def task_done(status: Any) -> bool:
    text = clean_text(status)
    return any(marker in text for marker in ("已建表", "已采集", "已导出", "完成", "校验通过"))


def task_started(status: Any) -> bool:
    text = clean_text(status)
    return bool(text) and not any(marker in text for marker in ("失败", "未通过", "错误"))


def settlement_progress(db: dict[str, Any], period: str, as_of: datetime | None = None) -> dict[str, Any]:
    expected = company_order("settlement")
    uploads = [
        item for item in db.get("uploads", [])
        if item.get("reportType") == "settlement" and clean_text(item.get("period")) == period
    ]
    if as_of:
        uploads = [
            item for item in uploads
            if (parse_datetime_value(item.get("uploadedAt")) or datetime.max) <= as_of
        ]
    uploaded = {
        standard_company_name(item.get("company"))
        for item in uploads
        if clean_text(item.get("status")) == "校验通过"
    }
    uploaded = {company for company in uploaded if company}
    target = max(len(expected), 1)
    value = clamp_percent(len(uploaded) / target * 100)
    return {
        "key": "settlement",
        "label": "竣工结算报表",
        "value": value,
        "detail": f"{period} 已上传 {len(uploaded)}/{len(expected)} 家",
        "view": "auto-reports",
    }


def material_progress(db: dict[str, Any], as_of: datetime | None = None) -> dict[str, Any]:
    tasks = db.get("materialPriceTasks", [])
    if as_of:
        tasks = [
            task for task in tasks
            if (parse_datetime_value(task.get("updatedAt") or task.get("createdAt")) or datetime.max) <= as_of
        ]
    done = sum(1 for task in tasks if task_done(task.get("status")))
    started = sum(1 for task in tasks if task_started(task.get("status")))
    value = clamp_percent((done * 70 + started * 30) / max(len(tasks), 1)) if tasks else 0
    source_rows = 0
    source_path = DATA_DIR / "material_price_source.json"
    data = read_json_file(source_path, {})
    source_time = None
    if source_path.exists():
        source_time = parse_datetime_value(data.get("capturedAt")) if isinstance(data, dict) else None
        source_time = source_time or datetime.fromtimestamp(source_path.stat().st_mtime)
    if isinstance(data, dict) and source_path.exists() and (not as_of or source_time <= as_of):
        source_rows = len(data.get("rows") or [])
    return {
        "key": "material-price",
        "label": "材料价格比对",
        "value": value,
        "detail": f"任务 {done}/{len(tasks)} 个已建表，源数据 {source_rows} 条",
        "view": "material-price",
    }


def pm_warning_progress(as_of: datetime | None = None) -> dict[str, Any]:
    if as_of:
        source_times: list[datetime] = []
        for path in sorted(PM_WARNING_RESULT_DIR.glob("latest_*.json")):
            try:
                source_times.append(datetime.fromtimestamp(path.stat().st_mtime))
            except OSError:
                continue
        if not source_times or min(source_times) > as_of:
            return {
                "key": "pm-warning",
                "label": "PM红蓝预警",
                "value": 0,
                "detail": "截至当月尚未形成抓数结果",
                "view": "pm-warning",
                "source": "",
            }
    data = load_pm_warning_data("all")
    totals = data.get("totals") or {}
    in_progress = int(to_number(totals.get("inProgress") or totals.get("total")))
    warned = int(to_number(totals.get("red")) + to_number(totals.get("blue")))
    source = data.get("source") or {}
    has_source = bool(data.get("projects") or data.get("companies") or in_progress)
    value = clamp_percent(20 + min(in_progress, 120) / 120 * 55 + min(warned, 40) / 40 * 25) if has_source else 0
    return {
        "key": "pm-warning",
        "label": "PM红蓝预警",
        "value": value,
        "detail": f"在建 {in_progress} 项，红/蓝预警 {warned} 项",
        "view": "pm-warning",
        "source": source.get("displaySource") or source.get("updatedAt") or "",
    }


def blacklist_progress(db: dict[str, Any], as_of: datetime | None = None) -> dict[str, Any]:
    jobs = list_blacklist_codex_jobs()
    entities = db.get("blacklistEntities") or []
    result_entities = load_blacklist_result_entities()
    if as_of:
        jobs = [
            job for job in jobs
            if (parse_datetime_value(job.get("updatedAt") or job.get("createdAt")) or datetime.max) <= as_of
        ]
        entities = [
            entity for entity in entities
            if (parse_datetime_value(entity.get("updatedAt")) or datetime.max) <= as_of
        ]
        result_entities = [
            entity for entity in result_entities
            if (parse_datetime_value(entity.get("updatedAt")) or datetime.max) <= as_of
        ]
    done = sum(1 for job in jobs if task_done(job.get("status")))
    entities = len(entities) + len(result_entities)
    job_ratio = done / max(len(jobs), 1) * 50 if jobs else 0
    value = clamp_percent(min(entities, 30) / 30 * 50 + job_ratio) if entities or jobs else 0
    return {
        "key": "blacklist-query",
        "label": "黑名单关联查询",
        "value": value,
        "detail": f"企业库 {entities} 条，穿透任务 {done}/{len(jobs)} 个完成",
        "view": "blacklist-query",
    }


def price_library_progress(as_of: datetime | None = None) -> dict[str, Any]:
    rows = load_price_library_rows()
    if as_of:
        rows = [
            row for row in rows
            if (parse_datetime_value(row.get("updatedAt") or row.get("createdAt")) or datetime.max) <= as_of
        ]
    companies = {clean_text(row.get("company")) for row in rows if clean_text(row.get("company"))}
    items = {clean_text(row.get("item")) for row in rows if clean_text(row.get("item"))}
    value = clamp_percent(min(len(rows), 200) / 200 * 70 + min(len(items), 50) / 50 * 30) if rows else 0
    return {
        "key": "price-library",
        "label": "分包价格库",
        "value": value,
        "detail": f"已沉淀 {len(rows)} 条价格记录，覆盖 {len(companies)} 个分公司、{len(items)} 类施工内容",
        "view": "price-library",
    }


def build_workbench_modules(db: dict[str, Any], period: str, as_of: datetime | None = None) -> list[dict[str, Any]]:
    return [
        pm_warning_progress(as_of),
        material_progress(db, as_of),
        settlement_progress(db, period, as_of),
        price_library_progress(as_of),
        blacklist_progress(db, as_of),
        {
            "key": "collusion-check",
            "label": "围标串标识别",
            "value": 0,
            "detail": "暂未接入识别任务数据",
            "view": "collusion-check",
        },
    ]

def modules_overall(modules: list[dict[str, Any]]) -> int:
    return clamp_percent(sum(item["value"] for item in modules) / max(len(modules), 1))


def build_workbench_progress() -> dict[str, Any]:
    db = load_db()
    period = latest_month_key(*(item.get("period") for item in db.get("uploads", [])))
    modules = [
        *build_workbench_modules(db, period),
    ]
    overall = modules_overall(modules)
    trend = build_workbench_trend(db, period)
    return {
        "period": period,
        "overall": overall,
        "modules": modules,
        "trend": trend,
        "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


PLUGINS: dict[str, ReportPlugin] = {
    SettlementPlugin.key: SettlementPlugin(),
    RiskLevelPlugin.key: RiskLevelPlugin(),
}


class AppHandler(BaseHTTPRequestHandler):
    server_version = "ReportHub/0.1"

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/":
            try:
                server_port = int(self.server.server_address[1])
            except Exception:
                server_port = 0
            if server_port == 8899 and PM_WARNING_START_DEBUG_BROWSERS_ON_8899:
                for platform in PM_WARNING_START_DEBUG_PLATFORMS_ON_8899:
                    start_pm_warning_debug_browser_async(platform)
            self.serve_file(STATIC_DIR / "index.html", "text/html; charset=utf-8")
        elif parsed.path.startswith("/static/"):
            file_path = STATIC_DIR / parsed.path.removeprefix("/static/")
            content_type = "text/plain; charset=utf-8"
            if file_path.suffix == ".css":
                content_type = "text/css; charset=utf-8"
            elif file_path.suffix == ".js":
                content_type = "application/javascript; charset=utf-8"
            elif file_path.suffix == ".json":
                content_type = "application/json; charset=utf-8"
            self.serve_file(file_path, content_type)
        elif parsed.path == "/api/report-types":
            self.send_json([
                {"key": plugin.key, "name": plugin.name, "description": plugin.description}
                for plugin in PLUGINS.values()
            ])
        elif parsed.path == "/api/uploads":
            self.send_json(load_db()["uploads"])
        elif parsed.path == "/api/workbench-progress":
            self.send_json(build_workbench_progress())
        elif parsed.path == "/api/issues":
            self.handle_issues(parse_qs(parsed.query))
        elif parsed.path == "/api/missing-companies":
            self.handle_missing_companies(parse_qs(parsed.query))
        elif parsed.path == "/api/summary":
            params = parse_qs(parsed.query)
            self.handle_summary(params)
        elif parsed.path == "/api/export":
            params = parse_qs(parsed.query)
            self.handle_export(params)
        elif parsed.path == "/api/export-settled-details":
            params = parse_qs(parsed.query)
            self.handle_export_settled_details(params)
        elif parsed.path == "/api/material-price-tasks":
            self.handle_material_price_tasks()
        elif parsed.path == "/api/material-price-companies":
            self.send_json(material_price_companies())
        elif parsed.path == "/api/material-price-trend":
            self.send_json(build_material_price_dispersion())
        elif parsed.path == "/api/codex-material-jobs":
            self.send_json(list_material_price_codex_jobs())
        elif parsed.path == "/api/material-price-export":
            self.handle_material_price_export(parse_qs(parsed.query))
        elif parsed.path == "/api/price-library":
            self.handle_price_library(parse_qs(parsed.query))
        elif parsed.path == "/api/price-library-export":
            self.handle_price_library_export(parse_qs(parsed.query))
        elif parsed.path == "/api/price-library-template":
            self.handle_price_library_template()
        elif parsed.path == "/api/pm-warning-data":
            self.handle_pm_warning_data(parse_qs(parsed.query))
        elif parsed.path == "/api/pm-warning-capture-jobs":
            self.handle_pm_warning_capture_jobs(parse_qs(parsed.query))
        elif parsed.path == "/api/pm-warning-export":
            self.handle_pm_warning_export(parse_qs(parsed.query))
        elif parsed.path == "/api/blacklist-query":
            query = parse_qs(parsed.query).get("q", [""])[0]
            self.send_json(search_blacklist_entities(query))
        elif parsed.path == "/api/blacklist-export":
            self.handle_blacklist_export(parse_qs(parsed.query))
        elif parsed.path == "/api/blacklist-codex-jobs":
            self.send_json(list_blacklist_codex_jobs())
        elif parsed.path == "/api/company-relation":
            params = parse_qs(parsed.query)
            left = params.get("left", [""])[0]
            right = params.get("right", [""])[0]
            self.send_json(judge_company_relation(left, right))
        else:
            self.send_error(404, "Not found")

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/api/uploads":
            self.handle_upload()
        elif parsed.path == "/api/issues":
            self.handle_create_issue()
        elif parsed.path == "/api/material-price-tasks":
            self.handle_create_material_price_task()
        elif parsed.path == "/api/material-price-import":
            self.handle_import_material_price_task()
        elif parsed.path == "/api/price-library":
            self.handle_create_price_library_row()
        elif parsed.path == "/api/price-library-import":
            self.handle_import_price_library()
        elif parsed.path == "/api/price-library-import-pm-labor":
            self.handle_import_pm_labor_price_library()
        elif parsed.path == "/api/price-library-import-pm-professional":
            self.handle_import_pm_professional_price_library()
        elif parsed.path == "/api/price-library-auto-extract":
            self.handle_auto_extract_price_library()
        elif parsed.path == "/api/pm-warning-capture":
            self.handle_create_pm_warning_capture()
        elif parsed.path == "/api/blacklist-codex-jobs":
            self.handle_create_blacklist_codex_job()
        elif parsed.path == "/api/blacklist-result-template":
            self.handle_create_blacklist_result_template()
        else:
            self.send_error(404, "Not found")

    def do_PUT(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/api/price-library":
            self.handle_update_price_library_row(parse_qs(parsed.query))
        else:
            self.send_error(404, "Not found")

    def do_DELETE(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/api/uploads":
            self.handle_delete_upload(parse_qs(parsed.query))
        elif parsed.path == "/api/issues":
            self.handle_delete_issue(parse_qs(parsed.query))
        elif parsed.path == "/api/material-price-tasks":
            self.handle_delete_material_price_task(parse_qs(parsed.query))
        elif parsed.path == "/api/price-library":
            self.handle_delete_price_library_row(parse_qs(parsed.query))
        elif parsed.path == "/api/pm-warning-data":
            params = parse_qs(parsed.query)
            platform = clean_text(params.get("platform", ["old-pm"])[0]) or "old-pm"
            self.send_json(clear_pm_warning_data(platform))
        else:
            self.send_error(404, "Not found")

    def serve_file(self, file_path: Path, content_type: str) -> None:
        if not file_path.exists() or not file_path.is_file():
            self.send_error(404, "Not found")
            return
        body = file_path.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def send_json(self, payload: Any, status: int = 200) -> None:
        body = json.dumps(payload, ensure_ascii=False, default=json_default).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def handle_price_library(self, params: dict[str, list[str]]) -> None:
        self.send_json(price_library_payload(params))

    def handle_create_price_library_row(self) -> None:
        length = int(self.headers.get("Content-Length", "0") or "0")
        payload = parse_json_body(self.headers, self.rfile.read(length))
        row = normalize_price_library_row(payload)
        if not row.get("item") or not row.get("unit") or not row.get("price"):
            self.send_json({"error": "分项名称、单位和单价不能为空。"}, 400)
            return
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        row.update({
            "id": uuid.uuid4().hex,
            "createdAt": now,
            "updatedAt": now,
            "source": row.get("source") or "手工录入",
        })
        with price_library_connect() as conn:
            insert_price_library_row(conn, row)
        self.send_json(row)

    def handle_update_price_library_row(self, params: dict[str, list[str]]) -> None:
        if not can_edit_price_library(self.headers):
            self.send_json({"error": "当前账号没有价格库修改权限。"}, 403)
            return
        row_id = clean_text(params.get("id", [""])[0])
        if not row_id:
            self.send_json({"error": "价格记录编号不能为空。"}, 400)
            return
        length = int(self.headers.get("Content-Length", "0") or "0")
        payload = parse_json_body(self.headers, self.rfile.read(length))
        row = normalize_price_library_row(payload)
        if not row.get("item") or not row.get("unit") or not row.get("price"):
            self.send_json({"error": "分项名称、单位和单价不能为空。"}, 400)
            return
        with price_library_connect() as conn:
            updated = update_price_library_row(conn, row_id, row)
        if not updated:
            self.send_json({"error": "未找到要修改的价格记录。"}, 404)
            return
        self.send_json(updated)

    def handle_import_price_library(self) -> None:
        length = int(self.headers.get("Content-Length", "0") or "0")
        form = parse_multipart_form(self.headers, self.rfile.read(length))
        file_items = form.get("file")
        if file_items is not None and not isinstance(file_items, list):
            file_items = [file_items]
        if not file_items:
            self.send_json({"error": "请选择要导入的价格库文件。"}, 400)
            return
        category = normalize_price_library_category(form.get("category"))
        source = clean_text(form.get("source")) or "Excel导入"
        imported: list[dict[str, Any]] = []
        imported_sheets: set[str] = set()
        PRICE_LIBRARY_SOURCE_DIR.mkdir(parents=True, exist_ok=True)
        for file_item in file_items:
            filename = Path(clean_text(file_item.get("filename"))).name
            suffix = Path(filename).suffix.lower()
            if suffix not in {".xlsx", ".xls", ".json"}:
                self.send_json({"error": "价格库导入仅支持 .xlsx、.xls 和 .json 文件。"}, 400)
                return
            import_id = uuid.uuid4().hex
            stored_path = PRICE_LIBRARY_SOURCE_DIR / f"{import_id}_{filename}"
            with stored_path.open("wb") as f:
                f.write(file_item.get("content", b""))
            parse_path = convert_xls_to_xlsx(stored_path) if suffix == ".xls" else stored_path
            for row in read_price_library_file(parse_path):
                row["category"] = normalize_price_library_category(row.get("category") or category)
                row["source"] = row.get("source") or source
                row["id"] = row.get("id") or stable_uploaded_price_library_id(filename, row)
                row["createdAt"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                row["updatedAt"] = row["createdAt"]
                row["_sourceFile"] = filename
                if clean_text(row.get("_sourceSheet")):
                    imported_sheets.add(f"{filename}/{clean_text(row.get('_sourceSheet'))}")
                imported.append(row)
        if not imported:
            self.send_json({"error": "导入文件中未识别到有效价格明细。"}, 400)
            return
        import_stats = {"inserted": 0, "updated": 0}
        with price_library_connect() as conn:
            for row in imported:
                result = insert_price_library_row(conn, row)
                if result == "updated":
                    import_stats["updated"] += 1
                else:
                    import_stats["inserted"] += 1
        self.send_json({"imported": len(imported), **import_stats, "rows": imported, "sheets": sorted(imported_sheets), "sheetCount": len(imported_sheets)})

    def handle_import_pm_labor_price_library(self) -> None:
        rows, sources = read_pm_labor_price_sources()
        if not rows:
            self.send_json({"error": "未找到可导入的劳务价格抓取成果，请先完成劳务价格抓取任务。"}, 400)
            return
        import_stats = {"inserted": 0, "updated": 0}
        with price_library_connect() as conn:
            for row in rows:
                result = insert_price_library_row(conn, row)
                if result == "updated":
                    import_stats["updated"] += 1
                else:
                    import_stats["inserted"] += 1
        self.send_json({"imported": len(rows), **import_stats, "sources": sources})

    def handle_import_pm_professional_price_library(self) -> None:
        rows, sources = read_pm_professional_price_sources()
        if not rows:
            self.send_json({"error": "未找到可导入的专业分包价格抓取成果，请先完成专业分包价格抓取任务。"}, 400)
            return
        import_stats = {"inserted": 0, "updated": 0}
        with price_library_connect() as conn:
            for row in rows:
                result = insert_price_library_row(conn, row)
                if result == "updated":
                    import_stats["updated"] += 1
                else:
                    import_stats["inserted"] += 1
        self.send_json({"imported": len(rows), **import_stats, "sources": sources})

    def handle_auto_extract_price_library(self) -> None:
        length = int(self.headers.get("Content-Length", "0") or "0")
        payload = parse_json_body(self.headers, self.rfile.read(length))
        platform = clean_text(payload.get("platform")) or "all"
        if platform not in {"all", "big-pm", "old-pm"}:
            self.send_json({"error": "自动提取平台选择无效。"}, 400)
            return
        include_labor = payload.get("labor") is not False
        include_professional = payload.get("professional") is not False
        if not include_labor and not include_professional:
            self.send_json({"error": "请选择劳务分包或专业分包。"}, 400)
            return

        import_stats = {"inserted": 0, "updated": 0}
        imported_rows: list[dict[str, Any]] = []
        summaries: list[dict[str, Any]] = []

        def add_bridge_rows(platform_key: str, platform_name: str, kind_name: str, rows: list[dict[str, Any]], sources: list[str]) -> None:
            prepared_rows: list[dict[str, Any]] = []
            for row in rows:
                source = clean_text(row.get("source")) or kind_name
                prepared_rows.append({
                    **row,
                    "source": f"自动提取/{platform_name}/{source}",
                    "_sourceFile": clean_text(row.get("_sourceFile")) or "Codex桥接",
                })
            imported_rows.extend(prepared_rows)
            summaries.append({
                "platform": platform_key,
                "name": platform_name,
                "kind": kind_name,
                "imported": len(prepared_rows),
                "sources": sources,
                "message": f"{platform_name}{kind_name}：读取 {len(prepared_rows)} 条",
            })

        if platform in {"all", "big-pm"}:
            if include_labor:
                labor_rows, labor_sources = read_pm_labor_price_sources()
                add_bridge_rows("big-pm", "PM平台", "劳务分包", labor_rows, labor_sources)
            if include_professional:
                professional_rows, professional_sources = read_pm_professional_price_sources()
                add_bridge_rows("big-pm", "PM平台", "专业分包", professional_rows, professional_sources)

        if platform in {"all", "old-pm"}:
            summaries.append({
                "platform": "old-pm",
                "name": "四版平台",
                "kind": "全部",
                "imported": 0,
                "sources": [],
                "message": "四版平台：端口已预留，直连提取待接入",
            })

        with price_library_connect() as conn:
            for row in imported_rows:
                result = insert_price_library_row(conn, row)
                if result == "updated":
                    import_stats["updated"] += 1
                else:
                    import_stats["inserted"] += 1

        if not imported_rows:
            message = "未读取到可入库的价格数据。"
        else:
            message = f"自动提取已处理 {len(imported_rows)} 条价格数据。"
        self.send_json({
            "imported": len(imported_rows),
            **import_stats,
            "platforms": summaries,
            "message": message,
        })

    def handle_price_library_template(self) -> None:
        EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        output_path = EXPORT_DIR / "分包价格库_标准导入模板.xlsx"
        write_price_library_template_workbook(output_path)
        body = output_path.read_bytes()
        encoded_name = quote(output_path.name)
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", f'attachment; filename="{encoded_name}"')
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def handle_price_library_export(self, params: dict[str, list[str]]) -> None:
        rows = load_price_library_rows(params)
        output_path = EXPORT_DIR / f"分包价格库_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        write_price_library_workbook(rows, output_path)
        body = output_path.read_bytes()
        encoded_name = quote(output_path.name)
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", f'attachment; filename="{encoded_name}"')
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def handle_delete_price_library_row(self, params: dict[str, list[str]]) -> None:
        row_id = clean_text(params.get("id", [""])[0])
        if not row_id:
            self.send_json({"error": "价格记录编号不能为空。"}, 400)
            return
        with price_library_connect() as conn:
            before = conn.execute("SELECT COUNT(*) FROM price_library").fetchone()[0]
            conn.execute("DELETE FROM price_library WHERE id = ?", (row_id,))
            after = conn.execute("SELECT COUNT(*) FROM price_library").fetchone()[0]
        self.send_json({"deleted": before - after})

    def handle_pm_warning_data(self, params: dict[str, list[str]]) -> None:
        platform = clean_text(params.get("platform", ["old-pm"])[0]) or "old-pm"
        period = clean_text(params.get("period", [""])[0])
        self.send_json(load_pm_warning_data(platform, period))

    def handle_pm_warning_capture_jobs(self, params: dict[str, list[str]]) -> None:
        job_id = clean_text(params.get("id", [""])[0])
        jobs = list_pm_warning_jobs(job_id)
        if job_id:
            if not jobs:
                self.send_json({"error": "未找到红蓝预警抓数任务。"}, 404)
                return
            self.send_json(jobs[0])
            return
        self.send_json(jobs)

    def handle_create_pm_warning_capture(self) -> None:
        length = int(self.headers.get("Content-Length", "0") or "0")
        payload = parse_json_body(self.headers, self.rfile.read(length))
        platform = pm_warning_safe_platform(payload.get("platform") or "old-pm")
        cdp_ready, cdp_status = ensure_pm_warning_debug_browser(platform)
        if not cdp_ready:
            self.send_json({"error": cdp_status.get("message", f"{pm_warning_platform_name(platform)}浏览器连接失败。"), "cdpStatus": cdp_status}, 400)
            return
        period = clean_text(payload.get("period")) or datetime.now().strftime("%Y-%m")
        periods = pm_warning_period_window(period)
        job_id = uuid.uuid4().hex
        result_path = pm_warning_latest_path(platform)
        platform_name = pm_warning_platform_name(platform)
        job = {
            "id": job_id,
            "type": "pm-warning-capture",
            "platform": platform,
            "platformName": platform_name,
            "period": period,
            "periods": periods,
            "status": "等待抓数",
            "message": f"任务已创建，等待连接已登录{platform_name}页面。",
            "createdAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "resultPath": str(result_path),
            "summaryPath": str(PM_WARNING_RESULT_DIR / f"red_blue_company_summary_{platform}.csv"),
            "detailPath": str(PM_WARNING_RESULT_DIR / f"red_blue_projects_{platform}.csv"),
            "totals": {},
        }
        job_path = write_pm_warning_job(job)
        started, message = start_pm_warning_worker(job, job_path)
        job["message"] = message
        if not started:
            job["status"] = "启动失败"
        write_pm_warning_job(job)
        self.send_json(job, 202 if started else 400)

    def handle_pm_warning_export(self, params: dict[str, list[str]]) -> None:
        platform = clean_text(params.get("platform", ["old-pm"])[0]) or "old-pm"
        period = clean_text(params.get("period", [""])[0]) or datetime.now().strftime("%Y-%m")
        safe_period = re.sub(r"[^0-9-]", "", period) or "all"
        output_path = EXPORT_DIR / f"PM平台红蓝预警报表_{safe_period}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xls"
        ok, message = export_pm_warning_report(platform, period, output_path)
        if not ok:
            self.send_json({"error": message}, 500)
            return
        body = output_path.read_bytes()
        ascii_name = f"pm-warning-{safe_period}.xls"
        encoded_name = quote(output_path.name)
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.ms-excel")
        self.send_header("Content-Disposition", f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{encoded_name}")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def handle_upload(self) -> None:
        length = int(self.headers.get("Content-Length", "0") or "0")
        form = parse_multipart_form(self.headers, self.rfile.read(length))
        report_type = clean_text(form.get("reportType"))
        period = clean_text(form.get("period"))
        company = standard_company_name(clean_text(form.get("company")))
        file_items = form.get("file")
        if file_items is not None and not isinstance(file_items, list):
            file_items = [file_items]
        if report_type not in PLUGINS or not period or not file_items:
            self.send_json({"error": "报表类型、期次和文件均不能为空。"}, 400)
            return
        records = []
        for file_item in file_items:
            filename = Path(clean_text(file_item.get("filename"))).name
            file_company = company or company_from_filename(filename)
            suffix = Path(filename).suffix.lower()
            if suffix not in {".xlsx", ".xls"}:
                records.append(self.failed_upload_record(report_type, period, file_company, filename, "", "当前版本支持 .xlsx 和 .xls 文件。"))
                continue
            upload_id = uuid.uuid4().hex
            stored_path = UPLOAD_DIR / f"{upload_id}_{filename}"
            with stored_path.open("wb") as f:
                f.write(file_item.get("content", b""))
            try:
                parse_path = convert_xls_to_xlsx(stored_path)
                result = PLUGINS[report_type].parse(parse_path, file_company)
                if report_type == "risk-level" and not file_company:
                    detected_names = {standard_company_name(row.get("company", "")) for row in result.rows if standard_company_name(row.get("company", ""))}
                    if len(detected_names) == 1:
                        file_company = next(iter(detected_names))
                    elif len(detected_names) > 1:
                        raise ValueError("该文件包含多个单位，请上传单个基层单位报表，或使用文件名标明单位。")
                    else:
                        raise ValueError("无法从文件名或表内单位名称识别公司。")
                for row in result.rows:
                    row["company"] = file_company or standard_company_name(row.get("company", ""))
                    row["_source_path"] = str(parse_path)
                    row["_source_file"] = filename
                detected_company = file_company or standard_company_name(clean_text(result.rows[0].get("company"))) if result.rows else file_company
                status = "校验通过" if not result.errors else "校验失败"
                record = {
                    "id": upload_id,
                    "reportType": report_type,
                    "period": period,
                    "company": detected_company,
                    "filename": filename,
                    "storedPath": str(parse_path),
                    "originalPath": str(stored_path),
                    "uploadedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "status": status,
                    "rows": result.rows,
                    "errors": result.errors,
                    "warnings": result.warnings,
                    "sheetNames": result.sheet_names,
                }
            except Exception as exc:
                record = self.failed_upload_record(report_type, period, file_company, filename, str(stored_path), f"解析失败：{exc}", upload_id)
            records.append(record)
        db = load_db()
        for record in reversed(records):
            db["uploads"].insert(0, record)
        save_db(db)
        self.send_json(records)

    def failed_upload_record(self, report_type: str, period: str, company: str, filename: str, stored_path: str, error: str, upload_id: str | None = None) -> dict[str, Any]:
        return {
            "id": upload_id or uuid.uuid4().hex,
            "reportType": report_type,
            "period": period,
            "company": company,
            "filename": filename,
            "storedPath": stored_path,
            "uploadedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "status": "校验失败",
            "rows": [],
            "errors": [error],
            "warnings": [],
            "sheetNames": [],
        }

    def selected_rows(self, params: dict[str, list[str]]) -> tuple[ReportPlugin | None, str, list[dict[str, Any]]]:
        report_type = params.get("type", [""])[0]
        period = params.get("period", [""])[0]
        plugin = PLUGINS.get(report_type)
        rows: list[dict[str, Any]] = []
        if plugin:
            for upload in self.selected_uploads(report_type, period):
                rows.extend(upload.get("rows", []))
        if report_type == "settlement":
            rows = unique_settlement_rows(rows)
        return plugin, period, rows

    def selected_uploads(self, report_type: str, period: str) -> list[dict[str, Any]]:
        latest_by_company: dict[str, dict[str, Any]] = {}
        for upload in load_db().get("uploads", []):
            if upload.get("reportType") != report_type:
                continue
            if period and upload.get("period") != period:
                continue
            if upload.get("errors"):
                continue
            company = standard_company_name(upload.get("company", "")) or upload.get("company", "")
            if not company:
                company = upload.get("filename", "")
            if company not in latest_by_company:
                latest_by_company[company] = upload
        return sorted(
            latest_by_company.values(),
            key=lambda upload: company_sort_key(
                standard_company_name(upload.get("company", "")) or upload.get("company", ""),
                report_type,
            ),
        )

    def handle_delete_upload(self, params: dict[str, list[str]]) -> None:
        upload_id = params.get("id", [""])[0]
        report_type = params.get("type", [""])[0]
        period = params.get("period", [""])[0]
        company = params.get("company", [""])[0]
        clear_all = params.get("scope", [""])[0] == "all"
        db = load_db()
        kept = []
        deleted = []
        for upload in db.get("uploads", []):
            should_delete = False
            if clear_all:
                should_delete = True
            elif upload_id:
                should_delete = upload.get("id") == upload_id
            elif report_type and period and not company:
                should_delete = upload.get("reportType") == report_type and upload.get("period") == period
            elif report_type and period and company:
                should_delete = (
                    upload.get("reportType") == report_type
                    and upload.get("period") == period
                    and upload.get("company") == company
                )
            if should_delete:
                deleted.append(upload)
            else:
                kept.append(upload)
        db["uploads"] = kept
        save_db(db)
        for upload in deleted:
            self.try_delete_upload_file(upload.get("storedPath", ""))
            self.try_delete_upload_file(upload.get("originalPath", ""))
        self.send_json({"deleted": len(deleted)})

    def handle_issues(self, params: dict[str, list[str]]) -> None:
        report_type = params.get("type", [""])[0]
        period = params.get("period", [""])[0]
        issues = []
        for issue in load_db().get("issues", []):
            if report_type and issue.get("reportType") != report_type:
                continue
            if period and issue.get("period") != period:
                continue
            issues.append(issue)
        self.send_json(issues)

    def handle_create_issue(self) -> None:
        length = int(self.headers.get("Content-Length", "0") or "0")
        payload = parse_json_body(self.headers, self.rfile.read(length))
        report_type = clean_text(payload.get("reportType"))
        period = clean_text(payload.get("period"))
        issue_type = clean_text(payload.get("issueType")) or "其他"
        company = standard_company_name(clean_text(payload.get("company"))) or clean_text(payload.get("company"))
        description = clean_text(payload.get("description"))
        if not report_type or not period or not description:
            self.send_json({"error": "报表类型、期次和问题说明不能为空。"}, 400)
            return
        issue = {
            "id": uuid.uuid4().hex,
            "reportType": report_type,
            "period": period,
            "issueType": issue_type,
            "company": company,
            "description": description,
            "status": "待处理",
            "createdAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
        db = load_db()
        db.setdefault("issues", []).insert(0, issue)
        save_db(db)
        self.send_json(issue)

    def handle_delete_issue(self, params: dict[str, list[str]]) -> None:
        issue_id = params.get("id", [""])[0]
        db = load_db()
        before = len(db.get("issues", []))
        db["issues"] = [issue for issue in db.get("issues", []) if issue.get("id") != issue_id]
        save_db(db)
        self.send_json({"deleted": before - len(db["issues"])})

    def handle_missing_companies(self, params: dict[str, list[str]]) -> None:
        report_type = params.get("type", [""])[0]
        period = params.get("period", [""])[0]
        uploaded: set[str] = set()
        for upload in self.selected_uploads(report_type, period):
            company = standard_company_name(upload.get("company", ""))
            if company:
                uploaded.add(company)
        expected = company_order(report_type)
        missing = [company for company in expected if company not in uploaded]
        self.send_json({"expected": expected, "uploaded": sorted(uploaded, key=lambda company: company_sort_key(company, report_type)), "missing": missing})

    def try_delete_upload_file(self, stored_path: str) -> None:
        try:
            path = Path(stored_path).resolve()
            upload_root = UPLOAD_DIR.resolve()
            if upload_root in path.parents and path.exists():
                path.unlink()
        except OSError:
            pass

    def handle_summary(self, params: dict[str, list[str]]) -> None:
        plugin, period, rows = self.selected_rows(params)
        if not plugin:
            self.send_json({"error": "未知报表类型。"}, 400)
            return
        if isinstance(plugin, SettlementPlugin):
            plugin.sync_month_settlement_marker_issues(rows, period)
        self.send_json(plugin.aggregate(rows))

    def handle_export(self, params: dict[str, list[str]]) -> None:
        plugin, period, rows = self.selected_rows(params)
        if not plugin:
            self.send_json({"error": "未知报表类型。"}, 400)
            return
        if isinstance(plugin, SettlementPlugin):
            plugin.sync_month_settlement_marker_issues(rows, period)
        is_settlement = isinstance(plugin, SettlementPlugin)
        if is_settlement:
            output_path = plugin.export_result_package(rows, period or "全部")
        else:
            output_path = EXPORT_DIR / f"{plugin.key}_{period or 'all'}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
            plugin.export(rows, output_path, period or "全部")
        body = output_path.read_bytes()
        encoded_name = output_path.name.encode("utf-8").decode("latin-1", errors="ignore")
        self.send_response(200)
        self.send_header(
            "Content-Type",
            "application/zip" if is_settlement else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.send_header("Content-Disposition", f'attachment; filename="{encoded_name}"')
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def handle_export_settled_details(self, params: dict[str, list[str]]) -> None:
        report_type = params.get("type", [""])[0]
        if report_type != "settlement":
            self.send_json({"error": "定案明细仅支持竣工结算报表。"}, 400)
            return
        plugin, period, rows = self.selected_rows(params)
        if not isinstance(plugin, SettlementPlugin):
            self.send_json({"error": "未知报表类型。"}, 400)
            return
        output_path = EXPORT_DIR / f"settlement_settled_details_{period or 'all'}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        plugin.export_settled_details(rows, output_path, period or "")
        body = output_path.read_bytes()
        encoded_name = output_path.name.encode("utf-8").decode("latin-1", errors="ignore")
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", f'attachment; filename="{encoded_name}"')
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def handle_material_price_tasks(self) -> None:
        db = load_db()
        changed = False
        active_capture_statuses = {
            "后台采集中",
            "连接平台中",
            "等待平台登录",
            "等待已登录浏览器",
            "等待平台页面",
            "平台采集中",
            "四源采集中",
            "集采源采集中",
            "材料中标信息采集中",
            "等待后台采集",
        }
        build_candidate_id = next(
            (
                task.get("id")
                for task in db.get("materialPriceTasks", [])
                if task.get("status") in {"已导出源数据", "正在生成表格", "已采集"}
            ),
            "",
        )
        for task in db.get("materialPriceTasks", []):
            existing_output = Path(task.get("outputPath", ""))
            if task.get("status") == "已建表" and existing_output.exists():
                continue
            if task.get("status") in {"已导出源数据", "正在生成表格"} or task.get("sourcePaths") or task.get("sourceDir"):
                if build_candidate_id and task.get("id") != build_candidate_id and task.get("status") not in active_capture_statuses:
                    continue
                if task.get("status") in active_capture_statuses and not task.get("sourcePaths"):
                    continue
                if has_material_bid_info_source(task) and task.get("status") in {"已采集", "已建表"}:
                    rows = load_material_price_rows(task)
                    if rows:
                        output_path = Path(task.get("outputPath") or "")
                        if not output_path.exists() or not output_path.name.startswith("material_bid_skill_"):
                            output_path = EXPORT_DIR / f"material_bid_skill_{task.get('startDate','')}_{task.get('endDate','')}_{task.get('id','')[:8]}.xlsx"
                        stats = write_platform_material_skill_workbook(task, rows, output_path)
                        task["rows"] = len(rows)
                        task["comparisonRows"] = stats.get("average", 0)
                        task["outputPath"] = str(output_path)
                        task["downloadContent"] = "平台材料中标信息生成的技能四表"
                        task["status"] = "已建表"
                        task["message"] = f"已按材料价格技能自动建表：源记录{len(rows)}条，表三{stats.get('details', 0)}行，表二{stats.get('average', 0)}项，表一{stats.get('high', 0)}项，表四{stats.get('dongtai', 0)}项。"
                        changed = True
                        continue
                result = build_bid_workbook_from_task_sources(task)
                if result.get("built"):
                    stats = result.get("stats", {})
                    task["rows"] = stats.get("details", 0)
                    task["comparisonRows"] = stats.get("average", 0)
                    task["outputPath"] = str(result.get("outputPath"))
                    task["downloadContent"] = "小程序按集采平台源表生成的技能四表"
                    validation_errors = stats.get("validationErrors") or []
                    if validation_errors:
                        task["status"] = "校验未通过"
                        task["message"] = f"集采源表已自动导出并成表，但技能校验未通过：{'；'.join(validation_errors[:5])}"
                    else:
                        task["status"] = "已建表"
                        task["message"] = f"集采源表已自动导出，小程序已按技能成表：表三{stats.get('details', 0)}行，表二{stats.get('average', 0)}行，表一{stats.get('high', 0)}行，表四{stats.get('dongtai', 0)}行；技能校验通过。"
                    changed = True
                    continue
                if result.get("missing") and task.get("status") not in {"四源采集中", "集采源采集中", "等待四源导出适配", "等待集采导出适配", "等待平台登录", "连接平台中"}:
                    task["status"] = "缺少源数据"
                    task["downloadContent"] = f"不可下载：缺少{'、'.join(result.get('missing') or [])}"
                    task["message"] = f"已发现部分源文件，但缺少：{'、'.join(result.get('missing') or [])}。"
                    changed = True
                    continue
            if task.get("status") not in {"待采集", "后台采集中", "连接平台中", "等待平台登录", "等待已登录浏览器", "等待平台页面", "等待平台采集适配", "等待后台采集", "平台采集中", "已采集", "四源采集中", "集采源采集中", "等待四源导出适配", "等待集采导出适配", "材料中标信息采集中", "等待材料中标信息导出适配"}:
                continue
            if task.get("source") == "集采平台":
                continue
            rows = load_material_price_rows(task)
            if not rows:
                continue
            comparison_rows = build_material_comparison_rows(rows)
            if task.get("rows") == len(rows) and task.get("comparisonRows") == len(comparison_rows) and task.get("status") == "已采集":
                continue
            task["rows"] = len(rows)
            task["comparisonRows"] = len(comparison_rows)
            task["status"] = "缺少源数据"
            task["downloadContent"] = "不可下载：缺少中标价格对比分析、中标价格查询、招标结果台帐"
            task["message"] = f"已采集到材料采购价格对比页 {len(rows)} 条价格记录，但这不是技能要求的中标价格横向对比源数据，不能据此生成正式四表。"
            changed = True
        if changed:
            save_db(db)
        self.send_json(db.get("materialPriceTasks", []))

    def handle_create_material_price_task(self) -> None:
        length = int(self.headers.get("Content-Length", "0") or "0")
        payload = parse_json_body(self.headers, self.rfile.read(length))
        start_date = clean_text(payload.get("startDate"))
        end_date = clean_text(payload.get("endDate"))
        company = clean_text(payload.get("company")) or "集团"
        if company == "全集团":
            company = "集团"
        if company not in material_price_companies():
            self.send_json({"error": "请选择有效的分公司。"}, 400)
            return
        if not start_date or not end_date:
            self.send_json({"error": "开始时间和结束时间不能为空。"}, 400)
            return
        if start_date > end_date:
            self.send_json({"error": "开始时间不能晚于结束时间。"}, 400)
            return
        task = {
            "id": uuid.uuid4().hex,
            "company": company,
            "startDate": start_date,
            "endDate": end_date,
            "keyword": clean_text(payload.get("keyword")),
            "source": "集采平台",
            "status": "后台采集中",
            "createdAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "rows": 0,
            "comparisonRows": 0,
            "outputPath": "",
            "downloadContent": "小程序按自动导出的集采平台源表生成技能四表",
            "message": "",
        }
        task["name"] = material_price_task_name(task)
        job_path = write_material_price_codex_job(task)
        task["codexJobPath"] = str(job_path)
        started, worker_message = start_material_platform_worker(task, job_path)
        task["message"] = f"已发布集采平台源表采集任务；采集完成后由小程序按技能自动建表。{worker_message}"
        if not started:
            task["status"] = "等待后台采集"
        output_path = EXPORT_DIR / f"material_bid_skill_{start_date}_{end_date}_{task['id'][:8]}.xlsx"
        task["outputPath"] = str(output_path)
        db = load_db()
        db.setdefault("materialPriceTasks", []).insert(0, task)
        save_db(db)
        self.send_json(task)

    def handle_import_material_price_task(self) -> None:
        length = int(self.headers.get("Content-Length", "0") or "0")
        form = parse_multipart_form(self.headers, self.rfile.read(length))
        start_date = clean_text(form.get("startDate"))
        end_date = clean_text(form.get("endDate"))
        company = clean_text(form.get("company")) or "集团"
        if company == "全集团":
            company = "集团"
        file_items = form.get("file")
        if file_items is not None and not isinstance(file_items, list):
            file_items = [file_items]
        if company not in material_price_companies():
            self.send_json({"error": "请选择有效的分公司。"}, 400)
            return
        if not start_date or not end_date:
            self.send_json({"error": "开始时间和结束时间不能为空。"}, 400)
            return
        if start_date > end_date:
            self.send_json({"error": "开始时间不能晚于结束时间。"}, 400)
            return
        if not file_items:
            self.send_json({"error": "请选择要导入的数据文件。"}, 400)
            return
        import_id = uuid.uuid4().hex
        stored_paths: list[Path] = []
        filenames: list[str] = []
        for file_item in file_items:
            filename = Path(clean_text(file_item.get("filename"))).name
            suffix = Path(filename).suffix.lower()
            if suffix not in {".xlsx", ".xls", ".json"}:
                self.send_json({"error": "仅支持 .xlsx、.xls 和 .json 数据文件。"}, 400)
                return
            stored_path = UPLOAD_DIR / f"{import_id}_{filename}"
            with stored_path.open("wb") as f:
                f.write(file_item.get("content", b""))
            stored_paths.append(stored_path)
            filenames.append(filename)
        sources = build_bid_source_sets(stored_paths)
        recognized_bid_sources = any(sources.get(key) for key in ("comparison", "query", "ledger", "pm"))
        missing_bid_sources = missing_bid_source_labels(sources)
        if recognized_bid_sources and missing_bid_sources:
            self.send_json({"error": f"平台源数据不完整，缺少：{'、'.join(missing_bid_sources)}。请同时导入集采平台的中标价格对比分析、中标价格查询、招标结果台帐。"}, 400)
            return
        if recognized_bid_sources and not missing_bid_sources:
            task = {
                "id": import_id,
                "company": company,
                "startDate": start_date,
                "endDate": end_date,
                "keyword": "",
                "source": "导入平台源数据",
                "status": "已建表",
                "createdAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "rows": 0,
                "comparisonRows": 0,
                "outputPath": "",
                "downloadContent": "小程序按集采平台源表生成的技能四表",
                "message": f"已导入平台源数据：{', '.join(filenames)}。",
            }
            task["name"] = material_price_task_name(task)
            output_path = EXPORT_DIR / f"bid_price_import_{start_date}_{end_date}_{task['id'][:8]}.xlsx"
            stats = write_bid_price_workbook(task, sources, output_path)
            task["rows"] = stats.get("details", 0)
            task["comparisonRows"] = stats.get("average", 0)
            task["outputPath"] = str(output_path)
            validation_errors = stats.get("validationErrors") or []
            if validation_errors:
                task["status"] = "校验未通过"
                task["message"] += f" 表三{stats.get('details', 0)}行，表二{stats.get('average', 0)}行，表一{stats.get('high', 0)}行，表四{stats.get('dongtai', 0)}行；校验问题：{'；'.join(validation_errors[:5])}"
            else:
                task["message"] += f" 表三{stats.get('details', 0)}行，表二{stats.get('average', 0)}行，表一{stats.get('high', 0)}行，表四{stats.get('dongtai', 0)}行；技能校验通过。"
            db = load_db()
            db.setdefault("materialPriceTasks", []).insert(0, task)
            save_db(db)
            self.send_json(task)
            return
        try:
            imported_rows = [normalize_material_row(row) for row in read_material_price_file(stored_paths[0])]
        except Exception as exc:
            self.send_json({"error": f"导入数据解析失败：{exc}"}, 400)
            return
        if company != "集团":
            for row in imported_rows:
                if not row.get("company"):
                    row["company"] = company
        imported_rows = [row for row in imported_rows if row.get("material")]
        if not imported_rows:
            self.send_json({"error": "导入文件中未识别到材料采购明细。"}, 400)
            return
        source_path = save_material_price_source(imported_rows, filenames[0])
        task = {
            "id": import_id,
            "company": company,
            "startDate": start_date,
            "endDate": end_date,
            "keyword": "",
            "source": "导入数据",
            "status": "已建表",
            "createdAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "rows": 0,
            "comparisonRows": 0,
            "outputPath": "",
            "downloadContent": "普通材料价格横向对比表",
            "message": f"已导入 {filenames[0]}，数据源：{source_path.name}。",
        }
        task["name"] = material_price_task_name(task)
        rows = load_material_price_rows(task)
        task["rows"] = len(rows)
        task["comparisonRows"] = len(build_material_comparison_rows(rows))
        output_path = EXPORT_DIR / f"material_price_import_{start_date}_{end_date}_{task['id'][:8]}.xlsx"
        export_material_price_task(task, rows, output_path)
        task["outputPath"] = str(output_path)
        db = load_db()
        db.setdefault("materialPriceTasks", []).insert(0, task)
        save_db(db)
        self.send_json(task)

    def handle_material_price_export(self, params: dict[str, list[str]]) -> None:
        task_id = params.get("id", [""])[0]
        db = load_db()
        task = next((item for item in db.get("materialPriceTasks", []) if item.get("id") == task_id), None)
        if not task:
            self.send_json({"error": "未找到材料价格任务。"}, 404)
            return
        output_path = Path(task.get("outputPath", ""))
        if task.get("source") == "导入平台源数据" and output_path.exists():
            body = output_path.read_bytes()
            encoded_name = output_path.name.encode("utf-8").decode("latin-1", errors="ignore")
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", f'attachment; filename="{encoded_name}"')
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        if task.get("sourcePaths") or task.get("sourceDir"):
            if has_material_bid_info_source(task) and task.get("status") in {"已采集", "已建表"}:
                rows = load_material_price_rows(task)
                if rows:
                    output_path = Path(task.get("outputPath") or EXPORT_DIR / f"material_bid_skill_{task.get('startDate','')}_{task.get('endDate','')}_{task.get('id','')[:8]}.xlsx")
                    if not output_path.exists() or not output_path.name.startswith("material_bid_skill_"):
                        output_path = EXPORT_DIR / f"material_bid_skill_{task.get('startDate','')}_{task.get('endDate','')}_{task.get('id','')[:8]}.xlsx"
                    stats = write_platform_material_skill_workbook(task, rows, output_path)
                    task["rows"] = len(rows)
                    task["comparisonRows"] = stats.get("average", 0)
                    task["outputPath"] = str(output_path)
                    task["downloadContent"] = "平台材料中标信息生成的技能四表"
                    task["status"] = "已建表"
                    task["message"] = f"已按材料价格技能自动建表：源记录{len(rows)}条，表三{stats.get('details', 0)}行，表二{stats.get('average', 0)}项，表一{stats.get('high', 0)}项，表四{stats.get('dongtai', 0)}项。"
                    save_db(db)
                    body = output_path.read_bytes()
                    encoded_name = output_path.name.encode("utf-8").decode("latin-1", errors="ignore")
                    self.send_response(200)
                    self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                    self.send_header("Content-Disposition", f'attachment; filename="{encoded_name}"')
                    self.send_header("Content-Length", str(len(body)))
                    self.end_headers()
                    self.wfile.write(body)
                    return
            result = build_bid_workbook_from_task_sources(task)
            if result.get("built"):
                stats = result.get("stats", {})
                output_path = Path(result.get("outputPath"))
                task["rows"] = stats.get("details", 0)
                task["comparisonRows"] = stats.get("average", 0)
                task["outputPath"] = str(output_path)
                task["downloadContent"] = "小程序按集采平台源表生成的技能四表"
                validation_errors = stats.get("validationErrors") or []
                if validation_errors:
                    task["status"] = "校验未通过"
                    task["message"] = f"集采源表已自动导出并成表，但技能校验未通过：{'；'.join(validation_errors[:5])}"
                else:
                    task["status"] = "已建表"
                    task["message"] = f"集采源表已自动导出，小程序已按技能成表：表三{stats.get('details', 0)}行，表二{stats.get('average', 0)}行，表一{stats.get('high', 0)}行，表四{stats.get('dongtai', 0)}行；技能校验通过。"
                save_db(db)
            else:
                missing = result.get("missing") or []
                task["status"] = "缺少源数据"
                task["downloadContent"] = f"不可下载：缺少{'、'.join(missing)}"
                task["message"] = f"自动成表缺少源数据：{'、'.join(missing)}。"
                save_db(db)
                self.send_json({"error": task["message"]}, 409)
                return
            if output_path.exists():
                body = output_path.read_bytes()
                encoded_name = output_path.name.encode("utf-8").decode("latin-1", errors="ignore")
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f'attachment; filename="{encoded_name}"')
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
                return
        if task.get("source") == "集采平台":
            task["status"] = "缺少源数据"
            task["downloadContent"] = "不可下载：缺少中标价格对比分析、中标价格查询、招标结果台帐"
            task["message"] = "当前任务尚未取得完整集采源表。技能要求取得集采平台中标价格对比分析、中标价格查询、招标结果台帐后生成正式四表。"
            save_db(db)
            self.send_json({"error": task["message"]}, 409)
            return
        rows = load_material_price_rows(task)
        if rows:
            task["rows"] = len(rows)
            task["comparisonRows"] = len(build_material_comparison_rows(rows))
            task["status"] = "已建表"
            if task.get("source") == "导入数据":
                task["downloadContent"] = "普通材料价格横向对比表"
                task["message"] = "已按导入数据生成普通材料采购价格横向对比表。"
                output_path = EXPORT_DIR / f"material_price_import_{task.get('startDate','')}_{task.get('endDate','')}_{task.get('id','')[:8]}.xlsx"
                export_material_price_task(task, rows, output_path)
            else:
                task["downloadContent"] = "小程序按集采平台源表生成的技能四表"
                output_path = EXPORT_DIR / f"bid_price_platform_{task.get('startDate','')}_{task.get('endDate','')}_{task.get('id','')[:8]}.xlsx"
                stats = write_platform_material_skill_workbook(task, rows, output_path)
                task["message"] = f"已按技能四表结构重新建表：表三{stats.get('details', 0)}行，表二{stats.get('average', 0)}行，表一{stats.get('high', 0)}行。"
            task["outputPath"] = str(output_path)
            save_db(db)
        elif not output_path.exists():
            export_material_price_task(task, rows, output_path)
        body = output_path.read_bytes()
        encoded_name = output_path.name.encode("utf-8").decode("latin-1", errors="ignore")
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", f'attachment; filename="{encoded_name}"')
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def handle_blacklist_export(self, params: dict[str, list[str]]) -> None:
        query = params.get("q", [""])[0]
        result = search_blacklist_entities(query)
        output_path = EXPORT_DIR / f"blacklist_relation_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
        export_blacklist_result(result, output_path)
        body = output_path.read_bytes()
        encoded_name = output_path.name.encode("utf-8").decode("latin-1", errors="ignore")
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", f'attachment; filename="{encoded_name}"')
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def handle_create_blacklist_codex_job(self) -> None:
        length = int(self.headers.get("Content-Length", "0") or "0")
        payload = parse_json_body(self.headers, self.rfile.read(length))
        query = clean_text(payload.get("query"))
        if not query:
            self.send_json({"error": "请先输入要穿透查询的企业、法人或信用代码。"}, 400)
            return
        job = {
            "id": uuid.uuid4().hex,
            "query": query,
            "status": "待天眼查处理",
            "createdAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "estimatedAt": "",
            "jobPath": "",
            "resultPath": "",
            "message": "已生成天眼查穿透查询任务，正在启动天眼查执行器。",
        }
        job["estimatedAt"] = estimate_blacklist_finish(job["createdAt"])
        job_path = write_blacklist_codex_job(job)
        result_path = blacklist_result_path(job["id"])
        job["jobPath"] = str(job_path)
        job["resultPath"] = str(result_path)
        db = load_db()
        db.setdefault("blacklistCodexJobs", []).insert(0, job)
        save_db(db)
        if start_blacklist_tianyancha_worker(job["id"]):
            job["status"] = "启动天眼查"
            job["message"] = "已启动天眼查执行器；如新 Edge 窗口要求登录或验证，请完成后再次点击“天眼查穿透”。"
            db = load_db()
            for item in db.get("blacklistCodexJobs", []):
                if item.get("id") == job["id"]:
                    item.update({"status": job["status"], "message": job["message"]})
            save_db(db)
        else:
            job["message"] = "任务已生成，但天眼查执行器启动失败，请检查 Node 运行时。"
        self.send_json(job)

    def handle_create_blacklist_result_template(self) -> None:
        length = int(self.headers.get("Content-Length", "0") or "0")
        payload = parse_json_body(self.headers, self.rfile.read(length))
        job_id = clean_text(payload.get("id"))
        if not job_id:
            self.send_json({"error": "任务编号不能为空。"}, 400)
            return
        db = load_db()
        job = next((item for item in db.get("blacklistCodexJobs", []) if item.get("id") == job_id), None)
        if not job:
            self.send_json({"error": "未找到天眼查穿透任务。"}, 404)
            return
        output_path = write_blacklist_result_template(job)
        job["status"] = "待填写结果"
        job["message"] = "回写模板已生成，请填入天眼查穿透结果。"
        job["resultPath"] = str(output_path)
        save_db(db)
        self.send_json({"id": job_id, "resultPath": str(output_path), "status": job["status"], "message": job["message"]})

    def handle_delete_material_price_task(self, params: dict[str, list[str]]) -> None:
        task_id = params.get("id", [""])[0]
        if not task_id:
            self.send_json({"error": "任务编号不能为空。"}, 400)
            return
        db = load_db()
        tasks = db.get("materialPriceTasks", [])
        deleted = [task for task in tasks if task.get("id") == task_id]
        db["materialPriceTasks"] = [task for task in tasks if task.get("id") != task_id]
        save_db(db)
        for task in deleted:
            self.try_delete_export_file(task.get("outputPath", ""))
        self.send_json({"deleted": len(deleted)})

    def try_delete_export_file(self, stored_path: str) -> None:
        try:
            path = Path(stored_path).resolve()
            export_root = EXPORT_DIR.resolve()
            if export_root in path.parents and path.exists():
                path.unlink()
        except OSError:
            pass


def main() -> None:
    ensure_dirs()
    port = int(sys.argv[1]) if len(sys.argv) > 1 else 8899
    host = sys.argv[2] if len(sys.argv) > 2 else "127.0.0.1"
    server = ThreadingHTTPServer((host, port), AppHandler)
    print(f"ReportHub running at http://{host}:{port}")
    if port == 8899 and PM_WARNING_START_DEBUG_BROWSERS_ON_8899:
        def start_debug_browsers() -> None:
            for message in start_pm_warning_debug_browsers_for_server(port):
                print(message)

        threading.Thread(target=start_debug_browsers, name="pm-warning-debug-browsers", daemon=True).start()
    server.serve_forever()


if __name__ == "__main__":
    main()
